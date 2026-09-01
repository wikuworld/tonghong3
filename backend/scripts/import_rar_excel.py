#!/usr/bin/env python3
"""Import RAR Excel (Recapitulation Account Receivable) ke tabel invoices."""

from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

# Allow running from repo root or backend/
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.db import SessionLocal
from app.models import Invoice
from app.services.accounts_receivable_report import invoice_remark, parse_payment_term_days


def _cell_date(val) -> date | None:
    if val is None or val == "" or str(val).strip() == " ":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()[:10]
    try:
        return date.fromisoformat(s)
    except ValueError:
        return None


def _dec(val) -> Decimal | None:
    if val is None or val == "":
        return None
    try:
        return Decimal(str(val))
    except Exception:
        return None


def _str(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s or None


def _map_row(row: tuple) -> dict | None:
    if not row or row[0] in (None, ""):
        return None
    tax_no = _str(row[1])
    inv_no = _str(row[11])
    if not tax_no and not inv_no:
        return None
    term_raw = row[7]
    term_days = None
    if term_raw is not None and str(term_raw).strip() not in ("", " "):
        try:
            term_days = int(float(term_raw))
        except (TypeError, ValueError):
            term_days = parse_payment_term_days(str(term_raw))
    remark = _str(row[15]) or ""
    pay_amt = _dec(row[13])
    total = _dec(row[10])
    payment_status = "paid" if remark.upper() == "FINISH PAYMENT" else "unpaid"
    return {
        "tax_invoice_no": tax_no,
        "customer_name": _str(row[2]),
        "tax_invoice_date": _cell_date(row[3]),
        "receipt_no": _str(row[4]),
        "received_by": _str(row[5]),
        "received_date": _cell_date(row[6]),
        "term_of_payment": str(term_days) if term_days is not None else None,
        "due_date": _cell_date(row[8]),
        "total_amount": total,
        "invoice_number": inv_no,
        "po_number": _str(row[12]),
        "payment_amount": pay_amt,
        "payment_date": _cell_date(row[14]),
        "payment_remark": remark or None,
        "payment_status": payment_status,
        "nopen": _str(row[16]),
        "nomor_aju": _str(row[17]),
        "credit_note_no": _str(row[18]),
        "credit_note_date": _cell_date(row[19]),
        "credit_note_amount": _dec(row[20]),
        "oia_export_usd": _dec(row[22]),
        "invoice_date": _cell_date(row[3]),
    }


def import_rar(path: Path, *, dry_run: bool = False) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    stats = {"read": 0, "inserted": 0, "updated": 0, "skipped": 0}
    db = SessionLocal()
    try:
        for row in ws.iter_rows(min_row=4, values_only=True):
            mapped = _map_row(row)
            if not mapped:
                continue
            stats["read"] += 1
            tax_no = mapped.get("tax_invoice_no")
            inv_no = mapped.get("invoice_number")
            existing = None
            if tax_no:
                existing = db.scalars(select(Invoice).where(Invoice.tax_invoice_no == tax_no)).first()
            if existing is None and inv_no:
                existing = db.scalars(select(Invoice).where(Invoice.invoice_number == inv_no)).first()
            if dry_run:
                if existing:
                    stats["updated"] += 1
                else:
                    stats["inserted"] += 1
                continue
            if existing:
                for k, v in mapped.items():
                    if v is not None:
                        setattr(existing, k, v)
                stats["updated"] += 1
            else:
                db.add(Invoice(**mapped))
                stats["inserted"] += 1
            db.flush()
            if stats["read"] % 500 == 0:
                pass
        if not dry_run:
            db.commit()
    finally:
        wb.close()
        db.close()
    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description="Import RAR Excel ke invoices")
    parser.add_argument(
        "path",
        nargs="?",
        default=str(Path(__file__).resolve().parents[2] / "dokumentasi"),
        help="Path ke file .xlsx atau folder dokumentasi",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    p = Path(args.path)
    if p.is_dir():
        matches = list(p.glob("*RAR*2026*.xlsx"))
        if not matches:
            print("File RAR tidak ditemukan di", p)
            return 1
        p = matches[0]
    if not p.exists():
        print("File tidak ada:", p)
        return 1
    stats = import_rar(p, dry_run=args.dry_run)
    mode = "DRY-RUN" if args.dry_run else "IMPORT"
    print(f"{mode} {p.name}: {stats}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
