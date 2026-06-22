from io import BytesIO

from openpyxl import load_workbook

from app.routers.marketing import _outstanding_manual_workbook_bytes


def _sample_row(po_category: str, thi_order: str, lot: str, po_cust: str) -> dict:
    return {
        "order_date": "2026-03-05",
        "thi_order": thi_order,
        "customer_name": "PT TEST",
        "brand": "NIKE",
        "order_type": "PROD",
        "no_po": po_cust,
        "po_customer_ref": po_cust,
        "po_category": po_category,
        "material_name": "Split Suede",
        "thickness": "1.2-1.4",
        "emboss": "",
        "softness_type": "SOFT",
        "color": "BLACK",
        "softness": "2.0-3.0",
        "qty_sf": "100",
        "unit_price": "1.33",
        "color_reference": "COLOR WINDOW",
        "color_standard": "",
        "backing": "",
        "model": "",
        "style_article": "",
        "rtd_text": "",
        "line_etd": None,
        "stock_wh_qty": "0",
        "balance_qty": "100",
        "system_reference": lot,
        "detail_remark": "auto remark po",
        "remark_marketing": "manual outstanding",
        "status_material": "Follow-up material: belum",
        "delivery_1_qty": "100",
        "delivery_1_date": "2026-03-10",
    }


def test_manual_workbook_has_category_sheets():
    rows = [
        _sample_row("production", "PS26030501", "MYCLR028260305001", "IBO123"),
        _sample_row("claim", "PSC26030501", "BYCMP033260305001", "CLAIM-001"),
        _sample_row("sample", "S26030501", "SYCPW010260305001", "DEV-001"),
        _sample_row("subcont", "SC26030501", "QYCMP033260305001", "SUB-001"),
        _sample_row("return", "PSR26030501", "QYCPA010260305001", "RET-PO"),
    ]
    data = _outstanding_manual_workbook_bytes(rows)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["PROD SUEDE", "CLAIM SUEDE", "SAMPLE SUEDE", "SUBCONT", "RETUR"]

    prod = wb["PROD SUEDE"]
    assert prod.max_row == 2
    headers = [c.value for c in prod[1]]
    assert "NO PO" in headers
    assert "STATUS MATERIAL" in headers
    assert "21th Qty" in headers

    claim = wb["CLAIM SUEDE"]
    claim_headers = [c.value for c in claim[1]]
    assert "PO CUST" in claim_headers
    assert "STATUS MATERIAL" not in claim_headers
    assert "5th Qty" in claim_headers
    assert "6th Qty" not in claim_headers
