"""Unit tests for Pending Payment (Report Mr. Bob) report."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

from app.services.pending_payment_report import (
    PENDING_PAYMENT_HEADERS,
    build_pending_payment_payload,
    normalize_customer_key,
    render_pending_payment_xlsx,
)


def _inv(**kwargs):
    defaults = {
        "id": 1,
        "customer_name": "PT. SHOETOWN LIGUNG INDONESIA",
        "tax_invoice_date": date(2026, 5, 30),
        "invoice_date": date(2026, 5, 30),
        "due_date": date(2026, 7, 30),
        "term_of_payment": "30",
        "total_amount": Decimal("5000"),
        "payment_amount": Decimal("737.92"),
        "payment_date": None,
        "payment_remark": "PROMISE PAY AT 08/14",
        "payment_status": "unpaid",
        "credit_note_amount": None,
        "oia_export_usd": None,
    }
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def test_normalize_customer_key_collapses_spaces():
    assert normalize_customer_key("PT. NIKOMAS  GEMILANG") == "PT. NIKOMAS GEMILANG"
    assert normalize_customer_key("PT\xa0SHOETOWN MUSTIKA") == "PT SHOETOWN MUSTIKA"


def test_build_pending_payment_payload_overdue_only():
    overdue = _inv()
    paid = _inv(
        id=2,
        customer_name="PT. METRO PEARL INDONESIA",
        payment_amount=Decimal("100"),
        total_amount=Decimal("100"),
        payment_status="paid",
    )
    db = MagicMock()
    db.scalars.return_value.all.return_value = [overdue, paid]

    payload = build_pending_payment_payload(db, as_of=date(2026, 8, 28), year=2026)
    assert payload["year"] == 2026
    assert payload["summary"]["overdue_invoice_count"] == 1
    shoetown = next(r for r in payload["rows"] if r["customer"] == "PT. SHOETOWN LIGUNG INDONESIA")
    assert shoetown["amount_usd"] == 4262.08
    assert shoetown["month"] == "05"
    assert shoetown["receivable_date"] == "30/07/2026"
    assert shoetown["remarks"] == "PROMISE PAY AT 08/14"
    metro = next(r for r in payload["rows"] if r["customer"] == "PT. METRO PEARL INDONESIA")
    assert metro["amount_usd"] == 0


def test_render_pending_payment_xlsx_layout():
    payload = build_pending_payment_payload(
        MagicMock(scalars=MagicMock(return_value=MagicMock(all=MagicMock(return_value=[])))),
        as_of=date(2026, 8, 28),
        year=2026,
    )
    data = render_pending_payment_xlsx(payload)
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    assert "Pending Payment 2026" in str(ws["A1"].value)
    headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]
    assert headers == PENDING_PAYMENT_HEADERS
    assert len(PENDING_PAYMENT_HEADERS) == 10
    assert ws.cell(row=4, column=2).value == "WEALTH VAST LTD"
    assert ws.cell(row=60, column=2).value == "總計 ："
