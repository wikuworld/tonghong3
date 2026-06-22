"""Sales contract / sales order ringkas — PDF (ReportLab) & Word (.docx)."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from html import escape
from typing import Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Customer, PoDetail, PoHeader

COMPANY_LINES = (
    "PT. TONG HONG TANNERY INDONESIA",
    "Kawasan Industri Modern II JL. Modern Industri XIV/W-1, Serang-Banten",
)


def _fmt_date(d: date | None) -> str:
    if not d:
        return "—"
    return d.isoformat()


def _dec(v: Decimal | None) -> str:
    if v is None:
        return ""
    return str(v.quantize(Decimal("0.01")))


def gather_sales_contract_ctx(db: Session, po_id: int) -> dict[str, Any]:
    ph = db.get(PoHeader, po_id)
    if not ph:
        raise ValueError("po_not_found")
    cust = db.get(Customer, ph.customer_id) if ph.customer_id else None
    details = list(
        db.scalars(select(PoDetail).where(PoDetail.po_id == po_id).order_by(PoDetail.id.asc())).all()
    )
    lines: list[dict[str, Any]] = []
    total = Decimal(0)
    for i, d in enumerate(details, 1):
        qty = d.quantity_order or Decimal(0)
        up = d.unit_price
        amt: Decimal | None = None
        if up is not None:
            amt = (qty * up).quantize(Decimal("0.01"))
            total += amt
        lines.append(
            {
                "no": i,
                "material": (d.material_name or "").strip(),
                "color": (d.color or "").strip(),
                "qty": qty,
                "unit_price": up,
                "amount": amt,
                "remark": (d.detail_remark or "").strip(),
            }
        )
    return {
        "po_number": (ph.po_number or str(ph.id)).strip(),
        "po_customer_ref": (ph.po_customer_ref or "").strip(),
        "order_date": ph.order_date,
        "order_date_fmt": _fmt_date(ph.order_date),
        "buyer": (ph.buyer or "").strip(),
        "customer_name": (cust.name if cust and cust.name else ph.customer_name_snapshot or "").strip() or "—",
        "billing_address": (ph.billing_address or "").strip(),
        "delivery_address": (ph.delivery_address or "").strip(),
        "lines": lines,
        "total_amount": total,
    }


def render_sales_contract_pdf(ctx: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(name="b", parent=styles["Normal"], fontSize=10, leading=13)
    small = ParagraphStyle(name="s", parent=styles["Normal"], fontSize=9, leading=11)
    story: list[Any] = []
    for ln in COMPANY_LINES:
        story.append(Paragraph(f"<b>{escape(ln)}</b>", body))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("<b>SALES CONTRACT / SALES ORDER</b>", styles["Heading2"]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"<b>PO No.:</b> {escape(ctx['po_number'])}", small))
    if ctx.get("po_customer_ref"):
        story.append(Paragraph(f"<b>Customer ref:</b> {escape(ctx['po_customer_ref'])}", small))
    story.append(Paragraph(f"<b>Date:</b> {escape(ctx['order_date_fmt'])}", small))
    story.append(Paragraph(f"<b>Buyer:</b> {escape(ctx.get('buyer') or '—')}", small))
    story.append(Paragraph(f"<b>Customer:</b> {escape(ctx['customer_name'])}", small))
    if ctx.get("billing_address"):
        story.append(Paragraph(f"<b>Bill to:</b> {escape(ctx['billing_address'])}", small))
    if ctx.get("delivery_address"):
        story.append(Paragraph(f"<b>Ship to:</b> {escape(ctx['delivery_address'])}", small))
    story.append(Spacer(1, 5 * mm))

    head = ["#", "Material", "Color", "Qty", "U.Price", "Amount", "Line rmk."]
    data: list[list[str]] = [head]
    for ln in ctx["lines"]:
        data.append(
            [
                str(ln["no"]),
                (ln["material"] or "")[:40],
                (ln["color"] or "")[:16],
                _dec(ln["qty"]),
                _dec(ln["unit_price"]) if ln["unit_price"] is not None else "",
                _dec(ln["amount"]) if ln["amount"] is not None else "",
                (ln["remark"] or "")[:20],
            ]
        )
    data.append(["", "", "", "", "TOTAL", _dec(ctx["total_amount"]), ""])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(tbl)
    story.append(Spacer(1, 10 * mm))
    story.append(Paragraph("Authorized buyer / supplier signature: _______________________", small))
    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def render_sales_contract_docx(ctx: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    d = Document()
    for ln in COMPANY_LINES:
        p = d.add_paragraph()
        r = p.add_run(ln)
        r.bold = True
        r.font.size = Pt(11)
    d.add_paragraph()
    t = d.add_paragraph()
    tr = t.add_run("SALES CONTRACT / SALES ORDER")
    tr.bold = True
    tr.font.size = Pt(14)
    d.add_paragraph(f"PO No.: {ctx['po_number']}")
    if ctx.get("po_customer_ref"):
        d.add_paragraph(f"Customer ref: {ctx['po_customer_ref']}")
    d.add_paragraph(f"Date: {ctx['order_date_fmt']}")
    d.add_paragraph(f"Buyer: {ctx.get('buyer') or '—'}")
    d.add_paragraph(f"Customer: {ctx['customer_name']}")
    if ctx.get("billing_address"):
        d.add_paragraph(f"Bill to: {ctx['billing_address']}")
    if ctx.get("delivery_address"):
        d.add_paragraph(f"Ship to: {ctx['delivery_address']}")
    d.add_paragraph()

    table = d.add_table(rows=1 + len(ctx["lines"]) + 1, cols=7)
    hdr = ["#", "Material", "Color", "Qty", "U.Price", "Amount", "Remark"]
    for j, h in enumerate(hdr):
        table.rows[0].cells[j].text = h
    for i, ln in enumerate(ctx["lines"], start=1):
        row = table.rows[i].cells
        row[0].text = str(ln["no"])
        row[1].text = ln["material"] or ""
        row[2].text = ln["color"] or ""
        row[3].text = _dec(ln["qty"])
        row[4].text = _dec(ln["unit_price"]) if ln["unit_price"] is not None else ""
        row[5].text = _dec(ln["amount"]) if ln["amount"] is not None else ""
        row[6].text = ln["remark"] or ""
    last = table.rows[len(ctx["lines"]) + 1].cells
    last[4].text = "TOTAL"
    last[5].text = _dec(ctx["total_amount"])
    d.add_paragraph()
    d.add_paragraph("Authorized buyer / supplier signature: _______________________")

    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf.getvalue()
