"""Normalisasi dan impor master item kategori PU / Suide dari sheet marketing."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import MasterItem
from app.services.subkon_master_items import (
    EMPTY_MARKERS,
    _clean_text,
    _detect_csv_delimiter,
    _pick_row_value,
)
from app.services.thickness_format import canonical_thickness_range

PU_CATEGORY = "PU"
SUIDE_CATEGORY = "Suide"
PU_SUEDE_SHEET = "PU, SUEDE"


@dataclass
class PuSuedeImportResult:
    imported: int
    inserted: int
    updated: int
    skipped: int
    pu: int
    suide: int

    def to_dict(self) -> dict[str, int]:
        return {
            "imported": self.imported,
            "inserted": self.inserted,
            "updated": self.updated,
            "skipped": self.skipped,
            "pu": self.pu,
            "suide": self.suide,
        }


def infer_pu_suede_category(material_name: str | None) -> str:
    name = str(material_name or "").strip().upper()
    if "SUEDE" in name:
        return SUIDE_CATEGORY
    return PU_CATEGORY


def resolve_pu_suede_category(
    row: dict[str, str | None],
    material_name: str | None,
) -> str:
    raw = _pick_row_value(row, "category", "kategori", "jenis", "")
    if raw:
        alias = raw.strip().lower()
        if alias in ("suide", "suede"):
            return SUIDE_CATEGORY
        if alias == "pu":
            return PU_CATEGORY
    return infer_pu_suede_category(material_name)


def map_pu_suede_source_row(row: dict[str, str | None]) -> dict[str, str | None]:
    return {
        "item_code": _pick_row_value(row, "item_code", "KODE BARANG", "kode barang"),
        "thickness_range": _pick_row_value(row, "thickness_range", "thickness", "THICKNESS", "ketebalan"),
        "material_name": _pick_row_value(row, "material_name", "MATERIAL NAME", "material name"),
        "color": _pick_row_value(row, "color", "COLOR", "warna"),
        "emboss": _pick_row_value(row, "emboss", "EMBOSS"),
        "category": _pick_row_value(row, "category", "kategori", "jenis", ""),
    }


def normalize_pu_suede_row(row: dict[str, str | None]) -> dict[str, str | None]:
    mapped = map_pu_suede_source_row(row)
    thickness_raw = _clean_text(mapped.get("thickness_range"))
    thickness = canonical_thickness_range(thickness_raw) if thickness_raw else None
    if thickness == "":
        thickness = None
    material_name = _clean_text(mapped.get("material_name"))
    return {
        "category": resolve_pu_suede_category(mapped, material_name),
        "item_code": _clean_text(mapped.get("item_code")),
        "material_name": material_name,
        "thickness_range": thickness,
        "color": _clean_text(mapped.get("color")),
        "emboss": _clean_text(mapped.get("emboss")),
    }


def iter_pu_suede_rows_from_csv(path: Path) -> Iterable[dict[str, str | None]]:
    delimiter = _detect_csv_delimiter(path)
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f, delimiter=delimiter):
            yield normalize_pu_suede_row(row)


def iter_pu_suede_rows_from_xlsx(
    path: Path,
    *,
    sheet_name: str = PU_SUEDE_SHEET,
) -> Iterable[dict[str, str | None]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' tidak ditemukan di {path}")
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = [str(cell or "").strip() for cell in next(rows, None) or ()]
    if not header:
        wb.close()
        return
    header_map = {name.strip().lower(): idx for idx, name in enumerate(header)}

    def col(*names: str) -> int | None:
        for name in names:
            idx = header_map.get(name.lower())
            if idx is not None:
                return idx
        return None

    code_idx = col("kode barang", "item_code", "item code")
    thickness_idx = col("thickness", "ketebalan", "thickness_range")
    material_idx = col("material name", "material_name", "material")
    emboss_idx = col("emboss")
    color_idx = col("color", "warna")

    for row in rows:
        if not row:
            continue

        def get(idx: int | None) -> str | None:
            if idx is None or idx >= len(row):
                return None
            val = row[idx]
            if val is None:
                return None
            text = str(val).strip()
            if text.lower() in EMPTY_MARKERS:
                return None
            return text

        source = {
            "item_code": get(code_idx),
            "thickness_range": get(thickness_idx),
            "material_name": get(material_idx),
            "emboss": get(emboss_idx),
            "color": get(color_idx),
        }
        if not source["item_code"]:
            continue
        yield normalize_pu_suede_row(source)
    wb.close()


def import_pu_suede_master_items(
    db: Session,
    rows: Iterable[dict[str, str | None]],
    *,
    dry_run: bool = False,
) -> PuSuedeImportResult:
    inserted = 0
    updated = 0
    skipped = 0
    pu = 0
    suide = 0
    seen_codes: set[str] = set()

    for item_data in rows:
        code = item_data.get("item_code")
        if not code:
            skipped += 1
            continue
        if code in seen_codes:
            skipped += 1
            continue
        seen_codes.add(code)

        if item_data.get("category") == SUIDE_CATEGORY:
            suide += 1
        else:
            pu += 1

        existing = db.scalars(select(MasterItem).where(MasterItem.item_code == code)).first()
        if existing:
            changed = False
            for field in ("category", "material_name", "thickness_range", "color", "emboss"):
                val = item_data.get(field)
                if val is not None and getattr(existing, field) != val:
                    setattr(existing, field, val)
                    changed = True
            if changed:
                updated += 1
            else:
                skipped += 1
            continue

        db.add(
            MasterItem(
                item_code=code,
                category=item_data.get("category"),
                material_name=item_data.get("material_name"),
                thickness_range=item_data.get("thickness_range"),
                color=item_data.get("color"),
                emboss=item_data.get("emboss"),
            )
        )
        inserted += 1

    if dry_run:
        db.rollback()
    else:
        db.commit()

    return PuSuedeImportResult(
        imported=inserted + updated,
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        pu=pu,
        suide=suide,
    )
