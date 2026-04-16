import io
from openpyxl import Workbook, load_workbook
from typing import BinaryIO

def export_to_excel(data: list[dict], columns: list[str], sheet_name: str = "Sheet1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    headers = [col.replace("_", " ").title() for col in columns]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    for row_idx, item in enumerate(data, 2):
        for col_idx, col in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=item.get(col, ""))
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def import_from_excel(file_content: bytes, columns: list[str]) -> list[dict]:
    wb = load_workbook(io.BytesIO(file_content))
    ws = wb.active
    
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        item = {}
        for i, col in enumerate(columns):
            item[col] = str(row[i]) if row[i] is not None else ""
        items.append(item)
    
    return items
