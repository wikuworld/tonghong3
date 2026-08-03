"""Export CSV (CR §3.10)."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.db import get_db
from app.models import PoDetail, PoHeader, User
from app.routers.auth import get_current_user
from app.services.mkt_workflow import stamp_rekap_printed_for_recap, sync_mkt_workflow_auto_advance

router = APIRouter(tags=["reports"])


class PoDailyRecapConfirmIn(BaseModel):
    """Konfirmasi setelah MKT benar-benar mencetak rekap (bukan unduh CSV)."""

    report_date: Optional[date] = None
    po_ids: Optional[list[int]] = None


def _thickness_display(d: PoDetail | None) -> str:
    if not d:
        return "—"
    if d.thickness_text and str(d.thickness_text).strip():
        return str(d.thickness_text).strip()
    a, b = d.thickness_min, d.thickness_max
    if a is not None and b is not None:
        return f"{a} - {b}"
    if a is not None or b is not None:
        return str(a if a is not None else b)
    return "—"


def _qty_float(q: Decimal | None) -> float:
    if q is None:
        return 0.0
    return float(q)


def _order_type_label(h: PoHeader) -> str:
    """Selaras lembar marketing / LAPORAN PO HARIAN."""
    if h.order_type_label and str(h.order_type_label).strip():
        return str(h.order_type_label).strip().upper()[:40]
    c = (h.po_category or "production").lower()
    alias = {
        "production": "PROD",
        "sample": "SAMPLE",
        "claim": "CLAIM",
        "return": "RETURN",
        "subcont": "SUBKON",
        "suede": "SUEDE",
    }
    return alias.get(c, c.upper()[:20])


def _pu_production_note(ln: PoDetail) -> str:
    vt = (ln.vir_softness_type or "").lower()
    if "vir" in vt or "pu" in vt:
        return "PU COATED (PU)"
    return ""


def _price_str(ln: PoDetail) -> str:
    if ln.unit_price is None:
        return ""
    return str(ln.unit_price)


def _header_matches_date_filters(
    h: PoHeader,
    date_from: Optional[date],
    date_to: Optional[date],
    date_field: str,
) -> bool:
    if date_from is None and date_to is None:
        return True

    def in_range(d: Optional[date]) -> bool:
        if d is None:
            return False
        if date_from is not None and d < date_from:
            return False
        if date_to is not None and d > date_to:
            return False
        return True

    order_ok = in_range(h.order_date)
    created_ok = False
    if h.created_at:
        ca = h.created_at
        cd: Optional[date] = ca.date() if isinstance(ca, datetime) else ca  # type: ignore[assignment]
        created_ok = in_range(cd)

    df = (date_field or "order").strip().lower()
    if df == "created":
        return created_ok
    if df == "either":
        return order_ok or created_ok
    return order_ok


def _header_matches_filters(
    h: PoHeader,
    date_from: Optional[date],
    date_to: Optional[date],
    customer: Optional[str],
    date_field: str = "order",
) -> bool:
    if not _header_matches_date_filters(h, date_from, date_to, date_field):
        return False
    if customer and str(customer).strip():
        q = str(customer).strip().lower()
        snap = (h.customer_name_snapshot or "").lower()
        code = (h.customer_code or "").lower()
        if not (q == snap or q == code or q in snap or q in code):
            return False
    return True


def _base_marketing_daily_row(h: PoHeader, ln: PoDetail) -> dict[str, Any]:
    cust = (h.customer_name_snapshot or h.customer_code or "").strip() or "—"
    po_no = h.po_number or f"SO-{h.id}"
    return {
        "poNo": po_no,
        "po_customer_ref": (h.po_customer_ref or "").strip(),
        "customer": cust,
        "brand": (h.brand_name_snapshot or "").strip(),
        "order_type": _order_type_label(h),
        "buyer": (h.buyer or "").strip(),
        "order_date": h.order_date.isoformat() if h.order_date else None,
        "thi_line_ref": (ln.thi_line_ref or "").strip(),
        "material": (ln.material_name or "").strip() or "—",
        "thickness": _thickness_display(ln),
        "emboss": (ln.emboss or "").strip(),
        "softness_type": (ln.vir_softness_type or "").strip(),
        "pu_production": _pu_production_note(ln),
        "color": (ln.color or "").strip(),
        "softness": (ln.softness or "").strip(),
        "unit_price": _price_str(ln),
        "color_reference": (ln.color_reference or "").strip(),
        "backing": (ln.backing or "").strip(),
        "model": (ln.model or "").strip(),
        "style_article": (ln.style_article or "").strip(),
        "line_etd": ln.line_etd.isoformat() if ln.line_etd else None,
        "qty": _qty_float(ln.quantity_order),
        "po_category": (h.po_category or "production").lower(),
        "status": h.status or "draft",
        "po_id": h.id,
        "detail_id": ln.id,
    }


def _build_marketing_daily_payload(
    db: Session,
    date_from: Optional[date],
    date_to: Optional[date],
    customer: Optional[str],
    date_field: str,
) -> dict[str, Any]:
    """Isi JSON + CSV laporan PO harian marketing."""
    headers = db.scalars(
        select(PoHeader)
        .options(selectinload(PoHeader.details))
        .order_by(PoHeader.id.desc())
    ).all()

    revise: list[dict[str, Any]] = []
    production_orders: list[dict[str, Any]] = []
    samples: list[dict[str, Any]] = []

    df = (date_field or "order").strip().lower()
    if df not in {"order", "created", "either"}:
        df = "order"

    for h in headers:
        if not _header_matches_filters(h, date_from, date_to, customer, df):
            continue
        for ln in h.details or []:
            if ln.copied_from_detail_id:
                old = db.get(PoDetail, ln.copied_from_detail_id)
                row = _base_marketing_daily_row(h, ln)
                row["oldThickness"] = _thickness_display(old)
                row["newThickness"] = _thickness_display(ln)
                revise.append(row)
                continue

            cat = (h.po_category or "production").lower()
            row = _base_marketing_daily_row(h, ln)
            row["oldThickness"] = ""
            row["newThickness"] = ""
            if cat == "sample":
                samples.append(row)
            else:
                production_orders.append(row)

    return {
        "report_title": "LAPORAN PO HARIAN MARKETING",
        "revise_thickness": revise,
        "production_orders": production_orders,
        "samples": samples,
    }


@router.get("/reports/po-order-description")
def report_po_order_description(
    db: Session = Depends(get_db),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    customer: Optional[str] = Query(None, description="Nama / kode customer (snapshot)"),
    date_field: str = Query(
        "order",
        description="Saring rentang tanggal: order (order_date), created (created_at header), either (salah satu)",
    ),
    _user=Depends(get_current_user),
) -> dict[str, Any]:
    """Data LAPORAN PO HARIAN MARKETING: revisi ketebalan, production order, sample — format lembar Excel."""
    _ = _user
    return _build_marketing_daily_payload(db, date_from, date_to, customer, date_field)


_MARKETING_CSV_COLS = [
    "section",
    "order_date",
    "thi_line_ref",
    "customer",
    "brand",
    "order_type",
    "po_no",
    "po_customer_ref",
    "material",
    "thickness",
    "old_thickness",
    "new_thickness",
    "emboss",
    "softness_type",
    "pu_production",
    "color",
    "softness",
    "qty",
    "unit_price",
    "color_reference",
    "backing",
    "model",
    "style_article",
    "line_etd",
    "buyer",
    "po_status",
    "po_category",
]

_MARKETING_XLSX_COLS = [
    ("No", None),
    ("Order Date", "order_date"),
    ("PO No", "poNo"),
    ("PO Cust", "po_customer_ref"),
    ("Customer", "customer"),
    ("Buyer", "buyer"),
    ("Order Type", "order_type"),
    ("THI/Lot", "thi_line_ref"),
    ("Material", "material"),
    ("Thickness", "thickness"),
    ("Old Thickness", "oldThickness"),
    ("New Thickness", "newThickness"),
    ("Color", "color"),
    ("Softness", "softness"),
    ("Qty", "qty"),
    ("Unit Price", "unit_price"),
    ("Color Ref", "color_reference"),
    ("Backing", "backing"),
    ("Model", "model"),
    ("Style/Article", "style_article"),
    ("ETD", "line_etd"),
]


def _row_to_csv_tuple(section: str, row: dict[str, Any]) -> tuple:
    return (
        section,
        row.get("order_date") or "",
        row.get("thi_line_ref") or "",
        row.get("customer") or "",
        row.get("brand") or "",
        row.get("order_type") or "",
        row.get("poNo") or "",
        row.get("po_customer_ref") or "",
        row.get("material") or "",
        row.get("thickness") or "",
        row.get("oldThickness") or "",
        row.get("newThickness") or "",
        row.get("emboss") or "",
        row.get("softness_type") or "",
        row.get("pu_production") or "",
        row.get("color") or "",
        row.get("softness") or "",
        row.get("qty", ""),
        row.get("unit_price") or "",
        row.get("color_reference") or "",
        row.get("backing") or "",
        row.get("model") or "",
        row.get("style_article") or "",
        row.get("line_etd") or "",
        row.get("buyer") or "",
        row.get("status") or "",
        row.get("po_category") or "",
    )


def _append_marketing_section_xlsx(ws, start_row: int, title: str, rows: list[dict[str, Any]]) -> int:
    ws.cell(row=start_row, column=1, value=title)
    ws.cell(row=start_row, column=1).font = Font(bold=True, size=12)
    row_idx = start_row + 1

    for col_idx, (label, _) in enumerate(_MARKETING_XLSX_COLS, start=1):
        c = ws.cell(row=row_idx, column=col_idx, value=label)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    row_idx += 1

    total_qty = 0.0
    for i, item in enumerate(rows, start=1):
        for col_idx, (_, key) in enumerate(_MARKETING_XLSX_COLS, start=1):
            value = i if key is None else item.get(key, "")
            ws.cell(row=row_idx, column=col_idx, value=value)
            if key == "qty":
                try:
                    total_qty += float(value or 0)
                except Exception:
                    pass
        row_idx += 1

    ws.cell(row=row_idx, column=1, value=f"TOTAL {title}")
    ws.cell(row=row_idx, column=1).font = Font(bold=True)
    qty_col_idx = next(i for i, (_, k) in enumerate(_MARKETING_XLSX_COLS, start=1) if k == "qty")
    ws.cell(row=row_idx, column=qty_col_idx, value=round(total_qty, 2))
    ws.cell(row=row_idx, column=qty_col_idx).font = Font(bold=True)
    return row_idx + 2


@router.get("/reports/po-daily-marketing.csv")
def export_po_daily_marketing_csv(
    db: Session = Depends(get_db),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    customer: Optional[str] = Query(None),
    date_field: str = Query("order"),
    _user=Depends(get_current_user),
) -> Response:
    """Ekspor CSV (UTF-8 BOM) selaras LAPORAN PO HARIAN MARKETING — tiga blok + kolom lengkap."""
    _ = _user
    payload = _build_marketing_daily_payload(db, date_from, date_to, customer, date_field)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_MARKETING_CSV_COLS)
    for row in payload["revise_thickness"]:
        w.writerow(_row_to_csv_tuple("REVISE_THICKNESS", row))
    for row in payload["production_orders"]:
        w.writerow(_row_to_csv_tuple("PRODUCTION_ORDER", row))
    for row in payload["samples"]:
        w.writerow(_row_to_csv_tuple("SAMPLE", row))

    fn = f"laporan_po_harian_marketing_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.csv"
    body = "\ufeff" + buf.getvalue()
    return Response(
        content=body,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


@router.get("/reports/po-daily-marketing.xlsx")
def export_po_daily_marketing_xlsx(
    db: Session = Depends(get_db),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    customer: Optional[str] = Query(None),
    date_field: str = Query("order"),
    _user=Depends(get_current_user),
) -> Response:
    """Ekspor XLSX LAPORAN PO HARIAN MARKETING (revise thickness, production order, sample)."""
    _ = _user
    payload = _build_marketing_daily_payload(db, date_from, date_to, customer, date_field)
    wb = Workbook()
    ws = wb.active
    ws.title = "PO Daily Marketing"

    ws["A1"] = payload.get("report_title") or "LAPORAN PO HARIAN MARKETING"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated at: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"

    row_idx = 4
    row_idx = _append_marketing_section_xlsx(ws, row_idx, "REVISE THICKNESS", payload.get("revise_thickness") or [])
    row_idx = _append_marketing_section_xlsx(ws, row_idx, "PRODUCTION ORDER", payload.get("production_orders") or [])
    _append_marketing_section_xlsx(ws, row_idx, "SAMPLE", payload.get("samples") or [])

    for idx, _ in enumerate(_MARKETING_XLSX_COLS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 16

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fn = f"laporan_po_harian_marketing_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


def _headers_for_daily_recap(db: Session, d: date) -> list[PoHeader]:
    return list(
        db.scalars(
            select(PoHeader)
            .options(selectinload(PoHeader.details))
            .where(
                or_(
                    PoHeader.order_date == d,
                    func.date(PoHeader.created_at) == d,
                )
            )
            .order_by(PoHeader.id.desc())
        ).all()
    )


def _headers_for_recap_by_ids(db: Session, po_ids: list[int]) -> list[PoHeader]:
    if not po_ids:
        return []
    wanted = {int(x) for x in po_ids}
    rows = list(
        db.scalars(
            select(PoHeader)
            .options(selectinload(PoHeader.details))
            .where(PoHeader.id.in_(wanted))
            .order_by(PoHeader.id.desc())
        ).all()
    )
    return rows


def _payload_from_headers(headers: list[PoHeader], *, report_date: Optional[date] = None) -> dict[str, Any]:
    """Bangun payload HTML/print dari daftar header — tanpa stamp status."""
    orders: list[dict[str, Any]] = []
    lines: list[dict[str, Any]] = []
    for h in headers:
        cust = (h.customer_name_snapshot or h.customer_code or "").strip() or "—"
        orders.append(
            {
                "po_id": h.id,
                "so_number": h.po_number or f"SO-{h.id}",
                "customer": cust,
                "customer_id": h.customer_id,
                "buyer": (h.buyer or "").strip(),
                "brand": (h.brand_name_snapshot or "").strip(),
                "order_type": _order_type_label(h),
                "status": h.status or "draft",
                "mkt_workflow_stage": h.mkt_workflow_stage or "",
                "order_date": h.order_date.isoformat() if h.order_date else None,
                "etd": h.etd.isoformat() if h.etd else None,
                "po_customer_ref": (h.po_customer_ref or "").strip(),
                "rekap_printed_at": h.rekap_printed_at.isoformat() if h.rekap_printed_at else None,
                "line_count": len(h.details or []),
            }
        )
        for ln in h.details or []:
            row = _base_marketing_daily_row(h, ln)
            row["po_id"] = h.id
            lines.append(row)

    pending_ids = [o["po_id"] for o in orders if not o["rekap_printed_at"]]
    d = report_date
    if d is None and headers:
        h0 = headers[0]
        d = h0.order_date or (h0.created_at.date() if h0.created_at else None)
    return {
        "report_title": "REKAP PO / SO MARKETING",
        "report_date": d.isoformat() if d else None,
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "order_count": len(orders),
        "line_count": len(lines),
        "pending_print_count": len(pending_ids),
        "pending_po_ids": pending_ids,
        "orders": orders,
        "lines": lines,
    }


def _build_po_daily_recap_payload(db: Session, d: date) -> dict[str, Any]:
    """Payload HTML/print rekap harian — tanpa mengubah status print."""
    return _payload_from_headers(_headers_for_daily_recap(db, d), report_date=d)


@router.get("/reports/po-daily-recap")
def get_po_daily_recap(
    report_date: Optional[date] = Query(None, alias="date"),
    po_id: Optional[int] = Query(None, description="Rekap satu SO (prioritas di atas filter tanggal)"),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """Data rekap PO/SO untuk dialog cetak HTML — tidak mengubah tahap/status."""
    _ = _user
    if po_id is not None:
        return _payload_from_headers(_headers_for_recap_by_ids(db, [po_id]))
    return _build_po_daily_recap_payload(db, report_date or date.today())


@router.post("/reports/po-daily-recap/confirm-print")
def confirm_po_daily_recap_print(
    body: PoDailyRecapConfirmIn,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """Stamp rekap_printed setelah MKT mengonfirmasi cetak dari dialog HTML."""
    if body.po_ids:
        headers = _headers_for_recap_by_ids(db, list(body.po_ids))
        d = body.report_date or date.today()
        if headers and headers[0].order_date:
            d = headers[0].order_date
    else:
        d = body.report_date or date.today()
        headers = _headers_for_daily_recap(db, d)
    stamped = stamp_rekap_printed_for_recap(db, headers, actor=user.username)
    for h in headers:
        # Setelah konfirmasi cetak: hanya naik sampai rekap_printed.
        # Jangan loncat with_ppic/outstanding hanya karena assign PPIC lama.
        sync_mkt_workflow_auto_advance(
            db,
            h.id,
            actor=user.username,
            stop_at_stage="rekap_printed",
        )
    db.commit()
    return {
        "ok": True,
        "report_date": d.isoformat() if isinstance(d, date) else str(d),
        "stamped_count": stamped,
        "po_ids": [h.id for h in headers],
    }


@router.get("/reports/po-daily-recap.csv")
def export_po_daily_recap_csv(
    report_date: Optional[date] = Query(None, alias="date"),
    po_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> Response:
    """Export CSV rekap (arsip) — tidak menandai status print; gunakan dialog cetak HTML."""
    _ = user
    d = report_date or date.today()
    if po_id is not None:
        headers = _headers_for_recap_by_ids(db, [po_id])
        if headers and headers[0].order_date:
            d = headers[0].order_date
    else:
        headers = _headers_for_daily_recap(db, d)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "so_number",
            "customer_id",
            "mkt_workflow_stage",
            "status",
            "order_date",
            "etd",
            "buyer",
            "rekap_printed_at",
        ]
    )
    for h in headers:
        w.writerow(
            [
                h.po_number or "",
                h.customer_id or "",
                h.mkt_workflow_stage or "",
                h.status or "",
                h.order_date or "",
                h.etd or "",
                h.buyer or "",
                h.rekap_printed_at.isoformat() if h.rekap_printed_at else "",
            ]
        )
    fn = f"rekap_po_{d.isoformat()}_{datetime.utcnow().strftime('%H%M')}.csv"
    return Response(
        content=buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


def _delivery_flat_columns(ln: PoDetail) -> list:
    """12 × (qty, date) selaras `03_Master_Data_Marketing.csv`."""
    by_seq = {d.sequence_no: d for d in (ln.delivery_schedules or [])}
    out = []
    for s in range(1, 13):
        d = by_seq.get(s)
        out.append(d.qty if d and d.qty is not None else "")
        out.append(d.delivery_date.isoformat() if d and d.delivery_date else "")
    return out


@router.get("/reports/sales-orders-export.csv")
def export_sales_orders_csv(
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> Response:
    headers = db.scalars(select(PoHeader).order_by(PoHeader.id.desc())).all()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "so_number",
            "customer_id",
            "status",
            "order_date",
            "etd",
            "billing_address",
            "delivery_address",
            "marketing_excel_date_serial",
            "po_customer_ref",
            "order_type_label",
            "brand_name_snapshot",
            "customer_name_snapshot",
        ]
    )
    for h in headers:
        w.writerow(
            [
                h.po_number or "",
                h.customer_id or "",
                h.status or "",
                h.order_date or "",
                h.etd or "",
                (h.billing_address or "").replace("\n", " ")[:300],
                (h.delivery_address or "").replace("\n", " ")[:300],
                h.marketing_excel_date_serial or "",
                (h.po_customer_ref or "").replace("\n", " ")[:500],
                h.order_type_label or "",
                h.brand_name_snapshot or "",
                h.customer_name_snapshot or "",
            ]
        )
    return Response(
        content=buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="sales_orders.csv"',
        },
    )


@router.get("/reports/sales-order-lines-export.csv")
def export_so_lines_csv(
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> Response:
    lines = db.scalars(
        select(PoDetail)
        .options(selectinload(PoDetail.delivery_schedules))
        .order_by(PoDetail.id.desc())
    ).all()
    buf = io.StringIO()
    w = csv.writer(buf)
    base_cols = [
        "so_id",
        "line_id",
        "master_item_id",
        "copied_from_detail_id",
        "thi_line_ref",
        "material",
        "thickness_min",
        "thickness_max",
        "thickness_text",
        "emboss",
        "vir_softness_type",
        "color",
        "softness",
        "qty",
        "line_etd",
        "etd_excel_serial",
        "unit_price",
        "color_reference",
        "backing",
        "model",
        "style_article",
        "rtd_text",
        "stock_wh_qty",
        "balance_qty",
        "system_reference",
        "detail_remark",
    ]
    del_cols = []
    for i in range(1, 13):
        del_cols.extend([f"delivery_{i}_qty", f"delivery_{i}_date"])
    w.writerow(base_cols + del_cols)
    for ln in lines:
        w.writerow(
            [
                ln.po_id or "",
                ln.id,
                ln.master_item_id or "",
                ln.copied_from_detail_id or "",
                ln.thi_line_ref or "",
                ln.material_name or "",
                ln.thickness_min or "",
                ln.thickness_max or "",
                ln.thickness_text or "",
                ln.emboss or "",
                ln.vir_softness_type or "",
                ln.color or "",
                ln.softness or "",
                ln.quantity_order or "",
                ln.line_etd or "",
                ln.etd_excel_serial or "",
                ln.unit_price or "",
                ln.color_reference or "",
                ln.backing or "",
                ln.model or "",
                ln.style_article or "",
                ln.rtd_text or "",
                ln.stock_wh_qty or "",
                ln.balance_qty or "",
                ln.system_reference or "",
                (ln.detail_remark or "").replace("\n", " ")[:500],
            ]
            + _delivery_flat_columns(ln)
        )
    return Response(
        content=buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="sales_order_lines.csv"',
        },
    )
