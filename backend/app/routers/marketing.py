"""Endpoint format lembar marketing (outstanding) dari data PO aktual."""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal
from enum import Enum
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.db import get_db
from app.models import (
    Customer,
    DeliveryDetail,
    DeliveryHeader,
    Invoice,
    InvoiceDetail,
    MarketingOutstandingSnapshot,
    MasterItem,
    PackingHeader,
    PoDetail,
    PoDetailDelivery,
    PoHeader,
    ProductionDetail,
)
from app.routers.auth import get_current_user, require_roles
from app.schemas import (
    MarketingDeliveryLineOut,
    MarketingDeliveryLinesPageOut,
    MarketingDeliveryPlanningPageOut,
    MarketingDeliveryPlanningSlotOut,
    MarketingDeliveryTodayOut,
    MarketingLineEtdTodayOut,
    MarketingOutstandingLineOut,
    MarketingOutstandingPageOut,
    MarketingOutstandingSnapshotCreate,
    MarketingOutstandingSnapshotDetailOut,
    MarketingOutstandingSnapshotListItem,
    OutstandingPackingOptionOut,
    OutstandingPackingOptionsOut,
    OutstandingScheduleBatchIn,
    OutstandingScheduleBatchOut,
    OutstandingShipBatchIn,
    OutstandingShipBatchOut,
)
from app.services.document_numbers import invoice_display_number, surat_jalan_display_number
from app.services.mkt_workflow import MKT_STAGES_ORDER, mkt_stage_index, normalize_mkt_stage
from app.services.mrp_po_alloc import enrich_outstanding_mrp_alloc
from app.services.outstanding_actions import (
    apply_outstanding_schedule_batch,
    apply_outstanding_ship_batch,
    packing_options_for_details,
)
router = APIRouter(tags=["marketing"])

# Selaras lembar manual «DATA OUTSTANDING ORDER» (PROD 21 gelombang, claim/sample 5, retur 3).
OUTSTANDING_DELIVERY_SLOTS_PROD = 21
OUTSTANDING_DELIVERY_SLOTS_CLAIM_SAMPLE = 5
OUTSTANDING_DELIVERY_SLOTS_RETURN = 3
PO_MAX_DELIVERY_SCHEDULES = 21


def _ordinal_delivery_label(n: int) -> str:
    if n == 1:
        return "1st"
    if n == 2:
        return "2nd"
    if n == 3:
        return "3rd"
    return f"{n}th"


def _delivery_export_column_pairs(max_slots: int) -> list[tuple[str, str]]:
    cols: list[tuple[str, str]] = []
    for i in range(1, max_slots + 1):
        lab = _ordinal_delivery_label(i)
        cols.append((f"delivery_{i}_qty", f"{lab} Qty"))
        cols.append((f"delivery_{i}_date", f"{lab} Date"))
    return cols


def _parse_iso_date_optional(s: Optional[str]) -> Optional[date]:
    if not s or not str(s).strip():
        return None
    try:
        return date.fromisoformat(str(s).strip()[:10])
    except ValueError:
        return None


class BalanceFilter(str, Enum):
    """Filter baris berdasarkan saldo baris PO (kolom balance_qty)."""

    all = "all"
    exclude_zero = "exclude_zero"
    positive_only = "positive_only"


class ShipmentStatusFilter(str, Enum):
    """Filter baris berdasarkan sisa kirim (balance terhitung atau tersimpan)."""

    all = "all"
    sent = "sent"
    not_sent = "not_sent"


class MaterialStatusFilter(str, Enum):
    """Status material versi marketing: shipped/sent, not yet shipped, cancelled, all."""

    all = "all"
    sent = "sent"
    not_sent = "not_sent"
    cancelled = "cancelled"


class DeliveryLinesStatusFilter(str, Enum):
    """Filter status header delivery pada daftar baris delivery."""

    all = "all"
    draft = "draft"
    sent = "sent"


class PackingLinesStatusFilter(str, Enum):
    """Filter status packing terkait delivery."""

    all = "all"
    draft = "draft"
    done = "done"


class InvplPrintedFilter(str, Enum):
    """Filter baris berdasarkan INVPL sudah ditandai dicetak atau belum."""

    all = "all"
    yes = "yes"
    no = "no"


def _passes_balance_filter(ln: PoDetail, mode: BalanceFilter) -> bool:
    b = ln.balance_qty
    if mode == BalanceFilter.all:
        return True
    if mode == BalanceFilter.exclude_zero:
        if b is None:
            return True
        return b != 0
    if mode == BalanceFilter.positive_only:
        if b is None:
            return False
        return b > 0
    return True


def _passes_balance_value(balance: Decimal | None, mode: BalanceFilter) -> bool:
    b = balance
    if mode == BalanceFilter.all:
        return True
    if mode == BalanceFilter.exclude_zero:
        if b is None:
            return True
        return b != 0
    if mode == BalanceFilter.positive_only:
        if b is None:
            return False
        return b > 0
    return True


def _dec_str(v: Decimal | None) -> str:
    if v is None:
        return ""
    return str(v)


def _delivered_qty_by_po_detail(db: Session) -> dict[int, Decimal]:
    """Qty terkirim per baris SO: jalur produksi + DeliveryDetail.po_detail_id langsung (MRP)."""
    delivered: dict[int, Decimal] = {}
    rows = db.execute(
        select(ProductionDetail.po_detail_id, func.sum(DeliveryDetail.qty))
        .join(DeliveryDetail, DeliveryDetail.production_detail_id == ProductionDetail.id)
        .group_by(ProductionDetail.po_detail_id)
    ).all()
    for pod_id, total in rows:
        if pod_id is None:
            continue
        delivered[int(pod_id)] = total or Decimal(0)
    direct_rows = db.execute(
        select(DeliveryDetail.po_detail_id, func.sum(DeliveryDetail.qty))
        .where(
            DeliveryDetail.po_detail_id.is_not(None),
            or_(
                DeliveryDetail.production_detail_id.is_(None),
                DeliveryDetail.production_detail_id == 0,
            ),
        )
        .group_by(DeliveryDetail.po_detail_id)
    ).all()
    for pod_id, total in direct_rows:
        if pod_id is None:
            continue
        key = int(pod_id)
        delivered[key] = (delivered.get(key) or Decimal(0)) + (total or Decimal(0))
    return delivered


def _planning_slot_ship_status(
    header_stage: Optional[str],
    delivered: Decimal,
    order_qty: Decimal,
) -> str:
    """planned | partial | shipped | too_early — berdasarkan tahap MKT + sisa kirim baris."""
    stage = normalize_mkt_stage(header_stage)
    if mkt_stage_index(stage) < mkt_stage_index("in_outstanding"):
        return "too_early"
    order = order_qty or Decimal(0)
    sent = delivered or Decimal(0)
    balance = order - sent
    if order > 0 and balance <= 0:
        return "shipped"
    if sent > 0:
        return "partial"
    return "planned"


def _thickness_display(ln: PoDetail) -> str:
    if ln.thickness_text and str(ln.thickness_text).strip():
        return str(ln.thickness_text).strip()
    a, b = ln.thickness_min, ln.thickness_max
    if a is not None and b is not None:
        return f"{a} - {b}"
    if a is not None or b is not None:
        return str(a if a is not None else b)
    return ""


def _delivery_line_amount(
    dd: DeliveryDetail,
    pod: Optional[PoDetail],
    invd: Optional[InvoiceDetail],
) -> Decimal | None:
    if invd is not None and invd.amount is not None:
        return invd.amount
    if dd.qty is not None and dd.qty > 0 and pod is not None and pod.unit_price is not None:
        return (dd.qty * pod.unit_price).quantize(Decimal("0.01"))
    return None


def _order_type_marketing(h: PoHeader) -> str:
    if h.order_type_label and str(h.order_type_label).strip():
        return str(h.order_type_label).strip().upper()[:40]
    c = (h.po_category or "production").lower()
    alias = {
        "production": "PROD",
        "sample": "SAMPLE",
        "claim": "CLAIM",
        "return": "RETURN",
        "subcont": "SUBKON",
        "suede": "SUEDE",
    }
    return alias.get(c, c.upper()[:20])


def _normalize_material_group_filter(raw: Optional[str]) -> Optional[str]:
    """Query `material_group`: pu | suede. Kosong = semua."""
    v = (raw or "").strip().lower()
    if not v:
        return None
    if v in {"suede", "suide"}:
        return "suede"
    if v == "pu":
        return "pu"
    raise HTTPException(status_code=400, detail="material_group harus pu atau suede")


def _classify_line_material_group(
    h: PoHeader,
    ln: PoDetail,
    master: MasterItem | None = None,
) -> str:
    """Kategori material baris: 'pu' atau 'suede' (tanpa raise)."""
    if master is not None:
        cat = (master.category or "").strip().lower()
        if cat in {"suide", "suede"}:
            return "suede"
        if cat == "pu":
            return "pu"
    vt = (getattr(ln, "vir_softness_type", None) or "").strip().lower()
    if "suede" in vt or "suide" in vt:
        return "suede"
    if "pu" in vt:
        return "pu"
    mn = (getattr(ln, "material_name", None) or "").strip().lower()
    if "suede" in mn or "suide" in mn:
        return "suede"
    if "pu" in mn:
        return "pu"
    if (getattr(h, "po_category", None) or "").strip().lower() == "suede":
        return "suede"
    return "pu"


def _normalize_mkt_stage_min(raw: Optional[str]) -> Optional[str]:
    v = (raw or "").strip()
    if not v:
        return None
    s = normalize_mkt_stage(v)
    if s not in MKT_STAGES_ORDER:
        raise HTTPException(status_code=400, detail="mkt_stage_min tidak valid")
    return s


def _passes_mkt_stage_min(header_stage: Optional[str], min_stage: Optional[str]) -> bool:
    if not min_stage:
        return True
    return mkt_stage_index(header_stage) >= mkt_stage_index(min_stage)


def _po_q_matches(header: Any, needle: Optional[str]) -> bool:
    """Cari no. SO (`po_number`) atau nomor PO customer (`po_customer_ref`)."""
    q = (needle or "").strip().lower()
    if not q:
        return True
    po_number = (getattr(header, "po_number", None) or "").strip().lower()
    po_cust = (getattr(header, "po_customer_ref", None) or "").strip().lower()
    return q in po_number or q in po_cust


def _prepare_outstanding_rows(
    db: Session,
    *args: Any,
    mkt_stage_min: Optional[str] = None,
    **kwargs: Any,
) -> list[dict[str, Any]]:
    rows = _query_outstanding_rows(db, *args, **kwargs)
    min_stage = _normalize_mkt_stage_min(mkt_stage_min)
    if min_stage:
        rows = [r for r in rows if _passes_mkt_stage_min(r.get("mkt_stage"), min_stage)]
    enrich_outstanding_mrp_alloc(rows)
    return rows


def _status_material_text(h: PoHeader, _ln: PoDetail, production_status: str) -> str:
    """Gabungan status produksi baris + follow-up material header (mirip kolom STATUS MATERIAL lembar Excel)."""
    parts: list[str] = []
    st = (production_status or "").strip()
    if st:
        parts.append(f"Produksi: {st}")
    stage = normalize_mkt_stage(getattr(h, "mkt_workflow_stage", None))
    if getattr(h, "material_followup_done", False) or mkt_stage_index(stage) >= mkt_stage_index(
        "delivery_arranged"
    ):
        parts.append("Follow-up material: selesai")
    elif stage == "material_followup":
        parts.append("Follow-up material: berjalan")
    else:
        parts.append("Follow-up material: belum")
    fu = (h.material_followup_remark or "").strip()
    if fu:
        parts.append(fu[:240])
    return " | ".join(parts) if parts else "—"


def _remark_system_text(ln: PoDetail, delivered: Decimal, balance: Decimal) -> str:
    """Ringkasan otomatis dari angka baris & ETD (berubah saat input PO / kiriman berubah)."""
    parts: list[str] = []
    if ln.quantity_order is not None:
        parts.append(f"Order {_dec_str(ln.quantity_order)} SF")
    parts.append(f"Terkirim {_dec_str(delivered)}")
    parts.append(f"Saldo {_dec_str(balance)}")
    if ln.stock_wh_qty is not None:
        parts.append(f"StokWH {_dec_str(ln.stock_wh_qty)}")
    if ln.line_etd:
        parts.append(f"ETD {ln.line_etd.isoformat()}")
    return " | ".join(parts)


def _delivery_slots(ln: PoDetail, max_slots: int = OUTSTANDING_DELIVERY_SLOTS_PROD) -> dict[str, Any]:
    by_seq: dict[int, Any] = {}
    for d in ln.delivery_schedules or []:
        by_seq[int(d.sequence_no)] = d
    out: dict[str, Any] = {}
    for i in range(1, max_slots + 1):
        d = by_seq.get(i)
        out[f"delivery_{i}_qty"] = _dec_str(d.qty) if d and d.qty is not None else ""
        out[f"delivery_{i}_date"] = d.delivery_date.isoformat() if d and d.delivery_date else ""
    return out


def _schedule_summary(ln: PoDetail, delivered: Decimal) -> dict[str, str]:
    """Ringkasan rencana jadwal vs order vs aktual terkirim."""
    total = Decimal(0)
    for d in ln.delivery_schedules or []:
        if d.qty is not None:
            total += d.qty
    order = ln.quantity_order
    out = {
        "schedule_qty_total": _dec_str(total) if total else "",
        "schedule_unallocated_qty": "",
        "actual_vs_schedule_qty": "",
    }
    if order is not None:
        if total:
            out["schedule_unallocated_qty"] = _dec_str(order - total)
            out["actual_vs_schedule_qty"] = _dec_str(delivered - total)
        else:
            out["schedule_unallocated_qty"] = _dec_str(order)
    return out


def _latest_production_status_by_pod(db: Session, pod_ids: set[int]) -> dict[int, str]:
    if not pod_ids:
        return {}
    rows = db.execute(
        select(ProductionDetail.po_detail_id, ProductionDetail.id, ProductionDetail.status)
        .where(ProductionDetail.po_detail_id.in_(pod_ids))
        .order_by(ProductionDetail.id.asc())
    ).all()
    out: dict[int, str] = {}
    for pod_id, _pid, status in rows:
        if pod_id is None:
            continue
        out[int(pod_id)] = (status or "").strip()
    return out


def _line_to_row(
    h: PoHeader,
    ln: PoDetail,
    delivered: Decimal,
    balance: Decimal,
    production_status: str,
) -> dict[str, Any]:
    brand = (h.brand_name_snapshot or "").strip() or "—"
    cust = (h.customer_name_snapshot or h.customer_code or "").strip() or "—"
    thi_order = (h.po_number or "").strip() or "—"
    system_no = (ln.system_reference or ln.thi_line_ref or "").strip()
    row: dict[str, Any] = {
        "order_date": h.order_date.isoformat() if h.order_date else None,
        "thi_order": thi_order,
        "customer_name": cust,
        "brand": brand,
        "order_type": _order_type_marketing(h),
        "no_po": (h.po_customer_ref or "").strip() or "—",
        "po_customer_ref": (h.po_customer_ref or "").strip() or "",
        "po_category": (h.po_category or "production").strip().lower(),
        "factory_code": ((h.factory_code or h.customer_code or "").strip()),
        "material_name": (ln.material_name or "").strip() or "—",
        "thickness": _thickness_display(ln),
        "emboss": (ln.emboss or "").strip() or "",
        "softness_type": (ln.vir_softness_type or "").strip() or "",
        "color": (ln.color or "").strip() or "",
        "softness": (ln.softness or "").strip() or "",
        "qty_sf": _dec_str(ln.quantity_order),
        "unit_price": _dec_str(ln.unit_price),
        "color_reference": (ln.color_reference or "").strip() or "",
        "color_standard": (ln.color_standard or "").strip() or "",
        "backing": (ln.backing or "").strip() or "",
        "model": (ln.model or "").strip() or "",
        "style_article": (ln.style_article or "").strip() or "",
        "rtd_text": (ln.rtd_text or "").strip() or "",
        "line_etd": ln.line_etd.isoformat() if ln.line_etd else None,
        "stock_wh_qty": _dec_str(ln.stock_wh_qty),
        "balance_qty": _dec_str(balance),
        "system_reference": system_no,
        "detail_remark": (ln.detail_remark or "").strip() or "",
        "remark_system": _remark_system_text(ln, delivered, balance),
        "remark_marketing": (ln.marketing_outstanding_remark or "").strip() or "",
        "material_followup_remark_header": (h.material_followup_remark or "").strip(),
        "material_followup_done_header": bool(h.material_followup_done),
        "status_material": _status_material_text(h, ln, production_status),
        "delivered_qty": _dec_str(delivered),
        "production_status": production_status or "",
        "po_id": h.id,
        "po_detail_id": ln.id,
        "po_status": h.status or "",
        "mkt_stage": h.mkt_workflow_stage or "",
    }
    row.update(_delivery_slots(ln))
    row.update(_schedule_summary(ln, delivered))
    return row


def _query_outstanding_rows(
    db: Session,
    balance_filter: BalanceFilter = BalanceFilter.exclude_zero,
    computed_balance: bool = True,
    po_q: Optional[str] = None,
    shipment_status: ShipmentStatusFilter = ShipmentStatusFilter.all,
    material_status: Optional[MaterialStatusFilter] = None,
    factory_code: Optional[str] = None,
    po_category: Optional[str] = None,
    material_group: Optional[str] = None,
    order_date_from: Optional[date] = None,
    order_date_to: Optional[date] = None,
    po_detail_ids: Optional[list[int]] = None,
) -> list[dict[str, Any]]:
    """Baris dari PO yang belum completed (aliran outstanding marketing)."""
    wanted_details = {int(x) for x in (po_detail_ids or []) if x is not None}
    group_needle = _normalize_material_group_filter(material_group)
    headers = db.scalars(
        select(PoHeader)
        .options(
            selectinload(PoHeader.details).selectinload(PoDetail.delivery_schedules)
        )
        .order_by(PoHeader.id.desc())
    ).all()

    masters_by_id: dict[int, MasterItem] = {}
    if group_needle:
        master_ids = {
            int(ln.master_item_id)
            for h in headers
            for ln in (h.details or [])
            if getattr(ln, "master_item_id", None)
        }
        if master_ids:
            for mi in db.scalars(select(MasterItem).where(MasterItem.id.in_(master_ids))).all():
                masters_by_id[int(mi.id)] = mi

    delivered_by_po_detail: dict[int, Decimal] = {}
    if computed_balance:
        delivered_by_po_detail = _delivered_qty_by_po_detail(db)

    po_needle = (po_q or "").strip().lower()
    factory_needle = (factory_code or "").strip().lower()
    cat_needle = (po_category or "").strip().lower()
    pending: list[tuple[PoHeader, PoDetail, Decimal, Decimal]] = []

    for h in headers:
        if order_date_from is not None or order_date_to is not None:
            od = h.order_date
            if od is None:
                continue
            if order_date_from is not None and od < order_date_from:
                continue
            if order_date_to is not None and od > order_date_to:
                continue
        if (h.status or "").lower() == "completed":
            continue
        ms = material_status or MaterialStatusFilter.all
        if ms == MaterialStatusFilter.cancelled and (h.status or "").strip().lower() not in {"cancelled", "canceled"}:
            continue
        if ms != MaterialStatusFilter.cancelled and (h.status or "").strip().lower() in {"cancelled", "canceled"}:
            # default: cancel tidak ikut outstanding kecuali diminta
            continue
        if factory_needle:
            if factory_needle not in ((h.factory_code or "").strip().lower()):
                continue
        if cat_needle:
            stored_cat = (h.po_category or "").strip().lower()
            if cat_needle == "production":
                if stored_cat not in {"", "production", "suede"}:
                    continue
            elif cat_needle != stored_cat:
                continue
        if po_needle and not _po_q_matches(h, po_needle):
            continue
        for ln in h.details or []:
            if wanted_details and int(ln.id) not in wanted_details:
                continue
            if group_needle:
                mid = getattr(ln, "master_item_id", None)
                master = masters_by_id.get(int(mid)) if mid else None
                if _classify_line_material_group(h, ln, master) != group_needle:
                    continue
            if computed_balance:
                delivered = delivered_by_po_detail.get(int(ln.id), Decimal(0))
                qty = ln.quantity_order or Decimal(0)
                bal = qty - delivered
                if not _passes_balance_value(bal, balance_filter):
                    continue
                eff = material_status.value if material_status is not None else shipment_status.value
                if eff == ShipmentStatusFilter.sent.value and bal > 0:
                    continue
                if eff == ShipmentStatusFilter.not_sent.value and bal <= 0:
                    continue
                pending.append((h, ln, delivered, bal))
            else:
                if not _passes_balance_filter(ln, balance_filter):
                    continue
                bstore = ln.balance_qty
                eff = material_status.value if material_status is not None else shipment_status.value
                if eff == ShipmentStatusFilter.sent.value:
                    if bstore is None or bstore > 0:
                        continue
                elif eff == ShipmentStatusFilter.not_sent.value:
                    if bstore is not None and bstore <= 0:
                        continue
                delivered = Decimal(0)
                bal = ln.balance_qty if ln.balance_qty is not None else Decimal(0)
                pending.append((h, ln, delivered, bal))

    pod_ids = {int(ln.id) for _, ln, _, _ in pending}
    prod_status = _latest_production_status_by_pod(db, pod_ids)

    out: list[dict[str, Any]] = []
    for h, ln, delivered, bal in pending:
        out.append(_line_to_row(h, ln, delivered, bal, prod_status.get(int(ln.id), "")))
    return out


# Kolom inti outstanding (tanpa gelombang kirim).
_OUTSTANDING_CORE_COLUMNS: list[tuple[str, str]] = [
    ("order_date", "DATE"),
    ("thi_order", "THI ORDER"),
    ("customer_name", "CUSTOMER NAME"),
    ("brand", "BRAND"),
    ("order_type", "ORDER TYPE"),
    ("material_name", "MATERIAL NAME"),
    ("thickness", "THICKNESS"),
    ("emboss", "EMBOSS"),
    ("softness_type", "SOFTNESS TYPE"),
    ("color", "COLOR"),
    ("softness", "SOFTNESS"),
    ("qty_sf", "QTY (S/F)"),
    ("unit_price", "PRICE"),
    ("color_reference", "COLOR S/W"),
    ("color_standard", "COLOR STD"),
    ("backing", "BACKING"),
    ("model", "MODEL"),
    ("style_article", "ARTICLE"),
    ("rtd_text", "RTA"),
    ("line_etd", "REPLY ETD"),
    ("stock_wh_qty", "STOCK WH"),
    ("balance_qty", "BALANCE"),
    ("system_reference", "SYSTEM NO"),
]

# Lembar PROD — NO PO + STATUS MATERIAL + 21 gelombang (mirip PROD SUEDE manual).
_PROD_OUTSTANDING_COLUMNS: list[tuple[str, str]] = (
    _OUTSTANDING_CORE_COLUMNS[:5]
    + [("no_po", "NO PO")]
    + _OUTSTANDING_CORE_COLUMNS[5:]
    + [
        ("detail_remark", "REMARK PO"),
        ("remark_marketing", "REMARK"),
        ("status_material", "STATUS MATERIAL"),
        ("mkt_stage", "TAHAP MKT"),
        ("mrp_alloc_status", "ALOKASI MRP"),
        ("mrp_qty_remaining", "SISA RENCANA"),
    ]
    + _delivery_export_column_pairs(OUTSTANDING_DELIVERY_SLOTS_PROD)
)

# Lembar CLAIM / SAMPLE / SUBCONT — PO CUST, tanpa STATUS MATERIAL, 5 gelombang.
_CLAIM_SAMPLE_OUTSTANDING_COLUMNS: list[tuple[str, str]] = (
    _OUTSTANDING_CORE_COLUMNS[:5]
    + [("po_customer_ref", "PO CUST")]
    + _OUTSTANDING_CORE_COLUMNS[5:]
    + [
        ("detail_remark", "REMARK PO"),
        ("remark_marketing", "REMARK"),
        ("mkt_stage", "TAHAP MKT"),
        ("mrp_alloc_status", "ALOKASI MRP"),
        ("mrp_qty_remaining", "SISA RENCANA"),
    ]
    + _delivery_export_column_pairs(OUTSTANDING_DELIVERY_SLOTS_CLAIM_SAMPLE)
)

# Kolom unduhan Excel gabungan (API legacy) — 21 gelombang + kolom audit internal.
_OUTSTANDING_XLS_COLUMNS: list[tuple[str, str]] = (
    _OUTSTANDING_CORE_COLUMNS[:5]
    + [("no_po", "NO PO"), ("po_customer_ref", "PO CUST")]
    + _OUTSTANDING_CORE_COLUMNS[5:]
    + [
        ("remark_system", "REMARK (SISTEM)"),
        ("remark_marketing", "REMARK (MARKETING)"),
        ("detail_remark", "CATATAN BARIS PO"),
        ("status_material", "STATUS MATERIAL"),
        ("delivered_qty", "QTY TERKIRIM"),
        ("schedule_qty_total", "RENCANA TOTAL"),
        ("schedule_unallocated_qty", "BELUM DIRENCANAKAN"),
        ("actual_vs_schedule_qty", "AKTUAL vs RENCANA"),
        ("production_status", "STATUS PRODUKSI"),
        ("mkt_stage", "TAHAP MKT"),
        ("mrp_alloc_status", "ALOKASI MRP"),
        ("mrp_qty_remaining", "SISA RENCANA"),
    ]
    + _delivery_export_column_pairs(OUTSTANDING_DELIVERY_SLOTS_PROD)
)


# Lembar «RETUR» — mapping dari baris outstanding (kategori return) ke kolom mirip Excel.
_RETURN_XLS_COLUMNS: list[tuple[str, str]] = [
    ("return_date", "RETURN DATE"),
    ("system_no_blank", "SYSTEM NO"),
    ("customer_name", "CUSTOMER NAME"),
    ("po_display", "PO NUMBER"),
    ("material_name", "MATERIAL NAME"),
    ("thickness", "THICKNESS"),
    ("thi_softness_col", "VIR SOFTNESS"),
    ("emboss", "EMBOSS"),
    ("color", "COLOR"),
    ("order_type", "SAMPLE/PROD"),
    ("brand", "BRAND"),
    ("qty_return", "QTY RETURN"),
    ("roll", "ROLL"),
    ("treatment", "TREATMENT"),
    ("reason_line", "REASON"),
    ("remark_followup", "REMARK"),
    ("confirmed_by", "CONFIRMED BY"),
    ("stock_wh_qty", "STOCK WH"),
    ("balance_qty", "BALANCE"),
    ("system_reference", "SYSTEM NUMBER"),
    ("detail_remark", "REMARK BARIS"),
    ("delivery_1_qty", "1st Qty"),
    ("delivery_1_date", "1st Date"),
    ("delivery_2_qty", "2nd Qty"),
    ("delivery_2_date", "2nd Date"),
    ("delivery_3_qty", "3rd Qty"),
    ("delivery_3_date", "3rd Date"),
]


def _outstanding_row_to_retur_export_row(r: dict[str, Any]) -> dict[str, Any]:
    """Flatkan satu baris outstanding ke skema export lembar RETUR."""
    po_cust = (r.get("po_customer_ref") or "").strip()
    no_po = (r.get("no_po") or "").strip()
    po_disp = po_cust or no_po
    thi = (r.get("thi_order") or "").strip()
    if thi in ("", "—"):
        thi = ""
    st = (r.get("softness_type") or "").strip()
    thi_soft = " | ".join(x for x in (thi, st) if x) or ""
    return {
        "return_date": r.get("order_date"),
        "system_no_blank": "",
        "customer_name": r.get("customer_name"),
        "po_display": po_disp,
        "material_name": r.get("material_name"),
        "thickness": r.get("thickness"),
        "thi_softness_col": thi_soft,
        "emboss": r.get("emboss"),
        "color": r.get("color"),
        "order_type": r.get("order_type"),
        "brand": r.get("brand"),
        "qty_return": r.get("balance_qty"),
        "roll": "",
        "treatment": "",
        "reason_line": (r.get("remark_marketing") or r.get("detail_remark") or "").strip() or "",
        "remark_followup": (r.get("material_followup_remark_header") or "").strip() or "",
        "confirmed_by": "",
        "stock_wh_qty": r.get("stock_wh_qty"),
        "balance_qty": r.get("balance_qty"),
        "system_reference": r.get("system_reference"),
        "detail_remark": (r.get("detail_remark") or "").strip() or "",
        "delivery_1_qty": r.get("delivery_1_qty"),
        "delivery_1_date": r.get("delivery_1_date"),
        "delivery_2_qty": r.get("delivery_2_qty"),
        "delivery_2_date": r.get("delivery_2_date"),
        "delivery_3_qty": r.get("delivery_3_qty"),
        "delivery_3_date": r.get("delivery_3_date"),
    }


def _return_rows_to_xlsx_bytes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "RETUR"
    keys = [k for k, _ in _RETURN_XLS_COLUMNS]
    ws.append([label for _, label in _RETURN_XLS_COLUMNS])
    for r in rows:
        er = _outstanding_row_to_retur_export_row(r)
        ws.append([er.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _outstanding_rows_to_xlsx_bytes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Outstanding"
    keys = [k for k, _ in _OUTSTANDING_XLS_COLUMNS]
    ws.append([label for _, label in _OUTSTANDING_XLS_COLUMNS])
    for r in rows:
        ws.append([r.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _write_outstanding_sheet(
    wb: Workbook,
    *,
    title: str,
    columns: list[tuple[str, str]],
    rows: list[dict[str, Any]],
    row_mapper: Optional[Any] = None,
    active: bool = False,
) -> None:
    keys = [k for k, _ in columns]
    if active:
        ws = wb.active
        ws.title = title[:31]
    else:
        ws = wb.create_sheet(title[:31])
    ws.append([label for _, label in columns])
    for r in rows:
        src = row_mapper(r) if row_mapper else r
        ws.append([src.get(k) for k in keys])


def _outstanding_manual_workbook_bytes(rows: list[dict[str, Any]]) -> bytes:
    """Workbook multi-lembar selaras file manual marketing (PROD / CLAIM / SAMPLE / SUBCONT / RETUR)."""
    wb = Workbook()
    by_cat: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        cat = (r.get("po_category") or "production").strip().lower()
        by_cat[cat].append(r)

    sheets: list[tuple[str, str, list[tuple[str, str]], Optional[Any]]] = [
        ("PROD SUEDE", "production", _PROD_OUTSTANDING_COLUMNS, None),
        ("CLAIM SUEDE", "claim", _CLAIM_SAMPLE_OUTSTANDING_COLUMNS, None),
        ("SAMPLE SUEDE", "sample", _CLAIM_SAMPLE_OUTSTANDING_COLUMNS, None),
        ("SUBCONT", "subcont", _CLAIM_SAMPLE_OUTSTANDING_COLUMNS, None),
        ("RETUR", "return", _RETURN_XLS_COLUMNS, _outstanding_row_to_retur_export_row),
    ]
    first = True
    for sheet_title, category, columns, mapper in sheets:
        cat_rows = by_cat.get(category, [])
        _write_outstanding_sheet(
            wb,
            title=sheet_title,
            columns=columns,
            rows=cat_rows,
            row_mapper=mapper,
            active=first,
        )
        first = False

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# Kolom unduhan Excel — selaras lembar «DELIVERY LIST DARI SISTEM» (label Indonesia).
_DELIVERY_LIST_XLS_COLUMNS: list[tuple[str, str]] = [
    ("surat_jalan_no", "NO SJ"),
    ("delivery_date", "TGL KIRIM"),
    ("factory_name", "FACTORY NAME"),
    ("system_no", "SYSTEM NO"),
    ("po_date", "PO DATE"),
    ("po_number", "PO NUMBER"),
    ("lot_no", "LOT NO"),
    ("thickness", "THICKNESS"),
    ("color", "COLOR"),
    ("material_type", "TYPE"),
    ("material_name", "MATERIAL NAME"),
    ("qty_po", "QTY PO"),
    ("unit_price", "UNIT PRICE"),
    ("qty_roll", "ROLL"),
    ("qty_pcs", "PCS"),
    ("qty_sf", "SF"),
    ("gross_weight", "GW"),
    ("net_weight", "NW"),
    ("amount", "AMOUNT"),
    ("color_std", "COLOR STD"),
    ("emboss", "EMBOSS"),
    ("buyer", "BUYER"),
    ("style_article", "STYLE"),
    ("model", "MODEL"),
    ("invoice_number", "INV NUMBER"),
    ("pi_no", "PI_NO"),
    ("remark", "REMARK"),
    ("order_type", "ORDER TYPE"),
    ("delivery_status", "STATUS DELIVERY"),
    ("packing_status", "STATUS PACKING"),
    ("invpl_printed", "INVPL SUDAH CETAK"),
    ("delivery_zone", "ZONA KIRIM"),
]


def _delivery_list_rows_to_xlsx_bytes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Delivery List"
    keys = [k for k, _ in _DELIVERY_LIST_XLS_COLUMNS]
    ws.append([label for _, label in _DELIVERY_LIST_XLS_COLUMNS])
    for r in rows:
        ws.append([r.get(k) for k in keys])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _delivery_detail_to_row(
    dd: DeliveryDetail,
    dh: DeliveryHeader,
    pd: Optional[ProductionDetail],
    pod: Optional[PoDetail],
    ph: Optional[PoHeader],
    cust: Optional[Customer],
    inv: Optional[Invoice],
    invd: Optional[InvoiceDetail],
    pkh: Optional[PackingHeader],
) -> dict[str, Any]:
    factory = ""
    if cust and (cust.name or "").strip():
        factory = (cust.name or "").strip()
    elif ph and (ph.customer_name_snapshot or ph.customer_code or "").strip():
        factory = (ph.customer_name_snapshot or ph.customer_code or "").strip()
    elif dh.destination:
        factory = (dh.destination or "")[:255]

    system_no = ""
    if ph and (ph.po_number or "").strip():
        system_no = (ph.po_number or "").strip()
    elif pod and (pod.system_reference or "").strip():
        system_no = (pod.system_reference or "").strip()

    lot_no = ""
    if pd and (pd.lot_number or "").strip():
        lot_no = (pd.lot_number or "").strip()
    elif pod and (pod.thi_line_ref or "").strip():
        lot_no = (pod.thi_line_ref or "").strip()

    sj_raw = (dh.delivery_number or "").strip()
    surat_jalan_no = surat_jalan_display_number(sj_raw, ph) if sj_raw else "—"
    inv_raw = (inv.invoice_number or "").strip() if inv else ""
    invoice_number = invoice_display_number(inv_raw, ph) if inv_raw else ""

    gw: Decimal | None = None
    nw: Decimal | None = None
    if invd is not None:
        if invd.gross_weight is not None:
            gw = invd.gross_weight
        if invd.net_weight is not None:
            nw = invd.net_weight
    if nw is None and dd.weight is not None:
        nw = dd.weight

    amt = _delivery_line_amount(dd, pod, invd)

    return {
        "surat_jalan_no": surat_jalan_no or "—",
        "delivery_date": dh.delivery_date.isoformat() if dh.delivery_date else None,
        "factory_name": factory or "—",
        "system_no": system_no or "—",
        "po_date": ph.order_date.isoformat() if ph and ph.order_date else None,
        "po_number": (ph.po_customer_ref or "").strip() if ph else "",
        "lot_no": lot_no,
        "thickness": _thickness_display(pod) if pod else "",
        "color": (pod.color or "").strip() if pod else "",
        "material_type": (pod.vir_softness_type or "").strip() if pod else "",
        "material_name": (pod.material_name or "").strip() if pod else "",
        "qty_po": _dec_str(pod.quantity_order) if pod else "",
        "unit_price": _dec_str(pod.unit_price) if pod else "",
        "qty_roll": dd.qty_roll if dd.qty_roll is not None else None,
        "qty_pcs": dd.qty_pcs if dd.qty_pcs is not None else None,
        "qty_sf": _dec_str(dd.qty),
        "gross_weight": _dec_str(gw) if gw is not None else "",
        "net_weight": _dec_str(nw) if nw is not None else "",
        "amount": _dec_str(amt) if amt is not None else "",
        "color_std": ((pod.color_standard or pod.color_reference) or "").strip() if pod else "",
        "emboss": (pod.emboss or "").strip() if pod else "",
        "buyer": ((ph.brand_name_snapshot or ph.buyer or "").strip() if ph else ""),
        "style_article": (pod.style_article or "").strip() if pod else "",
        "model": (pod.model or "").strip() if pod else "",
        "invoice_number": invoice_number,
        "pi_no": (ph.po_customer_ref or "").strip() if ph else "",
        "remark": ((pod.detail_remark or "").strip() if pod else ""),
        "order_type": _order_type_marketing(ph) if ph else "",
        "factory_code": ((ph.factory_code or ph.customer_code or "").strip() if ph else "")
        or ((cust.code or "").strip() if cust else ""),
        "destination": (dh.destination or "").strip() if dh.destination else "",
        "delivery_id": dh.id,
        "delivery_detail_id": dd.id,
        "delivery_status": dh.status or "",
        "packing_id": dh.packing_id,
        "packing_status": (pkh.status or "").strip() if pkh else "",
        "invpl_printed": bool(pkh and pkh.invpl_printed_at is not None),
        "delivery_zone": (dh.delivery_zone or "").strip() if dh.delivery_zone else "",
        "source_mrp": bool(getattr(dh, "mrp_fg_id", None) or (pkh and (pkh.source or "") == "mrp")),
    }


def _query_delivery_lines_rows(
    db: Session,
    q: Optional[str] = None,
    delivery_status: DeliveryLinesStatusFilter = DeliveryLinesStatusFilter.all,
    packing_status: PackingLinesStatusFilter = PackingLinesStatusFilter.all,
    invpl_printed: InvplPrintedFilter = InvplPrintedFilter.all,
    delivery_date_from: Optional[date] = None,
    delivery_date_to: Optional[date] = None,
    factory_code: Optional[str] = None,
    po_category: Optional[str] = None,
    customer: Optional[str] = None,
) -> list[dict[str, Any]]:
    """Baris delivery + join ke produksi & PO (format lembar DELIVERY LIST)."""
    stmt = (
        select(
            DeliveryDetail,
            DeliveryHeader,
            ProductionDetail,
            PoDetail,
            PoHeader,
            Customer,
            PackingHeader,
        )
        .join(DeliveryHeader, DeliveryDetail.delivery_id == DeliveryHeader.id)
        .outerjoin(ProductionDetail, DeliveryDetail.production_detail_id == ProductionDetail.id)
        .outerjoin(PoDetail, ProductionDetail.po_detail_id == PoDetail.id)
        .outerjoin(PoHeader, PoDetail.po_id == PoHeader.id)
        .outerjoin(Customer, DeliveryHeader.customer_id == Customer.id)
        .outerjoin(PackingHeader, DeliveryHeader.packing_id == PackingHeader.id)
        .order_by(DeliveryHeader.id.desc(), DeliveryDetail.id.asc())
    )
    if delivery_date_from is not None:
        stmt = stmt.where(
            DeliveryHeader.delivery_date.is_not(None),
            DeliveryHeader.delivery_date >= delivery_date_from,
        )
    if delivery_date_to is not None:
        stmt = stmt.where(
            DeliveryHeader.delivery_date.is_not(None),
            DeliveryHeader.delivery_date <= delivery_date_to,
        )
    fc = (factory_code or "").strip()
    if fc:
        like_fc = f"%{fc}%"
        stmt = stmt.where(
            or_(
                PoHeader.id.is_(None),
                PoHeader.factory_code.ilike(like_fc),
                PoHeader.customer_code.ilike(like_fc),
                Customer.code.ilike(like_fc),
            )
        )
    pc = (po_category or "").strip().lower()
    if pc:
        if pc == "production":
            stmt = stmt.where(
                or_(
                    PoHeader.id.is_(None),
                    PoHeader.po_category.in_(("production", "suede")),
                    PoHeader.po_category.is_(None),
                    PoHeader.po_category == "",
                )
            )
        else:
            stmt = stmt.where(or_(PoHeader.id.is_(None), PoHeader.po_category == pc))
    cq = (customer or "").strip()
    if cq:
        like_cust = f"%{cq}%"
        stmt = stmt.where(
            or_(
                Customer.name.ilike(like_cust),
                Customer.code.ilike(like_cust),
                PoHeader.customer_name_snapshot.ilike(like_cust),
            )
        )
    raw = db.execute(stmt).all()
    if not raw:
        return []

    delivery_ids = {row[1].id for row in raw}
    inv_by_delivery: dict[int, Invoice] = {}
    if delivery_ids:
        inv_rows = db.scalars(
            select(Invoice).where(Invoice.delivery_id.in_(delivery_ids))
        ).all()
        for inv in inv_rows:
            did = inv.delivery_id
            if did is None:
                continue
            if did not in inv_by_delivery:
                inv_by_delivery[did] = inv

    inv_ids = [i.id for i in inv_by_delivery.values()]
    details_by_invoice: dict[int, list[InvoiceDetail]] = defaultdict(list)
    if inv_ids:
        idets = db.scalars(
            select(InvoiceDetail)
            .where(InvoiceDetail.invoice_id.in_(inv_ids))
            .order_by(InvoiceDetail.invoice_id.asc(), InvoiceDetail.id.asc())
        ).all()
        for d in idets:
            details_by_invoice[int(d.invoice_id)].append(d)

    line_index: dict[int, int] = defaultdict(int)
    out: list[dict[str, Any]] = []
    for dd, dh, pd, pod, ph, cust, pkh in raw:
        # Jalur MRP: DeliveryDetail.po_detail_id langsung (tanpa ProductionDetail)
        if pod is None and getattr(dd, "po_detail_id", None):
            pod = db.get(PoDetail, int(dd.po_detail_id))
            if pod and ph is None and pod.po_id:
                ph = db.get(PoHeader, pod.po_id)
        inv = inv_by_delivery.get(dh.id)
        invd: InvoiceDetail | None = None
        if inv:
            lst = details_by_invoice.get(int(inv.id), [])
            li = line_index[dh.id]
            if li < len(lst):
                invd = lst[li]
        line_index[dh.id] += 1
        out.append(_delivery_detail_to_row(dd, dh, pd, pod, ph, cust, inv, invd, pkh))

    if delivery_status != DeliveryLinesStatusFilter.all:
        want = delivery_status.value
        out = [r for r in out if (r.get("delivery_status") or "").lower() == want]
    if packing_status != PackingLinesStatusFilter.all:
        want = packing_status.value
        out = [r for r in out if (r.get("packing_status") or "").lower() == want]
    if invpl_printed == InvplPrintedFilter.yes:
        out = [r for r in out if r.get("invpl_printed") is True]
    elif invpl_printed == InvplPrintedFilter.no:
        out = [r for r in out if r.get("invpl_printed") is not True]

    if q and str(q).strip():
        qq = str(q).strip().lower()

        def _match(row: dict[str, Any]) -> bool:
            for v in row.values():
                if v is None:
                    continue
                if qq in str(v).lower():
                    return True
            return False

        out = [r for r in out if _match(r)]
    return out


def _parse_po_detail_ids(raw: Optional[str]) -> Optional[list[int]]:
    if not raw or not str(raw).strip():
        return None
    out: list[int] = []
    for part in str(raw).split(","):
        part = part.strip()
        if part.isdigit():
            out.append(int(part))
    return out or None


@router.get("/marketing/outstanding-lines", response_model=MarketingOutstandingPageOut)
def marketing_outstanding_lines(
    db: Session = Depends(get_db),
    balance_filter: BalanceFilter = Query(
        default=BalanceFilter.exclude_zero,
        description="all: semua baris; exclude_zero: sembunyikan balance tepat 0; positive_only: hanya balance > 0",
    ),
    computed_balance: bool = Query(True, description="Hitung balance dari qty_order - qty_terkirim (delivery)"),
    po_q: Optional[str] = Query(
        None,
        description="Substring pencarian nomor PO (case-insensitive)",
    ),
    shipment_status: ShipmentStatusFilter = Query(
        default=ShipmentStatusFilter.all,
        description="all | sent (balance 0) | not_sent (balance > 0)",
    ),
    material_status: Optional[MaterialStatusFilter] = Query(
        default=None,
        description="Override shipment_status: all | sent | not_sent | cancelled",
    ),
    factory_code: Optional[str] = Query(None, description="Filter factory code (substring, case-insensitive)"),
    po_category: Optional[str] = Query(None, description="Filter kategori PO: production/sample/claim/return/subcont/suede"),
    material_group: Optional[str] = Query(None, description="Filter jenis material: pu | suede"),
    mkt_stage_min: Optional[str] = Query(None, description="Minimal tahap MKT, mis. in_outstanding"),
    order_date_from: Optional[str] = Query(None, description="Tanggal order PO header: dari (YYYY-MM-DD)"),
    order_date_to: Optional[str] = Query(None, description="Tanggal order PO header: sampai (YYYY-MM-DD)"),
    po_detail_ids: Optional[str] = Query(None, description="CSV po_detail_id — batasi ke baris terpilih"),
    skip: int = Query(0, ge=0, description="Offset for pagination"),
    limit: int = Query(100, ge=1, le=1000, description="Limit for pagination"),
    _user=Depends(get_current_user),
) -> MarketingOutstandingPageOut:
    rows = _prepare_outstanding_rows(
        db,
        balance_filter,
        computed_balance,
        po_q=po_q,
        shipment_status=shipment_status,
        material_status=material_status,
        factory_code=factory_code,
        po_category=po_category,
        material_group=material_group,
        mkt_stage_min=mkt_stage_min,
        order_date_from=_parse_iso_date_optional(order_date_from),
        order_date_to=_parse_iso_date_optional(order_date_to),
        po_detail_ids=_parse_po_detail_ids(po_detail_ids),
    )
    total = len(rows)
    paginated = rows[skip : skip + limit]
    return MarketingOutstandingPageOut(
        data=[MarketingOutstandingLineOut.model_validate(r) for r in paginated],
        total=total,
    )


@router.get("/marketing/outstanding-lines.csv")
def marketing_outstanding_lines_csv(
    db: Session = Depends(get_db),
    balance_filter: BalanceFilter = Query(default=BalanceFilter.exclude_zero),
    computed_balance: bool = Query(True),
    po_q: Optional[str] = Query(None),
    shipment_status: ShipmentStatusFilter = Query(default=ShipmentStatusFilter.all),
    material_status: Optional[MaterialStatusFilter] = Query(default=None),
    factory_code: Optional[str] = Query(None),
    po_category: Optional[str] = Query(None),
    material_group: Optional[str] = Query(None),
    mkt_stage_min: Optional[str] = Query(None),
    order_date_from: Optional[str] = Query(None),
    order_date_to: Optional[str] = Query(None),
    po_detail_ids: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    rows = _prepare_outstanding_rows(
        db,
        balance_filter,
        computed_balance,
        po_q=po_q,
        shipment_status=shipment_status,
        material_status=material_status,
        factory_code=factory_code,
        po_category=po_category,
        material_group=material_group,
        mkt_stage_min=mkt_stage_min,
        order_date_from=_parse_iso_date_optional(order_date_from),
        order_date_to=_parse_iso_date_optional(order_date_to),
        po_detail_ids=_parse_po_detail_ids(po_detail_ids),
    )
    buf = io.StringIO()
    if not rows:
        w = csv.writer(buf)
        w.writerow(["order_date", "thi_order", "customer_name", "no_po", "material_name"])
    else:
        fieldnames = list(rows[0].keys())
        w = csv.DictWriter(buf, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    return Response(
        content=buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="marketing_outstanding.csv"'},
    )


@router.get("/marketing/outstanding-lines.xlsx")
def marketing_outstanding_lines_xlsx(
    db: Session = Depends(get_db),
    balance_filter: BalanceFilter = Query(default=BalanceFilter.exclude_zero),
    computed_balance: bool = Query(True),
    po_q: Optional[str] = Query(None),
    shipment_status: ShipmentStatusFilter = Query(default=ShipmentStatusFilter.all),
    material_status: Optional[MaterialStatusFilter] = Query(default=None),
    factory_code: Optional[str] = Query(None),
    po_category: Optional[str] = Query(None),
    material_group: Optional[str] = Query(None),
    mkt_stage_min: Optional[str] = Query(None),
    order_date_from: Optional[str] = Query(None),
    order_date_to: Optional[str] = Query(None),
    po_detail_ids: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    rows = _prepare_outstanding_rows(
        db,
        balance_filter,
        computed_balance,
        po_q=po_q,
        shipment_status=shipment_status,
        material_status=material_status,
        factory_code=factory_code,
        po_category=po_category,
        material_group=material_group,
        mkt_stage_min=mkt_stage_min,
        order_date_from=_parse_iso_date_optional(order_date_from),
        order_date_to=_parse_iso_date_optional(order_date_to),
        po_detail_ids=_parse_po_detail_ids(po_detail_ids),
    )
    data = _outstanding_rows_to_xlsx_bytes(rows)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="marketing_outstanding.xlsx"'},
    )


@router.get("/marketing/outstanding-return-lines.xlsx")
def marketing_outstanding_return_lines_xlsx(
    db: Session = Depends(get_db),
    balance_filter: BalanceFilter = Query(default=BalanceFilter.exclude_zero),
    computed_balance: bool = Query(True),
    po_q: Optional[str] = Query(None),
    shipment_status: ShipmentStatusFilter = Query(default=ShipmentStatusFilter.all),
    material_status: Optional[MaterialStatusFilter] = Query(default=None),
    factory_code: Optional[str] = Query(None),
    material_group: Optional[str] = Query(None),
    mkt_stage_min: Optional[str] = Query(None),
    order_date_from: Optional[str] = Query(None),
    order_date_to: Optional[str] = Query(None),
    po_detail_ids: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    """Export lembar RETUR (hanya `po_category=return`) — kolom diselaraskan contoh Excel marketing."""
    rows = _prepare_outstanding_rows(
        db,
        balance_filter,
        computed_balance,
        po_q=po_q,
        shipment_status=shipment_status,
        material_status=material_status,
        factory_code=factory_code,
        po_category="return",
        material_group=material_group,
        mkt_stage_min=mkt_stage_min,
        order_date_from=_parse_iso_date_optional(order_date_from),
        order_date_to=_parse_iso_date_optional(order_date_to),
        po_detail_ids=_parse_po_detail_ids(po_detail_ids),
    )
    data = _return_rows_to_xlsx_bytes(rows)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="marketing_outstanding_retur.xlsx"'},
    )


@router.get("/marketing/outstanding-manual.xlsx")
def marketing_outstanding_manual_workbook_xlsx(
    db: Session = Depends(get_db),
    balance_filter: BalanceFilter = Query(default=BalanceFilter.exclude_zero),
    computed_balance: bool = Query(True),
    po_q: Optional[str] = Query(None),
    shipment_status: ShipmentStatusFilter = Query(default=ShipmentStatusFilter.all),
    material_status: Optional[MaterialStatusFilter] = Query(default=None),
    factory_code: Optional[str] = Query(None),
    material_group: Optional[str] = Query(None),
    mkt_stage_min: Optional[str] = Query(None),
    order_date_from: Optional[str] = Query(None),
    order_date_to: Optional[str] = Query(None),
    po_detail_ids: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    """Unduh workbook multi-lembar (PROD / CLAIM / SAMPLE / SUBCONT / RETUR) seperti Excel manual marketing."""
    rows = _prepare_outstanding_rows(
        db,
        balance_filter,
        computed_balance,
        po_q=po_q,
        shipment_status=shipment_status,
        material_status=material_status,
        factory_code=factory_code,
        po_category=None,
        material_group=material_group,
        mkt_stage_min=mkt_stage_min,
        order_date_from=_parse_iso_date_optional(order_date_from),
        order_date_to=_parse_iso_date_optional(order_date_to),
        po_detail_ids=_parse_po_detail_ids(po_detail_ids),
    )
    data = _outstanding_manual_workbook_bytes(rows)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="data_outstanding_order.xlsx"'},
    )


@router.get("/marketing/delivery-lines", response_model=MarketingDeliveryLinesPageOut)
def marketing_delivery_lines(
    db: Session = Depends(get_db),
    q: Optional[str] = Query(None, description="Saring baris (substring, case-insensitive)"),
    delivery_status: DeliveryLinesStatusFilter = Query(
        default=DeliveryLinesStatusFilter.all,
        description="all | draft | sent — status header delivery",
    ),
    packing_status: PackingLinesStatusFilter = Query(
        default=PackingLinesStatusFilter.all,
        description="all | draft | done — status packing terkait",
    ),
    invpl_printed: InvplPrintedFilter = Query(
        default=InvplPrintedFilter.all,
        description="all | yes | no — INVPL sudah ditandai dicetak",
    ),
    delivery_date_from: Optional[str] = Query(
        None,
        description="Filter tgl kirim header SJ: dari (YYYY-MM-DD)",
    ),
    delivery_date_to: Optional[str] = Query(
        None,
        description="Filter tgl kirim header SJ: sampai (YYYY-MM-DD)",
    ),
    factory_code: Optional[str] = Query(None, description="Substring kode pabrik (PO / customer)"),
    po_category: Optional[str] = Query(None, description="Kategori PO: production/sample/claim/return/subcont/suede"),
    customer: Optional[str] = Query(None, description="Substring nama atau kode customer"),
    skip: int = Query(0, ge=0, description="Offset for pagination"),
    limit: int = Query(100, ge=1, le=1000, description="Limit for pagination"),
    _user=Depends(get_current_user),
) -> MarketingDeliveryLinesPageOut:
    rows = _query_delivery_lines_rows(
        db,
        q=q,
        delivery_status=delivery_status,
        packing_status=packing_status,
        invpl_printed=invpl_printed,
        delivery_date_from=_parse_iso_date_optional(delivery_date_from),
        delivery_date_to=_parse_iso_date_optional(delivery_date_to),
        factory_code=factory_code,
        po_category=po_category,
        customer=customer,
    )
    total = len(rows)
    paginated = rows[skip : skip + limit]
    return MarketingDeliveryLinesPageOut(
        data=[MarketingDeliveryLineOut.model_validate(r) for r in paginated],
        total=total,
    )


@router.get("/marketing/delivery-lines.csv")
def marketing_delivery_lines_csv(
    db: Session = Depends(get_db),
    q: Optional[str] = Query(None),
    delivery_status: DeliveryLinesStatusFilter = Query(default=DeliveryLinesStatusFilter.all),
    packing_status: PackingLinesStatusFilter = Query(default=PackingLinesStatusFilter.all),
    invpl_printed: InvplPrintedFilter = Query(default=InvplPrintedFilter.all),
    delivery_date_from: Optional[str] = Query(None),
    delivery_date_to: Optional[str] = Query(None),
    factory_code: Optional[str] = Query(None),
    po_category: Optional[str] = Query(None),
    customer: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    rows = _query_delivery_lines_rows(
        db,
        q=q,
        delivery_status=delivery_status,
        packing_status=packing_status,
        invpl_printed=invpl_printed,
        delivery_date_from=_parse_iso_date_optional(delivery_date_from),
        delivery_date_to=_parse_iso_date_optional(delivery_date_to),
        factory_code=factory_code,
        po_category=po_category,
        customer=customer,
    )
    buf = io.StringIO()
    if not rows:
        w = csv.writer(buf)
        w.writerow(
            [
                "surat_jalan_no",
                "delivery_date",
                "factory_name",
                "system_no",
                "po_number",
            ]
        )
    else:
        fieldnames = list(rows[0].keys())
        w = csv.DictWriter(buf, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    return Response(
        content=buf.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="marketing_delivery_list.csv"'},
    )


@router.get("/marketing/delivery-lines.xlsx")
def marketing_delivery_lines_xlsx(
    db: Session = Depends(get_db),
    q: Optional[str] = Query(None),
    delivery_status: DeliveryLinesStatusFilter = Query(default=DeliveryLinesStatusFilter.all),
    packing_status: PackingLinesStatusFilter = Query(default=PackingLinesStatusFilter.all),
    invpl_printed: InvplPrintedFilter = Query(default=InvplPrintedFilter.all),
    delivery_date_from: Optional[str] = Query(None),
    delivery_date_to: Optional[str] = Query(None),
    factory_code: Optional[str] = Query(None),
    po_category: Optional[str] = Query(None),
    customer: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> Response:
    rows = _query_delivery_lines_rows(
        db,
        q=q,
        delivery_status=delivery_status,
        packing_status=packing_status,
        invpl_printed=invpl_printed,
        delivery_date_from=_parse_iso_date_optional(delivery_date_from),
        delivery_date_to=_parse_iso_date_optional(delivery_date_to),
        factory_code=factory_code,
        po_category=po_category,
        customer=customer,
    )
    data = _delivery_list_rows_to_xlsx_bytes(rows)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="delivery_list_sistem.xlsx"'},
    )


def _parse_balance_filter(val: str) -> BalanceFilter:
    try:
        return BalanceFilter(str(val).strip())
    except ValueError:
        return BalanceFilter.exclude_zero


def _parse_shipment_status(val: str) -> ShipmentStatusFilter:
    try:
        return ShipmentStatusFilter(str(val).strip())
    except ValueError:
        return ShipmentStatusFilter.all


def _parse_material_status_optional(val: Optional[str]) -> Optional[MaterialStatusFilter]:
    if not val or not str(val).strip():
        return None
    try:
        return MaterialStatusFilter(str(val).strip())
    except ValueError:
        return None


def _query_delivery_planning_slots(
    db: Session,
    date_from: date,
    date_to: date,
    factory_code: Optional[str] = None,
    mkt_stage_min: Optional[str] = None,
    po_q: Optional[str] = None,
    *,
    as_of: Optional[date] = None,
) -> list[dict[str, Any]]:
    fc = (factory_code or "").strip().lower()
    min_stage = _normalize_mkt_stage_min(mkt_stage_min)
    stmt = (
        select(PoDetailDelivery, PoDetail, PoHeader)
        .join(PoDetail, PoDetailDelivery.po_detail_id == PoDetail.id)
        .join(PoHeader, PoDetail.po_id == PoHeader.id)
        .where(PoDetailDelivery.delivery_date.is_not(None))
        .where(PoDetailDelivery.delivery_date >= date_from)
        .where(PoDetailDelivery.delivery_date <= date_to)
        .order_by(
            PoDetailDelivery.delivery_date.asc(),
            PoHeader.id.desc(),
            PoDetail.id.asc(),
            PoDetailDelivery.sequence_no.asc(),
        )
    )
    raw = db.execute(stmt).all()
    delivered_map = _delivered_qty_by_po_detail(db)
    ref = as_of or date.today()
    out: list[dict[str, Any]] = []
    for pdd, pod, ph in raw:
        st = (ph.status or "").lower()
        if st in ("completed", "cancelled", "canceled"):
            continue
        if pdd.qty is not None and pdd.qty <= 0:
            continue
        if fc and fc not in ((ph.factory_code or "").strip().lower()):
            continue
        if not _po_q_matches(ph, po_q):
            continue
        stage = normalize_mkt_stage(getattr(ph, "mkt_workflow_stage", None))
        if min_stage and not _passes_mkt_stage_min(stage, min_stage):
            continue
        order = pod.quantity_order or Decimal(0)
        delivered = delivered_map.get(int(pod.id), Decimal(0))
        balance = order - delivered
        ship_status = _planning_slot_ship_status(stage, delivered, order)
        slot_date = pdd.delivery_date
        out.append(
            {
                "delivery_date": slot_date.isoformat() if slot_date else None,
                "sequence_no": pdd.sequence_no,
                "qty": _dec_str(pdd.qty) if pdd.qty is not None else "",
                "po_id": ph.id,
                "po_number": (ph.po_number or "").strip(),
                "po_customer_ref": (ph.po_customer_ref or "").strip(),
                "customer_name": (ph.customer_name_snapshot or ph.customer_code or "").strip(),
                "brand": (ph.brand_name_snapshot or ph.buyer or "").strip(),
                "factory_code": (ph.factory_code or "").strip(),
                "material_name": (pod.material_name or "").strip(),
                "thi_order": (ph.po_number or "").strip(),
                "thi_line_ref": (pod.thi_line_ref or "").strip(),
                "lot_no": (pod.thi_line_ref or "").strip(),
                "po_detail_id": pod.id,
                "line_etd": pod.line_etd.isoformat() if pod.line_etd else None,
                "mkt_stage": stage,
                "order_qty": _dec_str(order),
                "delivered_qty": _dec_str(delivered),
                "balance_qty": _dec_str(balance),
                "ship_status": ship_status,
                "overdue": bool(
                    slot_date
                    and slot_date < ref
                    and ship_status in {"planned", "partial"}
                ),
            }
        )
    return out


def _query_overdue_planning_slots(
    db: Session,
    before: date,
    factory_code: Optional[str] = None,
    mkt_stage_min: Optional[str] = None,
    po_q: Optional[str] = None,
) -> list[dict[str, Any]]:
    start = date(max(before.year - 2, 2000), 1, 1)
    if start >= before:
        return []
    until = before - timedelta(days=1)
    slots = _query_delivery_planning_slots(
        db,
        start,
        until,
        factory_code=factory_code,
        mkt_stage_min=mkt_stage_min,
        po_q=po_q,
        as_of=before,
    )
    return [s for s in slots if s.get("overdue")]


def _query_line_etd_on_date(db: Session, on_date: date, factory_code: Optional[str] = None) -> list[dict[str, Any]]:
    fc = (factory_code or "").strip().lower()
    stmt = (
        select(PoDetail, PoHeader)
        .join(PoHeader, PoDetail.po_id == PoHeader.id)
        .where(PoDetail.line_etd == on_date)
        .order_by(PoHeader.id.desc(), PoDetail.id.asc())
    )
    out: list[dict[str, Any]] = []
    for pod, ph in db.execute(stmt).all():
        st = (ph.status or "").lower()
        if st in ("completed", "cancelled", "canceled"):
            continue
        if fc and fc not in ((ph.factory_code or "").strip().lower()):
            continue
        cust = (ph.customer_name_snapshot or ph.customer_code or "").strip() or "—"
        out.append(
            {
                "line_etd": pod.line_etd.isoformat() if pod.line_etd else None,
                "po_id": ph.id,
                "po_number": (ph.po_number or "").strip(),
                "factory_code": (ph.factory_code or "").strip(),
                "customer_name": cust,
                "material_name": (pod.material_name or "").strip(),
                "thi_order": (ph.po_number or "").strip(),
                "lot_no": (pod.thi_line_ref or "").strip(),
                "po_detail_id": pod.id,
            }
        )
    return out


def _delivery_lines_for_header_date(db: Session, on_date: date) -> list[dict[str, Any]]:
    rows = _query_delivery_lines_rows(db)
    dstr = on_date.isoformat()
    return [r for r in rows if r.get("delivery_date") and str(r.get("delivery_date")).startswith(dstr)]


@router.post(
    "/marketing/outstanding-snapshots",
    response_model=MarketingOutstandingSnapshotListItem,
)
def create_marketing_outstanding_snapshot(
    body: MarketingOutstandingSnapshotCreate,
    db: Session = Depends(get_db),
    user=Depends(require_roles("admin", "marketing")),
) -> MarketingOutstandingSnapshotListItem:
    f = body.filters
    rows = _prepare_outstanding_rows(
        db,
        _parse_balance_filter(f.balance_filter),
        f.computed_balance,
        po_q=f.po_q,
        shipment_status=_parse_shipment_status(f.shipment_status),
        material_status=_parse_material_status_optional(f.material_status),
        factory_code=f.factory_code,
        po_category=f.po_category,
        material_group=f.material_group,
        mkt_stage_min=f.mkt_stage_min,
        order_date_from=_parse_iso_date_optional(f.order_date_from),
        order_date_to=_parse_iso_date_optional(f.order_date_to),
    )
    snap = MarketingOutstandingSnapshot(
        label=(body.label or "").strip() or None,
        filter_params_json=f.model_dump(),
        rows_json=rows,
        row_count=len(rows),
        created_by_username=user.username,
    )
    db.add(snap)
    db.commit()
    db.refresh(snap)
    return MarketingOutstandingSnapshotListItem.model_validate(snap)


@router.get("/marketing/outstanding-snapshots", response_model=list[MarketingOutstandingSnapshotListItem])
def list_marketing_outstanding_snapshots(
    db: Session = Depends(get_db),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    _user=Depends(get_current_user),
) -> list[MarketingOutstandingSnapshotListItem]:
    rows = db.scalars(
        select(MarketingOutstandingSnapshot)
        .order_by(MarketingOutstandingSnapshot.id.desc())
        .offset(skip)
        .limit(limit)
    ).all()
    return [MarketingOutstandingSnapshotListItem.model_validate(r) for r in rows]


@router.get(
    "/marketing/outstanding-snapshots/{snapshot_id}",
    response_model=MarketingOutstandingSnapshotDetailOut,
)
def get_marketing_outstanding_snapshot(
    snapshot_id: int,
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> MarketingOutstandingSnapshotDetailOut:
    row = db.get(MarketingOutstandingSnapshot, snapshot_id)
    if not row:
        raise HTTPException(status_code=404, detail="Snapshot tidak ditemukan")
    return MarketingOutstandingSnapshotDetailOut.model_validate(row)


@router.get("/marketing/delivery-planning", response_model=MarketingDeliveryPlanningPageOut)
def marketing_delivery_planning(
    db: Session = Depends(get_db),
    date_from: Optional[str] = Query(None, description="YYYY-MM-DD"),
    date_to: Optional[str] = Query(None, description="YYYY-MM-DD"),
    factory_code: Optional[str] = Query(None),
    mkt_stage_min: Optional[str] = Query(None, description="Kosong = semua tahap; in_outstanding = siap outstanding ke atas"),
    po_q: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> MarketingDeliveryPlanningPageOut:
    today = date.today()
    df = _parse_iso_date_optional(date_from) or today
    dt = _parse_iso_date_optional(date_to) or today
    if df > dt:
        df, dt = dt, df
    slots = _query_delivery_planning_slots(
        db,
        df,
        dt,
        factory_code=factory_code,
        mkt_stage_min=mkt_stage_min,
        po_q=po_q,
        as_of=today,
    )
    return MarketingDeliveryPlanningPageOut(
        date_from=df.isoformat(),
        date_to=dt.isoformat(),
        data=[MarketingDeliveryPlanningSlotOut.model_validate(s) for s in slots],
        total=len(slots),
    )


@router.get("/marketing/delivery-today", response_model=MarketingDeliveryTodayOut)
def marketing_delivery_today(
    db: Session = Depends(get_db),
    factory_code: Optional[str] = Query(None),
    on_date: Optional[str] = Query(None, description="Hari acuan (default hari ini), YYYY-MM-DD"),
    mkt_stage_min: Optional[str] = Query(None),
    po_q: Optional[str] = Query(None),
    _user=Depends(get_current_user),
) -> MarketingDeliveryTodayOut:
    d = _parse_iso_date_optional(on_date) or date.today()
    planned = _query_delivery_planning_slots(
        db,
        d,
        d,
        factory_code=factory_code,
        mkt_stage_min=mkt_stage_min,
        po_q=po_q,
        as_of=d,
    )
    overdue = _query_overdue_planning_slots(
        db,
        before=d,
        factory_code=factory_code,
        mkt_stage_min=mkt_stage_min,
        po_q=po_q,
    )
    line_etd = _query_line_etd_on_date(db, d, factory_code=factory_code)
    sj_lines = _delivery_lines_for_header_date(db, d)
    return MarketingDeliveryTodayOut(
        as_of=d.isoformat(),
        planned_delivery_slots=[MarketingDeliveryPlanningSlotOut.model_validate(s) for s in planned],
        overdue_slots=[MarketingDeliveryPlanningSlotOut.model_validate(s) for s in overdue],
        lines_reply_etd_today=[MarketingLineEtdTodayOut.model_validate(s) for s in line_etd],
        surat_jalan_lines=[MarketingDeliveryLineOut.model_validate(s) for s in sj_lines],
    )


@router.post("/marketing/outstanding-schedule-batch", response_model=OutstandingScheduleBatchOut)
def marketing_outstanding_schedule_batch(
    body: OutstandingScheduleBatchIn,
    db: Session = Depends(get_db),
    user=Depends(require_roles("admin", "marketing")),
) -> OutstandingScheduleBatchOut:
    """Jadwalkan delivery untuk satu/banyak baris outstanding (lintas SO)."""
    result = apply_outstanding_schedule_batch(db, body, actor=user.username)
    return OutstandingScheduleBatchOut.model_validate(result)


@router.get("/marketing/outstanding-packing-options", response_model=OutstandingPackingOptionsOut)
def marketing_outstanding_packing_options(
    po_detail_ids: str = Query(..., description="CSV po_detail_id"),
    db: Session = Depends(get_db),
    _user=Depends(get_current_user),
) -> OutstandingPackingOptionsOut:
    """Daftar packing list (MRP/QC OK) siap kirim untuk baris outstanding terpilih."""
    ids: list[int] = []
    for part in str(po_detail_ids).split(","):
        part = part.strip()
        if part.isdigit():
            ids.append(int(part))
    rows = packing_options_for_details(db, ids)
    return OutstandingPackingOptionsOut(
        data=[OutstandingPackingOptionOut.model_validate(r) for r in rows],
        total=len(rows),
    )


@router.post("/marketing/outstanding-ship-batch", response_model=OutstandingShipBatchOut)
def marketing_outstanding_ship_batch(
    body: OutstandingShipBatchIn,
    db: Session = Depends(get_db),
    user=Depends(require_roles("admin", "marketing", "warehouse")),
) -> OutstandingShipBatchOut:
    """Susun delivery dari packing list yang dipilih per baris outstanding."""
    result = apply_outstanding_ship_batch(db, body, actor=user.username)
    return OutstandingShipBatchOut.model_validate(result)
