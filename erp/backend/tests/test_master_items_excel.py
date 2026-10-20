"""Export/import Excel master items — round-trip, anti-duplikat, update parsial."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import MasterItem
from app.services.master_items_excel import (
    EXPORT_COLUMNS,
    export_master_items_template_xlsx,
    export_master_items_xlsx,
    import_master_items_xlsx,
)


def _xlsx_from_rows(rows: list[dict], *, headers: list[str] | None = None) -> bytes:
    headers = headers or list(EXPORT_COLUMNS)
    wb = Workbook()
    ws = wb.active
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(headers, 1):
            key = h if h in EXPORT_COLUMNS else h
            # allow title headers in input rows keyed by EXPORT_COLUMNS
            if h in EXPORT_COLUMNS:
                ws.cell(r_idx, c_idx, row.get(h, ""))
            else:
                # map title back
                snake = h.lower().replace(" ", "_")
                ws.cell(r_idx, c_idx, row.get(snake, row.get(h, "")))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_db(items: list[MasterItem] | None = None):
    items = list(items or [])
    db = MagicMock()

    def scalars(stmt):
        result = MagicMock()
        text = str(stmt).lower()
        if "master_items" in text or True:
            result.all.return_value = list(items)
            result.first.return_value = items[0] if items else None
        return result

    db.scalars.side_effect = scalars

    def add(obj):
        items.append(obj)

    def flush():
        for it in items:
            if getattr(it, "id", None) is None:
                it.id = len(items)

    db.add.side_effect = add
    db.flush.side_effect = flush
    return db, items


def test_export_uses_snake_case_headers():
    row = MasterItem(
        item_code="FG1",
        category="PU",
        material_name="Leather",
        thickness_range="1.2-1.4",
        color="BLACK (00A)",
        emboss="R-8",
    )
    db, _ = _make_db([row])
    data = export_master_items_xlsx(db)
    wb = load_workbook(BytesIO(data))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS


def test_template_headers_only():
    data = export_master_items_template_xlsx()
    wb = load_workbook(BytesIO(data))
    assert wb.active.max_row == 1
    assert [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)] == EXPORT_COLUMNS


def test_import_create_then_reimport_skips_unchanged():
    db, items = _make_db([])
    content = _xlsx_from_rows(
        [
            {
                "category": "PU",
                "item_code": "FG100",
                "material_name": "Virgin PU",
                "thickness_range": "1.2-1.4",
                "color": "BLACK (00A)",
                "emboss": "R-8",
            }
        ]
    )
    r1 = import_master_items_xlsx(db, content)
    assert r1.created == 1
    assert r1.updated == 0
    assert len(items) == 1
    assert items[0].item_code == "FG100"
    assert items[0].color_code  # extracted

    r2 = import_master_items_xlsx(db, content)
    assert r2.created == 0
    assert r2.updated == 0
    assert r2.skipped == 1


def test_import_duplicate_item_code_in_file_no_double_create():
    db, items = _make_db([])
    content = _xlsx_from_rows(
        [
            {
                "category": "PU",
                "item_code": "FG200",
                "material_name": "A",
                "thickness_range": "1.2-1.4",
                "color": "BLACK",
                "emboss": "",
            },
            {
                "category": "PU",
                "item_code": "FG200",
                "material_name": "B",
                "thickness_range": "1.2-1.4",
                "color": "WHITE",
                "emboss": "X",
            },
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.created == 1
    assert result.updated == 1
    assert len(items) == 1
    assert items[0].material_name == "B"
    assert items[0].color == "WHITE"


def test_import_partial_color_only_keeps_other_fields():
    existing = MasterItem(
        id=1,
        item_code="FG300",
        category="Suede",
        material_name="Split Suede",
        thickness_range="1.2-1.4",
        color="BLACK (00A)",
        emboss="Nappa",
    )
    db, items = _make_db([existing])
    content = _xlsx_from_rows(
        [
            {
                "category": "",
                "item_code": "FG300",
                "material_name": "",
                "thickness_range": "",
                "color": "BROWN (11B)",
                "emboss": "",
            }
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.updated == 1
    assert items[0].color == "BROWN (11B)"
    assert items[0].material_name == "Split Suede"
    assert items[0].thickness_range == "1.2-1.4"
    assert items[0].emboss == "Nappa"
    assert items[0].category == "Suede"


def test_import_missing_item_code_and_invalid_category():
    db, items = _make_db([])
    content = _xlsx_from_rows(
        [
            {
                "category": "PU",
                "item_code": "",
                "material_name": "X",
                "thickness_range": "",
                "color": "",
                "emboss": "",
            },
            {
                "category": "INVALID",
                "item_code": "FG400",
                "material_name": "Y",
                "thickness_range": "",
                "color": "",
                "emboss": "",
            },
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.created == 0
    assert len(items) == 0
    assert result.skipped == 2
    assert len(result.errors) == 2


def test_import_accepts_legacy_title_headers():
    db, items = _make_db([])
    titles = [c.replace("_", " ").title() for c in EXPORT_COLUMNS]
    content = _xlsx_from_rows(
        [
            {
                "category": "suide",
                "item_code": "FG500",
                "material_name": "Suede Mat",
                "thickness_range": "1.0-1.2",
                "color": "NAVY",
                "emboss": "",
            }
        ],
        headers=titles,
    )
    result = import_master_items_xlsx(db, content)
    assert result.created == 1
    assert items[0].category == "Suede"


def test_import_flags_blank_keeps_existing_and_true_sets_new():
    existing = MasterItem(
        id=1,
        item_code="FG700",
        category="PU",
        material_name="Virgin PU",
        thickness_range="1.2-1.4",
        flag_color_black=True,
    )
    db, items = _make_db([existing])
    content = _xlsx_from_rows(
        [
            {
                "category": "",
                "item_code": "FG700",
                "material_name": "",
                "thickness_range": "",
                "color": "",
                "emboss": "",
                "flag_color_black": "",
                "flag_metalic": "true",
            }
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.updated == 1
    assert items[0].flag_color_black is True
    assert items[0].flag_metalic is True
    assert not items[0].flag_emboss


def test_import_invalid_flag_skipped():
    db, items = _make_db([])
    content = _xlsx_from_rows(
        [
            {
                "category": "PU",
                "item_code": "FG800",
                "flag_color": "maybe",
            }
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.created == 0
    assert result.skipped == 1
    assert "flag addon" in result.errors[0].reason
    assert items == []


def test_import_new_requires_category():
    db, items = _make_db([])
    content = _xlsx_from_rows(
        [
            {
                "category": "",
                "item_code": "FG600",
                "material_name": "Only Mat",
                "thickness_range": "",
                "color": "",
                "emboss": "",
            }
        ]
    )
    result = import_master_items_xlsx(db, content)
    assert result.created == 0
    assert result.errors
    assert "category wajib" in result.errors[0].reason
