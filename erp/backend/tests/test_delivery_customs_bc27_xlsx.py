"""Export Excel BC 2.7 — multi SJ, customer sama, kolom kuning kosong, bobot MRP."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.models import (
    Customer,
    DeliveryDetail,
    DeliveryHeader,
    ExchangeRate,
    Invoice,
    InvoiceDetail,
    MasterItem,
    PackingHeader,
    PoDetail,
)
from app.services.customs_numbering import generate_nomor_aju, generate_nomor_kontrak
from app.services.delivery_customs_bc27_xlsx import (
    BAHANBAKU_YELLOW_KEYS,
    render_customs_bc27_xlsx,
    render_delivery_customs_bc27_xlsx,
    resolve_bruto_netto,
)

AJU_SEQ = "004407"


def _col_map(ws):
    return {ws.cell(1, c).value: c for c in range(1, (ws.max_column or 0) + 1) if ws.cell(1, c).value}


def _make_db(*, headers, customer, details_by_delivery, invoices, rates, pods, masters, packings=None):
    db = MagicMock()
    packings = packings or {}

    def scalars(stmt):
        result = MagicMock()
        entity = None
        try:
            entity = stmt.column_descriptions[0]["entity"]
        except Exception:
            entity = None
        text = str(stmt).lower()
        if entity is DeliveryDetail or "delivery_details" in text:
            found = None
            try:
                for crit in stmt._where_criteria:  # type: ignore[attr-defined]
                    right = getattr(crit, "right", None)
                    val = getattr(right, "value", None)
                    if val is not None and int(val) in details_by_delivery:
                        found = int(val)
                        break
            except Exception:
                found = None
            if found is not None:
                result.all.return_value = details_by_delivery.get(found, [])
            else:
                flat = []
                for rows in details_by_delivery.values():
                    flat.extend(rows)
                result.all.return_value = flat
        elif entity is Invoice or "invoices" in text:
            result.all.return_value = invoices
        elif entity is InvoiceDetail or "invoice_details" in text:
            result.all.return_value = []
        elif entity is ExchangeRate or "exchange_rates" in text:
            result.all.return_value = rates
        else:
            result.all.return_value = []
        return result

    db.scalars.side_effect = scalars

    header_map = {int(h.id): h for h in headers}

    def _get(model, pk):
        pk = int(pk)
        if model is DeliveryHeader:
            return header_map.get(pk)
        if model is Customer:
            return customer if customer and int(customer.id) == pk else None
        if model is PoDetail:
            return pods.get(pk)
        if model is MasterItem:
            return masters.get(pk)
        if model is PackingHeader:
            return packings.get(pk)
        return None

    db.get.side_effect = _get
    return db


def test_generate_nomor_aju_format():
    assert generate_nomor_aju(date(2026, 10, 2), AJU_SEQ) == "000027TOP41720261002004407"
    assert generate_nomor_aju(date(2026, 10, 2), "004410") == "000027TOP41720261002004410"


def test_generate_nomor_kontrak_format():
    assert generate_nomor_kontrak("SC006", date(2026, 10, 2)) == "002/KPJB-THI/SC006/X/2026"
    assert generate_nomor_kontrak("HI024", date(2026, 10, 6)) == "006/KPJB-THI/HI024/X/2026"


def test_render_bc27_single_fills_red_mappings_yellow_empty():
    dh = SimpleNamespace(
        id=4407,
        delivery_number="THI-M261001012",
        customer_id=9,
        delivery_date=date(2026, 10, 2),
        driver_name="Budi",
        vehicle_number="B 1234 XX",
        packing_id=None,
    )
    cust = SimpleNamespace(
        id=9,
        code="SC006",
        name="SELALU CINTA INDONESIA",
        address="JL. LINGKAR SELATAN",
        billing_address=None,
        tax_id="0729241489505000000000",
        skep_kb_number="193/WBC.10/2020",
        customs_office_code="060800",
        default_currency="USD",
        reference_quotation=None,
    )
    pod = SimpleNamespace(
        id=1,
        master_item_id=2,
        material_name="VIRGIN HYDRO PU LEATHER",
        thickness_text="1.2-1.4MM",
        color="BLACK (00A)",
        emboss="R-8 M",
        unit_price=Decimal("1.20"),
        system_reference=None,
    )
    master = SimpleNamespace(id=2, item_code="FG00071", exim_material_code=None)
    dd = SimpleNamespace(
        id=1,
        po_detail_id=1,
        production_detail_id=None,
        packing_detail_id=None,
        qty=Decimal("100"),
        qty_roll=10,
        qty_pcs=0,
        weight=Decimal("12.5"),
    )
    inv = SimpleNamespace(
        id=3,
        invoice_number="THIN-M261001012",
        invoice_date=date(2026, 10, 2),
        delivery_id=4407,
    )
    rate = SimpleNamespace(idr_per_usd=Decimal("17863"), effective_date=date(2026, 10, 1))

    db = _make_db(
        headers=[dh],
        customer=cust,
        details_by_delivery={4407: [dd]},
        invoices=[inv],
        rates=[rate],
        pods={1: pod},
        masters={2: master},
    )

    data, name = render_delivery_customs_bc27_xlsx(db, 4407, aju_seq=AJU_SEQ)
    assert name.endswith("_bc27.xlsx")
    assert data[:2] == b"PK"

    wb = load_workbook(BytesIO(data))
    header = {
        wb["HEADER"].cell(1, c).value: wb["HEADER"].cell(2, c).value
        for c in range(1, wb["HEADER"].max_column + 1)
    }
    assert header["NOMOR AJU"] == "000027TOP41720261002004407"
    assert header["KODE DOKUMEN"] == "27"
    assert header["KODE KANTOR TUJUAN"] == "060800"
    assert header["CIF"] == 120.0
    assert header["NDPBM"] == 17863.0
    assert header["PPN PAJAK"] == pytest.approx(120.0 * 17863 * 0.11, rel=1e-6)
    assert header["NOMOR DAFTAR"] is None
    assert header["NETTO"] == 12.5
    assert header["BRUTO"] == 12.5

    bmap = _col_map(wb["BARANG"])
    assert wb["BARANG"].cell(2, bmap["KODE BARANG"]).value == "FG00071"
    assert wb["BARANG"].cell(2, bmap["HS"]).value is None
    assert wb["BARANG"].cell(2, bmap["JUMLAH SATUAN"]).value == 100.0
    assert wb["BARANG"].cell(2, bmap["CIF"]).value == 120.0

    bbmap = _col_map(wb["BAHANBAKU"])
    for key in BAHANBAKU_YELLOW_KEYS:
        assert wb["BAHANBAKU"].cell(2, bbmap[key]).value is None, key
    assert wb["BAHANBAKU"].cell(2, bbmap["JUMLAH SATUAN"]).value == 100.0

    btmap = _col_map(wb["BAHANBAKUTARIF"])
    assert wb["BAHANBAKUTARIF"].cell(2, btmap["NILAI FASILITAS"]).value is None

    dmap = _col_map(wb["DOKUMEN"])
    codes = [wb["DOKUMEN"].cell(r, dmap["KODE DOKUMEN"]).value for r in range(2, wb["DOKUMEN"].max_row + 1)]
    assert codes == ["380", "630", "315"]
    kontrak_rows = [
        wb["DOKUMEN"].cell(r, dmap["NOMOR DOKUMEN"]).value
        for r in range(2, wb["DOKUMEN"].max_row + 1)
        if wb["DOKUMEN"].cell(r, dmap["KODE DOKUMEN"]).value == "315"
    ]
    assert kontrak_rows == ["002/KPJB-THI/SC006/X/2026"]


def test_render_bc27_multi_merge_same_customer():
    dh1 = SimpleNamespace(
        id=4407,
        delivery_number="THI-M261001012",
        customer_id=9,
        delivery_date=date(2026, 10, 2),
        driver_name="Budi",
        vehicle_number="B 1",
        packing_id=None,
    )
    dh2 = SimpleNamespace(
        id=4408,
        delivery_number="THI-M261001013",
        customer_id=9,
        delivery_date=date(2026, 10, 3),
        driver_name="Budi",
        vehicle_number="B 1",
        packing_id=None,
    )
    cust = SimpleNamespace(
        id=9,
        code="SC006",
        name="SELALU CINTA INDONESIA",
        address="JL. LINGKAR",
        billing_address=None,
        tax_id="0729",
        skep_kb_number="193/WBC.10/2020",
        customs_office_code="060800",
        default_currency="USD",
        reference_quotation=None,
    )
    pod = SimpleNamespace(
        id=1,
        master_item_id=2,
        material_name="LEATHER",
        thickness_text="1.2MM",
        color="BLACK",
        emboss="",
        unit_price=Decimal("2"),
        system_reference=None,
    )
    master = SimpleNamespace(id=2, item_code="FG1", exim_material_code=None)
    dd1 = SimpleNamespace(
        id=1,
        po_detail_id=1,
        production_detail_id=None,
        packing_detail_id=None,
        qty=Decimal("10"),
        qty_roll=2,
        qty_pcs=0,
        weight=Decimal("1"),
    )
    dd2 = SimpleNamespace(
        id=2,
        po_detail_id=1,
        production_detail_id=None,
        packing_detail_id=None,
        qty=Decimal("20"),
        qty_roll=3,
        qty_pcs=0,
        weight=Decimal("2"),
    )
    invs = [
        SimpleNamespace(id=1, invoice_number="INV-1", invoice_date=date(2026, 10, 2), delivery_id=4407),
        SimpleNamespace(id=2, invoice_number="INV-2", invoice_date=date(2026, 10, 3), delivery_id=4408),
    ]
    rate = SimpleNamespace(idr_per_usd=Decimal("10000"), effective_date=date(2026, 10, 1))

    db = _make_db(
        headers=[dh1, dh2],
        customer=cust,
        details_by_delivery={4407: [dd1], 4408: [dd2]},
        invoices=invs,
        rates=[rate],
        pods={1: pod},
        masters={2: master},
    )

    data, name = render_customs_bc27_xlsx(db, [4408, 4407], aju_seq=AJU_SEQ)
    assert name.startswith("000027TOP41720261002004407")
    wb = load_workbook(BytesIO(data))

    header = {
        wb["HEADER"].cell(1, c).value: wb["HEADER"].cell(2, c).value
        for c in range(1, wb["HEADER"].max_column + 1)
    }
    assert header["CIF"] == 60.0
    assert header["BRUTO"] == 3.0
    assert header["NETTO"] == 3.0
    assert header["NOMOR AJU"] == "000027TOP41720261002004407"

    bmap = _col_map(wb["BARANG"])
    assert wb["BARANG"].max_row == 3
    assert wb["BARANG"].cell(2, bmap["SERI BARANG"]).value == 1
    assert wb["BARANG"].cell(3, bmap["SERI BARANG"]).value == 2

    dmap = _col_map(wb["DOKUMEN"])
    codes = [wb["DOKUMEN"].cell(r, dmap["KODE DOKUMEN"]).value for r in range(2, wb["DOKUMEN"].max_row + 1)]
    assert codes.count("380") == 2
    assert codes.count("630") == 2
    assert codes.count("315") == 1

    kmap = _col_map(wb["KEMASAN"])
    assert wb["KEMASAN"].cell(2, kmap["JUMLAH KEMASAN"]).value == 5


def test_render_bc27_rejects_different_customers():
    dh1 = SimpleNamespace(
        id=1,
        delivery_number="A",
        customer_id=9,
        delivery_date=date(2026, 10, 1),
        driver_name=None,
        vehicle_number=None,
        packing_id=None,
    )
    dh2 = SimpleNamespace(
        id=2,
        delivery_number="B",
        customer_id=10,
        delivery_date=date(2026, 10, 1),
        driver_name=None,
        vehicle_number=None,
        packing_id=None,
    )
    cust = SimpleNamespace(
        id=9,
        code="A001",
        name="A",
        address="",
        billing_address=None,
        tax_id=None,
        skep_kb_number=None,
        customs_office_code=None,
        default_currency="USD",
        reference_quotation=None,
    )
    db = _make_db(
        headers=[dh1, dh2],
        customer=cust,
        details_by_delivery={1: [], 2: []},
        invoices=[],
        rates=[],
        pods={},
        masters={},
    )
    with pytest.raises(ValueError, match="delivery_customer_mismatch"):
        render_customs_bc27_xlsx(db, [1, 2], aju_seq=AJU_SEQ)


def test_render_bc27_rejects_invalid_aju_seq():
    dh = SimpleNamespace(
        id=1,
        delivery_number="A",
        customer_id=9,
        delivery_date=date(2026, 10, 1),
        driver_name=None,
        vehicle_number=None,
        packing_id=None,
    )
    cust = SimpleNamespace(
        id=9,
        code="SC006",
        name="A",
        address="",
        billing_address=None,
        tax_id=None,
        skep_kb_number=None,
        customs_office_code=None,
        default_currency="USD",
        reference_quotation=None,
    )
    db = _make_db(
        headers=[dh],
        customer=cust,
        details_by_delivery={1: []},
        invoices=[],
        rates=[],
        pods={},
        masters={},
    )
    with pytest.raises(ValueError, match="aju_seq_invalid"):
        render_customs_bc27_xlsx(db, [1], aju_seq="44")


def test_resolve_bruto_from_mrp_packing_ratio(monkeypatch):
    dh = SimpleNamespace(id=10, packing_id=55, customer_id=1)
    pk = SimpleNamespace(id=55, mrp_packing_id=99)
    db = MagicMock()
    db.get.side_effect = lambda model, pk_id: pk if model is PackingHeader and int(pk_id) == 55 else None

    monkeypatch.setattr(
        "app.services.delivery_customs_bc27_xlsx.fetch_mrp_packing_weights",
        lambda mid: (Decimal("1157.9"), Decimal("1136.3")),
    )
    monkeypatch.setattr(
        "app.services.delivery_customs_bc27_xlsx._invoice_weight_totals",
        lambda *_a, **_k: (None, None),
    )

    lines = [{"nw": Decimal("500"), "gw": Decimal("500"), "qty_roll": 1, "amount": Decimal("1")}]
    nw, gw = resolve_bruto_netto(db, [dh], lines)
    assert nw == Decimal("500")
    # 500 * (1157.9 / 1136.3) ≈ 509.50
    assert gw == Decimal("509.50")


def test_render_bc27_uses_mrp_bruto_ratio(monkeypatch):
    dh = SimpleNamespace(
        id=4407,
        delivery_number="THI-M261001012",
        customer_id=9,
        delivery_date=date(2026, 10, 2),
        driver_name="Budi",
        vehicle_number="B 1",
        packing_id=55,
    )
    pk = SimpleNamespace(id=55, mrp_packing_id=99)
    cust = SimpleNamespace(
        id=9,
        code="SC006",
        name="CUST",
        address="ADDR",
        billing_address=None,
        tax_id="1",
        skep_kb_number="S",
        customs_office_code="060800",
        default_currency="USD",
        reference_quotation=None,
    )
    pod = SimpleNamespace(
        id=1,
        master_item_id=2,
        material_name="LEATHER",
        thickness_text="1.2MM",
        color="BLACK",
        emboss="",
        unit_price=Decimal("1"),
        system_reference=None,
    )
    master = SimpleNamespace(id=2, item_code="FG1", exim_material_code=None)
    dd = SimpleNamespace(
        id=1,
        po_detail_id=1,
        production_detail_id=None,
        packing_detail_id=None,
        qty=Decimal("10"),
        qty_roll=1,
        qty_pcs=0,
        weight=Decimal("100"),
    )
    rate = SimpleNamespace(idr_per_usd=Decimal("10000"), effective_date=date(2026, 10, 1))

    db = _make_db(
        headers=[dh],
        customer=cust,
        details_by_delivery={4407: [dd]},
        invoices=[],
        rates=[rate],
        pods={1: pod},
        masters={2: master},
        packings={55: pk},
    )
    monkeypatch.setattr(
        "app.services.delivery_customs_bc27_xlsx.fetch_mrp_packing_weights",
        lambda mid: (Decimal("110"), Decimal("100")),
    )

    data, _name = render_customs_bc27_xlsx(db, [4407], aju_seq=AJU_SEQ)
    wb = load_workbook(BytesIO(data))
    header = {
        wb["HEADER"].cell(1, c).value: wb["HEADER"].cell(2, c).value
        for c in range(1, wb["HEADER"].max_column + 1)
    }
    assert header["NETTO"] == 100.0
    assert header["BRUTO"] == 110.0

    bmap = _col_map(wb["BARANG"])
    assert wb["BARANG"].cell(2, bmap["NETTO"]).value == 100.0
    assert wb["BARANG"].cell(2, bmap["BRUTO"]).value == 110.0
