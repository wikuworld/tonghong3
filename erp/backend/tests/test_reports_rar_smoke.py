"""HTTP smoke tests for accounts receivable (RAR) report API."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.services.accounts_receivable_report import RAR_HEADERS
from tests.http_fixtures import (
    clear_api_overrides,
    clear_auth_user,
    make_api_client,
    set_auth_user,
)

_PREFIX = "/api/v1"


def _sample_invoice(**kwargs):
    defaults = {
        "id": 1,
        "tax_invoice_no": "07002600001575917",
        "customer_name": "PT TEST",
        "tax_invoice_date": date(2026, 1, 5),
        "invoice_date": date(2026, 1, 5),
        "receipt_no": "002/01/2025",
        "received_by": "SETO",
        "received_date": date(2026, 1, 5),
        "term_of_payment": "30",
        "due_date": date(2026, 2, 4),
        "total_amount": Decimal("6871.94"),
        "invoice_number": "THIN-M251229010",
        "po_number": "PO-TEST",
        "payment_amount": Decimal("6871.94"),
        "payment_date": date(2026, 3, 2),
        "payment_remark": None,
        "payment_status": "paid",
        "nopen": "000.192",
        "nomor_aju": "00002721147220260105000013",
        "credit_note_no": None,
        "credit_note_date": None,
        "credit_note_amount": None,
        "oia_export_usd": None,
        "delivery_id": 1,
        "customer_id": 1,
        "po_id": 1,
        "packing_list_number": "PL-1",
        "customer_address": "",
        "customer_phone": "",
        "description": "",
        "bank_name": "",
        "account_number": "",
        "bank_address": "",
        "bank_phone": "",
        "total_in_words": "",
        "payment_bank_account": "",
        "payment_proof_url": "",
        "created_at": date(2026, 1, 5),
    }
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _rar_db(*invoices):
    db = MagicMock()
    db.scalars.return_value.all.return_value = list(invoices)
    db.scalar.return_value = Decimal("6087.6")
    return db


def _xlsx_headers(content: bytes) -> list[str]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    return [str(c.value or "") for c in ws[3]]


@pytest.fixture(autouse=True)
def _cleanup_overrides():
    yield
    clear_api_overrides()


def test_rar_json_requires_auth():
    client = make_api_client(_rar_db(_sample_invoice()))
    clear_auth_user()
    resp = client.get(f"{_PREFIX}/reports/accounts-receivable")
    assert resp.status_code == 401


def test_rar_json_ok_authenticated():
    client = make_api_client(_rar_db(_sample_invoice()))
    set_auth_user("marketing")
    resp = client.get(f"{_PREFIX}/reports/accounts-receivable", params={"year": 2026, "limit": 10})
    assert resp.status_code == 200
    body = resp.json()
    assert body["report_title"] == "RECAPITULATION ACCOUNT RECEIVABLE 2026"
    assert body["year"] == 2026
    assert body["summary"]["count"] == 1
    assert body["summary"]["finish_payment_count"] == 1
    assert len(body["rows"]) == 1
    assert body["rows"][0]["customer_name"] == "PT TEST"
    assert "rar_rows" not in body


def test_rar_json_filter_unpaid():
    paid = _sample_invoice()
    unpaid = _sample_invoice(
        id=2,
        tax_invoice_no="07002600001575918",
        payment_amount=None,
        payment_date=None,
        payment_status="unpaid",
    )
    client = make_api_client(_rar_db(paid, unpaid))
    set_auth_user("finance")
    resp = client.get(
        f"{_PREFIX}/reports/accounts-receivable",
        params={"payment_status": "unpaid", "limit": 50},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["summary"]["count"] == 1
    assert body["summary"]["not_yet_count"] == 1
    assert body["rows"][0]["tax_invoice_no"] == "07002600001575918"


def test_rar_xlsx_forbidden_marketing():
    client = make_api_client(_rar_db(_sample_invoice()))
    set_auth_user("marketing")
    resp = client.get(f"{_PREFIX}/reports/accounts-receivable.xlsx", params={"year": 2026})
    assert resp.status_code == 403


def test_rar_xlsx_ok_finance():
    client = make_api_client(_rar_db(_sample_invoice()))
    set_auth_user("finance")
    resp = client.get(f"{_PREFIX}/reports/accounts-receivable.xlsx", params={"year": 2026})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert len(resp.content) > 500
    headers = _xlsx_headers(resp.content)
    assert headers == RAR_HEADERS
    assert len(headers) == 24


def test_rar_xlsx_ok_admin():
    client = make_api_client(_rar_db(_sample_invoice()))
    set_auth_user("admin")
    resp = client.get(f"{_PREFIX}/reports/accounts-receivable.xlsx")
    assert resp.status_code == 200
    assert _xlsx_headers(resp.content) == RAR_HEADERS
