"""Batal SO: MRP dicek terpisah dari timeout, push sebelum commit, ulang aman."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.routers.marketing import (
    _DELIVERY_LIST_XLS_COLUMNS,
    _delivery_detail_to_row,
    _delivery_list_rows_to_xlsx_bytes,
)
from app.schemas import MarketingDeliveryLineOut
from app.services.business_rules import validate_po_status_transition
from app.services.mrp_po_alloc import MrpAllocUnreachable, require_mrp_alloc_snapshot
from app.services.mrp_po_sync import map_po_detail
from app.services.po_cancel import cancel_sales_order


def _po(**kwargs):
    data = {
        "id": 1,
        "status": "draft",
        "po_number": "SO-1",
        "cancel_reason": None,
        "cancelled_at": None,
        "cancelled_by": None,
    }
    data.update(kwargs)
    return SimpleNamespace(**data)


def _db(po):
    db = MagicMock()
    db.get.return_value = po
    return db


def _guards(**kwargs):
    defaults = {
        "_active_delivery_id": None,
        "_open_invoice_id": None,
        "_production_detail_id": None,
        "_ppic_in_progress_id": None,
        "_detail_ids": [11],
        "_pending_ppic": [],
    }
    defaults.update(kwargs)
    return defaults


def _run(db, po_status_reason="salah input", **guards):
    values = _guards(**guards)
    pending = values.pop("_pending_ppic")
    with (
        patch("app.services.po_cancel._active_delivery_id", return_value=values["_active_delivery_id"]),
        patch("app.services.po_cancel._open_invoice_id", return_value=values["_open_invoice_id"]),
        patch("app.services.po_cancel._production_detail_id", return_value=values["_production_detail_id"]),
        patch("app.services.po_cancel._ppic_in_progress_id", return_value=values["_ppic_in_progress_id"]),
        patch("app.services.po_cancel._detail_ids", return_value=values["_detail_ids"]),
        patch("app.services.po_cancel._pending_ppic", return_value=pending),
        patch("app.services.po_cancel.log_change") as log_change,
        patch("app.services.po_cancel.create_notification_for_roles") as notify,
        patch("app.services.po_cancel.require_mrp_alloc_snapshot") as snapshot,
        patch("app.services.po_cancel.sync_po_details_to_mrp") as sync,
    ):
        if "snapshot" in guards:
            snapshot.side_effect = guards["snapshot"]
        elif "snapshot_return" in guards:
            snapshot.return_value = guards["snapshot_return"]
        if guards.get("sync_error"):
            sync.side_effect = RuntimeError("mrp down")
        row = cancel_sales_order(db, 1, reason=po_status_reason, actor="mkt")
    return row, snapshot, sync, notify, log_change


def test_draft_cancel_skips_mrp_and_closes_pending_ppic():
    po = _po()
    pending = SimpleNamespace(id=7, status="pending")
    db = _db(po)
    row, snapshot, sync, notify, _log = _run(db, _pending_ppic=[pending])
    assert row.status == "cancelled"
    assert row.cancel_reason == "salah input"
    assert row.cancelled_by == "mkt"
    assert pending.status == "cancelled"
    snapshot.assert_not_called()
    sync.assert_not_called()
    db.commit.assert_called_once()
    notify.assert_called_once()
    assert notify.call_args.kwargs["route_query"] == {
        "material_status": "cancelled",
        "po_q": "SO-1",
    }


def test_confirmed_without_alloc_pushes_before_commit():
    po = _po(status="confirmed")
    db = _db(po)
    order = []
    with (
        patch("app.services.po_cancel._active_delivery_id", return_value=None),
        patch("app.services.po_cancel._open_invoice_id", return_value=None),
        patch("app.services.po_cancel._production_detail_id", return_value=None),
        patch("app.services.po_cancel._ppic_in_progress_id", return_value=None),
        patch("app.services.po_cancel._detail_ids", return_value=[11]),
        patch("app.services.po_cancel._pending_ppic", return_value=[]),
        patch("app.services.po_cancel.log_change"),
        patch("app.services.po_cancel.create_notification_for_roles"),
        patch(
            "app.services.po_cancel.require_mrp_alloc_snapshot",
            return_value={
                11: {"mrp_alloc_status": "outstanding", "mrp_qty_processed": "0", "mrp_qty_wip": "0"}
            },
        ),
        patch("app.services.po_cancel.sync_po_details_to_mrp", side_effect=lambda *a, **k: order.append(po.status)),
    ):
        db.commit.side_effect = lambda: order.append("commit")
        cancel_sales_order(db, 1, reason="batal", actor="mkt")
    assert order == ["cancelled", "commit"]


def test_retry_skips_push_when_mrp_already_cancelled():
    po = _po(status="confirmed")
    db = _db(po)
    _row, _snapshot, sync, _notify, _log = _run(
        db,
        snapshot_return={
            11: {"mrp_alloc_status": "cancelled", "mrp_qty_processed": "0", "mrp_qty_wip": "0"}
        },
    )
    sync.assert_not_called()
    db.commit.assert_called_once()
    assert po.status == "cancelled"


def test_missing_mrp_row_still_pushes():
    po = _po(status="production")
    db = _db(po)
    _row, _snapshot, sync, _notify, _log = _run(db, snapshot_return={})
    sync.assert_called_once()
    assert sync.call_args.kwargs["raise_on_error"] is True


def test_qty_or_wip_blocks_cancel():
    po = _po(status="confirmed")
    db = _db(po)
    with pytest.raises(HTTPException) as exc:
        _run(
            db,
            snapshot_return={
                11: {"mrp_alloc_status": "partial", "mrp_qty_processed": "1", "mrp_qty_wip": "0"}
            },
        )
    assert exc.value.status_code == 400
    db.commit.assert_not_called()
    assert po.status == "confirmed"


def test_unreachable_mrp_blocks_cancel():
    po = _po(status="completed")
    db = _db(po)
    with pytest.raises(HTTPException) as exc:
        _run(db, snapshot=MrpAllocUnreachable("timeout"))
    assert "tidak terjangkau" in exc.value.detail
    db.commit.assert_not_called()
    assert po.status == "completed"


def test_active_delivery_invoice_production_and_ppic_block():
    for key, detail_bit in (
        ("_active_delivery_id", "surat jalan"),
        ("_open_invoice_id", "invoice"),
        ("_production_detail_id", "produksi"),
        ("_ppic_in_progress_id", "PPIC"),
    ):
        po = _po(status="confirmed")
        db = _db(po)
        with pytest.raises(HTTPException) as exc:
            _run(db, **{key: 9})
        assert exc.value.status_code == 400
        assert detail_bit in exc.value.detail
        db.commit.assert_not_called()


def test_already_cancelled_is_idempotent():
    po = _po(status="cancelled", cancel_reason="lama")
    db = _db(po)
    row, snapshot, sync, notify, _log = _run(db)
    assert row is po
    assert po.cancel_reason == "lama"
    snapshot.assert_not_called()
    sync.assert_not_called()
    notify.assert_not_called()
    db.commit.assert_not_called()


def test_empty_reason_rejected():
    with pytest.raises(HTTPException) as exc:
        cancel_sales_order(_db(_po()), 1, reason="   ", actor="mkt")
    assert exc.value.status_code == 400


def test_push_failure_rolls_back_erp():
    po = _po(status="confirmed")
    db = _db(po)
    with pytest.raises(HTTPException):
        _run(
            db,
            snapshot_return={
                11: {"mrp_alloc_status": "outstanding", "mrp_qty_processed": "0", "mrp_qty_wip": "0"}
            },
            sync_error=True,
        )
    db.rollback.assert_called_once()
    db.commit.assert_not_called()


def test_put_cannot_enter_or_leave_cancelled():
    with pytest.raises(HTTPException):
        validate_po_status_transition("confirmed", "cancelled")
    with pytest.raises(HTTPException):
        validate_po_status_transition("cancelled", "confirmed")
    validate_po_status_transition("cancelled", "cancelled")


def test_mkt_advance_rejects_cancelled():
    from app.services.mkt_workflow import validate_mkt_advance

    db = MagicMock()
    db.get.return_value = SimpleNamespace(status="cancelled", mkt_workflow_stage="input_system")
    with pytest.raises(HTTPException) as exc:
        validate_mkt_advance(db, 1, "rekap_printed")
    assert "dibatalkan" in exc.value.detail


def test_map_po_detail_marks_cancelled_without_price():
    detail = SimpleNamespace(
        id=11,
        po=SimpleNamespace(
            po_number="SO-1",
            po_customer_ref="PO-1",
            status="cancelled",
            customer_name_snapshot="Buyer",
            customer_code="B",
            customer_id=1,
        ),
        material_name="PU",
        style_article="",
        color_swatch_recipe=None,
        quantity_order=Decimal("10"),
        balance_qty=Decimal("10"),
        thi_line_ref="LOT",
        system_reference="",
        color="BLACK",
        emboss="",
        thickness_text="1.0",
        thickness_min=None,
        thickness_max=None,
        unit_price=Decimal("2.50"),
    )
    payload = map_po_detail(detail)
    assert payload["cancelled"] is True
    assert "unit_price" not in payload


def test_require_snapshot_empty_list_is_success_and_timeout_raises():
    response = MagicMock()
    response.status_code = 200
    response.content = b"[]"
    response.json.return_value = []
    client = MagicMock()
    client.get.return_value = response
    client.__enter__.return_value = client
    with patch("app.services.mrp_po_alloc.httpx.Client", return_value=client):
        assert require_mrp_alloc_snapshot([11]) == {}

    broken = MagicMock()
    broken.__enter__.side_effect = TimeoutError("timeout")
    with patch("app.services.mrp_po_alloc.httpx.Client", return_value=broken):
        with pytest.raises(MrpAllocUnreachable):
            require_mrp_alloc_snapshot([11])


def test_delivery_list_appends_foc_and_category_columns():
    labels = [label for _, label in _DELIVERY_LIST_XLS_COLUMNS]
    assert labels[-2:] == ["FOC/BERBAYAR", "KATEGORI"]
    assert "ORDER TYPE" in labels
    assert labels.index("ORDER TYPE") < labels.index("FOC/BERBAYAR")
    raw = _delivery_list_rows_to_xlsx_bytes(
        [{"order_type": "PROD", "charge_type": "FOC", "po_category_label": "SAMPLE"}]
    )
    sheet = load_workbook(BytesIO(raw)).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers[-2:] == ["FOC/BERBAYAR", "KATEGORI"]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert values[headers.index("FOC/BERBAYAR")] == "FOC"
    assert values[headers.index("KATEGORI")] == "SAMPLE"
    assert values[headers.index("ORDER TYPE")] == "PROD"


def test_delivery_row_charge_and_category_do_not_change_amount():
    dd = SimpleNamespace(id=3, qty=Decimal("2"), qty_roll=1, qty_pcs=2, weight=None, returned_qty=None)
    dh = SimpleNamespace(
        id=8,
        delivery_number="SJ-1",
        delivery_date=date(2026, 4, 1),
        destination="",
        status="draft",
        is_deleted=False,
        packing_id=None,
        delivery_zone="local",
        is_foc=True,
        mrp_fg_id=None,
    )
    pod = SimpleNamespace(
        system_reference="",
        thi_line_ref="LOT1",
        quantity_order=Decimal("2"),
        unit_price=Decimal("3.50"),
        color="",
        vir_softness_type="",
        material_name="PU",
        color_standard="",
        color_reference="",
        emboss="",
        style_article="",
        model="",
        detail_remark="",
        thickness_text="",
        thickness_min=None,
        thickness_max=None,
    )
    ph = SimpleNamespace(
        po_number="SO-1",
        customer_name_snapshot="Pabrik",
        customer_code="P",
        order_date=date(2026, 4, 1),
        po_customer_ref="REF",
        brand_name_snapshot="",
        buyer="",
        order_type_label="",
        po_category="sample",
        factory_code="",
    )
    row = _delivery_detail_to_row(dd, dh, None, pod, ph, None, None, None, None)
    assert row["charge_type"] == "FOC"
    assert row["is_foc"] is True
    assert MarketingDeliveryLineOut.model_validate(row).is_foc is True
    assert row["po_category_label"] == "SAMPLE"
    assert row["order_type"] == "SAMPLE"
    assert row["unit_price"] == "3.50"
    assert row["amount"] == "7.00"
