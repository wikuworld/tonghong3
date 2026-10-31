"""Export/import Excel master CSI Color."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import CsiColorMaster
from app.services.csi_color_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_csi_colors_template_xlsx,
    export_csi_colors_xlsx,
    import_csi_colors_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    for name in (
        "csi_color_masters",
        "buyer_price_lists",
        "brand_buyer_maps",
        "buyer_codes",
    ):
        if name in text:
            return name
    return ""


def _make_db(rows: list[CsiColorMaster] | None = None):
    rows = list(rows or [])
    db = MagicMock()

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        if table == "csi_color_masters":
            result.all.return_value = list(rows)
        else:
            result.all.return_value = []
        result.first.return_value = None
        return result

    db.scalars.side_effect = scalars
    db.add.side_effect = rows.append
    return db, rows


def _row(**kwargs) -> CsiColorMaster:
    data = {
        "buyer_code": "NIKE",
        "color_code": "00A",
        "color_name": "BLACK",
        "is_active": True,
    }
    data.update(kwargs)
    return CsiColorMaster(**data)


def test_export_headers_and_values():
    row = _row(id=7, source_file="csi.xlsx")
    db, _ = _make_db([row])
    wb = load_workbook(BytesIO(export_csi_colors_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 7
    assert wb.active.cell(2, 2).value == "NIKE"
    assert wb.active.cell(2, 3).value == "00A"
    assert wb.active.cell(2, 5).value == "csi.xlsx"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_csi_colors_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert wb.active.cell(2, 1).value is None


def test_import_creates_and_updates_by_id():
    existing = _row(id=3, color_name="BLACK")
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "buyer_code", "color_code", "color_name"],
        [
            {"id": 3, "buyer_code": "nike", "color_code": "00A", "color_name": "BLACK REV"},
            {"id": "", "buyer_code": "ADIDAS", "color_code": "A0QM", "color_name": "CORE BLACK"},
        ],
    )
    result = import_csi_colors_xlsx(db, content)
    assert result.updated == 1
    assert result.created == 1
    assert result.errors == []
    assert existing.color_name == "BLACK REV"
    assert existing.buyer_code == "NIKE"
    assert rows[-1].color_code == "A0QM"


def test_import_skips_unchanged_and_unknown_buyer():
    existing = _row(id=1)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "buyer", "color_code", "color_name"],
        [
            {"id": 1, "buyer": "NIKE", "color_code": "00A", "color_name": "BLACK"},
            {"id": "", "buyer": "ZZZ", "color_code": "X", "color_name": "UNKNOWN"},
        ],
    )
    result = import_csi_colors_xlsx(db, content)
    assert result.updated == 0
    assert result.created == 0
    assert result.skipped == 2
    assert len(rows) == 1
    assert any("ZZZ" in err.reason for err in result.errors)


def test_reimport_same_key_does_not_duplicate():
    existing = _row(id=4)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["buyer_code", "color_code", "color_name"],
        [
            {"buyer_code": "NIKE", "color_code": "00A", "color_name": "BLACK REV"},
            {"buyer_code": "nike", "color_code": "00A", "color_name": "BLACK REV"},
        ],
    )
    first = import_csi_colors_xlsx(db, content)
    second = import_csi_colors_xlsx(db, content)
    assert first.created == 0
    assert first.updated == 1
    assert first.skipped == 1
    assert second.updated == 0
    assert len(rows) == 1
    assert existing.color_name == "BLACK REV"


def test_import_rejects_missing_header():
    db, _ = _make_db()
    content = _xlsx(["color_code"], [{"color_code": "00A"}])
    result = import_csi_colors_xlsx(db, content)
    assert result.created == 0
    assert result.errors[0].row == 1
