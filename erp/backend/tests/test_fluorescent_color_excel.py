"""Export/import Excel master Fluorescent Color."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import FluorescentColorMaster
from app.services.fluorescent_color_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_fluorescent_colors_template_xlsx,
    export_fluorescent_colors_xlsx,
    import_fluorescent_colors_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    for name in (
        "fluorescent_color_masters",
        "buyer_price_lists",
        "brand_buyer_maps",
        "buyer_codes",
        "brands",
    ):
        if name in text:
            return name
    return ""


def _make_db(rows: list[FluorescentColorMaster] | None = None):
    rows = list(rows or [])
    db = MagicMock()

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        if table == "fluorescent_color_masters":
            result.all.return_value = list(rows)
            result.first.return_value = rows[0] if rows else None
        else:
            result.all.return_value = []
            result.first.return_value = None
        return result

    db.scalars.side_effect = scalars

    def add(obj):
        rows.append(obj)

    db.add.side_effect = add
    return db, rows


def _row(**kwargs) -> FluorescentColorMaster:
    data = {
        "buyer_code": "NB",
        "season_label": "SS26",
        "season_year": 2026,
        "season_type": "SS",
        "color_code": "FL-01",
        "color_name": "VOLT",
        "color_class": "fluorescent",
        "is_active": True,
    }
    data.update(kwargs)
    return FluorescentColorMaster(**data)


def test_export_headers_and_values():
    row = _row(id=7, source_file="nb.csv")
    db, _ = _make_db([row])
    wb = load_workbook(BytesIO(export_fluorescent_colors_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 7
    assert wb.active.cell(2, 2).value == "NB"
    assert wb.active.cell(2, 7).value == "VOLT"
    assert wb.active.cell(2, 10).value == "1"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_fluorescent_colors_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert wb.active.cell(2, 1).value is None


def test_import_creates_and_updates_by_id():
    existing = _row(id=3, color_class="fluorescent")
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "buyer_code", "color_name", "color_code", "season_label", "season_year", "season_type", "color_class"],
        [
            {
                "id": 3,
                "buyer_code": "nb",
                "color_name": "VOLT",
                "color_code": "FL-01",
                "season_label": "SS26",
                "season_year": 2026,
                "season_type": "SS",
                "color_class": "metallic",
            },
            {
                "id": "",
                "buyer_code": "ADIDAS",
                "color_name": "SOLAR RED",
                "color_code": "M-02",
                "season_label": "FW25",
                "season_year": 2025,
                "season_type": "FW",
                "color_class": "metallic",
            },
        ],
    )
    result = import_fluorescent_colors_xlsx(db, content)
    assert result.updated == 1
    assert result.created == 1
    assert result.errors == []
    assert existing.color_class == "metallic"
    assert existing.buyer_code == "NB"
    assert existing.updated_at is not None
    assert rows[-1].color_name == "SOLAR RED"
    assert rows[-1].buyer_code == "ADIDAS"


def test_import_skips_unchanged_and_unknown_buyer():
    existing = _row(id=1)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "buyer", "color_name", "color_code", "season_label", "color_class"],
        [
            {
                "id": 1,
                "buyer": "NB",
                "color_name": "VOLT",
                "color_code": "FL-01",
                "season_label": "SS26",
                "color_class": "fluorescent",
            },
            {
                "id": "",
                "buyer": "ZZZ",
                "color_name": "UNKNOWN",
                "color_code": "X",
                "season_label": "SS26",
                "color_class": "fluorescent",
            },
        ],
    )
    result = import_fluorescent_colors_xlsx(db, content)
    assert result.updated == 0
    assert result.created == 0
    assert result.skipped == 2
    assert len(rows) == 1
    assert any("ZZZ" in err.reason for err in result.errors)


def test_reimport_same_key_does_not_duplicate():
    existing = _row(id=4, color_class="fluorescent")
    db, rows = _make_db([existing])
    content = _xlsx(
        ["buyer_code", "color_name", "color_code", "season_label", "color_class"],
        [
            {
                "buyer_code": "NB",
                "color_name": "VOLT",
                "color_code": "FL-01",
                "season_label": "SS26",
                "color_class": "metallic",
            },
            {
                "buyer_code": "nb",
                "color_name": "VOLT",
                "color_code": "FL-01",
                "season_label": "SS26",
                "color_class": "metallic",
            },
        ],
    )
    first = import_fluorescent_colors_xlsx(db, content)
    second = import_fluorescent_colors_xlsx(db, content)
    assert first.created == 0
    assert first.updated == 1
    assert first.skipped == 1
    assert second.created == 0
    assert second.updated == 0
    assert len(rows) == 1
    assert existing.color_class == "metallic"


def test_duplicate_rows_in_one_file_insert_once():
    db, rows = _make_db([])
    content = _xlsx(
        ["buyer_code", "color_name", "color_code", "season_label", "color_class"],
        [
            {
                "buyer_code": "NB",
                "color_name": "VOLT",
                "color_code": "",
                "season_label": "SS26",
                "color_class": "fluorescent",
            },
            {
                "buyer_code": "NB",
                "color_name": "VOLT",
                "color_code": "",
                "season_label": "SS26",
                "color_class": "metallic",
            },
        ],
    )
    result = import_fluorescent_colors_xlsx(db, content)
    assert result.created == 1
    assert result.updated == 1
    assert len(rows) == 1
    assert rows[0].color_class == "metallic"


def test_import_rejects_missing_header():
    db, _ = _make_db()
    content = _xlsx(["color_code"], [{"color_code": "FL-01"}])
    result = import_fluorescent_colors_xlsx(db, content)
    assert result.created == 0
    assert result.errors[0].row == 1


def test_invalid_color_class_skips_row():
    db, rows = _make_db([])
    content = _xlsx(
        ["buyer_code", "color_name", "color_class"],
        [{"buyer_code": "NB", "color_name": "VOLT", "color_class": "pearl"}],
    )
    result = import_fluorescent_colors_xlsx(db, content)
    assert result.created == 0
    assert result.skipped == 1
    assert len(rows) == 0
    assert any("color_class" in err.reason for err in result.errors)
