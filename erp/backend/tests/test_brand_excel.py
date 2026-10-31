"""Export/import Excel master brand."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import Brand, BrandBuyerMap, Customer, CustomerBrandMapping
from app.services.brand_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_brands_template_xlsx,
    export_brands_xlsx,
    import_brands_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    for name in (
        "customer_brand_mappings",
        "brand_buyer_maps",
        "buyer_codes",
        "buyer_price_lists",
        "customers",
        "brands",
    ):
        if name in text:
            return name
    return ""


def _make_db(
    brands: list[Brand] | None = None,
    customers: list[Customer] | None = None,
    mappings: list[CustomerBrandMapping] | None = None,
    buyers: list[BrandBuyerMap] | None = None,
):
    brands = list(brands or [])
    customers = list(customers or [])
    mappings = list(mappings or [])
    buyers = list(buyers or [])
    added: list = []
    db = MagicMock()
    seq = {"n": 20}

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        result.first.return_value = None
        if table == "brands":
            result.all.return_value = list(brands)
        elif table == "customers":
            result.all.return_value = list(customers)
        elif table == "customer_brand_mappings":
            result.all.return_value = list(mappings)
        elif table == "brand_buyer_maps":
            result.all.return_value = list(buyers)
        else:
            result.all.return_value = []
        return result

    def flush():
        for obj in added:
            if isinstance(obj, Brand) and obj.id is None:
                seq["n"] += 1
                obj.id = seq["n"]

    db.scalars.side_effect = scalars
    db.add.side_effect = added.append
    db.flush.side_effect = flush
    db.delete.side_effect = lambda obj: None
    return db, added


def test_export_headers_and_relations():
    brand = Brand(id=4, code="NB", name="NB")
    customer = Customer(id=8, code="AD002", name="ADIS")
    db, _ = _make_db(
        [brand],
        [customer],
        [CustomerBrandMapping(id=1, customer_id=8, brand_id=4, is_default=False)],
        [BrandBuyerMap(id=1, brand_id=4, buyer_code="NB")],
    )
    wb = load_workbook(BytesIO(export_brands_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 4
    assert wb.active.cell(2, 2).value == "NB"
    assert wb.active.cell(2, 4).value == "AD002"
    assert wb.active.cell(2, 5).value == "NB"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_brands_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert wb.active.cell(2, 1).value is None


def test_import_creates_brand_customer_and_buyer():
    db, added = _make_db([], [Customer(id=8, code="AD002", name="ADIS")])
    content = _xlsx(
        ["name", "customer_codes", "buyer_codes"],
        [{"name": "Puma", "customer_codes": "ad002", "buyer_codes": "puma"}],
    )
    result = import_brands_xlsx(db, content)
    assert result.created == 1
    assert result.errors == []
    brands = [obj for obj in added if isinstance(obj, Brand)]
    maps = [obj for obj in added if isinstance(obj, CustomerBrandMapping)]
    buyers = [obj for obj in added if isinstance(obj, BrandBuyerMap)]
    assert brands[0].name == "Puma"
    assert brands[0].code == "Puma"
    assert maps[0].customer_id == 8
    assert maps[0].brand_id == brands[0].id
    assert buyers[0].buyer_code == "PUMA"


def test_import_updates_then_skips_same_name():
    existing = Brand(id=4, code="NB", name="NB")
    db, added = _make_db([existing])
    headers = ["id", "name"]
    row = {"id": 4, "name": "New Balance"}
    first = import_brands_xlsx(db, _xlsx(headers, [row]))
    assert first.updated == 1
    assert existing.name == "New Balance"
    assert existing.code == "New Balance"
    assert added == []
    second = import_brands_xlsx(db, _xlsx(headers, [row]))
    assert second.skipped == 1
    assert second.updated == 0


def test_import_rejects_unknown_customer():
    db, added = _make_db()
    content = _xlsx(
        ["name", "customer_codes"],
        [{"name": "NB", "customer_codes": "ZZZ"}],
    )
    result = import_brands_xlsx(db, content)
    assert result.created == 0
    assert result.errors
    assert added == []


def test_import_missing_header():
    db, _ = _make_db()
    result = import_brands_xlsx(db, _xlsx(["buyer_codes"], [{"buyer_codes": "NB"}]))
    assert result.total_rows == 0
    assert result.errors
    assert "name" in result.errors[0].reason
