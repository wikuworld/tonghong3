"""Export/import Excel list harga buyer."""

from datetime import date
from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import Brand, BrandBuyerMap, BuyerPriceList, BuyerPriceSheetRow
from app.services.buyer_price_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_buyer_prices_template_xlsx,
    export_buyer_prices_xlsx,
    import_buyer_prices_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    for name in (
        "buyer_price_sheet_rows",
        "buyer_price_lists",
        "brand_buyer_maps",
        "brands",
    ):
        if name in text:
            return name
    return ""


def _make_db(
    lists: list[BuyerPriceList] | None = None,
    sheets: list[BuyerPriceSheetRow] | None = None,
    brands: list[Brand] | None = None,
    maps: list[BrandBuyerMap] | None = None,
):
    lists = list(lists or [])
    sheets = list(sheets or [])
    brands = list(brands or [])
    maps = list(maps or [])
    added: list = []
    db = MagicMock()
    seq = {"n": 80}

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        if table == "buyer_price_lists":
            result.all.return_value = list(lists)
        elif table == "buyer_price_sheet_rows":
            result.all.return_value = list(sheets)
        elif table == "brands":
            result.all.return_value = list(brands)
        elif table == "brand_buyer_maps":
            result.all.return_value = list(maps)
        else:
            result.all.return_value = []
        return result

    def flush():
        for obj in added:
            if isinstance(obj, BuyerPriceList) and obj.id is None:
                seq["n"] += 1
                obj.id = seq["n"]

    db.scalars.side_effect = scalars
    db.add.side_effect = added.append
    db.flush.side_effect = flush
    return db, added


def _brand() -> Brand:
    return Brand(id=1, code="NB", name="New Balance")


def _map() -> BrandBuyerMap:
    return BrandBuyerMap(id=1, brand_id=1, buyer_code="NB")


def _price_list(**kwargs) -> BuyerPriceList:
    data = {
        "id": 5,
        "buyer_code": "NB",
        "brand_id": 1,
        "currency": "USD",
        "title": "Harga 2026",
        "period_label": "2026",
        "period_start": date(2026, 1, 1),
        "period_end": date(2026, 12, 31),
        "outside_adjust_type": "nominal",
        "discount_modifier_enabled": False,
        "discount_modifier_type": "nominal",
    }
    data.update(kwargs)
    return BuyerPriceList(**data)


def _sheet(**kwargs) -> BuyerPriceSheetRow:
    data = {
        "id": 9,
        "list_id": 5,
        "row_number": 0,
        "cells": {
            "ITEM": "PU",
            "THICKNESS": "1.0-1.2",
            "COLOR": "",
            "COLOR WHITE": 10,
        },
    }
    data.update(kwargs)
    return BuyerPriceSheetRow(**data)


def test_export_headers_and_price_line():
    db, _ = _make_db([_price_list()], [_sheet()], [_brand()], [_map()])
    wb = load_workbook(BytesIO(export_buyer_prices_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 5
    assert wb.active.cell(2, 2).value == "NB"
    assert wb.active.cell(2, 3).value == "NB"
    assert wb.active.cell(2, 14).value == "PU"
    assert wb.active.cell(2, 17).value == 10


def test_export_list_without_sheet_rows():
    db, _ = _make_db([_price_list()], [], [_brand()], [_map()])
    wb = load_workbook(BytesIO(export_buyer_prices_xlsx(db)))
    assert wb.active.cell(2, 1).value == 5
    assert wb.active.cell(2, 14).value is None


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_buyer_prices_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert "list_id" not in headers
    assert wb.active.cell(2, 1).value is None


def test_import_creates_list_and_line():
    db, added = _make_db([], [], [_brand()], [_map()])
    content = _xlsx(
        ["brand_code", "buyer_code", "period_start", "period_end", "item", "color_white"],
        [{"brand_code": "nb", "buyer_code": "nb", "period_start": "2027-01-01", "period_end": "2027-12-31", "item": "PU", "color_white": 8}],
    )
    result = import_buyer_prices_xlsx(db, content)
    assert result.created == 1
    assert result.errors == []
    created_lists = [obj for obj in added if isinstance(obj, BuyerPriceList)]
    created_rows = [obj for obj in added if isinstance(obj, BuyerPriceSheetRow)]
    assert len(created_lists) == 1
    assert created_lists[0].buyer_code == "NB"
    assert created_lists[0].brand_id == 1
    assert created_rows[0].list_id == created_lists[0].id
    assert created_rows[0].cells["ITEM"] == "PU"
    assert created_rows[0].cells["COLOR WHITE"] == 8


def test_import_updates_then_skips_same_line():
    sheet = _sheet()
    db, added = _make_db([_price_list()], [sheet], [_brand()], [_map()])
    headers = ["list_id", "brand_code", "buyer_code", "period_start", "period_end", "item", "thickness", "color_white"]
    row = {
        "list_id": 5,
        "brand_code": "NB",
        "buyer_code": "NB",
        "period_start": "2026-01-01",
        "period_end": "2026-12-31",
        "item": "PU",
        "thickness": "1.0-1.2",
        "color_white": 12,
    }
    first = import_buyer_prices_xlsx(db, _xlsx(headers, [row]))
    assert first.updated == 1
    assert first.created == 0
    assert sheet.cells["COLOR WHITE"] == 12
    assert added == []
    second = import_buyer_prices_xlsx(db, _xlsx(headers, [row]))
    assert second.skipped == 1
    assert second.updated == 0
    assert second.created == 0


def test_import_rejects_unknown_brand_and_overlap():
    db, added = _make_db([_price_list()], [], [_brand()], [_map()])
    content = _xlsx(
        ["brand_code", "buyer_code", "period_start", "period_end", "item", "color_white"],
        [
            {"brand_code": "ZZZ", "buyer_code": "NB", "period_start": "2027-01-01", "period_end": "2027-06-30", "item": "PU", "color_white": 1},
            {"brand_code": "NB", "buyer_code": "NB", "period_start": "2026-06-01", "period_end": "2026-08-01", "item": "PU", "color_white": 1},
        ],
    )
    result = import_buyer_prices_xlsx(db, content)
    assert result.created == 0
    assert len(result.errors) == 2
    assert added == []


def test_import_missing_header():
    db, _ = _make_db()
    result = import_buyer_prices_xlsx(db, _xlsx(["item"], [{"item": "PU"}]))
    assert result.total_rows == 0
    assert result.errors
    assert "brand_code" in result.errors[0].reason
