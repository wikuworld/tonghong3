"""Export/import Excel master Color Window."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import ColorWindowMaster
from app.services.color_window_excel import (
    EXPORT_COLUMNS,
    export_color_windows_template_xlsx,
    export_color_windows_xlsx,
    import_color_windows_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_db(rows: list[ColorWindowMaster] | None = None, brand_codes: list[str] | None = None):
    rows = list(rows or [])
    brand_codes = list(brand_codes or ["NB", "ADIDAS"])
    db = MagicMock()

    def scalars(stmt):
        result = MagicMock()
        text = str(stmt).lower()
        if "color_window" in text:
            result.all.return_value = list(rows)
        else:
            result.all.return_value = list(brand_codes)
        return result

    db.scalars.side_effect = scalars

    def add(obj):
        rows.append(obj)

    db.add.side_effect = add
    return db, rows


def test_export_headers_and_values():
    row = ColorWindowMaster(
        id=7,
        color_code="039-90-05",
        color_name="ANGORA COLORO",
        brand_code="NB",
        remark="C/W",
    )
    db, _ = _make_db([row])
    wb = load_workbook(BytesIO(export_color_windows_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, 8)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 7
    assert wb.active.cell(2, 2).value == "039-90-05"
    assert wb.active.cell(2, 4).value == "NB"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_color_windows_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, 5)]
    assert headers == ["color_code", "color_name", "brand_code", "remark"]
    assert wb.active.cell(2, 1).value is None


def test_import_creates_and_updates_by_id():
    existing = ColorWindowMaster(
        id=3,
        color_code="A0QM",
        color_name="CORE BLACK",
        brand_code="ADIDAS",
        remark=None,
    )
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "color_code", "color_name", "brand_code", "remark"],
        [
            {
                "id": 3,
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "brand_code": "adidas",
                "remark": "revisi",
            },
            {
                "id": "",
                "color_code": "039-90-05",
                "color_name": "ANGORA COLORO",
                "brand_code": "NB",
                "remark": "",
            },
        ],
    )
    result = import_color_windows_xlsx(db, content)
    assert result.updated == 1
    assert result.created == 1
    assert result.errors == []
    assert existing.remark == "revisi"
    assert existing.brand_code == "ADIDAS"
    assert existing.updated_at is not None
    assert rows[-1].color_code == "039-90-05"
    assert rows[-1].brand_code == "NB"


def test_import_skips_unchanged_and_unknown_brand():
    existing = ColorWindowMaster(
        id=1,
        color_code="AA2R",
        color_name="GREY ONE",
        brand_code="ADIDAS",
        remark=None,
    )
    db, rows = _make_db([existing])
    content = _xlsx(
        ["color_code", "color_name", "brand", "remark", "id"],
        [
            {"id": 1, "color_code": "AA2R", "color_name": "GREY ONE", "brand": "ADIDAS", "remark": ""},
            {"id": "", "color_code": "ZZ", "color_name": "UNKNOWN", "brand": "PUMA", "remark": ""},
        ],
    )
    result = import_color_windows_xlsx(db, content)
    assert result.updated == 0
    assert result.created == 0
    assert result.skipped == 2
    assert len(rows) == 1
    assert any("PUMA" in err.reason for err in result.errors)


def test_reimport_same_brand_and_code_does_not_duplicate():
    existing = ColorWindowMaster(
        id=4,
        color_code="039-90-05",
        color_name="ANGORA COLORO",
        brand_code="NB",
        remark=None,
    )
    db, rows = _make_db([existing])
    content = _xlsx(
        ["color_code", "color_name", "brand_code", "remark"],
        [
            {
                "color_code": "039-90-05",
                "color_name": "ANGORA COLORO",
                "brand_code": "NB",
                "remark": "salinan",
            },
            {
                "color_code": "039-90-05",
                "color_name": "ANGORA COLORO",
                "brand_code": "nb",
                "remark": "salinan",
            },
        ],
    )
    first = import_color_windows_xlsx(db, content)
    second = import_color_windows_xlsx(db, content)
    assert first.created == 0
    assert first.updated == 1
    assert first.skipped == 1
    assert second.created == 0
    assert second.updated == 0
    assert len(rows) == 1
    assert existing.remark == "salinan"


def test_duplicate_rows_in_one_file_insert_once():
    db, rows = _make_db([])
    content = _xlsx(
        ["color_code", "color_name", "brand_code", "remark"],
        [
            {"color_code": "A0QM", "color_name": "CORE BLACK", "brand_code": "ADIDAS", "remark": ""},
            {"color_code": "a0qm", "color_name": "CORE BLACK", "brand_code": "ADIDAS", "remark": ""},
        ],
    )
    result = import_color_windows_xlsx(db, content)
    assert result.created == 1
    assert result.updated == 1
    assert len(rows) == 1
    assert rows[0].color_code == "a0qm"


def test_import_rejects_missing_header():
    db, _ = _make_db()
    content = _xlsx(["color_code"], [{"color_code": "A0QM"}])
    result = import_color_windows_xlsx(db, content)
    assert result.created == 0
    assert result.errors[0].row == 1
