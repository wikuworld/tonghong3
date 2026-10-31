"""Export/import Excel master Color Swatch."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import ColorSwatchRecipe, CsiColorMaster, Customer
from app.services.color_swatch_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_color_swatches_template_xlsx,
    export_color_swatches_xlsx,
    import_color_swatches_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    for name in (
        "csi_color_masters",
        "color_swatch_recipes",
        "buyer_price_lists",
        "brand_buyer_maps",
        "buyer_codes",
        "customers",
        "brands",
    ):
        if name in text:
            return name
    return ""


def _make_db(
    rows: list[ColorSwatchRecipe] | None = None,
    customers: list[Customer] | None = None,
    csi_rows: list[CsiColorMaster] | None = None,
):
    rows = list(rows or [])
    customers = list(customers or [])
    csi_rows = list(csi_rows or [])
    db = MagicMock()

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        if table == "color_swatch_recipes":
            result.all.return_value = list(rows)
            result.first.return_value = rows[0] if rows else None
        elif table == "customers":
            result.all.return_value = list(customers)
            result.first.return_value = customers[0] if customers else None
        elif table == "csi_color_masters":
            result.all.return_value = list(csi_rows)
            result.first.return_value = csi_rows[0] if csi_rows else None
        else:
            result.all.return_value = []
            result.first.return_value = None
        return result

    db.scalars.side_effect = scalars

    def add(obj):
        rows.append(obj)

    db.add.side_effect = add
    return db, rows


def _recipe(**kwargs) -> ColorSwatchRecipe:
    data = {
        "factory_key": "ADIS",
        "buyer_code": "ADIDAS",
        "color_code": "A0QM",
        "color_name": "CORE BLACK",
        "emboss": "AR-167",
        "color_standard": "SW1",
        "material_name": "GENERIC PU",
        "is_active": True,
    }
    data.update(kwargs)
    return ColorSwatchRecipe(**data)


def test_export_headers_and_values():
    row = _recipe(id=7, customer_id=5, csi_color_master_id=9, color_code="039-90-05")
    customer = Customer(id=5, code="SL043", name="Suela")
    csi = CsiColorMaster(id=9, buyer_code="ADIDAS", color_code="CSI-9", color_name="CORE BLACK")
    db, _ = _make_db([row], customers=[customer], csi_rows=[csi])
    wb = load_workbook(BytesIO(export_color_swatches_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 7
    assert wb.active.cell(2, 3).value == "SL043"
    assert wb.active.cell(2, 5).value == "039-90-05"
    assert wb.active.cell(2, 13).value == "CSI-9"
    assert wb.active.cell(2, 14).value == "1"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_color_swatches_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert wb.active.cell(2, 1).value is None


def test_import_creates_and_updates_by_id():
    existing = _recipe(id=3, color_name="CORE BLACK", backing=None)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "factory_key", "buyer_code", "color_code", "color_name", "emboss", "color_standard", "material_name", "backing"],
        [
            {
                "id": 3,
                "factory_key": "adis",
                "buyer_code": "adidas",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "AR-167",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
                "backing": "revisi",
            },
            {
                "id": "",
                "factory_key": "ADIS",
                "buyer_code": "NB",
                "color_code": "039-90-05",
                "color_name": "ANGORA COLORO",
                "emboss": "",
                "color_standard": "SW2",
                "material_name": "SUEDE",
                "backing": "",
            },
        ],
    )
    result = import_color_swatches_xlsx(db, content)
    assert result.updated == 1
    assert result.created == 1
    assert result.errors == []
    assert existing.backing == "revisi"
    assert existing.buyer_code == "ADIDAS"
    assert existing.updated_at is not None
    assert rows[-1].color_code == "039-90-05"
    assert rows[-1].buyer_code == "NB"
    assert rows[-1].factory_key == "ADIS"


def test_import_skips_unchanged_and_unknown_buyer():
    existing = _recipe(id=1)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["id", "factory_key", "buyer", "color_code", "color_name", "emboss", "color_standard", "material_name"],
        [
            {
                "id": 1,
                "factory_key": "ADIS",
                "buyer": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "AR-167",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
            },
            {
                "id": "",
                "factory_key": "ADIS",
                "buyer": "ZZZ",
                "color_code": "ZZ",
                "color_name": "UNKNOWN",
                "emboss": "",
                "color_standard": "SW9",
                "material_name": "SUEDE",
            },
        ],
    )
    result = import_color_swatches_xlsx(db, content)
    assert result.updated == 0
    assert result.created == 0
    assert result.skipped == 2
    assert len(rows) == 1
    assert any("ZZZ" in err.reason for err in result.errors)


def test_reimport_same_key_does_not_duplicate():
    existing = _recipe(id=4, backing=None)
    db, rows = _make_db([existing])
    content = _xlsx(
        ["factory_key", "buyer_code", "color_code", "color_name", "emboss", "color_standard", "material_name", "backing"],
        [
            {
                "factory_key": "ADIS",
                "buyer_code": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "AR-167",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
                "backing": "salinan",
            },
            {
                "factory_key": "ADIS",
                "buyer_code": "adidas",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "AR-167",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
                "backing": "salinan",
            },
        ],
    )
    first = import_color_swatches_xlsx(db, content)
    second = import_color_swatches_xlsx(db, content)
    assert first.created == 0
    assert first.updated == 1
    assert first.skipped == 1
    assert second.created == 0
    assert second.updated == 0
    assert len(rows) == 1
    assert existing.backing == "salinan"


def test_duplicate_rows_in_one_file_insert_once():
    db, rows = _make_db([])
    content = _xlsx(
        ["factory_key", "buyer_code", "color_code", "color_name", "emboss", "color_standard", "material_name"],
        [
            {
                "factory_key": "ADIS",
                "buyer_code": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
            },
            {
                "factory_key": "ADIS",
                "buyer_code": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK REV",
                "emboss": "",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
            },
        ],
    )
    result = import_color_swatches_xlsx(db, content)
    assert result.created == 1
    assert result.updated == 1
    assert len(rows) == 1
    assert rows[0].color_name == "CORE BLACK REV"


def test_import_rejects_missing_header():
    db, _ = _make_db()
    content = _xlsx(["color_code"], [{"color_code": "A0QM"}])
    result = import_color_swatches_xlsx(db, content)
    assert result.created == 0
    assert result.errors[0].row == 1


def test_empty_color_standard_on_create_uses_suggestion():
    db, rows = _make_db([])
    content = _xlsx(
        ["factory_key", "buyer_code", "color_code", "color_name", "emboss", "color_standard", "material_name"],
        [
            {
                "factory_key": "ADIS",
                "buyer_code": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "",
                "color_standard": "",
                "material_name": "GENERIC PU",
            }
        ],
    )
    result = import_color_swatches_xlsx(db, content)
    assert result.created == 1
    assert result.errors == []
    assert rows[0].color_standard
    assert rows[0].color_standard.startswith("TMAD002")


def test_unknown_csi_color_code_skips_row():
    db, rows = _make_db([], csi_rows=[])
    content = _xlsx(
        [
            "factory_key",
            "buyer_code",
            "color_code",
            "color_name",
            "emboss",
            "color_standard",
            "material_name",
            "csi_color_code",
        ],
        [
            {
                "factory_key": "ADIS",
                "buyer_code": "ADIDAS",
                "color_code": "A0QM",
                "color_name": "CORE BLACK",
                "emboss": "",
                "color_standard": "SW1",
                "material_name": "GENERIC PU",
                "csi_color_code": "NOPE",
            }
        ],
    )
    result = import_color_swatches_xlsx(db, content)
    assert result.created == 0
    assert result.skipped == 1
    assert len(rows) == 0
    assert any("NOPE" in err.reason for err in result.errors)
