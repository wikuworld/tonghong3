"""Workbook Data Intelligent: isian unmapping, kolom terkunci, dan nama CSI."""

from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook
from io import BytesIO

from app.models import BrandBuyerMap, CsiColorMaster, CustomerBrandMapping
from app.services.data_intelligent import (
    SHEET_ORDER,
    _apply_brand_buyer,
    _apply_csi_baru,
    _apply_customer_brand,
    _apply_item_code,
    _apply_nama_beda,
    _apply_swatch_csi,
    name_gap_kind,
    parse_workbook,
    render_workbook,
    sheets_from_web,
)


def test_name_gap_ignores_fluorescent_and_abbreviation():
    assert name_gap_kind("UNIVERSITY RED (FLUORESCENT)", "UNIVERSITY RED") == "same"
    assert name_gap_kind("LT OREWOOD BROWN", "LT OREWOOD") == "prefix"
    assert name_gap_kind("PURE PLATINUM", "PURE LATINUM") == "different"
    assert name_gap_kind("RED", "DARK MAROON") == "different"


def _added(db):
    found = []

    def _add(obj):
        found.append(obj)

    db.add.side_effect = _add
    db.flush.side_effect = lambda: None
    return found


def test_customer_brand_filled_only_when_brand_exists():
    db = MagicMock()
    added = _added(db)
    customer = SimpleNamespace(id=6, code="DS026", name="Dean")
    brand = SimpleNamespace(id=3, code="NIKE")
    expected = {"6": {"customer_id": 6, "customer_code": "DS026", "customer_name": "Dean", "brand_code": ""}}
    rows = [
        {"__row__": 2, "customer_id": "6", "customer_code": "DS026", "customer_name": "Dean", "brand_code": "NOPE"},
        {"__row__": 3, "customer_id": "6", "customer_code": "DS026", "customer_name": "Dean", "brand_code": "NIKE"},
        {"__row__": 4, "customer_id": "6", "customer_code": "DS026", "customer_name": "Dean", "brand_code": ""},
    ]
    applied, rejected = [], []
    _apply_customer_brand(db, rows, expected, {6: customer}, {"NIKE": brand}, set(), applied, rejected)
    assert [row["row"] for row in applied] == [3]
    assert rejected[0]["reason"] == "Kode brand tidak ada di master"
    assert len(added) == 1
    assert isinstance(added[0], CustomerBrandMapping)
    assert added[0].customer_id == 6 and added[0].brand_id == 3


def test_customer_brand_rejects_locked_change_and_existing_pair():
    db = MagicMock()
    added = _added(db)
    customer = SimpleNamespace(id=6, code="DS026", name="Dean")
    brand = SimpleNamespace(id=1, code="NIKE")
    expected = {"6": {"customer_id": 6, "customer_code": "DS026", "customer_name": "Dean", "brand_code": ""}}
    rows = [
        {"__row__": 2, "customer_id": "6", "customer_code": "CHANGED", "customer_name": "Dean", "brand_code": "NIKE"},
        {"__row__": 3, "customer_id": "6", "customer_code": "DS026", "customer_name": "Dean", "brand_code": "NIKE"},
    ]
    applied, rejected = [], []
    _apply_customer_brand(db, rows, expected, {6: customer}, {"NIKE": brand}, {(6, 1)}, applied, rejected)
    assert applied == []
    assert added == []
    assert rejected[0]["reason"] == "Kolom terkunci berubah"
    assert rejected[1]["reason"] == "Relasi customer-brand sudah ada"


def test_brand_with_buyer_is_not_overwritten():
    db = MagicMock()
    added = _added(db)
    brand = SimpleNamespace(id=6, code="ON", name="ON CLOUD")
    expected = {"6": {"brand_id": 6, "brand_code": "ON", "brand_name": "ON CLOUD", "buyer_code": ""}}
    rows = [{"__row__": 2, "brand_id": "6", "brand_code": "ON", "brand_name": "ON CLOUD", "buyer_code": "NIKE"}]
    applied, rejected = [], []
    _apply_brand_buyer(db, rows, expected, {6: brand}, {"NIKE"}, {6}, applied, rejected)
    assert applied == []
    assert added == []
    assert rejected[0]["reason"] == "Brand sudah punya buyer"


def test_swatch_links_only_same_buyer_csi():
    recipe = SimpleNamespace(id=9, buyer_code="NB", csi_color_master_id=None, color_name="WHITE")
    csi = SimpleNamespace(id=4, buyer_code="NIKE", color_code="100", is_active=True, color_name="WHITE")
    expected = {
        "9": {
            "recipe_id": 9,
            "factory_key": "SAI",
            "buyer_code": "NB",
            "color_code": "100",
            "color_name": "WHITE",
            "material_name": "PU",
            "emboss": "AR",
            "csi_color_code": "",
        }
    }
    record = {**expected["9"], "__row__": 2, "recipe_id": "9", "csi_color_code": "100"}
    applied, rejected = [], []
    _apply_swatch_csi([record], expected, {9: recipe}, {("NIKE", "100"): csi}, applied, rejected)
    assert applied == []
    assert recipe.csi_color_master_id is None
    assert rejected[0]["reason"] == "Kode CSI tidak ada untuk buyer ini"

    csi.buyer_code = "NB"
    applied, rejected = [], []
    _apply_swatch_csi([record], expected, {9: recipe}, {("NB", "100"): csi}, applied, rejected)
    assert recipe.csi_color_master_id == 4
    assert applied[0]["action"] == "swatch_csi"


def test_csi_baru_does_not_overwrite_existing_code():
    db = MagicMock()
    added = _added(db)
    existing = SimpleNamespace(id=1, buyer_code="NB", color_code="100", color_name="WHITE", is_active=True)
    rows = [
        {"__row__": 2, "buyer_code": "NB", "color_code": "100", "color_name": "LAIN"},
        {"__row__": 3, "buyer_code": "NB", "color_code": "200", "color_name": "BARU"},
        {"__row__": 4, "buyer_code": "", "color_code": "", "color_name": ""},
    ]

    def _flush():
        for obj in added:
            if getattr(obj, "id", None) is None:
                obj.id = 77

    db.flush.side_effect = _flush
    applied, rejected = [], []
    index = {("NB", "100"): existing}
    _apply_csi_baru(db, rows, {"NB"}, index, applied, rejected)
    assert existing.color_name == "WHITE"
    assert rejected[0]["reason"] == "CSI buyer dan kode ini sudah ada"
    assert len(added) == 1
    assert isinstance(added[0], CsiColorMaster)
    assert added[0].color_code == "200"
    assert index[("NB", "200")] is added[0]
    assert applied[0]["action"] == "csi_baru"


def test_nama_beda_decisions():
    recipe = SimpleNamespace(id=3, color_name="PURE PLATINUM", csi_color_master_id=8)
    csi = SimpleNamespace(id=8, color_name="PURE LATINUM")
    expected = {
        "3": {
            "recipe_id": 3,
            "buyer_code": "NIKE",
            "color_code": "04Z",
            "swatch_name": "PURE PLATINUM",
            "csi_id": 8,
            "csi_name": "PURE LATINUM",
            "keputusan": "",
        }
    }

    def row(decision):
        return {**{key: str(value) if key in {"recipe_id", "csi_id"} else value for key, value in expected["3"].items()}, "__row__": 2, "keputusan": decision}

    applied, rejected = [], []
    _apply_nama_beda([row("")], expected, {3: recipe}, {8: csi}, applied, rejected)
    assert applied == [] and rejected == []
    assert recipe.color_name == "PURE PLATINUM"
    assert csi.color_name == "PURE LATINUM"

    applied, rejected = [], []
    _apply_nama_beda([row("BIARKAN")], expected, {3: recipe}, {8: csi}, applied, rejected)
    assert applied == [] and rejected == []

    applied, rejected = [], []
    _apply_nama_beda([row("IKUT_CSI")], expected, {3: recipe}, {8: csi}, applied, rejected)
    assert recipe.color_name == "PURE LATINUM"
    assert applied[0]["action"] == "ikut_csi"

    recipe.color_name = "PURE PLATINUM"
    applied, rejected = [], []
    _apply_nama_beda([row("IKUT_SWATCH")], expected, {3: recipe}, {8: csi}, applied, rejected)
    assert csi.color_name == "PURE PLATINUM"
    assert applied[0]["action"] == "ikut_swatch"

    csi.color_name = "PURE LATINUM"
    applied, rejected = [], []
    _apply_nama_beda([row("SALAH")], expected, {3: recipe}, {8: csi}, applied, rejected)
    assert rejected[0]["reason"].startswith("Keputusan harus")
    assert csi.color_name == "PURE LATINUM"


def test_item_code_and_locked_column():
    item = SimpleNamespace(id=5, item_code="FG1", color_code="")
    expected = {
        "5": {
            "item_id": 5,
            "item_code": "FG1",
            "material_name": "PU",
            "color": "BLACK",
            "emboss": "AR",
            "thickness_range": "1.2-1.4",
            "color_code": "",
        }
    }
    bad = {**expected["5"], "__row__": 2, "item_id": "5", "material_name": "LAIN", "color_code": "A0QM"}
    applied, rejected = [], []
    _apply_item_code([bad], expected, {5: item}, applied, rejected)
    assert item.color_code == ""
    assert rejected[0]["reason"] == "Kolom terkunci berubah"

    good = {**expected["5"], "__row__": 3, "item_id": "5", "color_code": "a0qm"}
    applied, rejected = [], []
    _apply_item_code([good], expected, {5: item}, applied, rejected)
    assert item.color_code == "A0QM"
    assert applied[0]["action"] == "item_color_code"


def test_workbook_sheets_and_locked_cells():
    snapshot = {
        "counts": {},
        "sheets": {name: [] for name in SHEET_ORDER},
    }
    snapshot["sheets"]["Map_Customer_Brand"] = [
        {"customer_id": 6, "customer_code": "DS026", "customer_name": "Dean", "brand_code": ""}
    ]
    raw = render_workbook(snapshot)
    wb = load_workbook(BytesIO(raw))
    assert tuple(wb.sheetnames) == SHEET_ORDER
    ws = wb["Map_Customer_Brand"]
    assert ws.protection.sheet is True
    assert ws["A2"].protection.locked is True
    assert ws["D2"].protection.locked is False
    parsed = parse_workbook(raw)
    assert parsed["Map_Customer_Brand"][0]["customer_code"] == "DS026"
    assert parsed["Map_Customer_Brand"][0]["brand_code"] == ""
    assert "CSI_Baru" not in parsed or parsed["CSI_Baru"] == []


def test_brand_buyer_fill_uses_existing_buyer_only():
    db = MagicMock()
    added = _added(db)
    brand = SimpleNamespace(id=6, code="ON", name="ON CLOUD")
    expected = {"6": {"brand_id": 6, "brand_code": "ON", "brand_name": "ON CLOUD", "buyer_code": ""}}
    missing = [{"__row__": 2, "brand_id": "6", "brand_code": "ON", "brand_name": "ON CLOUD", "buyer_code": "NOPE"}]
    applied, rejected = [], []
    _apply_brand_buyer(db, missing, expected, {6: brand}, {"NIKE"}, set(), applied, rejected)
    assert added == []
    assert rejected[0]["reason"] == "Kode buyer tidak ada di master"

    ok = [{"__row__": 3, "brand_id": "6", "brand_code": "ON", "brand_name": "ON CLOUD", "buyer_code": "nike"}]
    applied, rejected = [], []
    _apply_brand_buyer(db, ok, expected, {6: brand}, {"NIKE"}, set(), applied, rejected)
    assert isinstance(added[0], BrandBuyerMap)
    assert added[0].buyer_code == "NIKE"
    assert applied[0]["action"] == "brand_buyer"


def test_web_payload_keeps_only_editable_sheets():
    parsed = sheets_from_web(
        {
            "sheets": {
                "Map_Customer_Brand": [
                    {"customer_id": 6, "customer_code": "DS026", "customer_name": "Dean", "brand_code": "NIKE"}
                ],
                "Item_Tanpa_Swatch": [{"item_id": 1, "color_code": "ZZ"}],
                "CSI_Baru": [{"buyer_code": "nb", "color_code": "200", "color_name": "Baru"}],
            }
        }
    )
    assert set(parsed) == {"Map_Customer_Brand", "CSI_Baru"}
    assert parsed["CSI_Baru"][0]["buyer_code"] == "nb"
    assert parsed["CSI_Baru"][0]["__row__"] == 2
    assert parsed["Map_Customer_Brand"][0]["brand_code"] == "NIKE"
