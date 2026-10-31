"""Export/import Excel master customer."""

from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook

from app.models import Customer, CustomerPic
from app.services.customer_excel import (
    EXPORT_COLUMNS,
    TEMPLATE_COLUMNS,
    export_customers_template_xlsx,
    export_customers_xlsx,
    import_customers_xlsx,
)


def _xlsx(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for i, header in enumerate(headers, 1):
        ws.cell(1, i, header)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, header in enumerate(headers, 1):
            ws.cell(r_idx, c_idx, row.get(header, ""))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stmt_table(stmt) -> str:
    text = str(stmt).lower()
    if "customer_pics" in text:
        return "customer_pics"
    if "customers" in text:
        return "customers"
    return ""


def _make_db(customers: list[Customer] | None = None, pics: list[CustomerPic] | None = None):
    customers = list(customers or [])
    pics = list(pics or [])
    added: list = []
    db = MagicMock()
    seq = {"n": 40}

    def scalars(stmt):
        result = MagicMock()
        table = _stmt_table(stmt)
        if table == "customers":
            result.all.return_value = list(customers)
        elif table == "customer_pics":
            result.all.return_value = list(pics)
        else:
            result.all.return_value = []
        return result

    def flush():
        for obj in added:
            if isinstance(obj, Customer) and obj.id is None:
                seq["n"] += 1
                obj.id = seq["n"]

    db.scalars.side_effect = scalars
    db.add.side_effect = added.append
    db.flush.side_effect = flush
    return db, added


def _customer(**kwargs) -> Customer:
    data = {
        "id": 3,
        "code": "AD002",
        "name": "ADIS",
        "is_active": True,
        "is_export_customer": False,
        "outside_banten": False,
        "discount_modifier_enabled": False,
        "discount_modifier_type": "nominal",
    }
    data.update(kwargs)
    return Customer(**data)


def test_export_headers_and_values():
    db, _ = _make_db([_customer(city="Tangerang")])
    wb = load_workbook(BytesIO(export_customers_xlsx(db)))
    headers = [wb.active.cell(1, c).value for c in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPORT_COLUMNS
    assert wb.active.cell(2, 1).value == 3
    assert wb.active.cell(2, 2).value == "AD002"
    assert wb.active.cell(2, 3).value == "ADIS"
    assert wb.active.cell(2, 8).value == "Tangerang"


def test_template_is_header_only():
    wb = load_workbook(BytesIO(export_customers_template_xlsx()))
    headers = [wb.active.cell(1, c).value for c in range(1, len(TEMPLATE_COLUMNS) + 1)]
    assert headers == TEMPLATE_COLUMNS
    assert "id" not in headers
    assert wb.active.cell(2, 1).value is None


def test_import_creates_customer_and_primary_contact():
    db, added = _make_db()
    content = _xlsx(
        ["code", "name", "contact_person", "phone"],
        [{"code": "hw036", "name": "Huawei", "contact_person": "Ani", "phone": "0812"}],
    )
    result = import_customers_xlsx(db, content)
    assert result.created == 1
    assert result.errors == []
    created = [obj for obj in added if isinstance(obj, Customer)]
    pics = [obj for obj in added if isinstance(obj, CustomerPic)]
    assert created[0].code == "hw036"
    assert created[0].name == "Huawei"
    assert created[0].is_active is True
    assert pics[0].customer_id == created[0].id
    assert pics[0].name == "Ani"
    assert pics[0].phone == "0812"


def test_import_updates_then_skips_same_row():
    existing = _customer()
    db, added = _make_db([existing])
    headers = ["id", "code", "name"]
    row = {"id": 3, "code": "ad002", "name": "ADIS UPDATED"}
    first = import_customers_xlsx(db, _xlsx(headers, [row]))
    assert first.updated == 1
    assert first.created == 0
    assert existing.name == "ADIS UPDATED"
    assert existing.code == "AD002"
    assert added == []
    second = import_customers_xlsx(db, _xlsx(headers, [row]))
    assert second.skipped == 1
    assert second.updated == 0


def test_import_rejects_missing_name_and_duplicate_code():
    other = _customer(id=9, code="ZZ001", name="Other")
    db, added = _make_db([_customer(), other])
    content = _xlsx(
        ["id", "code", "name"],
        [
            {"id": "", "code": "NEW1", "name": ""},
            {"id": 3, "code": "ZZ001", "name": "ADIS"},
        ],
    )
    result = import_customers_xlsx(db, content)
    assert result.created == 0
    assert result.updated == 0
    assert len(result.errors) == 2
    assert added == []


def test_import_missing_header():
    db, _ = _make_db()
    result = import_customers_xlsx(db, _xlsx(["code"], [{"code": "AD002"}]))
    assert result.total_rows == 0
    assert result.errors
    assert "name" in result.errors[0].reason
