"""Laporan piutang (RAR) — selaras Excel RAR 2026."""

from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.models import Invoice, InvoiceDetail

RAR_HEADERS = [
    "NO",
    "NO FAKTUR PAJAK",
    "CUSTOMER",
    "FAKTUR DATE",
    "NO TANDA TERIMA FAKTUR",
    "RECEIVED BY",
    "RECEIVED DATE",
    "PAYMENT TERM T/T DAYS",
    "MATURITY",
    "QTY",
    "AMOUNT FAKTUR (USD)",
    "NO. INVOICE",
    "PO NO.",
    "AMOUNT PAYMENT",
    "DATE PAYMENT",
    "REMARK",
    "NOPEN",
    "AJU",
    "NOTA RETUR/PEMBATALAN",
    "DATE",
    "AMOUNT",
    "DATE",
    "OIA EXPORT (USD)",
    "BALANCE AFTER DEDUCT",
]


def _dec(val: Any) -> Decimal:
    if val is None or val == "":
        return Decimal(0)
    try:
        return Decimal(str(val))
    except Exception:
        return Decimal(0)


def parse_payment_term_days(term: Optional[str]) -> Optional[int]:
    if not term:
        return None
    m = re.search(r"\d+", str(term))
    if not m:
        return None
    try:
        return int(m.group())
    except ValueError:
        return None


def payment_term_days_for_due(term: Optional[str]) -> Optional[int]:
    """Hari untuk hitung jatuh tempo. Angka dari term, atau 0 untuk 'advance'."""
    days = parse_payment_term_days(term)
    if days is not None:
        return days
    if term and re.search(r"advance", str(term), re.I):
        return 0
    return None


def resolve_invoice_payment_terms(
    *,
    customer_payment_terms: Optional[str] = None,
    term_of_payment: Optional[str] = None,
    due_date: Optional[date] = None,
    base_date: Optional[date] = None,
) -> tuple[Optional[str], Optional[date]]:
    """Isi term dari customer bila kosong; hitung due_date = base_date + N hari bila kosong."""
    term = (term_of_payment or "").strip() or None
    if not term:
        term = (customer_payment_terms or "").strip() or None

    due = due_date
    if due is None and base_date is not None and term:
        days = payment_term_days_for_due(term)
        if days is not None:
            due = base_date + timedelta(days=days)
    return term, due


def invoice_faktur_date(inv: Invoice) -> Optional[date]:
    return inv.tax_invoice_date or inv.invoice_date


def invoice_maturity(inv: Invoice) -> Optional[date]:
    if inv.due_date:
        return inv.due_date
    base = invoice_faktur_date(inv)
    days = payment_term_days_for_due(inv.term_of_payment)
    if base and days is not None:
        return base + timedelta(days=days)
    return None


def invoice_qty(db: Session, inv: Invoice) -> Decimal:
    total = db.scalar(
        select(func.coalesce(func.sum(InvoiceDetail.qty), 0)).where(
            InvoiceDetail.invoice_id == inv.id
        )
    )
    return _dec(total)


def invoice_balance(inv: Invoice) -> Decimal:
    total = _dec(inv.total_amount)
    paid = _dec(inv.payment_amount)
    retur = _dec(inv.credit_note_amount)
    oia = _dec(inv.oia_export_usd)
    return total - paid - retur - oia


def invoice_remark(inv: Invoice) -> str:
    if inv.payment_remark and str(inv.payment_remark).strip():
        return str(inv.payment_remark).strip()
    bal = invoice_balance(inv)
    if bal <= 0 and (inv.payment_amount is not None or inv.payment_status == "paid"):
        return "FINISH PAYMENT"
    if inv.payment_status == "paid" and bal <= 0:
        return "FINISH PAYMENT"
    if inv.payment_date and bal <= 0:
        return "FINISH PAYMENT"
    return "NOT YET"


def invoice_to_rar_row(inv: Invoice, *, seq: int, qty: Decimal) -> dict[str, Any]:
    bal = invoice_balance(inv)
    term_days = parse_payment_term_days(inv.term_of_payment)
    maturity = invoice_maturity(inv)
    faktur_dt = invoice_faktur_date(inv)
    pay_amt = inv.payment_amount
    if pay_amt is None and inv.payment_status == "paid" and inv.total_amount is not None:
        pay_amt = inv.total_amount
    return {
        "no": seq,
        "tax_invoice_no": inv.tax_invoice_no or "",
        "customer": inv.customer_name or "",
        "faktur_date": faktur_dt.isoformat() if faktur_dt else "",
        "receipt_no": inv.receipt_no or "",
        "received_by": inv.received_by or "",
        "received_date": inv.received_date.isoformat() if inv.received_date else "",
        "payment_term_days": term_days,
        "maturity": maturity.isoformat() if maturity else "",
        "qty": float(qty),
        "amount_faktur_usd": float(_dec(inv.total_amount)),
        "invoice_number": inv.invoice_number or "",
        "po_number": inv.po_number or "",
        "amount_payment": float(_dec(pay_amt)) if pay_amt is not None else None,
        "date_payment": inv.payment_date.isoformat() if inv.payment_date else "",
        "remark": invoice_remark(inv),
        "nopen": inv.nopen or "",
        "nomor_aju": inv.nomor_aju or "",
        "credit_note_no": inv.credit_note_no or "",
        "credit_note_date": inv.credit_note_date.isoformat() if inv.credit_note_date else "",
        "credit_note_amount": float(_dec(inv.credit_note_amount)) if inv.credit_note_amount else None,
        "credit_note_date_2": inv.credit_note_date.isoformat() if inv.credit_note_date else "",
        "oia_export_usd": float(_dec(inv.oia_export_usd)) if inv.oia_export_usd else None,
        "balance_after_deduct": float(bal),
        "payment_status": inv.payment_status,
        "invoice_id": inv.id,
    }


_INVOICE_OUT_FIELDS = (
    "id",
    "delivery_id",
    "customer_id",
    "po_id",
    "invoice_number",
    "packing_list_number",
    "po_number",
    "customer_name",
    "customer_address",
    "customer_phone",
    "invoice_date",
    "description",
    "term_of_payment",
    "bank_name",
    "account_number",
    "bank_address",
    "bank_phone",
    "total_amount",
    "total_in_words",
    "due_date",
    "tax_invoice_no",
    "tax_invoice_date",
    "receipt_no",
    "received_by",
    "received_date",
    "nopen",
    "nomor_aju",
    "bc_date",
    "payment_date",
    "payment_bank_account",
    "payment_proof_url",
    "payment_status",
    "payment_amount",
    "payment_remark",
    "credit_note_no",
    "credit_note_date",
    "credit_note_amount",
    "oia_export_usd",
    "created_at",
)


def enrich_invoice_out(inv: Invoice, db: Session) -> dict[str, Any]:
    qty = invoice_qty(db, inv)
    bal = invoice_balance(inv)
    data: dict[str, Any] = {}
    for key in _INVOICE_OUT_FIELDS:
        v = getattr(inv, key, None)
        if isinstance(v, Decimal):
            data[key] = float(v)
        elif isinstance(v, date):
            data[key] = v.isoformat()
        elif isinstance(v, datetime):
            data[key] = v.isoformat()
        else:
            data[key] = v
    data["qty_total"] = float(qty)
    data["balance_after_deduct"] = float(bal)
    data["remark"] = invoice_remark(inv)
    data["payment_term_days"] = parse_payment_term_days(inv.term_of_payment)
    return data


def _invoice_date_filter(inv: Invoice, date_from: Optional[date], date_to: Optional[date]) -> bool:
    fd = invoice_faktur_date(inv)
    if fd is None:
        return date_from is None and date_to is None
    if date_from and fd < date_from:
        return False
    if date_to and fd > date_to:
        return False
    return True


def _query_invoices(
    db: Session,
    *,
    date_from: Optional[date],
    date_to: Optional[date],
    customer: Optional[str],
    payment_status: str,
    year: Optional[int],
) -> list[Invoice]:
    stmt = (
        select(Invoice)
        .options(selectinload(Invoice.customer))
        .where(Invoice.voided_at.is_(None))
        .order_by(Invoice.tax_invoice_date.asc(), Invoice.id.asc())
    )
    if customer:
        like = f"%{customer.strip()}%"
        stmt = stmt.where(Invoice.customer_name.ilike(like))
    if year:
        stmt = stmt.where(
            or_(
                func.year(Invoice.tax_invoice_date) == year,
                func.year(Invoice.invoice_date) == year,
            )
        )
    if date_from:
        stmt = stmt.where(
            or_(
                Invoice.tax_invoice_date >= date_from,
                Invoice.invoice_date >= date_from,
            )
        )
    if date_to:
        stmt = stmt.where(
            or_(
                Invoice.tax_invoice_date <= date_to,
                Invoice.invoice_date <= date_to,
            )
        )
    rows = list(db.scalars(stmt).all())
    if payment_status == "paid":
        rows = [r for r in rows if invoice_remark(r) == "FINISH PAYMENT"]
    elif payment_status == "unpaid":
        rows = [r for r in rows if invoice_remark(r) != "FINISH PAYMENT"]
    return rows


def build_rar_payload(
    db: Session,
    *,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    customer: Optional[str] = None,
    payment_status: str = "all",
    year: Optional[int] = None,
    limit: int = 500,
    offset: int = 0,
) -> dict[str, Any]:
    all_rows = _query_invoices(
        db,
        date_from=date_from,
        date_to=date_to,
        customer=customer,
        payment_status=payment_status,
        year=year,
    )
    report_year = year or (date_from.year if date_from else date.today().year)
    rar_rows: list[dict[str, Any]] = []
    total_qty = Decimal(0)
    total_faktur = Decimal(0)
    total_payment = Decimal(0)
    for idx, inv in enumerate(all_rows, start=1):
        qty = invoice_qty(db, inv)
        row = invoice_to_rar_row(inv, seq=idx, qty=qty)
        rar_rows.append(row)
        total_qty += qty
        total_faktur += _dec(inv.total_amount)
        pay = row.get("amount_payment")
        if pay is not None:
            total_payment += _dec(pay)

    page = all_rows[offset : offset + limit]
    enriched = [enrich_invoice_out(inv, db) for inv in page]

    return {
        "report_title": f"RECAPITULATION ACCOUNT RECEIVABLE {report_year}",
        "year": report_year,
        "summary": {
            "count": len(all_rows),
            "qty_total": float(total_qty),
            "amount_faktur_usd": float(total_faktur),
            "amount_payment_usd": float(total_payment),
            "not_yet_count": sum(1 for r in rar_rows if r["remark"] != "FINISH PAYMENT"),
            "finish_payment_count": sum(1 for r in rar_rows if r["remark"] == "FINISH PAYMENT"),
        },
        "rows": enriched,
        "rar_rows": rar_rows,
        "total": len(all_rows),
        "limit": limit,
        "offset": offset,
    }


def render_rar_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    year = payload.get("year") or date.today().year
    summary = payload.get("summary") or {}
    rows = payload.get("rar_rows") or []

    ws["A1"] = payload.get("report_title") or f"RECAPITULATION ACCOUNT RECEIVABLE {year}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["J2"] = summary.get("qty_total", 0)
    ws["K2"] = summary.get("amount_faktur_usd", 0)
    ws["N2"] = summary.get("amount_payment_usd", 0)

    for col_idx, header in enumerate(RAR_HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=4):
        values = [
            row.get("no"),
            row.get("tax_invoice_no"),
            row.get("customer"),
            row.get("faktur_date") or None,
            row.get("receipt_no"),
            row.get("received_by"),
            row.get("received_date") or None,
            row.get("payment_term_days"),
            row.get("maturity") or None,
            row.get("qty"),
            row.get("amount_faktur_usd"),
            row.get("invoice_number"),
            row.get("po_number"),
            row.get("amount_payment"),
            row.get("date_payment") or None,
            row.get("remark"),
            row.get("nopen"),
            row.get("nomor_aju"),
            row.get("credit_note_no"),
            row.get("credit_note_date") or None,
            row.get("credit_note_amount"),
            row.get("credit_note_date_2") or None,
            row.get("oia_export_usd"),
            row.get("balance_after_deduct"),
        ]
        for col_idx, val in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
