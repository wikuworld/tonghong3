"""Impor folder CSV `LIST_HARGA_ALL_BUYER_extracted` → buyer_price_lists + baris JSON."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from app.services.buyer_price_aliases import (
    canonical_price_block,
    normalize_sheet_cells,
    split_material_and_thickness,
)

PERIODE_RE = re.compile(
    r"PERIODE\s+(\d{4}/\d{2}/\d{2})\s*[-–]\s*(\d{4}/\d{2}/\d{2})",
    re.IGNORECASE,
)


def filename_to_buyer(name: str) -> Optional[str]:
    m = re.match(r"^\d+_(.+)\.csv$", name.strip(), re.I)
    if not m:
        return None
    return m.group(1).upper()


def parse_period_from_rows(rows: list[list[str]]) -> tuple[str, Optional[date], Optional[date]]:
    label = ""
    start_d = end_d = None
    for row in rows[:20]:
        line = ",".join(row)
        if "PERIODE" in line.upper():
            label = line.strip()
            m = PERIODE_RE.search(line)
            if m:
                try:
                    start_d = datetime.strptime(m.group(1), "%Y/%m/%d").date()
                    end_d = datetime.strptime(m.group(2), "%Y/%m/%d").date()
                except ValueError:
                    pass
            break
    return label, start_d, end_d


def row_nonempty(row: list[str]) -> bool:
    return any((c or "").strip() for c in row)


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _fmt_num(v: Optional[float]) -> str:
    if v is None:
        return "0"
    # Simpan ringkas agar UI input-number tetap nyaman dibaca.
    if abs(v - int(v)) < 1e-9:
        return str(int(v))
    return f"{v:.6f}".rstrip("0").rstrip(".")


STANDARD_COLUMN_KEYS = [
    "ITEM",
    "THICKNESS",
    "COLOR",
    "COLOR WHITE",
    "COLOR BLACK",
    "COLOR COLOR",
    "FLOURASENT",
    "METALIC",
    "EMBOSS",
    "SPESIAL EMBOSS",
]


def _addon(total: Optional[float], base: Optional[float]) -> str:
    if total is None or base is None:
        return "0"
    return _fmt_num(max(0.0, float(total) - float(base)))


def build_price_cells(
    *,
    item: str,
    thickness: str = "",
    block: str = "",
    white: Optional[float],
    black: Optional[float] = None,
    color: Optional[float] = None,
    fluo: Optional[float] = None,
    metallic: Optional[float] = None,
    metallic_is_addon: bool = False,
) -> Optional[dict[str, Any]]:
    name, thick = split_material_and_thickness(item, thickness or _thickness_from_text(item))
    if not name or white is None:
        return None
    metal = _fmt_num(metallic) if metallic_is_addon and metallic is not None else _addon(metallic, white)
    if metallic_is_addon and metallic is None:
        metal = "0"
    return {
        "ITEM": name,
        "THICKNESS": thick,
        "COLOR": canonical_price_block(block),
        "COLOR WHITE": _fmt_num(white),
        "COLOR BLACK": _addon(black, white),
        "COLOR COLOR": _addon(color, white),
        "FLOURASENT": _addon(fluo, white),
        "METALIC": metal,
        "EMBOSS": "0",
        "SPESIAL EMBOSS": "0",
    }


def normalize_buyer_cells(cells: dict[str, Any]) -> Optional[dict[str, Any]]:
    """
    Normalisasi berbagai format CSV buyer ke shape standar UI:
    ITEM, THICKNESS, COLOR, COLOR WHITE, COLOR BLACK, COLOR COLOR, FLOURASENT, METALIC, EMBOSS, SPESIAL EMBOSS
    """
    upper = {str(k or "").strip().upper(): v for k, v in (cells or {}).items()}

    # Format "LIST HARGA ALL BUYER_extracted": MATERIAL NAME + PRICE + METALLIC/...
    if "MATERIAL NAME" in upper and "PRICE" in upper:
        item = str(upper.get("MATERIAL NAME") or "").strip()
        thickness = str(upper.get("THICKNESS") or "").strip()
        base = _to_float(upper.get("PRICE"))
        metallic_total = _to_float(upper.get("METALLIC/PEARLIZED/FLUORESCENT"))
        if not item or base is None:
            return None
        metallic_addon = 0.0
        if metallic_total is not None:
            metallic_addon = max(0.0, metallic_total - base)
        return normalize_sheet_cells({
            "ITEM": item,
            "THICKNESS": thickness,
            "COLOR": "",
            "COLOR WHITE": _fmt_num(base),
            "COLOR BLACK": "0",
            "COLOR COLOR": "0",
            "FLOURASENT": "0",
            "METALIC": _fmt_num(metallic_addon),
            "EMBOSS": "0",
            "SPESIAL EMBOSS": "0",
        })

    # Sudah dalam format standar UI.
    if "ITEM" in upper and "COLOR WHITE" in upper:
        return normalize_sheet_cells({
            "ITEM": str(upper.get("ITEM") or "").strip(),
            "THICKNESS": str(upper.get("THICKNESS") or "").strip(),
            "COLOR": str(upper.get("COLOR") or "").strip(),
            "COLOR WHITE": str(upper.get("COLOR WHITE") or "0").strip() or "0",
            "COLOR BLACK": str(upper.get("COLOR BLACK") or "0").strip() or "0",
            "COLOR COLOR": str(upper.get("COLOR COLOR") or "0").strip() or "0",
            "FLOURASENT": str(upper.get("FLOURASENT") or "0").strip() or "0",
            "METALIC": str(upper.get("METALIC") or "0").strip() or "0",
            "EMBOSS": str(upper.get("EMBOSS") or "0").strip() or "0",
            "SPESIAL EMBOSS": str(upper.get("SPESIAL EMBOSS") or "0").strip() or "0",
        })

    return None


def _extract_thickness(material_name: str) -> str:
    text = str(material_name or "").strip()
    m = re.search(r"(\d+(?:\.\d+)?\s*-\s*\d+(?:\.\d+)?)\s*MM?$", text, re.IGNORECASE)
    if not m:
        return ""
    return m.group(1).replace(" ", "")


def _first_numeric(values: list[str]) -> Optional[float]:
    for val in values:
        num = _to_float(val)
        if num is not None:
            return num
    return None


def parse_nike_block_rows(raw: list[list[str]], header_idx: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    current: Optional[dict[str, Any]] = None

    def flush_current() -> None:
        nonlocal current
        if not current:
            return
        item = str(current.get("ITEM") or "").strip()
        if not item:
            current = None
            return
        white = _to_float(current.get("WHITE"))
        black = _to_float(current.get("BLACK"))
        color = _to_float(current.get("COLOR_RAW"))
        metallic = _to_float(current.get("METALLIC_RAW"))
        fluo = _to_float(current.get("FLUORESCENT_RAW"))

        base = white
        if base is None:
            for candidate in (black, color, metallic, fluo):
                if candidate is not None:
                    base = candidate
                    break
        if base is None:
            current = None
            return

        built = build_price_cells(
            item=item,
            thickness=str(current.get("THICKNESS") or "").strip(),
            white=base,
            black=black,
            color=color,
            fluo=fluo,
            metallic=metallic,
        )
        if built:
            rows.append(built)
        current = None

    for row in raw[header_idx + 1 :]:
        if not row_nonempty(row):
            continue
        no_val = (row[0] if len(row) > 0 else "").strip()
        material_cell = (row[1] if len(row) > 1 else "").strip()
        numeric = _first_numeric([row[i].strip() for i in range(2, min(len(row), 6))])

        # Baris awal item material.
        if no_val.isdigit() and material_cell:
            flush_current()
            current = {
                "ITEM": material_cell,
                "THICKNESS": _extract_thickness(material_cell),
            }
            continue

        # Baris breakdown warna di bawah item aktif.
        if current and not no_val and material_cell:
            label = material_cell.upper()
            if label == "WHITE":
                current["WHITE"] = numeric
            elif label == "BLACK":
                current["BLACK"] = numeric
            elif label == "COLOR":
                current["COLOR_RAW"] = numeric
            elif "METALLIC" in label or "PEARL" in label or "ANODIZED" in label:
                current["METALLIC_RAW"] = numeric
            elif "FLUORESCENT" in label:
                current["FLUORESCENT_RAW"] = numeric

    flush_current()
    return rows


def parse_price_block_rows(raw: list[list[str]], header_idx: int, headers: list[str]) -> Optional[list[dict[str, Any]]]:
    """
    Parser untuk format blok:
    - baris item utama: NO terisi (1,2,3...), MATERIAL NAME berisi nama material
    - baris turunan: NO kosong, MATERIAL NAME berisi WHITE/BLACK/COLOR/FLUORESCENT...
    """
    try:
        idx_no = next(i for i, h in enumerate(headers) if h.strip().upper() == "NO")
        idx_material = next(i for i, h in enumerate(headers) if h.strip().upper() == "MATERIAL NAME")
        idx_price = next(i for i, h in enumerate(headers) if h.strip().upper() == "PRICE")
    except StopIteration:
        return None

    idx_metallic: Optional[int] = None
    for i, h in enumerate(headers):
        up = h.strip().upper()
        if "METALLIC" in up or "PEARL" in up or "ANODIZED" in up:
            idx_metallic = i
            break

    def get_val(row: list[str], idx: int) -> str:
        return row[idx].strip() if idx < len(row) else ""

    # Deteksi apakah benar format blok warna.
    continuation_count = 0
    for row in raw[header_idx + 1 :]:
        if not row_nonempty(row):
            continue
        no_val = get_val(row, idx_no)
        mat = get_val(row, idx_material).upper()
        if (not no_val) and mat in {"WHITE", "BLACK", "COLOR", "FLUORESCENT"}:
            continuation_count += 1
            if continuation_count >= 2:
                break
    if continuation_count < 2:
        return None

    rows: list[dict[str, Any]] = []
    current: Optional[dict[str, Any]] = None

    def flush_current() -> None:
        nonlocal current
        if not current:
            return
        item = str(current.get("ITEM") or "").strip()
        if not item:
            current = None
            return

        white = _to_float(current.get("WHITE"))
        black = _to_float(current.get("BLACK"))
        fluo = _to_float(current.get("FLUORESCENT"))
        metallic_total = _to_float(current.get("METALLIC_TOTAL"))
        color_price = _to_float(current.get("COLOR_PRICE"))

        base = white
        if base is None:
            for candidate in (black, color_price, fluo, metallic_total):
                if candidate is not None:
                    base = candidate
                    break
        if base is None:
            current = None
            return

        metallic_ref = color_price if color_price is not None else base
        metallic_addon = 0.0
        if metallic_total is not None:
            metallic_addon = max(0.0, metallic_total - metallic_ref)

        built = build_price_cells(
            item=item,
            thickness=str(current.get("THICKNESS") or "").strip(),
            white=base,
            black=black,
            color=color_price,
            fluo=fluo,
            metallic=metallic_addon,
            metallic_is_addon=True,
        )
        if built:
            rows.append(built)
        current = None

    for row in raw[header_idx + 1 :]:
        if not row_nonempty(row):
            continue
        no_val = get_val(row, idx_no)
        material = get_val(row, idx_material)
        price_val = _to_float(get_val(row, idx_price))
        metallic_val = _to_float(get_val(row, idx_metallic)) if idx_metallic is not None else None
        material_up = material.upper()

        if no_val.isdigit() and material:
            flush_current()
            current = {"ITEM": material, "THICKNESS": _extract_thickness(material)}
            continue

        if current and not no_val and material:
            if material_up == "WHITE":
                current["WHITE"] = price_val
            elif material_up == "BLACK":
                current["BLACK"] = price_val
            elif material_up == "COLOR":
                current["COLOR_PRICE"] = price_val
                if metallic_val is not None:
                    current["METALLIC_TOTAL"] = metallic_val
            elif "FLUORESCENT" in material_up:
                current["FLUORESCENT"] = price_val
                if metallic_val is not None:
                    current["METALLIC_TOTAL"] = metallic_val
            elif any(x in material_up for x in ("METALLIC", "PEARL", "ANODIZED")):
                current["METALLIC_TOTAL"] = price_val

    flush_current()
    return rows


def parse_buyer_csv(path: Path) -> dict[str, Any]:
    with open(path, newline="", encoding="utf-8") as f:
        raw = list(csv.reader(f))
    title = (raw[0][0] if raw and raw[0] else "").strip()
    period_label, ps, pe = parse_period_from_rows(raw)

    header_idx: Optional[int] = None
    headers: list[str] = []
    for i, row in enumerate(raw):
        if not row:
            continue
        c0 = (row[0] or "").strip().upper()
        if c0 == "NO" or c0.startswith("NO"):
            header_idx = i
            headers = [(c or "").strip() for c in row]
            break

    data_rows: list[dict[str, Any]] = []
    if header_idx is not None:
        upper_headers = {h.strip().upper() for h in headers}
        if "PRICE COMMON" in upper_headers and "MATERIAL NAME" in upper_headers:
            data_rows = parse_nike_block_rows(raw, header_idx)
            return {
                "title": title,
                "period_label": period_label,
                "period_start": ps,
                "period_end": pe,
                "header_row_index": header_idx,
                "csv_line_count": len(raw),
                "data_rows": data_rows,
            }
        if "MATERIAL NAME" in upper_headers and "PRICE" in upper_headers:
            block_rows = parse_price_block_rows(raw, header_idx, headers)
            if block_rows is not None:
                return {
                    "title": title,
                    "period_label": period_label,
                    "period_start": ps,
                    "period_end": pe,
                    "header_row_index": header_idx,
                    "csv_line_count": len(raw),
                    "data_rows": block_rows,
                }
        for row in raw[header_idx + 1 :]:
            if not row_nonempty(row):
                continue
            d: dict[str, Any] = {}
            for k, h in enumerate(headers):
                key = h if h else f"_col{k}"
                val = row[k].strip() if k < len(row) else ""
                d[key] = val
            normalized = normalize_buyer_cells(d)
            if normalized is not None:
                data_rows.append(normalized)
    else:
        for row in raw:
            if not row_nonempty(row):
                continue
            normalized = normalize_buyer_cells(
                {f"col_{k}": row[k] if k < len(row) else "" for k in range(len(row))}
            )
            if normalized is not None:
                data_rows.append(normalized)

    return {
        "title": title,
        "period_label": period_label,
        "period_start": ps,
        "period_end": pe,
        "header_row_index": header_idx,
        "csv_line_count": len(raw),
        "data_rows": data_rows,
    }


def persist_buyer_folder(
    db: Any,
    folder: Path,
    source_label: Optional[str],
) -> list[Any]:
    """Timpa list hanya bila buyer dan periode sama; periode lain tetap utuh."""
    from app.models import BuyerPriceList, BuyerPriceSheetRow
    from app.services.buyer_price_period import find_list_same_period, unique_brand_id_for_buyer

    if not folder.is_dir():
        raise ValueError(f"Bukan folder: {folder}")

    out: list[Any] = []
    for p in sorted(folder.glob("*.csv")):
        if p.name.upper() == "README.MD":
            continue
        buyer = filename_to_buyer(p.name)
        if not buyer:
            continue
        parsed = parse_buyer_csv(p)
        data_rows = parsed["data_rows"]

        existing = find_list_same_period(
            db, buyer, parsed["period_start"], parsed["period_end"]
        )
        if existing:
            db.delete(existing)
            db.flush()

        lst = BuyerPriceList(
            buyer_code=buyer,
            brand_id=unique_brand_id_for_buyer(db, buyer),
            title=parsed["title"] or None,
            period_label=parsed["period_label"] or None,
            period_start=parsed["period_start"],
            period_end=parsed["period_end"],
            currency="USD",
            meta_json={
                "filename": p.name,
                "source_label": source_label or folder.name,
                "header_row_index": parsed["header_row_index"],
                "csv_line_count": parsed["csv_line_count"],
                "data_row_count": len(data_rows),
            },
        )
        db.add(lst)
        db.flush()
        for ri, cells in enumerate(data_rows):
            db.add(BuyerPriceSheetRow(list_id=lst.id, row_number=ri, cells=cells))
        out.append(lst)
    db.flush()
    for o in out:
        db.refresh(o)
    return out


def _xlsx_num(val: Any) -> Optional[float]:
    if val is None:
        return None
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return float(val)
    s = str(val).strip().replace(",", "")
    if not s or s in {"-", "–"}:
        return None
    if s.startswith("+"):
        s = s[1:].strip()
    try:
        return float(s)
    except ValueError:
        return None


def _xlsx_metallic_addon(val: Any) -> Optional[float]:
    if val is None:
        return None
    s = str(val).strip()
    if s.startswith("+"):
        return _xlsx_num(s)
    return None


def _thickness_from_text(*parts: Any) -> str:
    blob = " ".join(str(p or "") for p in parts)
    m = re.search(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)", blob)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = re.search(r"(\d+(?:\.\d+)?)\s*MM\b", blob, re.I)
    if m:
        return m.group(1)
    m = re.search(r"\b(\d+\.\d+)\b", blob)
    if m:
        return m.group(1)
    return ""


def _is_note_material(name: str) -> bool:
    up = str(name or "").strip().upper()
    return up.startswith("SPECIAL DIRECT") or up.startswith("SPECIAL OIA")


def parse_adidas_quotation_xlsx(path: Path, sheet_name: str = "ADCI") -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {sheet_name} tidak ada di {path.name}")
    rows = list(wb[sheet_name].iter_rows(values_only=True))
    wb.close()

    header_idx = None
    for i, row in enumerate(rows[:20]):
        vals = [str(x).strip().upper() if x is not None else "" for x in (row or ())]
        if "MATERIAL NAME" in vals and "THICKNESS" in vals:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("Header Material Name tidak ditemukan")

    data_rows: list[dict[str, Any]] = []
    notes: list[str] = []
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        first = str(row[0] or "").strip()
        if first.upper().startswith("REMARK"):
            break
        no = row[1] if len(row) > 1 else None
        if not (isinstance(no, (int, float)) or (isinstance(no, str) and str(no).strip().isdigit())):
            if first.startswith("1.") or "Period of validity" in first:
                notes.append(first)
            continue
        name = str(row[5] or "").strip() if len(row) > 5 else ""
        if not name:
            continue
        name = name.lstrip("'\"")
        thick = _thickness_from_text(row[6] if len(row) > 6 else "", name)
        white = _xlsx_num(row[8] if len(row) > 8 else None)
        built = build_price_cells(
            item=name,
            thickness=thick,
            block="",
            white=white,
            black=_xlsx_num(row[9] if len(row) > 9 else None),
            color=_xlsx_num(row[10] if len(row) > 10 else None),
            fluo=_xlsx_num(row[11] if len(row) > 11 else None),
            metallic=_xlsx_metallic_addon(row[12] if len(row) > 12 else None),
            metallic_is_addon=True,
        )
        if built:
            data_rows.append(built)
    return {
        "title": "QUOTATION ADIDAS FW26 (FEB-JUL 2026)",
        "period_label": "PERIODE 2026/02/01 - 2026/07/31",
        "period_start": date(2026, 2, 1),
        "period_end": date(2026, 7, 31),
        "data_rows": data_rows,
        "notes": notes,
        "canonical_sheet": sheet_name,
    }


def parse_nike_tempel_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    notes: list[str] = []
    data_rows: list[dict[str, Any]] = []

    blocks = (
        (1, 2, "COMMON"),
        (7, 8, "SPECIAL"),
        (14, 15, "OUTSIDE BANTEN"),
        (20, 21, "OUTSIDE SPECIAL"),
    )
    for r in range(4, 32):
        for name_col, price_col, block in blocks:
            name = ws.cell(r, name_col).value
            if name is None or not str(name).strip():
                continue
            name_s = str(name).strip()
            if _is_note_material(name_s):
                notes.append(name_s)
                continue
            prices = [_xlsx_num(ws.cell(r, price_col + i).value) for i in range(5)]
            if all(p is None for p in prices):
                continue
            white, black, color, metallic, fluo = prices
            if white is None:
                continue
            built = build_price_cells(
                item=name_s,
                thickness=_thickness_from_text(name_s),
                block=block,
                white=white,
                black=black,
                color=color,
                fluo=fluo,
                metallic=metallic,
            )
            if built:
                data_rows.append(built)
    wb.close()
    return {
        "title": "UPDATE PRICE NIKE (JUL-DES 2026)",
        "period_label": "JUL-DES 2026",
        "period_start": date(2026, 7, 1),
        "period_end": date(2026, 12, 31),
        "data_rows": data_rows,
        "notes": notes,
    }


def _replace_buyer_list(
    db: Any,
    *,
    buyer_code: str,
    parsed: dict[str, Any],
    source_file: str,
    kind: str,
    extra_meta: Optional[dict[str, Any]] = None,
) -> Any:
    from app.models import BuyerPriceList, BuyerPriceSheetRow
    from app.services.buyer_price_period import find_list_same_period, unique_brand_id_for_buyer
    from app.services.buyer_registry import ensure_buyer_code

    ensure_buyer_code(db, buyer_code)
    existing = find_list_same_period(
        db,
        buyer_code,
        parsed.get("period_start"),
        parsed.get("period_end"),
    )
    if existing:
        db.delete(existing)
        db.flush()
    meta = {
        "source_file": source_file,
        "kind": kind,
        "data_row_count": len(parsed["data_rows"]),
        "column_keys": list(STANDARD_COLUMN_KEYS),
        "notes": parsed.get("notes") or [],
    }
    if extra_meta:
        meta.update(extra_meta)
    lst = BuyerPriceList(
        buyer_code=buyer_code,
        brand_id=unique_brand_id_for_buyer(db, buyer_code),
        title=parsed.get("title"),
        period_label=parsed.get("period_label"),
        period_start=parsed.get("period_start"),
        period_end=parsed.get("period_end"),
        currency="USD",
        meta_json=meta,
    )
    db.add(lst)
    db.flush()
    for ri, cells in enumerate(parsed["data_rows"]):
        db.add(BuyerPriceSheetRow(list_id=lst.id, row_number=ri, cells=cells))
    db.flush()
    db.refresh(lst)
    return lst


def _first_existing(paths: list[Path]) -> Optional[Path]:
    for p in paths:
        if p.is_file():
            return p
    return None


def persist_official_buyer_workbooks(db: Any, doc_dir: Path) -> list[Any]:
    """Isi ulang list ADIDAS (sheet ADCI) dan NIKE (4 blok, tanpa kolom kanan)."""
    adidas_path = _first_existing(
        [
            doc_dir / "1._QUOTATION_ADIDAS_FW26_FEB-JUL_2026_ (1).xlsx",
            doc_dir / "1._QUOTATION_ADIDAS_FW26_FEB-JUL_2026_ (2).xlsx",
        ]
    )
    nike_path = _first_existing(
        [
            doc_dir / "HARGA_TEMPEL_NIKE_JUL_-_DES_2026_UPDATE (1).xlsx",
            doc_dir / "HARGA_TEMPEL_NIKE_JUL_-_DES_2026_UPDATE (2).xlsx",
        ]
    )
    out: list[Any] = []
    if adidas_path:
        parsed = parse_adidas_quotation_xlsx(adidas_path)
        out.append(
            _replace_buyer_list(
                db,
                buyer_code="ADIDAS",
                parsed=parsed,
                source_file=adidas_path.name,
                kind="adidas_quotation",
                extra_meta={"canonical_sheet": parsed.get("canonical_sheet") or "ADCI"},
            )
        )
    if nike_path:
        parsed = parse_nike_tempel_xlsx(nike_path)
        out.append(
            _replace_buyer_list(
                db,
                buyer_code="NIKE",
                parsed=parsed,
                source_file=nike_path.name,
                kind="nike_tempel",
            )
        )
    return out


def list_to_api_dict(db: Any, list_id: int) -> dict[str, Any]:
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload

    from app.models import BuyerPriceList, MasterItem

    row = db.scalars(
        select(BuyerPriceList)
        .options(selectinload(BuyerPriceList.rows))
        .where(BuyerPriceList.id == list_id)
    ).first()
    if not row:
        return {}
    lines = sorted(row.rows, key=lambda r: r.row_number)
    rows_out: list[dict[str, Any]] = []
    item_codes = {
        str((r.cells or {}).get("ITEM") or "").strip()
        for r in lines
        if isinstance(r.cells, dict) and str((r.cells or {}).get("ITEM") or "").strip()
    }
    item_codes = {x for x in item_codes if x}
    material_by_code: dict[str, str] = {}
    if item_codes:
        mats = db.execute(
            select(MasterItem.item_code, MasterItem.material_name).where(MasterItem.item_code.in_(item_codes))
        ).all()
        material_by_code = {
            str(code or "").strip(): str(material or "").strip()
            for code, material in mats
            if str(code or "").strip() and str(material or "").strip()
        }

    for line in lines:
        cells = dict(line.cells or {}) if isinstance(line.cells, dict) else {}
        item_val = str(cells.get("ITEM") or "").strip()
        if item_val in material_by_code:
            cells["ITEM"] = material_by_code[item_val]
        rows_out.append(cells)

    brand_code = None
    brand_name = None
    if row.brand_id:
        from app.models import Brand

        brand = db.get(Brand, row.brand_id)
        if brand:
            brand_code = brand.code
            brand_name = brand.name

    return {
        "id": row.id,
        "buyer_code": row.buyer_code,
        "brand_id": row.brand_id,
        "brand_code": brand_code,
        "brand_name": brand_name,
        "title": row.title,
        "period_label": row.period_label,
        "period_start": row.period_start.isoformat() if row.period_start else None,
        "period_end": row.period_end.isoformat() if row.period_end else None,
        "currency": row.currency,
        "outside_price_adjust": float(row.outside_price_adjust) if row.outside_price_adjust is not None else None,
        "outside_adjust_type": row.outside_adjust_type,
        "discount_modifier_enabled": bool(getattr(row, "discount_modifier_enabled", False)),
        "discount_modifier": float(row.discount_modifier) if row.discount_modifier is not None else None,
        "discount_modifier_type": row.discount_modifier_type,
        "meta_json": row.meta_json,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "rows": rows_out,
        "row_count": len(lines),
    }
