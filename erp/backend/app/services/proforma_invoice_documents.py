"""Proforma Invoice (PI) — PDF (ReportLab) & Word (.docx) dari data PO."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from html import escape
from typing import Any, Optional

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Customer, CustomerPic, PoDetail, PoHeader

COMPANY_LINES = (
    "PT. TONG HONG TANNERY INDONESIA",
    "Kawasan Industri Modern II Jl. Modern Industri XIV/W-1, Ds. Barengkok Kec. KibinSerang - Banten Indonesia",
    "Telp. 0254-8480870, Fax. 0254-8480743, E-Mail Address: thin.pu02@tong-hong.id",
)

BANK_LINES = (
    "BANK CTBC INDONESIA",
    "A/C NO. : 2020357701 (USD)",
    "Bank Add: TAMARA CENTER 16th FLOOR, JL. JENDRAL SUDIRMAN KAV.24 JAKARTA",
    "Tel: 62-21-5237788/7245, Fax : 021-2500622",
)

ROMAN_MONTHS = ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII")

DEFAULT_PAYMENT_TERMS = "BY T/T 30 DAYS"
DEFAULT_UNIT = "SF"


def _fmt_date(d: date | None) -> str:
    if not d:
        return "—"
    return d.isoformat()


def _fmt_etd(d: date | None) -> str:
    if not d:
        return ""
    return f"{d.month}-{d.day}"


def _dec_money(v: Decimal | None) -> str:
    if v is None:
        return ""
    return f"{v.quantize(Decimal('0.01')):,.2f}"


def _dec_qty(v: Decimal | None) -> str:
    if v is None:
        return ""
    return f"{v.quantize(Decimal('0.01')):,.2f}"


def _dec_price(v: Decimal | None) -> str:
    if v is None:
        return ""
    return f"{v.quantize(Decimal('0.01')):,.2f}"


def _thickness_text(pod: PoDetail) -> str:
    from app.services.thickness_format import display_po_thickness_mm

    return display_po_thickness_mm(pod)


def _pi_category_label(ph: PoHeader, details: list[PoDetail]) -> str:
    if ph.order_type_label and str(ph.order_type_label).strip():
        return str(ph.order_type_label).strip().upper().replace(" ", "-")
    cat = (ph.po_category or "production").lower()
    mat = (details[0].material_name or "").lower() if details else ""
    if cat == "suede" or (cat == "production" and "suede" in mat):
        return "PROD-SUEDE"
    mapping = {
        "production": "PROD",
        "sample": "SAMPLE",
        "claim": "CLAIM",
        "return": "RETURN",
        "subcont": "SUBCON",
        "suede": "PROD-SUEDE",
    }
    return mapping.get(cat, cat.upper())


def _pi_sequence(db: Session, ph: PoHeader) -> int:
    od = ph.order_date or date.today()
    month_start = od.replace(day=1)
    if od.month == 12:
        month_end = date(od.year + 1, 1, 1)
    else:
        month_end = date(od.year, od.month + 1, 1)
    buyer = (ph.buyer or "").strip()
    q = select(func.count()).select_from(PoHeader).where(
        PoHeader.id <= ph.id,
        PoHeader.order_date >= month_start,
        PoHeader.order_date < month_end,
    )
    if buyer:
        q = q.where(PoHeader.buyer == buyer)
    return int(db.scalar(q) or 1)


def _build_pi_number(db: Session, ph: PoHeader, details: list[PoDetail]) -> str:
    od = ph.order_date or date.today()
    buyer = ((ph.buyer or ph.customer_code or "XX").strip()[:12]).upper()
    cat = _pi_category_label(ph, details)
    seq = _pi_sequence(db, ph)
    roman = ROMAN_MONTHS[od.month - 1]
    return f"{seq:03d}/PI/{buyer}/{cat}/{roman}/{od.year}"


def _resolve_salesperson(db: Session, customer_id: int | None) -> str:
    return "MARKETING"


def _resolve_attn(db: Session, cust: Customer | None) -> str:
    if not cust:
        return ""
    if cust.contact_person and str(cust.contact_person).strip():
        return str(cust.contact_person).strip().upper()
    pic = db.scalar(
        select(CustomerPic.name)
        .where(CustomerPic.customer_id == cust.id)
        .where(CustomerPic.name.isnot(None))
        .order_by(CustomerPic.id.asc())
        .limit(1)
    )
    return (pic or "").strip().upper()


def _resolve_payment_terms(cust: Customer | None, db: Session, customer_id: int | None) -> str:
    if cust and cust.payment_terms and str(cust.payment_terms).strip():
        return str(cust.payment_terms).strip().upper()
    return DEFAULT_PAYMENT_TERMS


def gather_proforma_invoice_ctx(
    db: Session,
    po_id: int,
    po_detail_ids: Optional[set[int] | list[int]] = None,
) -> dict[str, Any]:
    ph = db.get(PoHeader, po_id)
    if not ph:
        raise ValueError("po_not_found")
    cust = db.get(Customer, ph.customer_id) if ph.customer_id else None
    details = list(
        db.scalars(select(PoDetail).where(PoDetail.po_id == po_id).order_by(PoDetail.id.asc())).all()
    )
    wanted = {int(x) for x in (po_detail_ids or []) if x is not None}
    if wanted:
        details = [d for d in details if int(d.id) in wanted]
        if not details:
            raise ValueError("details_not_found")
    # PO NO. pada cetakan = No. PO Customer (fallback ke nomor SO internal).
    po_number = (ph.po_customer_ref or ph.po_number or str(ph.id)).strip()
    customer_name = (cust.name if cust and cust.name else ph.customer_name_snapshot or "").strip() or "—"
    ship_addr = (ph.delivery_address or (cust.address if cust else "") or ph.billing_address or "").strip()
    phone = (cust.phone if cust and cust.phone else "").strip()
    salesperson = _resolve_salesperson(db, ph.customer_id)
    attn = _resolve_attn(db, cust)
    payment_terms = _resolve_payment_terms(cust, db, ph.customer_id)

    lines: list[dict[str, Any]] = []
    total_qty = Decimal(0)
    total_amount = Decimal(0)
    for d in details:
        qty = d.quantity_order or Decimal(0)
        up = d.unit_price
        amt: Decimal | None = None
        if up is not None:
            amt = (qty * up).quantize(Decimal("0.01"))
            total_amount += amt
        total_qty += qty
        etd_src = d.line_etd or ph.etd
        lines.append(
            {
                "po_detail_id": int(d.id),
                "po_no": po_number,
                "item": (d.material_name or "").strip(),
                "thickness": _thickness_text(d),
                "emboss": (d.emboss or "").strip(),
                "color": (d.color or "").strip(),
                "qty": qty,
                "unit": DEFAULT_UNIT,
                "unit_price": up,
                "amount": amt,
                "etd": _fmt_etd(etd_src),
            }
        )

    pi_number = _build_pi_number(db, ph, details)
    po_list_label = po_number if len(lines) <= 1 else "AS LISTED BELLOW"
    ex_factory_label = "AS LISTED BELLOW" if lines else po_number

    return {
        "po_number": po_number,
        "pi_number": pi_number,
        "document_date": ph.order_date,
        "document_date_fmt": _fmt_date(ph.order_date),
        "customer_name": customer_name,
        "ship_address": ship_addr,
        "phone": phone,
        "attn": attn,
        "from_name": salesperson,
        "signature_name": salesperson,
        "po_list_label": po_list_label,
        "ex_factory_label": ex_factory_label,
        "payment_terms": payment_terms,
        "lines": lines,
        "total_qty": total_qty,
        "total_amount": total_amount,
        "unit": DEFAULT_UNIT,
        "po_detail_ids": [int(d.id) for d in details],
        "po_ids": [int(po_id)],
        "customer_id": ph.customer_id,
    }


def gather_proforma_invoice_ctx_multi(
    db: Session,
    po_ids: list[int],
    po_detail_ids: Optional[list[int]] = None,
) -> dict[str, Any]:
    """Gabung beberapa SO jadi 1 PI — wajib customer yang sama.

    Jika `po_detail_ids` diisi, hanya baris itu yang masuk dokumen (sesuai checklist outstanding).
    """
    ids = sorted({int(x) for x in po_ids if x is not None})
    wanted = {int(x) for x in (po_detail_ids or []) if x is not None} or None
    if not ids:
        raise ValueError("po_ids_required")
    if len(ids) == 1:
        return gather_proforma_invoice_ctx(db, ids[0], wanted)

    headers: list[PoHeader] = []
    for pid in ids:
        ph = db.get(PoHeader, pid)
        if not ph:
            raise ValueError("po_not_found")
        headers.append(ph)

    customer_ids = {(h.customer_id or 0) for h in headers}
    if len(customer_ids) > 1 or (0 in customer_ids and len(headers) > 1):
        raise ValueError("customer_mismatch")

    all_lines: list[dict[str, Any]] = []
    total_qty = Decimal(0)
    total_amount = Decimal(0)
    po_numbers: list[str] = []
    included_detail_ids: list[int] = []
    base: dict[str, Any] | None = None

    for h in headers:
        try:
            ctx = gather_proforma_invoice_ctx(db, h.id, wanted)
        except ValueError as e:
            if str(e) == "details_not_found" and wanted is not None:
                # SO ini tidak punya baris tercentang — lewati
                continue
            raise
        if base is None:
            base = dict(ctx)
        po_numbers.append(str(ctx["po_number"]))
        for ln in ctx["lines"]:
            all_lines.append(ln)
        total_qty += ctx["total_qty"] or Decimal(0)
        total_amount += ctx["total_amount"] or Decimal(0)
        included_detail_ids.extend(ctx.get("po_detail_ids") or [])

    if base is None or (wanted is not None and not all_lines):
        raise ValueError("details_not_found")

    po_label = ", ".join(po_numbers[:4])
    if len(po_numbers) > 4:
        po_label = f"{po_label} +{len(po_numbers) - 4}"

    base["po_number"] = po_label
    primary = headers[0]
    primary_details = list(
        db.scalars(select(PoDetail).where(PoDetail.po_id == primary.id).order_by(PoDetail.id.asc())).all()
    )
    if wanted:
        filtered = [d for d in primary_details if int(d.id) in wanted]
        if filtered:
            primary_details = filtered
    base["pi_number"] = _ensure_unique_pi_number(db, _build_pi_number(db, primary, primary_details))
    base["po_list_label"] = "AS LISTED BELLOW"
    base["ex_factory_label"] = "AS LISTED BELLOW"
    base["lines"] = all_lines
    base["total_qty"] = total_qty
    base["total_amount"] = total_amount
    base["po_ids"] = ids
    base["po_detail_ids"] = included_detail_ids
    base["customer_id"] = primary.customer_id
    return base


def _ensure_unique_pi_number(db: Session, pi_number: str) -> str:
    from app.models import ProformaInvoiceRecord

    base = (pi_number or "").strip() or "001/PI/XX/PROD/I/2026"
    candidate = base
    n = 1
    while n < 100:
        row = db.scalars(
            select(ProformaInvoiceRecord).where(ProformaInvoiceRecord.pi_number == candidate).limit(1)
        ).first()
        if not row:
            return candidate
        n += 1
        candidate = f"{base}-R{n}"
    return candidate


def persist_proforma_invoice_record(
    db: Session,
    ctx: dict[str, Any],
    *,
    actor: str | None = None,
) -> Any:
    """Simpan nomor PI resmi setelah generate dokumen."""
    from app.models import ProformaInvoiceRecord

    pi_number = _ensure_unique_pi_number(db, str(ctx.get("pi_number") or ""))
    ctx["pi_number"] = pi_number
    row = ProformaInvoiceRecord(
        pi_number=pi_number,
        customer_id=ctx.get("customer_id"),
        customer_name=str(ctx.get("customer_name") or "")[:255] or None,
        po_ids_json=list(ctx.get("po_ids") or []),
        po_detail_ids_json=list(ctx.get("po_detail_ids") or []),
        total_qty=ctx.get("total_qty"),
        total_amount=ctx.get("total_amount"),
        created_by=actor,
    )
    db.add(row)
    db.flush()
    return row


def _terms_paragraphs(ctx: dict[str, Any], style: ParagraphStyle) -> list[Any]:
    bank_block = "<br/>".join(escape(ln) for ln in BANK_LINES)
    terms = [
        f"1. <b>TERMS OF PAYMENT:</b> {escape(ctx['payment_terms'])}",
        f"2. <b>OUR ADVISING BANK:</b> {bank_block}",
        "3. <b>BENEFICIARY:</b> PT. TONG HONG TANNERY INDONESIA",
        f"4. <b>Ex-Factory DATED:</b> {escape(ctx['ex_factory_label'])}",
        "5. <b>PACKING:</b> BY ROLL",
        "6. <b>SHIPPING METHOD:</b> BY BOX CAR",
        "7. <b>TRADE TERM:</b> FRANCO FACTORY",
        "8. <b>DETAILS OF MATERIALS AS BELOW:</b>",
    ]
    return [Paragraph(t, style) for t in terms]


def render_proforma_invoice_pdf(ctx: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    center = ParagraphStyle(name="center", parent=styles["Normal"], alignment=1, fontSize=10, leading=12)
    center_bold = ParagraphStyle(name="center_b", parent=center, fontSize=11, leading=13)
    small = ParagraphStyle(name="pi_s", parent=styles["Normal"], fontSize=8.5, leading=11)
    cell = ParagraphStyle(name="pi_cell", parent=styles["Normal"], fontSize=7, leading=8.5)
    cell_r = ParagraphStyle(name="pi_cell_r", parent=cell, alignment=2)
    cell_c = ParagraphStyle(name="pi_cell_c", parent=cell, alignment=1)
    head_style = ParagraphStyle(
        name="pi_head",
        parent=styles["Normal"],
        fontSize=6.5,
        leading=8,
        alignment=1,
        fontName="Helvetica-Bold",
    )
    story: list[Any] = []

    for i, ln in enumerate(COMPANY_LINES):
        st = center_bold if i == 0 else center
        story.append(Paragraph(f"<b>{escape(ln)}</b>" if i == 0 else escape(ln), st))
    story.append(Spacer(1, 3 * mm))

    divider = Table([[""]], colWidths=[doc.width], rowHeights=[2])
    divider.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 2, colors.black)]))
    story.append(divider)
    story.append(Spacer(1, 4 * mm))

    left_block = [
        f"<b>TO:</b> {escape(ctx['customer_name'])}",
        f"<b>ADD:</b> {escape(ctx['ship_address'] or '—')}",
        f"<b>TELP.:</b> {escape(ctx['phone'] or '—')}",
        f"<b>ATTN:</b> {escape(ctx['attn'] or '—')}",
    ]
    right_block = [
        f"<b>FROM:</b> {escape(ctx['from_name'])}",
        f"<b>DATE:</b> {escape(ctx['document_date_fmt'])}",
        f"<b>PO NO.:</b> {escape(ctx['po_list_label'])}",
    ]
    addr_tbl = Table(
        [
            [
                Paragraph("<br/>".join(left_block), small),
                Paragraph("<br/>".join(right_block), small),
            ]
        ],
        colWidths=[doc.width * 0.58, doc.width * 0.42],
    )
    addr_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(addr_tbl)
    story.append(Spacer(1, 5 * mm))

    story.append(Paragraph("<b>PROFORMA INVOICE</b>", ParagraphStyle(name="title", parent=center_bold, fontSize=12)))
    story.append(Paragraph(f"<b>NO. {escape(ctx['pi_number'])}</b>", center_bold))
    story.append(Spacer(1, 4 * mm))

    for p in _terms_paragraphs(ctx, small):
        story.append(p)
    story.append(Spacer(1, 3 * mm))

    head_labels = [
        "PO NO",
        "ITEM",
        "THICKNESS",
        "EMBOSS",
        "COLOR",
        "QTY",
        "UNIT",
        "UNIT PRICE<br/>(USD)",
        "AMOUNT<br/>(USD)",
        "ETD",
    ]
    data: list[list[Any]] = [[Paragraph(h, head_style) for h in head_labels]]
    for ln in ctx["lines"]:
        data.append(
            [
                Paragraph(escape(str(ln["po_no"] or "")), cell),
                Paragraph(escape(str(ln["item"] or "")), cell),
                Paragraph(escape(str(ln["thickness"] or "")), cell_c),
                Paragraph(escape(str(ln["emboss"] or "")), cell),
                Paragraph(escape(str(ln["color"] or "")), cell),
                Paragraph(escape(_dec_qty(ln["qty"])), cell_r),
                Paragraph(escape(str(ln["unit"] or "")), cell_c),
                Paragraph(escape(_dec_price(ln["unit_price"])), cell_r),
                Paragraph(escape(_dec_money(ln["amount"])), cell_r),
                Paragraph(escape(str(ln["etd"] or "")), cell_c),
            ]
        )
    data.append(
        [
            Paragraph("<b>TOTAL</b>", cell),
            "",
            "",
            "",
            "",
            Paragraph(f"<b>{escape(_dec_qty(ctx['total_qty']))}</b>", cell_r),
            Paragraph(escape(str(ctx.get("unit") or "")), cell_c),
            "",
            Paragraph(f"<b>{escape(_dec_money(ctx['total_amount']))}</b>", cell_r),
            "",
        ]
    )

    # Lebar disetel agar COLOR/QTY tidak saling tabrak; teks wrap via Paragraph.
    col_widths = [
        doc.width * 0.11,  # PO
        doc.width * 0.13,  # ITEM
        doc.width * 0.08,  # THICKNESS
        doc.width * 0.08,  # EMBOSS
        doc.width * 0.16,  # COLOR (sering panjang)
        doc.width * 0.09,  # QTY
        doc.width * 0.05,  # UNIT
        doc.width * 0.10,  # PRICE
        doc.width * 0.11,  # AMOUNT
        doc.width * 0.09,  # ETD
    ]
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.90, 0.90, 0.90)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("SPAN", (0, -1), (4, -1)),
                ("BACKGROUND", (0, -1), (-1, -1), colors.Color(0.96, 0.96, 0.96)),
            ]
        )
    )
    story.append(tbl)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("BEST REGARDS,", small))
    story.append(Paragraph("Yours Truly,", small))
    story.append(Spacer(1, 12 * mm))
    story.append(Paragraph(f"<b>{escape(ctx['signature_name'])}</b>", small))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def render_proforma_invoice_docx(ctx: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = Document()
    for i, ln in enumerate(COMPANY_LINES):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(ln)
        r.bold = i == 0
        r.font.size = Pt(10 if i else 11)

    d.add_paragraph()

    left = d.add_table(rows=1, cols=2)
    left.rows[0].cells[0].text = "\n".join(
        [
            f"TO: {ctx['customer_name']}",
            f"ADD: {ctx['ship_address'] or '—'}",
            f"TELP.: {ctx['phone'] or '—'}",
            f"ATTN: {ctx['attn'] or '—'}",
        ]
    )
    left.rows[0].cells[1].text = "\n".join(
        [
            f"FROM: {ctx['from_name']}",
            f"DATE: {ctx['document_date_fmt']}",
            f"PO NO.: {ctx['po_list_label']}",
        ]
    )

    title = d.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("PROFORMA INVOICE")
    tr.bold = True
    tr.font.size = Pt(13)

    no_p = d.add_paragraph()
    no_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nr = no_p.add_run(f"NO. {ctx['pi_number']}")
    nr.bold = True

    d.add_paragraph(f"1. TERMS OF PAYMENT: {ctx['payment_terms']}")
    p2 = d.add_paragraph("2. OUR ADVISING BANK:")
    for bl in BANK_LINES:
        p2.add_run(f"\n{bl}")
    d.add_paragraph("3. BENEFICIARY: PT. TONG HONG TANNERY INDONESIA")
    d.add_paragraph(f"4. Ex-Factory DATED: {ctx['ex_factory_label']}")
    d.add_paragraph("5. PACKING: BY ROLL")
    d.add_paragraph("6. SHIPPING METHOD: BY BOX CAR")
    d.add_paragraph("7. TRADE TERM: FRANCO FACTORY")
    d.add_paragraph("8. DETAILS OF MATERIALS AS BELOW:")

    cols = ["PO NO", "ITEM", "THICKNESS", "EMBOSS", "COLOR", "QTY", "UNIT", "PRICE (USD)", "AMOUNT (USD)", "ETD"]
    table = d.add_table(rows=1 + len(ctx["lines"]) + 1, cols=len(cols))
    table.style = "Table Grid"
    for j, h in enumerate(cols):
        table.rows[0].cells[j].text = h
    for i, ln in enumerate(ctx["lines"], start=1):
        row = table.rows[i].cells
        row[0].text = ln["po_no"]
        row[1].text = ln["item"]
        row[2].text = ln["thickness"]
        row[3].text = ln["emboss"]
        row[4].text = ln["color"]
        row[5].text = _dec_qty(ln["qty"])
        row[6].text = ln["unit"]
        row[7].text = _dec_price(ln["unit_price"])
        row[8].text = _dec_money(ln["amount"])
        row[9].text = ln["etd"]
    last = table.rows[len(ctx["lines"]) + 1].cells
    last[0].text = "TOTAL"
    last[5].text = _dec_qty(ctx["total_qty"])
    last[6].text = ctx["unit"]
    last[8].text = _dec_money(ctx["total_amount"])

    d.add_paragraph()
    d.add_paragraph("BEST REGARDS,")
    d.add_paragraph("Yours Truly,")
    d.add_paragraph()
    sig = d.add_paragraph()
    sr = sig.add_run(ctx["signature_name"])
    sr.bold = True

    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf.getvalue()
