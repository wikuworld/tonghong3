"""Export/import Excel master Color Window.

Kunci unik impor: brand_code + color_code. Baris yang sudah ada diperbarui.
File yang diunggah ulang tidak menambah duplikat.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Brand, ColorWindowMaster
from app.services.color_window import normalize_color_window_fields
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "color_code",
    "color_name",
    "brand_code",
    "remark",
    "created_at",
    "updated_at",
]

IMPORT_FIELDS = ("color_code", "color_name", "brand_code", "remark")

_HEADER_ALIAS = {
    "id": "id",
    "color_code": "color_code",
    "kode": "color_code",
    "code": "color_code",
    "color_name": "color_name",
    "color": "color_name",
    "warna": "color_name",
    "brand_code": "brand_code",
    "brand": "brand_code",
    "remark": "remark",
    "catatan": "remark",
    "created_at": "created_at",
    "updated_at": "updated_at",
}


@dataclass
class ImportErrorRow:
    row: int
    color_code: str | None
    reason: str


@dataclass
class ColorWindowImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(e) for e in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    s = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in s:
        s = s.replace("__", "_")
    return s


def _cell_str(val: Any) -> str:
    if val is None or isinstance(val, bool):
        return "" if val is None else str(val).strip()
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, (int, float)) and 20000 < float(val) < 80000:
        stamp = datetime(1899, 12, 30) + timedelta(days=float(val))
        return stamp.strftime("%Y-%m-%d %H:%M:%S")
    return str(val).strip()


def _business_key(brand: str, code: str) -> tuple[str, str]:
    return ((brand or "").strip().upper(), (code or "").strip().upper())


def _parse_dt(raw: str) -> datetime | None:
    text = (raw or "").strip().replace("T", " ")
    if not text:
        return None
    if len(text) > 19:
        text = text[:19]
    for fmt, size in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%d/%m/%Y", 10),
    ):
        try:
            return datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
    return None


def _format_dt(value: datetime | None) -> str:
    if not value:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _workbook_bytes(ws_title: str, headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ws_title
    for col_idx, col in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=col)
    for row_idx, values in enumerate(data_rows, 2):
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_color_windows_xlsx(db: Session) -> bytes:
    rows = list(
        db.scalars(
            select(ColorWindowMaster).order_by(
                ColorWindowMaster.brand_code.asc(),
                ColorWindowMaster.color_code.asc(),
                ColorWindowMaster.id.asc(),
            )
        ).all()
    )
    data = [
        [
            row.id,
            row.color_code or "",
            row.color_name or "",
            row.brand_code or "",
            row.remark or "",
            _format_dt(row.created_at),
            _format_dt(row.updated_at),
        ]
        for row in rows
    ]
    return _workbook_bytes("Color Window", EXPORT_COLUMNS, data)


def export_color_windows_template_xlsx() -> bytes:
    return _workbook_bytes("Color Window", ["color_code", "color_name", "brand_code", "remark"], [])


def _parse_rows(content: bytes) -> tuple[list[dict[str, str]], list[ImportErrorRow]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel tidak bisa dibaca")]
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel kosong")]

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row or []):
        key = _HEADER_ALIAS.get(_normalize_header(cell))
        if key and key not in col_map.values():
            col_map[idx] = key

    missing = [name for name in ("color_code", "color_name", "brand_code") if name not in col_map.values()]
    if missing:
        return [], [
            ImportErrorRow(
                row=1,
                color_code=None,
                reason="Header wajib memuat kolom color_code, color_name, dan brand_code",
            )
        ]

    parsed: list[dict[str, str]] = []
    for excel_row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {key: "" for key in ("id", *IMPORT_FIELDS, "created_at", "updated_at")}
        for idx, key in col_map.items():
            if idx < len(row) and key in item:
                item[key] = _cell_str(row[idx])
        if not any(item.get(k) for k in ("id", *IMPORT_FIELDS)):
            continue
        item["_row"] = str(excel_row_num)
        item["_present"] = ",".join(col_map.values())
        parsed.append(item)
    return parsed, []


def _parse_id(raw: str) -> int | None:
    text = (raw or "").strip()
    if not text:
        return None
    try:
        value = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None
    return value if value > 0 else None


def import_color_windows_xlsx(db: Session, content: bytes) -> ColorWindowImportResult:
    result = ColorWindowImportResult()
    rows, parse_errors = _parse_rows(content)
    result.errors.extend(parse_errors)
    if parse_errors and not rows:
        return result

    existing = list(db.scalars(select(ColorWindowMaster)).all())
    by_id = {row.id: row for row in existing if row.id is not None}
    by_key: dict[tuple[str, str], ColorWindowMaster] = {}
    for row in sorted(existing, key=lambda item: item.id or 0):
        by_key.setdefault(_business_key(row.brand_code or "", row.color_code or ""), row)
    brand_codes = {
        str(code).strip().upper()
        for code in db.scalars(select(Brand.code)).all()
        if code and str(code).strip()
    }
    seen_ids: set[int] = set()

    for item in rows:
        result.total_rows += 1
        excel_row = int(item.pop("_row", "0") or 0)
        present = {part for part in (item.pop("_present", "") or "").split(",") if part}
        code_label = _cell_str(item.get("color_code")) or None
        raw_id = _cell_str(item.get("id"))
        row_id = _parse_id(raw_id)
        if raw_id and row_id is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=code_label, reason="id tidak valid")
            )
            continue
        if row_id is not None and row_id in seen_ids:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason=f"id {row_id} muncul lebih dari sekali di file",
                )
            )
            continue

        try:
            data = normalize_color_window_fields(
                {
                    "color_code": item.get("color_code"),
                    "color_name": item.get("color_name"),
                    "brand_code": item.get("brand_code"),
                    "remark": item.get("remark"),
                }
            )
        except ValueError as exc:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=code_label, reason=str(exc)))
            continue

        if data["brand_code"] not in brand_codes:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=data["color_code"],
                    reason=f"Brand '{data['brand_code']}' tidak ditemukan di master Brand",
                )
            )
            continue

        key = _business_key(data["brand_code"], data["color_code"])
        owner = by_key.get(key)
        target = by_id.get(row_id) if row_id is not None else None
        if owner is not None and target is not None and owner is not target:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=data["color_code"],
                    reason="brand + color_code sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue
        if target is None:
            target = owner
        if row_id is not None:
            seen_ids.add(row_id)

        created_at = _parse_dt(item.get("created_at") or "") if "created_at" in present else None
        updated_at = _parse_dt(item.get("updated_at") or "") if "updated_at" in present else None
        if "created_at" in present and item.get("created_at") and created_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=data["color_code"], reason="created_at tidak valid")
            )
            continue
        if "updated_at" in present and item.get("updated_at") and updated_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=data["color_code"], reason="updated_at tidak valid")
            )
            continue

        if target is None:
            created = ColorWindowMaster(
                color_code=data["color_code"],
                color_name=data["color_name"],
                brand_code=data["brand_code"],
                remark=data["remark"],
                created_at=created_at or app_now(),
                updated_at=updated_at,
            )
            db.add(created)
            by_key[key] = created
            result.created += 1
            continue

        old_key = _business_key(target.brand_code or "", target.color_code or "")
        if old_key != key and by_key.get(key) not in (None, target):
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=data["color_code"],
                    reason="brand + color_code sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue

        same_fields = all(getattr(target, field_name) == data[field_name] for field_name in IMPORT_FIELDS)
        same_created = created_at is None or _format_dt(target.created_at) == _format_dt(created_at)
        same_updated = updated_at is None or _format_dt(target.updated_at) == _format_dt(updated_at)
        if same_fields and same_created and same_updated:
            result.skipped += 1
            continue

        for field_name in IMPORT_FIELDS:
            setattr(target, field_name, data[field_name])
        if created_at is not None:
            target.created_at = created_at
        target.updated_at = updated_at if updated_at is not None else app_now()
        if old_key != key and by_key.get(old_key) is target:
            by_key.pop(old_key, None)
        by_key[key] = target
        result.updated += 1

    return result
