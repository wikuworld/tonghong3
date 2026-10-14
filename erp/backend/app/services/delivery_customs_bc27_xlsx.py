"""Export Excel bea cukai BC 2.7 (pengeluaran KB/TPB) dari satu/lebih delivery.

Format mengacu `dokumentasi/000027TOP41720261002004407 (1).xlsx`.
Label merah = mapping isi dari ERP; sel kuning = dikosongkan.
Multi-SJ wajib satu customer; digabung jadi satu workbook AJU.
"""

from __future__ import annotations

import logging
from copy import copy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, Sequence

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    Customer,
    DeliveryDetail,
    DeliveryHeader,
    ExchangeRate,
    Invoice,
    InvoiceDetail,
    MasterItem,
    PackingDetail,
    PackingHeader,
    PoDetail,
    ProductionDetail,
)
from app.services.customs_numbering import (
    generate_nomor_aju,
    generate_nomor_kontrak,
    normalize_aju_seq,
)
from app.timeutil import app_now

logger = logging.getLogger(__name__)

# Identitas perusahaan (isi tetap dari template contoh AJU).
TH_NAME = "TONG HONG TANNERY INDONESIA"
TH_ADDRESS = (
    "KAWASAN INDUSTRI MODERN CIKANDE JL MODERN INDUSTRI XIV, NO. W1, "
    "BARENGKOK, KIBIN, KAB. SERANG, BANTEN, 42185"
)
TH_NPWP = "0211472683401000000000"
TH_NIB = "8120104992265"
TH_SKEP = "KM-86/WBC.07/2019"
TH_SKEP_DATE = "04-16-2019 00:00:00"
TH_KODE_KANTOR = "050400"
TH_KOTA_PERNYATAAN = "SERANG"
TH_NAMA_PERNYATAAN = "CHIU PO HAO"
TH_JABATAN_PERNYATAAN = "DIREKTUR"
TH_KODE_JENIS_API = "02"
TH_KODE_STATUS = "LAINNYA"

KODE_DOKUMEN_BC27 = "27"
KODE_DOK_INVOICE = "380"
KODE_DOK_SJ = "630"
KODE_DOK_KONTRAK = "315"
KODE_SATUAN_FTK = "FTK"
KODE_KEMASAN_RO = "RO"
VERSI_TEMPLATE = "1.3"

# Kolom kuning di template anotasi — jangan diisi.
BAHANBAKU_YELLOW_KEYS = frozenset(
    {
        "NOMOR DAFTAR ASAL",
        "TANGGAL DAFTAR ASAL",
        "NOMOR AJU ASAL",
        "SERI BARANG ASAL",
        "NETTO",
        "BRUTO",
        "CIF",
        "CIF RUPIAH",
    }
)
BAHANBAKUTARIF_YELLOW_KEYS = frozenset({"NILAI FASILITAS"})

SHEET_ORDER = [
    "HEADER",
    "ENTITAS",
    "DOKUMEN",
    "PENGANGKUT",
    "KEMASAN",
    "KONTAINER",
    "KOMPONENBIAYA",
    "BARANG",
    "BARANGTARIF",
    "BARANGDOKUMEN",
    "BARANGENTITAS",
    "BARANGSPEKKHUSUS",
    "BARANGVD",
    "BAHANBAKU",
    "BAHANBAKUTARIF",
    "BAHANBAKUDOKUMEN",
    "PUNGUTAN",
    "JAMINAN",
    "BANKDEVISA",
    "VERSI",
    "RESPON",
]

_TEMPLATE_NAMES = (
    "000027TOP41720261002004407 (1).xlsx",
    "000027TOP41720261002004407.xlsx",
)


def _template_candidates() -> list[Path]:
    here = Path(__file__).resolve()
    roots = [
        here.parents[3] / "dokumentasi",
        here.parents[2] / "dokumentasi",
        Path("/app/dokumentasi"),
        Path("/home/developer/project_uat/dokumentasi"),
    ]
    out: list[Path] = []
    for root in roots:
        for name in _TEMPLATE_NAMES:
            out.append(root / name)
    return out


def resolve_bc27_template_path() -> Path | None:
    for p in _template_candidates():
        if p.is_file():
            return p
    return None


def _dec(val: Any) -> Decimal:
    if val is None or val == "":
        return Decimal("0")
    if isinstance(val, Decimal):
        return val
    return Decimal(str(val))


def _money2(val: Decimal) -> Decimal:
    return val.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _fmt_date(d: date | datetime | None) -> str | None:
    if d is None:
        return None
    if isinstance(d, datetime):
        d = d.date()
    return d.isoformat()


def _num(val: Decimal | int | float | None) -> float | int | None:
    if val is None:
        return None
    if isinstance(val, Decimal):
        if val == val.to_integral_value():
            return int(val)
        return float(val)
    return val


def _text(val: Any) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s or None


def _header_map(ws: Worksheet) -> dict[str, list[int]]:
    """Kolom 1-based per nama header (dukung nama duplikat)."""
    out: dict[str, list[int]] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        name = ws.cell(1, c).value
        if name is None:
            continue
        key = str(name).strip()
        if not key:
            continue
        out.setdefault(key, []).append(c)
    return out


def _write_row(ws: Worksheet, row_idx: int, headers: dict[str, list[int]], values: dict[str, Any]) -> None:
    used: dict[str, int] = {}
    for key, val in values.items():
        cols = headers.get(key) or []
        if not cols:
            continue
        occ = used.get(key, 0)
        if occ >= len(cols):
            continue
        ws.cell(row_idx, cols[occ], val)
        used[key] = occ + 1


def _append_row(ws: Worksheet, headers: dict[str, list[int]], values: dict[str, Any]) -> None:
    _write_row(ws, (ws.max_row or 1) + 1, headers, values)


def _copy_sheet_headers(src: Worksheet, dst: Worksheet) -> None:
    for c in range(1, (src.max_column or 0) + 1):
        src_cell = src.cell(1, c)
        dst_cell = dst.cell(1, c)
        dst_cell.value = src_cell.value
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.border = copy(src_cell.border)
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.number_format = src_cell.number_format


def _blank_workbook_from_template(template: Path) -> Any:
    src = load_workbook(template)
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name in SHEET_ORDER:
        if name not in src.sheetnames:
            wb.create_sheet(name)
            continue
        src_ws = src[name]
        ws = wb.create_sheet(name)
        _copy_sheet_headers(src_ws, ws)
    for name in src.sheetnames:
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            _copy_sheet_headers(src[name], ws)
    src.close()
    return wb


def _latest_kurs(db: Session, on_date: date | None = None) -> Decimal | None:
    q = select(ExchangeRate).order_by(ExchangeRate.effective_date.desc())
    rows = list(db.scalars(q).all())
    if not rows:
        return None
    if on_date is None:
        return _dec(rows[0].idr_per_usd)
    for r in rows:
        if r.effective_date and r.effective_date <= on_date:
            return _dec(r.idr_per_usd)
    return _dec(rows[0].idr_per_usd)


def _item_uraian(pod: PoDetail | None) -> str | None:
    if not pod:
        return None
    from app.services.thickness_format import display_po_thickness_mm

    thick = display_po_thickness_mm(pod)
    rest = " ".join(
        p
        for p in [
            (pod.material_name or "").strip(),
            (pod.emboss or "").strip(),
            (pod.color or "").strip(),
        ]
        if p
    )
    if thick and rest:
        return f"{thick} {rest}".strip()
    return rest or thick or None


def _mrp_base() -> str:
    return (settings.mrp_upstream_url or "http://127.0.0.1:8002").rstrip("/")


def _mrp_timeout() -> float:
    return max(float(settings.mrp_sync_timeout_seconds or 8.0), 20.0)


def fetch_mrp_packing_weights(mrp_packing_id: int) -> tuple[Decimal | None, Decimal | None]:
    """Ambil (gross_weight, net_weight) dari packing MRP. Gagal/network → (None, None)."""
    if not mrp_packing_id:
        return None, None
    base = _mrp_base()
    try:
        with httpx.Client(timeout=_mrp_timeout()) as client:
            r = client.get(f"{base}/api/mrp/packing/{int(mrp_packing_id)}")
            if r.status_code != 200:
                logger.warning("MRP packing/%s weights HTTP %s", mrp_packing_id, r.status_code)
                return None, None
            body = r.json() if r.content else {}
            packing = body.get("packing") if isinstance(body, dict) else None
            if not isinstance(packing, dict):
                packing = body if isinstance(body, dict) else {}
            gw_raw = packing.get("gross_weight")
            nw_raw = packing.get("net_weight")
            gw = _dec(gw_raw) if gw_raw not in (None, "") else None
            nw = _dec(nw_raw) if nw_raw not in (None, "") else None
            if gw is not None and gw <= 0:
                gw = None
            if nw is not None and nw <= 0:
                nw = None
            return gw, nw
    except httpx.HTTPError as exc:
        logger.warning("MRP packing/%s weights unavailable: %s", mrp_packing_id, exc)
        return None, None
    except Exception as exc:  # noqa: BLE001 — jangan gagalkan export Excel
        logger.warning("MRP packing/%s weights error: %s", mrp_packing_id, exc)
        return None, None


def _invoice_weight_totals(db: Session, delivery_id: int) -> tuple[Decimal | None, Decimal | None]:
    """Σ net/gross dari invoice_details untuk satu SJ (jika terisi)."""
    inv_ids = list(
        db.scalars(select(Invoice.id).where(Invoice.delivery_id == int(delivery_id))).all()
    )
    if not inv_ids:
        return None, None
    details = list(
        db.scalars(select(InvoiceDetail).where(InvoiceDetail.invoice_id.in_(inv_ids))).all()
    )
    if not details:
        return None, None
    nw = sum((_dec(d.net_weight) for d in details if d.net_weight is not None), Decimal("0"))
    gw = sum((_dec(d.gross_weight) for d in details if d.gross_weight is not None), Decimal("0"))
    return (nw if nw > 0 else None), (gw if gw > 0 else None)


def _mrp_packing_ids_for_headers(db: Session, headers: Sequence[DeliveryHeader]) -> list[int]:
    """Unique mrp_packing_id dari packing ERP yang tertaut SJ."""
    seen: set[int] = set()
    out: list[int] = []
    for dh in headers:
        if not dh.packing_id:
            continue
        pk = db.get(PackingHeader, int(dh.packing_id))
        if not pk or not pk.mrp_packing_id:
            continue
        mid = int(pk.mrp_packing_id)
        if mid in seen:
            continue
        seen.add(mid)
        out.append(mid)
    return out


def resolve_bruto_netto(
    db: Session,
    headers: Sequence[DeliveryHeader],
    lines: Sequence[dict[str, Any]],
) -> tuple[Decimal, Decimal]:
    """NETTO dari baris delivery (+invoice fallback); BRUTO dari rasio packing MRP / invoice GW.

    Prioritas BRUTO:
    1) Rasio Σ gross/net packing MRP (unik) × NETTO
    2) Rasio / total gross dari invoice_details
    3) Fallback BRUTO = NETTO
    """
    line_nw = sum((ln["nw"] or Decimal("0") for ln in lines), Decimal("0"))

    inv_nw_total = Decimal("0")
    inv_gw_total = Decimal("0")
    has_inv_nw = False
    has_inv_gw = False
    for dh in headers:
        inv_nw, inv_gw = _invoice_weight_totals(db, int(dh.id))
        if inv_nw is not None:
            inv_nw_total += inv_nw
            has_inv_nw = True
        if inv_gw is not None:
            inv_gw_total += inv_gw
            has_inv_gw = True

    total_nw = line_nw if line_nw > 0 else (inv_nw_total if has_inv_nw else Decimal("0"))

    mrp_gw_sum = Decimal("0")
    mrp_nw_sum = Decimal("0")
    has_mrp_gw = False
    has_mrp_nw = False
    for mid in _mrp_packing_ids_for_headers(db, headers):
        gw, nw = fetch_mrp_packing_weights(mid)
        if gw is not None:
            mrp_gw_sum += gw
            has_mrp_gw = True
        if nw is not None:
            mrp_nw_sum += nw
            has_mrp_nw = True

    total_gw: Decimal | None = None
    if has_mrp_gw and has_mrp_nw and mrp_nw_sum > 0 and total_nw > 0:
        total_gw = _money2(total_nw * (mrp_gw_sum / mrp_nw_sum))
    elif has_mrp_gw and total_nw > 0 and len(headers) == 1:
        # Satu SJ: boleh pakai bruto packing MRP langsung bila netto packing tak ada
        total_gw = mrp_gw_sum
    elif has_inv_gw and has_inv_nw and inv_nw_total > 0 and total_nw > 0:
        total_gw = _money2(total_nw * (inv_gw_total / inv_nw_total))
    elif has_inv_gw:
        total_gw = inv_gw_total

    if total_gw is None or total_gw <= 0:
        total_gw = total_nw

    return total_nw, total_gw


def _gather_lines(db: Session, delivery_id: int, *, seri_start: int = 1) -> list[dict[str, Any]]:
    details = list(
        db.scalars(
            select(DeliveryDetail)
            .where(DeliveryDetail.delivery_id == delivery_id)
            .order_by(DeliveryDetail.id.asc())
        ).all()
    )
    lines: list[dict[str, Any]] = []
    for idx, dd in enumerate(details):
        pod: PoDetail | None = None
        master: MasterItem | None = None
        pd: ProductionDetail | None = None
        if dd.po_detail_id:
            pod = db.get(PoDetail, int(dd.po_detail_id))
        if dd.production_detail_id:
            pd = db.get(ProductionDetail, int(dd.production_detail_id))
            if pod is None and pd and pd.po_detail_id:
                pod = db.get(PoDetail, int(pd.po_detail_id))
        if pod and pod.master_item_id:
            master = db.get(MasterItem, int(pod.master_item_id))

        qty = _dec(dd.qty)
        unit_price = _dec(pod.unit_price) if pod and pod.unit_price is not None else None
        amount = _money2(qty * unit_price) if unit_price is not None else None

        nw = _dec(dd.weight) if dd.weight is not None else None
        if nw is None and dd.packing_detail_id:
            pkd = db.get(PackingDetail, int(dd.packing_detail_id))
            if pkd and pkd.weight_net is not None:
                nw = _dec(pkd.weight_net)

        kode_barang = ""
        if master:
            kode_barang = (master.item_code or master.exim_material_code or "").strip()
        if not kode_barang and pod and pod.system_reference:
            kode_barang = pod.system_reference.strip()

        lines.append(
            {
                "seri": seri_start + idx,
                "delivery_id": delivery_id,
                "uraian": _item_uraian(pod),
                "kode_barang": kode_barang or None,
                "qty": qty,
                "qty_roll": int(dd.qty_roll or 0),
                "qty_pcs": int(dd.qty_pcs or 0),
                "nw": nw,
                "gw": nw,
                "amount": amount,
                "unit_price": unit_price,
            }
        )
    return lines


def _parse_delivery_ids(delivery_ids: Sequence[int] | int) -> list[int]:
    if isinstance(delivery_ids, int):
        ids = [delivery_ids]
    else:
        ids = [int(x) for x in delivery_ids]
    # unik, urut stabil
    seen: set[int] = set()
    out: list[int] = []
    for i in ids:
        if i in seen:
            continue
        seen.add(i)
        out.append(i)
    if not out:
        raise ValueError("delivery_ids_required")
    return out


def build_bc27_workbook(
    db: Session,
    delivery_ids: Sequence[int] | int,
    *,
    aju_seq: str | int,
) -> tuple[Any, str]:
    ids = _parse_delivery_ids(delivery_ids)
    aju_digits = normalize_aju_seq(aju_seq)

    headers: list[DeliveryHeader] = []
    for did in ids:
        dh = db.get(DeliveryHeader, did)
        if not dh:
            raise ValueError("delivery_not_found")
        headers.append(dh)

    customer_ids = {int(h.customer_id) for h in headers if h.customer_id is not None}
    if len(customer_ids) > 1:
        raise ValueError("delivery_customer_mismatch")
    if not customer_ids:
        raise ValueError("delivery_customer_missing")

    customer = db.get(Customer, next(iter(customer_ids)))
    invoices = list(
        db.scalars(
            select(Invoice)
            .where(Invoice.delivery_id.in_(ids))
            .order_by(Invoice.delivery_id.asc(), Invoice.id.asc())
        ).all()
    )

    lines: list[dict[str, Any]] = []
    seri = 1
    for dh in headers:
        chunk = _gather_lines(db, int(dh.id), seri_start=seri)
        lines.extend(chunk)
        seri += len(chunk)

    dates = [h.delivery_date for h in headers if h.delivery_date]
    dday = min(dates) if dates else app_now().date()
    nomor_aju = generate_nomor_aju(dday, aju_digits)
    cust_code = None
    if customer:
        cust_code = _text(customer.code)
    kontrak = generate_nomor_kontrak(cust_code, dday)
    kurs = _latest_kurs(db, dday)

    total_cif = sum((ln["amount"] or Decimal("0") for ln in lines), Decimal("0"))
    total_nw, total_gw = resolve_bruto_netto(db, headers, lines)
    total_roll = sum(int(ln["qty_roll"] or 0) for ln in lines)

    bruto_ratio: Decimal | None = None
    if total_nw > 0 and total_gw > 0:
        bruto_ratio = total_gw / total_nw

    cif_idr = _money2(total_cif * kurs) if kurs is not None else None
    ppn = _money2(cif_idr * Decimal("0.11")) if cif_idr is not None else None

    template = resolve_bc27_template_path()
    if template:
        wb = _blank_workbook_from_template(template)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        for name in SHEET_ORDER:
            wb.create_sheet(name)

    # --- HEADER ---
    ws = wb["HEADER"]
    hmap = _header_map(ws)
    if not hmap:
        for i, col in enumerate(
            [
                "NOMOR AJU",
                "KODE DOKUMEN",
                "KODE KANTOR",
                "KODE KANTOR TUJUAN",
                "KODE JENIS TPB",
                "KODE TUJUAN PENGIRIMAN",
                "KODE TUJUAN TPB",
                "CIF",
                "HARGA_PENYERAHAN",
                "NDPBM",
                "DASAR PENGENAAN PAJAK",
                "BRUTO",
                "NETTO",
                "KOTA PERNYATAAN",
                "TANGGAL PERNYATAAN",
                "NAMA PERNYATAAN",
                "JABATAN PERNYATAAN",
                "KODE VALUTA",
                "NOMOR DAFTAR",
                "TANGGAL DAFTAR",
                "PPN PAJAK",
                "PPNBM PAJAK",
                "TARIF PPN PAJAK",
            ],
            1,
        ):
            ws.cell(1, i, col)
        hmap = _header_map(ws)

    kantor_tujuan = _text(customer.customs_office_code) if customer else None
    valuta = "USD"
    if customer and _text(customer.default_currency):
        valuta = customer.default_currency.strip().upper()

    header_vals: dict[str, Any] = {
        "NOMOR AJU": nomor_aju,
        "KODE DOKUMEN": KODE_DOKUMEN_BC27,
        "KODE KANTOR": TH_KODE_KANTOR,
        "KODE KANTOR TUJUAN": kantor_tujuan,
        "KODE JENIS TPB": "1",
        "KODE TUJUAN PENGIRIMAN": "1",
        "KODE TUJUAN TPB": "1",
        "ASURANSI": 0,
        "NILAI BARANG": 0,
        "FREIGHT": 0,
        "FOB": 0,
        "BIAYA TAMBAHAN": 0,
        "BIAYA PENGURANG": 0,
        "VD": 0,
        "CIF": _num(total_cif) if total_cif else None,
        "HARGA_PENYERAHAN": _num(cif_idr),
        "NDPBM": _num(kurs),
        "DASAR PENGENAAN PAJAK": _num(cif_idr),
        "UANG MUKA": 0,
        "BRUTO": _num(total_gw) if total_gw else None,
        "NETTO": _num(total_nw) if total_nw else None,
        "VOLUME": 0,
        "KOTA PERNYATAAN": TH_KOTA_PERNYATAAN,
        "TANGGAL PERNYATAAN": _fmt_date(dday),
        "NAMA PERNYATAAN": TH_NAMA_PERNYATAAN,
        "JABATAN PERNYATAAN": TH_JABATAN_PERNYATAAN,
        "KODE VALUTA": valuta,
        "NOMOR DAFTAR": None,
        "TANGGAL DAFTAR": None,
        "PPN PAJAK": _num(ppn),
        "PPNBM PAJAK": 0,
        "TARIF PPN PAJAK": 11,
    }
    _append_row(ws, hmap, header_vals)

    # --- ENTITAS ---
    ws = wb["ENTITAS"]
    emap = _header_map(ws)
    _append_row(
        ws,
        emap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": 7,
            "KODE ENTITAS": "7",
            "KODE JENIS IDENTITAS": "6",
            "NOMOR IDENTITAS": TH_NPWP,
            "NAMA ENTITAS": TH_NAME,
            "ALAMAT ENTITAS": TH_ADDRESS,
            "KODE JENIS API": TH_KODE_JENIS_API,
            "KODE STATUS": TH_KODE_STATUS,
        },
    )
    _append_row(
        ws,
        emap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": 3,
            "KODE ENTITAS": "3",
            "KODE JENIS IDENTITAS": "6",
            "NOMOR IDENTITAS": TH_NPWP,
            "NAMA ENTITAS": TH_NAME,
            "ALAMAT ENTITAS": TH_ADDRESS,
            "NIB ENTITAS": TH_NIB,
            "KODE JENIS API": TH_KODE_JENIS_API,
            "KODE STATUS": TH_KODE_STATUS,
            "NOMOR IJIN ENTITAS": TH_SKEP,
            "TANGGAL IJIN ENTITAS": TH_SKEP_DATE,
        },
    )
    cust_name = _text(customer.name) if customer else None
    cust_addr = None
    if customer:
        cust_addr = _text(customer.address) or _text(customer.billing_address)
    cust_tax = _text(customer.tax_id) if customer else None
    cust_skep = _text(customer.skep_kb_number) if customer else None
    _append_row(
        ws,
        emap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": 8,
            "KODE ENTITAS": "8",
            "KODE JENIS IDENTITAS": "6",
            "NOMOR IDENTITAS": cust_tax,
            "NAMA ENTITAS": cust_name,
            "ALAMAT ENTITAS": cust_addr,
            "KODE JENIS API": TH_KODE_JENIS_API,
            "KODE STATUS": TH_KODE_STATUS,
            "NOMOR IJIN ENTITAS": cust_skep,
            "TANGGAL IJIN ENTITAS": None,
        },
    )

    # --- DOKUMEN ---
    ws = wb["DOKUMEN"]
    dmap = _header_map(ws)
    seri_dok = 1
    for inv in invoices:
        inv_day = inv.invoice_date
        if inv_day is None:
            dh_match = next((h for h in headers if int(h.id) == int(inv.delivery_id or 0)), None)
            inv_day = (dh_match.delivery_date if dh_match else None) or dday
        _append_row(
            ws,
            dmap,
            {
                "NOMOR AJU": nomor_aju,
                "SERI": seri_dok,
                "KODE DOKUMEN": KODE_DOK_INVOICE,
                "NOMOR DOKUMEN": _text(inv.invoice_number),
                "TANGGAL DOKUMEN": _fmt_date(inv_day),
            },
        )
        seri_dok += 1

    for dh in headers:
        sj_day = dh.delivery_date or dday
        _append_row(
            ws,
            dmap,
            {
                "NOMOR AJU": nomor_aju,
                "SERI": seri_dok,
                "KODE DOKUMEN": KODE_DOK_SJ,
                "NOMOR DOKUMEN": _text(dh.delivery_number),
                "TANGGAL DOKUMEN": _fmt_date(sj_day),
            },
        )
        seri_dok += 1

    _append_row(
        ws,
        dmap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": seri_dok,
            "KODE DOKUMEN": KODE_DOK_KONTRAK,
            "NOMOR DOKUMEN": kontrak,
            "TANGGAL DOKUMEN": _fmt_date(dday),
        },
    )

    # --- PENGANGKUT ---
    ws = wb["PENGANGKUT"]
    pmap = _header_map(ws)
    first = headers[0]
    drivers = {_text(h.driver_name) for h in headers}
    vehicles = {_text(h.vehicle_number) for h in headers}
    driver = _text(first.driver_name) if len(drivers) == 1 else None
    vehicle = _text(first.vehicle_number) if len(vehicles) == 1 else None
    _append_row(
        ws,
        pmap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": 1,
            "KODE CARA ANGKUT": None,
            "NAMA PENGANGKUT": driver or "-",
            "NOMOR PENGANGKUT": vehicle or "-",
        },
    )

    # --- KEMASAN ---
    ws = wb["KEMASAN"]
    kmap = _header_map(ws)
    _append_row(
        ws,
        kmap,
        {
            "NOMOR AJU": nomor_aju,
            "SERI": 1,
            "KODE KEMASAN": KODE_KEMASAN_RO,
            "JUMLAH KEMASAN": total_roll if total_roll > 0 else None,
        },
    )

    # --- BARANG ---
    ws = wb["BARANG"]
    bmap = _header_map(ws)
    for ln in lines:
        amt = ln["amount"]
        amt_idr = _money2(amt * kurs) if amt is not None and kurs is not None else None
        nw = ln["nw"]
        gw = None
        if nw is not None and bruto_ratio is not None:
            gw = _money2(nw * bruto_ratio)
        elif nw is not None:
            gw = nw
        _append_row(
            ws,
            bmap,
            {
                "NOMOR AJU": nomor_aju,
                "SERI BARANG": ln["seri"],
                "HS": None,
                "KODE BARANG": ln["kode_barang"],
                "URAIAN": ln["uraian"],
                "MEREK": "-",
                "TIPE": "-",
                "UKURAN": "-",
                "SPESIFIKASI LAIN": "-",
                "KODE SATUAN": KODE_SATUAN_FTK,
                "JUMLAH SATUAN": _num(ln["qty"]) if ln["qty"] else None,
                "NETTO": _num(nw),
                "BRUTO": _num(gw),
                "CIF": _num(amt),
                "CIF RUPIAH": _num(amt_idr),
                "NDPBM": _num(kurs),
                "FOB": 0,
                "NILAI TAMBAH": 0,
                "DISKON": 0,
                "HARGA PENYERAHAN": _num(amt_idr),
                "HARGA PEROLEHAN": 0,
                "HARGA EKSPOR": 0,
                "NILAI BARANG": 0,
                "NILAI JASA": 0,
            },
        )

    # --- BAHANBAKU (BOM belum ada; qty mengikuti invoice; kolom kuning kosong) ---
    ws = wb["BAHANBAKU"]
    bbmap = _header_map(ws)
    for ln in lines:
        row = {
            "NOMOR AJU": nomor_aju,
            "SERI BARANG": ln["seri"],
            "SERI BAHAN BAKU": 1,
            "KODE ASAL BAHAN BAKU": "0",
            "HS": None,
            "KODE BARANG": None,
            "URAIAN": None,
            "MEREK": "-",
            "TIPE": "-",
            "UKURAN": "-",
            "SPESIFIKASI LAIN": "-",
            "KODE SATUAN": KODE_SATUAN_FTK,
            "JUMLAH SATUAN": _num(ln["qty"]) if ln["qty"] else None,
            "KODE DOKUMEN ASAL": "23",
            "KODE KANTOR ASAL": TH_KODE_KANTOR,
            "NDPBM": _num(kurs),
            "HARGA PENYERAHAN": 0,
            "HARGA PEROLEHAN": 0,
            "NILAI JASA": 0,
            "SERI IZIN": 0,
        }
        for k in BAHANBAKU_YELLOW_KEYS:
            row[k] = None
        _append_row(ws, bbmap, row)

    # --- BAHANBAKUTARIF ---
    ws = wb["BAHANBAKUTARIF"]
    btmap = _header_map(ws)
    tarif_rows = [
        ("BM", "1", 0, "6", 100),
        ("PPH", "1", 2.5, "6", 100),
        ("PPN", "1", 11, "6", 100),
    ]
    for ln in lines:
        qty = _num(ln["qty"]) if ln["qty"] else None
        for kode_pungutan, kode_tarif, tarif, kode_fas, tarif_fas in tarif_rows:
            row = {
                "NOMOR AJU": nomor_aju,
                "SERI BARANG": ln["seri"],
                "SERI BAHAN BAKU": 1,
                "KODE ASAL BAHAN BAKU": "0",
                "KODE PUNGUTAN": kode_pungutan,
                "KODE TARIF": kode_tarif,
                "TARIF": tarif,
                "KODE FASILITAS": kode_fas,
                "TARIF FASILITAS": tarif_fas,
                "NILAI BAYAR": 0,
                "NILAI SUDAH DILUNASI": 0,
                "KODE SATUAN": KODE_SATUAN_FTK,
                "JUMLAH SATUAN": qty,
            }
            for k in BAHANBAKUTARIF_YELLOW_KEYS:
                row[k] = None
            _append_row(ws, btmap, row)

    # --- VERSI ---
    ws = wb["VERSI"]
    if ws.max_row < 1 or ws.cell(1, 1).value is None:
        ws.cell(1, 1, "VERSI")
    ws.cell(2, 1, VERSI_TEMPLATE)

    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in nomor_aju)[:48]
    filename = f"{safe}_bc27.xlsx"
    return wb, filename


def render_customs_bc27_xlsx(
    db: Session,
    delivery_ids: Sequence[int] | int,
    *,
    aju_seq: str | int,
) -> tuple[bytes, str]:
    wb, filename = build_bc27_workbook(db, delivery_ids, aju_seq=aju_seq)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), filename


def render_delivery_customs_bc27_xlsx(
    db: Session,
    delivery_id: int,
    *,
    aju_seq: str | int,
) -> tuple[bytes, str]:
    """Kompatibilitas: satu delivery_id."""
    return render_customs_bc27_xlsx(db, [delivery_id], aju_seq=aju_seq)
