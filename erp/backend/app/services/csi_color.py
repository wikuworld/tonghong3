"""CRUD helpers & impor katalog CSI."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CsiColorMaster
from app.services.buyer_registry import ensure_buyer_code
from app.services.color_code import csi_color_label

SHEET_BUYER = {
    "NIKE CSI": "NIKE",
    "UA CSI": "UA",
    "NB": "NB",
    "ADIDAS": "ADIDAS",
}


def apply_csi_fields(data: dict[str, Any]) -> dict[str, Any]:
    buyer = (data.get("buyer_code") or "").strip().upper()
    code = (data.get("color_code") or "").strip()
    name = (data.get("color_name") or "").strip()
    return {
        "buyer_code": buyer,
        "color_code": code,
        "color_name": name,
        "source_file": (data.get("source_file") or "").strip() or None,
        "is_active": bool(data.get("is_active", True)),
    }


def resolve_csi_by_buyer_code(
    db: Session, buyer_code: str, color_code: str
) -> CsiColorMaster | None:
    buyer = (buyer_code or "").strip().upper()
    code = (color_code or "").strip()
    if not buyer or not code:
        return None
    return db.scalars(
        select(CsiColorMaster).where(
            CsiColorMaster.buyer_code == buyer,
            CsiColorMaster.color_code == code,
            CsiColorMaster.is_active.is_(True),
        )
    ).first()


def parse_csi_workbook(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[dict[str, Any]] = []
    try:
        for sheet_name, buyer in SHEET_BUYER.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            header_idx = None
            for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
                vals = [str(x).strip().upper() if x is not None else "" for x in (row or ())]
                if "COLOR CODE" in vals and "COLOR NAME" in vals:
                    header_idx = i
                    break
            if header_idx is None:
                continue
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i <= header_idx:
                    continue
                if not row or row[1] is None:
                    continue
                code = str(row[1]).strip()
                name = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                if not code or code.upper() == "COLOR CODE":
                    continue
                out.append(
                    {
                        "buyer_code": buyer,
                        "color_code": code,
                        "color_name": name or code,
                        "source_file": path.name,
                        "is_active": True,
                    }
                )
    finally:
        wb.close()
    return out


def upsert_csi_rows(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    inserted = 0
    updated = 0
    for raw in rows:
        data = apply_csi_fields(raw)
        if not data["buyer_code"] or not data["color_code"]:
            continue
        ensure_buyer_code(db, data["buyer_code"])
        existing = db.scalars(
            select(CsiColorMaster).where(
                CsiColorMaster.buyer_code == data["buyer_code"],
                CsiColorMaster.color_code == data["color_code"],
            )
        ).first()
        if existing:
            existing.color_name = data["color_name"]
            existing.source_file = data["source_file"] or existing.source_file
            existing.is_active = data["is_active"]
            updated += 1
        else:
            db.add(CsiColorMaster(**data))
            inserted += 1
    db.flush()
    return {"inserted": inserted, "updated": updated, "total": inserted + updated, "warnings": []}


def apply_csi_to_swatch_fields(
    db: Session,
    *,
    csi_id: int | None,
    buyer_code: str,
    color_code: str,
    color_name: str,
) -> tuple[int | None, str, str]:
    """Jika CSI dipilih, isi kode/nama yang kosong dari CSI (override manual user tetap dihormati)."""
    code = (color_code or "").strip()
    name = (color_name or "").strip()
    if csi_id:
        csi = db.get(CsiColorMaster, csi_id)
        if csi and csi.is_active:
            return (
                int(csi.id),
                code or (csi.color_code or ""),
                name or (csi.color_name or ""),
            )
    if buyer_code and code:
        csi = resolve_csi_by_buyer_code(db, buyer_code, code)
        if csi:
            return int(csi.id), code, name
    return None, code, name


def csi_option_label(row: CsiColorMaster) -> str:
    return csi_color_label(row.color_name, row.color_code)
