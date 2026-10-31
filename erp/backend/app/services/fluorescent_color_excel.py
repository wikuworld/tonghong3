"""Export/import Excel master Fluorescent Color.

Kunci unik impor: buyer_code + season_label + color_code + color_name.
Baris yang sudah ada diperbarui. File yang diunggah ulang tidak menambah duplikat
dan tidak menonaktifkan baris di luar file.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Brand, FluorescentColorMaster
from app.services.buyer_registry import list_active_buyer_codes
from app.services.fluorescent_color import VALID_COLOR_CLASSES, normalize_buyer_code
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "buyer_code",
    "season_label",
    "season_year",
    "season_type",
    "color_code",
    "color_name",
    "color_class",
    "source_file",
    "is_active",
    "created_at",
    "updated_at",
]

TEMPLATE_COLUMNS = [
    "buyer_code",
    "color_name",
    "color_code",
    "season_label",
    "season_year",
    "season_type",
    "color_class",
    "is_active",
]

_ROW_FIELDS = (
    "id",
    "buyer_code",
    "season_label",
    "season_year",
    "season_type",
    "color_code",
    "color_name",
    "color_class",
    "source_file",
    "is_active",
    "created_at",
    "updated_at",
)

_REQUIRED_HEADERS = ("buyer_code", "color_name")

_HEADER_ALIAS = {
    "id": "id",
    "buyer_code": "buyer_code",
    "buyer": "buyer_code",
    "brand": "buyer_code",
    "season_label": "season_label",
    "season": "season_label",
    "season_year": "season_year",
    "year": "season_year",
    "tahun": "season_year",
    "season_type": "season_type",
    "type": "season_type",
    "tipe": "season_type",
    "color_code": "color_code",
    "kode": "color_code",
    "code": "color_code",
    "color_name": "color_name",
    "color": "color_name",
    "warna": "color_name",
    "nama": "color_name",
    "color_class": "color_class",
    "class": "color_class",
    "kelas": "color_class",
    "source_file": "source_file",
    "source": "source_file",
    "sumber": "source_file",
    "is_active": "is_active",
    "active": "is_active",
    "aktif": "is_active",
    "created_at": "created_at",
    "updated_at": "updated_at",
}

_TRUE_VALUES = {"1", "true", "yes", "y", "aktif", "active"}
_FALSE_VALUES = {"0", "false", "no", "n", "nonaktif", "inactive", "tidak"}


@dataclass
class ImportErrorRow:
    row: int
    color_code: str | None
    reason: str


@dataclass
class FluorescentImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(e) for e in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    s = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in s:
        s = s.replace("__", "_")
    return s


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val).strip()
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)) and 20000 < float(val) < 80000:
        stamp = datetime(1899, 12, 30) + timedelta(days=float(val))
        return stamp.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _blank(raw: Any) -> str | None:
    text = _text(raw)
    return text or None


def _business_key(buyer: str, season: str | None, code: str | None, name: str) -> tuple[str, str, str, str]:
    return (
        normalize_buyer_code(buyer),
        _text(season),
        _text(code),
        _text(name),
    )


def _key_from_row(row: FluorescentColorMaster) -> tuple[str, str, str, str]:
    return _business_key(row.buyer_code or "", row.season_label, row.color_code, row.color_name or "")


def _parse_dt(raw: str) -> datetime | None:
    text = (raw or "").strip().replace("T", " ")
    if not text:
        return None
    if len(text) > 19:
        text = text[:19]
    for fmt, size in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%d/%m/%Y", 10),
    ):
        try:
            return datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
    return None


def _format_dt(value: datetime | None) -> str:
    if not value:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _parse_id(raw: str) -> int | None:
    text = (raw or "").strip()
    if not text:
        return None
    try:
        value = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None
    return value if value > 0 else None


def _parse_year(raw: str) -> int | None:
    text = (raw or "").strip()
    if not text:
        return None
    try:
        value = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None
    if value < 1900 or value > 2200:
        return None
    return value


def _parse_active(raw: str) -> bool | None:
    text = (raw or "").strip().lower()
    if not text:
        return True
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def _parse_class(raw: str) -> str | None:
    text = (raw or "").strip().lower()
    if not text:
        return "fluorescent"
    if text in VALID_COLOR_CLASSES:
        return text
    return None


def _workbook_bytes(ws_title: str, headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ws_title
    for col_idx, col in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=col)
    for row_idx, values in enumerate(data_rows, 2):
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_fluorescent_colors_xlsx(db: Session) -> bytes:
    rows = list(
        db.scalars(
            select(FluorescentColorMaster).order_by(
                FluorescentColorMaster.buyer_code.asc(),
                FluorescentColorMaster.season_label.asc(),
                FluorescentColorMaster.color_name.asc(),
                FluorescentColorMaster.id.asc(),
            )
        ).all()
    )
    data = [
        [
            row.id,
            row.buyer_code or "",
            row.season_label or "",
            row.season_year if row.season_year is not None else "",
            row.season_type or "",
            row.color_code or "",
            row.color_name or "",
            row.color_class or "",
            row.source_file or "",
            "0" if row.is_active is False else "1",
            _format_dt(row.created_at),
            _format_dt(row.updated_at),
        ]
        for row in rows
    ]
    return _workbook_bytes("Fluorescent Color", EXPORT_COLUMNS, data)


def export_fluorescent_colors_template_xlsx() -> bytes:
    return _workbook_bytes("Fluorescent Color", TEMPLATE_COLUMNS, [])


def _parse_rows(content: bytes) -> tuple[list[dict[str, str]], list[ImportErrorRow]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel tidak bisa dibaca")]
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel kosong")]

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row or []):
        key = _HEADER_ALIAS.get(_normalize_header(cell))
        if key and key not in col_map.values():
            col_map[idx] = key

    if any(name not in col_map.values() for name in _REQUIRED_HEADERS):
        return [], [
            ImportErrorRow(
                row=1,
                color_code=None,
                reason="Header wajib memuat kolom buyer_code dan color_name",
            )
        ]

    parsed: list[dict[str, str]] = []
    for excel_row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {key: "" for key in _ROW_FIELDS}
        for idx, key in col_map.items():
            if idx < len(row) and key in item:
                item[key] = _cell_str(row[idx])
        if not any(item.get(k) for k in ("id", "buyer_code", "color_name", "color_code")):
            continue
        item["_row"] = str(excel_row_num)
        item["_present"] = ",".join(col_map.values())
        parsed.append(item)
    return parsed, []


def _brand_id_for_buyer(brands: list[Brand], buyer: str) -> int | None:
    if not buyer or buyer == "GLOBAL":
        return None
    for brand in brands:
        if _text(getattr(brand, "code", "")).upper() == buyer:
            brand_id = getattr(brand, "id", None)
            return int(brand_id) if brand_id is not None else None
    return None


def import_fluorescent_colors_xlsx(db: Session, content: bytes) -> FluorescentImportResult:
    result = FluorescentImportResult()
    rows, parse_errors = _parse_rows(content)
    result.errors.extend(parse_errors)
    if parse_errors and not rows:
        return result

    known_buyers = {code.strip().upper() for code in list_active_buyer_codes(db) if code}
    existing = list(db.scalars(select(FluorescentColorMaster)).all())
    brands = list(db.scalars(select(Brand)).all())
    by_id = {row.id: row for row in existing if row.id is not None}
    by_key: dict[tuple[str, str, str, str], FluorescentColorMaster] = {}
    for row in sorted(existing, key=lambda item: item.id or 0):
        by_key.setdefault(_key_from_row(row), row)
    seen_ids: set[int] = set()

    for item in rows:
        result.total_rows += 1
        excel_row = int(item.pop("_row", "0") or 0)
        present = {part for part in (item.pop("_present", "") or "").split(",") if part}
        code_label = _text(item.get("color_code")) or None
        raw_id = _text(item.get("id"))
        row_id = _parse_id(raw_id)
        if raw_id and row_id is None:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=code_label, reason="id tidak valid"))
            continue
        if row_id is not None and row_id in seen_ids:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason=f"id {row_id} muncul lebih dari sekali di file",
                )
            )
            continue

        buyer = normalize_buyer_code(item.get("buyer_code") or "")
        color_name = _text(item.get("color_name"))
        color_code = _blank(item.get("color_code"))
        season_label = _blank(item.get("season_label"))
        season_type = _blank(item.get("season_type")) if "season_type" in present else None
        if not _text(item.get("buyer_code")) or not color_name:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason="buyer_code dan color_name wajib diisi",
                )
            )
            continue
        if buyer not in known_buyers:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason=f"Buyer '{buyer}' tidak ditemukan di registry",
                )
            )
            continue

        color_class = _parse_class(item.get("color_class") or "")
        if color_class is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason="color_class harus fluorescent atau metallic",
                )
            )
            continue

        season_year: int | None = None
        if "season_year" in present and _text(item.get("season_year")):
            season_year = _parse_year(item.get("season_year") or "")
            if season_year is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason="season_year tidak valid",
                    )
                )
                continue

        is_active = True
        if "is_active" in present:
            parsed_active = _parse_active(item.get("is_active") or "")
            if parsed_active is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason="is_active tidak valid",
                    )
                )
                continue
            is_active = parsed_active

        created_at = _parse_dt(item.get("created_at") or "") if "created_at" in present else None
        updated_at = _parse_dt(item.get("updated_at") or "") if "updated_at" in present else None
        if "created_at" in present and _text(item.get("created_at")) and created_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=color_code or code_label, reason="created_at tidak valid")
            )
            continue
        if "updated_at" in present and _text(item.get("updated_at")) and updated_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=color_code or code_label, reason="updated_at tidak valid")
            )
            continue

        key = _business_key(buyer, season_label, color_code, color_name)
        owner = by_key.get(key)
        target = by_id.get(row_id) if row_id is not None else None
        if owner is not None and target is not None and owner is not target:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason="buyer + season + kode + nama sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue
        if target is None:
            target = owner
        if row_id is not None:
            seen_ids.add(row_id)

        source_file = _blank(item.get("source_file")) if "source_file" in present else None
        brand_id = _brand_id_for_buyer(brands, buyer)

        if target is None:
            created = FluorescentColorMaster(
                buyer_code=buyer,
                brand_id=brand_id,
                season_label=season_label,
                season_year=season_year if "season_year" in present else None,
                season_type=season_type if "season_type" in present else None,
                color_name=color_name,
                color_code=color_code,
                color_class=color_class,
                source_file=source_file,
                is_active=is_active if "is_active" in present else True,
                created_at=created_at or app_now(),
                updated_at=updated_at,
            )
            db.add(created)
            by_key[key] = created
            result.created += 1
            continue

        old_key = _key_from_row(target)
        if old_key != key and by_key.get(key) not in (None, target):
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason="buyer + season + kode + nama sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue

        same_type = True if "season_type" not in present else _text(target.season_type) == _text(season_type)
        same_year = True if "season_year" not in present else target.season_year == season_year
        same_source = True if "source_file" not in present else _text(target.source_file) == _text(source_file)
        same_active = True if "is_active" not in present else bool(target.is_active) == is_active
        same_created = created_at is None or _format_dt(target.created_at) == _format_dt(created_at)
        same_updated = updated_at is None or _format_dt(target.updated_at) == _format_dt(updated_at)
        same_fields = (
            normalize_buyer_code(target.buyer_code or "") == buyer
            and _text(target.season_label) == _text(season_label)
            and same_type
            and _text(target.color_code) == _text(color_code)
            and _text(target.color_name) == color_name
            and _text(target.color_class) == color_class
            and target.brand_id == brand_id
            and same_year
            and same_source
            and same_active
            and same_created
            and same_updated
        )
        if same_fields:
            result.skipped += 1
            continue

        target.buyer_code = buyer
        target.brand_id = brand_id
        target.season_label = season_label
        if "season_type" in present:
            target.season_type = season_type
        target.color_name = color_name
        target.color_code = color_code
        target.color_class = color_class
        if "season_year" in present:
            target.season_year = season_year
        if "source_file" in present:
            target.source_file = source_file
        if "is_active" in present:
            target.is_active = is_active
        if created_at is not None:
            target.created_at = created_at
        target.updated_at = updated_at if updated_at is not None else app_now()
        if old_key != key and by_key.get(old_key) is target:
            by_key.pop(old_key, None)
        by_key[key] = target
        result.updated += 1

    return result
