"""Export/import Excel master CSI Color.

Kunci unik impor: buyer_code + color_code.
Baris yang sudah ada diperbarui. File yang diunggah ulang tidak menambah duplikat
dan tidak menonaktifkan baris di luar file.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CsiColorMaster
from app.services.buyer_registry import list_active_buyer_codes
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "buyer_code",
    "color_code",
    "color_name",
    "source_file",
    "is_active",
    "created_at",
    "updated_at",
]

TEMPLATE_COLUMNS = [
    "buyer_code",
    "color_code",
    "color_name",
    "is_active",
]

_ROW_FIELDS = (
    "id",
    "buyer_code",
    "color_code",
    "color_name",
    "source_file",
    "is_active",
    "created_at",
    "updated_at",
)

_REQUIRED_HEADERS = ("buyer_code", "color_code", "color_name")

_HEADER_ALIAS = {
    "id": "id",
    "buyer_code": "buyer_code",
    "buyer": "buyer_code",
    "brand": "buyer_code",
    "color_code": "color_code",
    "kode": "color_code",
    "code": "color_code",
    "color_name": "color_name",
    "color": "color_name",
    "warna": "color_name",
    "nama": "color_name",
    "source_file": "source_file",
    "source": "source_file",
    "sumber": "source_file",
    "is_active": "is_active",
    "active": "is_active",
    "aktif": "is_active",
    "created_at": "created_at",
    "updated_at": "updated_at",
}

_TRUE_VALUES = {"1", "true", "yes", "y", "aktif", "active"}
_FALSE_VALUES = {"0", "false", "no", "n", "nonaktif", "inactive", "tidak"}


@dataclass
class ImportErrorRow:
    row: int
    color_code: str | None
    reason: str


@dataclass
class CsiColorImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(e) for e in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    s = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in s:
        s = s.replace("__", "_")
    return s


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val).strip()
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)) and 20000 < float(val) < 80000:
        stamp = datetime(1899, 12, 30) + timedelta(days=float(val))
        return stamp.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _blank(raw: Any) -> str | None:
    return _text(raw) or None


def _business_key(buyer: str, code: str) -> tuple[str, str]:
    return (_text(buyer).upper(), _text(code))


def _key_from_row(row: CsiColorMaster) -> tuple[str, str]:
    return _business_key(row.buyer_code or "", row.color_code or "")


def _parse_dt(raw: str) -> datetime | None:
    text = (raw or "").strip().replace("T", " ")
    if not text:
        return None
    if len(text) > 19:
        text = text[:19]
    for fmt, size in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%d/%m/%Y", 10),
    ):
        try:
            return datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
    return None


def _format_dt(value: datetime | None) -> str:
    if not value:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _parse_id(raw: str) -> int | None:
    text = (raw or "").strip()
    if not text:
        return None
    try:
        value = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None
    return value if value > 0 else None


def _parse_active(raw: str) -> bool | None:
    text = (raw or "").strip().lower()
    if not text:
        return True
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def _workbook_bytes(headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CSI Color"
    for col_idx, col in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=col)
    for row_idx, values in enumerate(data_rows, 2):
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_csi_colors_xlsx(db: Session) -> bytes:
    rows = list(
        db.scalars(
            select(CsiColorMaster).order_by(
                CsiColorMaster.buyer_code.asc(),
                CsiColorMaster.color_code.asc(),
                CsiColorMaster.id.asc(),
            )
        ).all()
    )
    data = [
        [
            row.id,
            row.buyer_code or "",
            row.color_code or "",
            row.color_name or "",
            row.source_file or "",
            "0" if row.is_active is False else "1",
            _format_dt(row.created_at),
            _format_dt(row.updated_at),
        ]
        for row in rows
    ]
    return _workbook_bytes(EXPORT_COLUMNS, data)


def export_csi_colors_template_xlsx() -> bytes:
    return _workbook_bytes(TEMPLATE_COLUMNS, [])


def _parse_rows(content: bytes) -> tuple[list[dict[str, str]], list[ImportErrorRow]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel tidak bisa dibaca")]
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel kosong")]

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row or []):
        key = _HEADER_ALIAS.get(_normalize_header(cell))
        if key and key not in col_map.values():
            col_map[idx] = key

    if any(name not in col_map.values() for name in _REQUIRED_HEADERS):
        return [], [
            ImportErrorRow(
                row=1,
                color_code=None,
                reason="Header wajib memuat kolom buyer_code, color_code, dan color_name",
            )
        ]

    parsed: list[dict[str, str]] = []
    for excel_row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {key: "" for key in _ROW_FIELDS}
        for idx, key in col_map.items():
            if idx < len(row) and key in item:
                item[key] = _cell_str(row[idx])
        if not any(item.get(k) for k in ("id", "buyer_code", "color_code", "color_name")):
            continue
        item["_row"] = str(excel_row_num)
        item["_present"] = ",".join(col_map.values())
        parsed.append(item)
    return parsed, []


def import_csi_colors_xlsx(db: Session, content: bytes) -> CsiColorImportResult:
    result = CsiColorImportResult()
    rows, parse_errors = _parse_rows(content)
    result.errors.extend(parse_errors)
    if parse_errors and not rows:
        return result

    known_buyers = {code.strip().upper() for code in list_active_buyer_codes(db) if code}
    existing = list(db.scalars(select(CsiColorMaster)).all())
    by_id = {row.id: row for row in existing if row.id is not None}
    by_key: dict[tuple[str, str], CsiColorMaster] = {}
    for row in sorted(existing, key=lambda item: item.id or 0):
        by_key.setdefault(_key_from_row(row), row)
    seen_ids: set[int] = set()

    for item in rows:
        result.total_rows += 1
        excel_row = int(item.pop("_row", "0") or 0)
        present = {part for part in (item.pop("_present", "") or "").split(",") if part}
        buyer = _text(item.get("buyer_code")).upper()
        color_code = _text(item.get("color_code"))
        color_name = _text(item.get("color_name"))
        code_label = color_code or None
        raw_id = _text(item.get("id"))
        row_id = _parse_id(raw_id)
        if raw_id and row_id is None:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=code_label, reason="id tidak valid"))
            continue
        if row_id is not None and row_id in seen_ids:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason=f"id {row_id} muncul lebih dari sekali di file",
                )
            )
            continue
        if not buyer or not color_code or not color_name:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason="buyer_code, color_code, dan color_name wajib diisi",
                )
            )
            continue
        if buyer not in known_buyers:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code,
                    reason=f"Buyer '{buyer}' tidak ditemukan di registry",
                )
            )
            continue

        is_active = True
        if "is_active" in present:
            parsed_active = _parse_active(item.get("is_active") or "")
            if parsed_active is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(row=excel_row, color_code=color_code, reason="is_active tidak valid")
                )
                continue
            is_active = parsed_active

        created_at = _parse_dt(item.get("created_at") or "") if "created_at" in present else None
        updated_at = _parse_dt(item.get("updated_at") or "") if "updated_at" in present else None
        if "created_at" in present and _text(item.get("created_at")) and created_at is None:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=color_code, reason="created_at tidak valid"))
            continue
        if "updated_at" in present and _text(item.get("updated_at")) and updated_at is None:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=color_code, reason="updated_at tidak valid"))
            continue

        key = _business_key(buyer, color_code)
        owner = by_key.get(key)
        target = by_id.get(row_id) if row_id is not None else None
        if owner is not None and target is not None and owner is not target:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code,
                    reason="buyer + color_code sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue
        if target is None:
            target = owner
        if row_id is not None:
            seen_ids.add(row_id)

        source_file = _blank(item.get("source_file")) if "source_file" in present else None

        if target is None:
            created = CsiColorMaster(
                buyer_code=buyer,
                color_code=color_code,
                color_name=color_name,
                source_file=source_file,
                is_active=is_active if "is_active" in present else True,
                created_at=created_at or app_now(),
                updated_at=updated_at,
            )
            db.add(created)
            by_key[key] = created
            result.created += 1
            continue

        old_key = _key_from_row(target)
        if old_key != key and by_key.get(key) not in (None, target):
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code,
                    reason="buyer + color_code sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue

        same_source = True if "source_file" not in present else _text(target.source_file) == _text(source_file)
        same_active = True if "is_active" not in present else bool(target.is_active) == is_active
        same_created = created_at is None or _format_dt(target.created_at) == _format_dt(created_at)
        same_updated = updated_at is None or _format_dt(target.updated_at) == _format_dt(updated_at)
        if (
            _text(target.buyer_code).upper() == buyer
            and _text(target.color_code) == color_code
            and _text(target.color_name) == color_name
            and same_source
            and same_active
            and same_created
            and same_updated
        ):
            result.skipped += 1
            continue

        target.buyer_code = buyer
        target.color_code = color_code
        target.color_name = color_name
        if "source_file" in present:
            target.source_file = source_file
        if "is_active" in present:
            target.is_active = is_active
        if created_at is not None:
            target.created_at = created_at
        target.updated_at = updated_at if updated_at is not None else app_now()
        if old_key != key and by_key.get(old_key) is target:
            by_key.pop(old_key, None)
        by_key[key] = target
        result.updated += 1

    return result
