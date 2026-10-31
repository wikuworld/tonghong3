"""Export/import Excel list harga buyer.

Satu baris Excel = satu baris sheet. Identitas list (brand, buyer, periode)
diulang di setiap baris supaya file yang diunduh bisa diimpor kembali.
Baris yang sudah ada diperbarui. Baris yang tidak ada di file tidak dihapus.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Brand, BrandBuyerMap, BuyerPriceList, BuyerPriceSheetRow
from app.services.buyer_price_aliases import normalize_sheet_cells
from app.services.buyer_price_import import STANDARD_COLUMN_KEYS

EXPORT_COLUMNS = [
    "list_id",
    "brand_code",
    "buyer_code",
    "currency",
    "title",
    "period_label",
    "period_start",
    "period_end",
    "outside_price_adjust",
    "outside_adjust_type",
    "discount_modifier_enabled",
    "discount_modifier",
    "discount_modifier_type",
    "item",
    "thickness",
    "color",
    "color_white",
    "color_black",
    "color_color",
    "flourasent",
    "metalic",
    "emboss",
    "spesial_emboss",
]

TEMPLATE_COLUMNS = [name for name in EXPORT_COLUMNS if name != "list_id"]

_REQUIRED_HEADERS = ("brand_code", "buyer_code", "period_start", "period_end")

_HEADER_ALIAS = {
    "list_id": "list_id",
    "id": "list_id",
    "brand_code": "brand_code",
    "brand": "brand_code",
    "buyer_code": "buyer_code",
    "buyer": "buyer_code",
    "currency": "currency",
    "title": "title",
    "period_label": "period_label",
    "period_start": "period_start",
    "period_end": "period_end",
    "outside_price_adjust": "outside_price_adjust",
    "outside_adjust_type": "outside_adjust_type",
    "discount_modifier_enabled": "discount_modifier_enabled",
    "discount_modifier": "discount_modifier",
    "discount_modifier_type": "discount_modifier_type",
    "item": "item",
    "thickness": "thickness",
    "color": "color",
    "color_white": "color_white",
    "color_black": "color_black",
    "color_color": "color_color",
    "flourasent": "flourasent",
    "fluorescent": "flourasent",
    "metalic": "metalic",
    "metallic": "metalic",
    "emboss": "emboss",
    "spesial_emboss": "spesial_emboss",
    "special_emboss": "spesial_emboss",
}

_SHEET_TEXT = (
    ("item", "ITEM"),
    ("thickness", "THICKNESS"),
    ("color", "COLOR"),
)

_SHEET_PRICES = (
    ("color_white", "COLOR WHITE"),
    ("color_black", "COLOR BLACK"),
    ("color_color", "COLOR COLOR"),
    ("flourasent", "FLOURASENT"),
    ("metalic", "METALIC"),
    ("emboss", "EMBOSS"),
    ("spesial_emboss", "SPESIAL EMBOSS"),
)

_ADJUST_ALIAS = {
    "nominal": "nominal",
    "amount": "nominal",
    "amt": "nominal",
    "percent": "percent",
    "percentage": "percent",
    "pct": "percent",
    "%": "percent",
}

_TRUE_VALUES = {"1", "true", "yes", "y", "ya", "aktif", "active"}
_FALSE_VALUES = {"0", "false", "no", "n", "tidak", "nonaktif", "inactive"}


@dataclass
class ImportErrorRow:
    row: int
    item: str | None
    reason: str


@dataclass
class BuyerPriceImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(err) for err in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    text = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in text:
        text = text.replace("__", "_")
    return _HEADER_ALIAS.get(text, text)


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _cell_str(val: Any, *, dates: bool = False) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if dates and isinstance(val, (int, float)) and not isinstance(val, bool):
        number = float(val)
        if 20000 < number < 80000:
            stamp = datetime(1899, 12, 30) + timedelta(days=number)
            return stamp.strftime("%Y-%m-%d")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    if isinstance(val, bool):
        return "1" if val else "0"
    return str(val).strip()


def _parse_date(raw: Any) -> date | None:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = _cell_str(raw, dates=True)
    if not text:
        return None
    for fmt, size in (("%Y-%m-%d", 10), ("%d/%m/%Y", 10), ("%Y/%m/%d", 10)):
        try:
            return datetime.strptime(text[:size], fmt).date()
        except ValueError:
            continue
    return None


def _parse_int(raw: Any) -> int | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _money(raw: Any) -> Decimal | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return Decimal(text).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def _same_money(left: Any, right: Decimal | None) -> bool:
    return _money(left) == right


def _adjust_type(raw: Any) -> str | None:
    text = _text(raw).lower()
    if not text:
        return "nominal"
    return _ADJUST_ALIAS.get(text)


def _parse_bool(raw: Any) -> bool | None:
    if isinstance(raw, bool):
        return raw
    text = _cell_str(raw).lower()
    if not text:
        return False
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def _price_number(raw: Any) -> float | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def _sheet_key(cells: dict) -> tuple[str, str, str]:
    norm = normalize_sheet_cells(cells)
    return (
        _text(norm.get("ITEM")).upper(),
        _text(norm.get("THICKNESS")).upper(),
        _text(norm.get("COLOR")).upper(),
    )


def _comparable_cells(cells: dict | None) -> dict[str, Any]:
    norm = normalize_sheet_cells(cells or {})
    out: dict[str, Any] = {
        "ITEM": _text(norm.get("ITEM")),
        "THICKNESS": _text(norm.get("THICKNESS")),
        "COLOR": _text(norm.get("COLOR")),
    }
    for _excel_key, sheet_key in _SHEET_PRICES:
        number = _price_number(norm.get(sheet_key))
        out[sheet_key] = "" if number is None else number
    return out


def _build_cells(data: dict[str, str]) -> dict[str, Any]:
    cells: dict[str, Any] = {}
    for excel_key, sheet_key in _SHEET_TEXT:
        cells[sheet_key] = data.get(excel_key, "")
    for excel_key, sheet_key in _SHEET_PRICES:
        number = _price_number(data.get(excel_key, ""))
        cells[sheet_key] = "" if number is None else number
    return normalize_sheet_cells(cells)


def _xlsx_bytes(headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Buyer Prices"
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for r_idx, row in enumerate(data_rows, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _export_price(cells: dict, key: str) -> Any:
    number = _price_number(cells.get(key))
    return None if number is None else number


def _list_values(price_list: BuyerPriceList, brand_code: str) -> list[Any]:
    return [
        price_list.id,
        brand_code,
        price_list.buyer_code or "",
        price_list.currency or "USD",
        price_list.title or "",
        price_list.period_label or "",
        price_list.period_start.isoformat() if isinstance(price_list.period_start, date) else "",
        price_list.period_end.isoformat() if isinstance(price_list.period_end, date) else "",
        float(price_list.outside_price_adjust) if price_list.outside_price_adjust is not None else None,
        price_list.outside_adjust_type or "",
        1 if price_list.discount_modifier_enabled else 0,
        float(price_list.discount_modifier) if price_list.discount_modifier is not None else None,
        price_list.discount_modifier_type or "",
    ]


def _line_values(cells: dict | None) -> list[Any]:
    norm = normalize_sheet_cells(cells or {})
    values: list[Any] = [
        _text(norm.get("ITEM")) or None,
        _text(norm.get("THICKNESS")) or None,
        _text(norm.get("COLOR")) or None,
    ]
    for _excel_key, sheet_key in _SHEET_PRICES:
        values.append(_export_price(norm, sheet_key))
    return values


def export_buyer_prices_template_xlsx() -> bytes:
    return _xlsx_bytes(TEMPLATE_COLUMNS, [])


def export_buyer_prices_xlsx(db: Session) -> bytes:
    lists = list(db.scalars(select(BuyerPriceList).order_by(BuyerPriceList.id.asc())).all())
    sheets = list(
        db.scalars(
            select(BuyerPriceSheetRow).order_by(
                BuyerPriceSheetRow.list_id.asc(),
                BuyerPriceSheetRow.row_number.asc(),
            )
        ).all()
    )
    brands = {row.id: (row.code or "") for row in db.scalars(select(Brand)).all()}
    by_list: dict[int, list[BuyerPriceSheetRow]] = {}
    for sheet in sheets:
        by_list.setdefault(sheet.list_id, []).append(sheet)

    data_rows: list[list[Any]] = []
    for price_list in lists:
        header = _list_values(price_list, brands.get(price_list.brand_id, ""))
        lines = by_list.get(price_list.id, [])
        if not lines:
            data_rows.append(header + _line_values({}))
            continue
        for sheet in lines:
            data_rows.append(header + _line_values(sheet.cells))
    return _xlsx_bytes(EXPORT_COLUMNS, data_rows)


def _read_table(content: bytes) -> tuple[list[str], list[tuple[int, dict[str, str]]], ImportErrorRow | None]:
    if not content:
        return [], [], ImportErrorRow(row=1, item=None, reason="File Excel kosong")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return [], [], ImportErrorRow(row=1, item=None, reason="File Excel tidak bisa dibaca")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], [], ImportErrorRow(row=1, item=None, reason="File Excel kosong")
    headers = [_normalize_header(cell) for cell in rows[0]]
    present = [name for name in headers if name]
    if any(name not in present for name in _REQUIRED_HEADERS):
        return [], [], ImportErrorRow(
            row=1,
            item=None,
            reason="Header wajib memuat kolom brand_code, buyer_code, period_start, dan period_end",
        )
    parsed: list[tuple[int, dict[str, str]]] = []
    for offset, raw in enumerate(rows[1:], start=2):
        data: dict[str, str] = {}
        empty = True
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = raw[idx] if idx < len(raw) else None
            as_date = header in {"period_start", "period_end"}
            text = _cell_str(value, dates=as_date)
            if text:
                empty = False
            data[header] = text
        if empty:
            continue
        parsed.append((offset, data))
    return present, parsed, None


def _index_brands(brands: list[Brand]) -> tuple[dict[str, Brand], dict[str, list[Brand]]]:
    by_code: dict[str, Brand] = {}
    by_name: dict[str, list[Brand]] = {}
    for brand in brands:
        code = _text(brand.code).upper()
        if code and code not in by_code:
            by_code[code] = brand
        name = _text(brand.name).upper()
        if name:
            by_name.setdefault(name, []).append(brand)
    return by_code, by_name


def _resolve_brand(token: str, by_code: dict[str, Brand], by_name: dict[str, list[Brand]]) -> Brand | None:
    key = token.strip().upper()
    if key in by_code:
        return by_code[key]
    named = by_name.get(key) or []
    if len(named) == 1:
        return named[0]
    return None


def _ranges_overlap(a_start: date, a_end: date, b_start: date, b_end: date) -> bool:
    return a_start <= b_end and b_start <= a_end


def _period_overlaps(
    lists: list[BuyerPriceList],
    *,
    brand_id: int,
    buyer_code: str,
    period_start: date,
    period_end: date,
    exclude: BuyerPriceList | None,
) -> bool:
    for other in lists:
        if exclude is not None and other is exclude:
            continue
        if other.brand_id != brand_id:
            continue
        if _text(other.buyer_code).upper() != buyer_code:
            continue
        if not isinstance(other.period_start, date) or not isinstance(other.period_end, date):
            continue
        if _ranges_overlap(period_start, period_end, other.period_start, other.period_end):
            return True
    return False


def _apply_header(target: BuyerPriceList, data: dict[str, str], present: set[str]) -> bool:
    changed = False
    if "currency" in present:
        currency = (data.get("currency") or "USD").upper()[:3] or "USD"
        if (target.currency or "") != currency:
            target.currency = currency
            changed = True
    if "title" in present:
        title = data.get("title") or None
        if (target.title or None) != title:
            target.title = title
            changed = True
    if "period_label" in present:
        label = data.get("period_label") or None
        if (target.period_label or None) != label:
            target.period_label = label
            changed = True
    if "outside_price_adjust" in present:
        amount = _money(data.get("outside_price_adjust"))
        if not _same_money(target.outside_price_adjust, amount):
            target.outside_price_adjust = amount
            changed = True
    if "outside_adjust_type" in present:
        kind = _adjust_type(data.get("outside_adjust_type")) or "nominal"
        if (target.outside_adjust_type or "") != kind:
            target.outside_adjust_type = kind
            changed = True
    if "discount_modifier_enabled" in present:
        enabled = _parse_bool(data.get("discount_modifier_enabled"))
        if enabled is None:
            enabled = False
        if bool(target.discount_modifier_enabled) != enabled:
            target.discount_modifier_enabled = enabled
            changed = True
    if "discount_modifier" in present:
        amount = _money(data.get("discount_modifier"))
        if not _same_money(target.discount_modifier, amount):
            target.discount_modifier = amount
            changed = True
    if "discount_modifier_type" in present:
        kind = _adjust_type(data.get("discount_modifier_type")) or "nominal"
        if (target.discount_modifier_type or "") != kind:
            target.discount_modifier_type = kind
            changed = True
    return changed


def _validate_types(data: dict[str, str], present: set[str]) -> str | None:
    for column in ("outside_adjust_type", "discount_modifier_type"):
        if column not in present or not data.get(column):
            continue
        if _adjust_type(data.get(column)) is None:
            return f"Nilai {column} harus nominal atau percent"
    if "discount_modifier_enabled" in present and data.get("discount_modifier_enabled"):
        if _parse_bool(data.get("discount_modifier_enabled")) is None:
            return "Nilai discount_modifier_enabled tidak dikenali"
    for column in ("outside_price_adjust", "discount_modifier"):
        if column in present and data.get(column) and _money(data.get(column)) is None:
            return f"Nilai {column} harus angka"
    return None


def _count(result: BuyerPriceImportResult, *, created: bool, updated: bool) -> None:
    if created:
        result.created += 1
    elif updated:
        result.updated += 1
    else:
        result.skipped += 1


def import_buyer_prices_xlsx(db: Session, content: bytes) -> BuyerPriceImportResult:
    present_list, parsed, header_error = _read_table(content)
    result = BuyerPriceImportResult()
    if header_error:
        result.errors.append(header_error)
        return result
    present = set(present_list)
    result.total_rows = len(parsed)

    lists = list(db.scalars(select(BuyerPriceList)).all())
    sheets = list(db.scalars(select(BuyerPriceSheetRow)).all())
    brands = list(db.scalars(select(Brand)).all())
    maps = list(db.scalars(select(BrandBuyerMap)).all())
    by_code, by_name = _index_brands(brands)
    buyers_by_brand: dict[int, set[str]] = {}
    for item in maps:
        buyers_by_brand.setdefault(item.brand_id, set()).add(_text(item.buyer_code).upper())

    by_id = {row.id: row for row in lists if row.id is not None}
    by_key: dict[tuple[int, str, date, date], BuyerPriceList] = {}
    for row in lists:
        if row.brand_id is None or not isinstance(row.period_start, date) or not isinstance(row.period_end, date):
            continue
        by_key[(row.brand_id, _text(row.buyer_code).upper(), row.period_start, row.period_end)] = row

    lines_by_list: dict[int, list[BuyerPriceSheetRow]] = {}
    for sheet in sheets:
        owner = by_id.get(sheet.list_id)
        if owner is None:
            continue
        lines_by_list.setdefault(id(owner), []).append(sheet)

    header_counted: set[int] = set()

    for excel_row, data in parsed:
        item_label = data.get("item") or None
        brand_token = data.get("brand_code", "")
        buyer = _text(data.get("buyer_code")).upper()
        start = _parse_date(data.get("period_start"))
        end = _parse_date(data.get("period_end"))
        if not brand_token:
            result.errors.append(ImportErrorRow(excel_row, item_label, "brand_code wajib diisi"))
            continue
        if not buyer:
            result.errors.append(ImportErrorRow(excel_row, item_label, "buyer_code wajib diisi"))
            continue
        if start is None or end is None:
            result.errors.append(ImportErrorRow(excel_row, item_label, "period_start dan period_end wajib tanggal yang valid"))
            continue
        if end < start:
            result.errors.append(ImportErrorRow(excel_row, item_label, "period_end tidak boleh sebelum period_start"))
            continue
        brand = _resolve_brand(brand_token, by_code, by_name)
        if brand is None:
            result.errors.append(ImportErrorRow(excel_row, item_label, f"Brand {brand_token} tidak ditemukan"))
            continue
        mapped = buyers_by_brand.get(brand.id, set())
        if buyer not in mapped:
            result.errors.append(ImportErrorRow(excel_row, item_label, f"Buyer {buyer} tidak terhubung ke brand {brand.code or brand_token}"))
            continue
        type_error = _validate_types(data, present)
        if type_error:
            result.errors.append(ImportErrorRow(excel_row, item_label, type_error))
            continue

        item_text = data.get("item", "")
        price_error = None
        if item_text:
            for column, _sheet_name in _SHEET_PRICES:
                raw_price = data.get(column, "")
                if raw_price and _price_number(raw_price) is None:
                    price_error = f"Nilai {column} harus angka"
                    break
        elif any(data.get(column) for column, _sheet_name in _SHEET_PRICES):
            price_error = "item wajib diisi jika ada nilai harga"
        if price_error:
            result.errors.append(ImportErrorRow(excel_row, item_label, price_error))
            continue

        list_id = _parse_int(data.get("list_id")) if "list_id" in present else None
        if "list_id" in present and data.get("list_id") and list_id is None:
            result.errors.append(ImportErrorRow(excel_row, item_label, "list_id harus angka"))
            continue

        target = by_id.get(list_id) if list_id is not None else None
        if list_id is not None and target is None:
            result.errors.append(ImportErrorRow(excel_row, item_label, f"List id {list_id} tidak ditemukan"))
            continue
        if target is not None and (target.brand_id != brand.id or _text(target.buyer_code).upper() != buyer):
            result.errors.append(ImportErrorRow(excel_row, item_label, "list_id tidak sesuai dengan brand atau buyer"))
            continue

        key = (brand.id, buyer, start, end)
        created_list = False
        period_changed = False
        if target is None:
            target = by_key.get(key)
        if target is None:
            if _period_overlaps(
                lists,
                brand_id=brand.id,
                buyer_code=buyer,
                period_start=start,
                period_end=end,
                exclude=None,
            ):
                result.errors.append(ImportErrorRow(excel_row, item_label, "Periode tumpang tindih dengan list harga yang sudah ada"))
                continue
            target = BuyerPriceList(
                buyer_code=buyer,
                brand_id=brand.id,
                currency=(data.get("currency") or "USD").upper()[:3] or "USD",
                title=data.get("title") or None,
                period_label=data.get("period_label") or None,
                period_start=start,
                period_end=end,
                outside_price_adjust=_money(data.get("outside_price_adjust")) if "outside_price_adjust" in present else None,
                outside_adjust_type=_adjust_type(data.get("outside_adjust_type")) if "outside_adjust_type" in present else "nominal",
                discount_modifier_enabled=_parse_bool(data.get("discount_modifier_enabled")) if "discount_modifier_enabled" in present else False,
                discount_modifier=_money(data.get("discount_modifier")) if "discount_modifier" in present else None,
                discount_modifier_type=_adjust_type(data.get("discount_modifier_type")) if "discount_modifier_type" in present else "nominal",
                meta_json={"column_keys": list(STANDARD_COLUMN_KEYS)},
            )
            if target.discount_modifier_enabled is None:
                target.discount_modifier_enabled = False
            db.add(target)
            db.flush()
            lists.append(target)
            by_key[key] = target
            if target.id is not None:
                by_id[target.id] = target
            created_list = True
        else:
            old_key = (
                target.brand_id,
                _text(target.buyer_code).upper(),
                target.period_start,
                target.period_end,
            )
            if old_key != key:
                other = by_key.get(key)
                if other is not None and other is not target:
                    result.errors.append(ImportErrorRow(excel_row, item_label, "Periode sudah dipakai list lain"))
                    continue
                if _period_overlaps(
                    lists,
                    brand_id=brand.id,
                    buyer_code=buyer,
                    period_start=start,
                    period_end=end,
                    exclude=target,
                ):
                    result.errors.append(ImportErrorRow(excel_row, item_label, "Periode tumpang tindih dengan list harga yang sudah ada"))
                    continue
                by_key.pop(old_key, None)
                target.period_start = start
                target.period_end = end
                by_key[key] = target
                period_changed = True

        header_changed = period_changed or (not created_list and _apply_header(target, data, present))
        if not item_text:
            _count(
                result,
                created=created_list,
                updated=header_changed and id(target) not in header_counted,
            )
            header_counted.add(id(target))
            continue

        cells = _build_cells(data)
        line_key = _sheet_key(cells)
        bucket = lines_by_list.setdefault(id(target), [])
        match = next((line for line in bucket if _sheet_key(line.cells or {}) == line_key), None)
        if match is None:
            numbers = [line.row_number for line in bucket if isinstance(line.row_number, int)]
            row_number = (max(numbers) + 1) if numbers else 0
            sheet = BuyerPriceSheetRow(list_id=target.id, row_number=row_number, cells=cells)
            db.add(sheet)
            bucket.append(sheet)
            result.created += 1
            header_counted.add(id(target))
            continue

        same_cells = _comparable_cells(match.cells) == _comparable_cells(cells)
        if not same_cells:
            merged = dict(match.cells or {})
            merged.update(cells)
            match.cells = merged
        _count(
            result,
            created=False,
            updated=(not same_cells) or (header_changed and id(target) not in header_counted),
        )
        if header_changed:
            header_counted.add(id(target))
    return result
