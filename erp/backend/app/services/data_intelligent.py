"""Peta relasi master dan isian unmapping untuk workbook Data Intelligent."""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Brand,
    BrandBuyerMap,
    BuyerCode,
    BuyerPriceList,
    ColorSwatchRecipe,
    CsiColorMaster,
    Customer,
    CustomerBrandMapping,
    FluorescentColorMaster,
    MasterItem,
    MaterialAlias,
)
from app.services.swatch_factory_map import (
    CUSTOMER_CODE_ALIASES,
    DEFAULT_FACTORY_CUSTOMER_CODE,
    extract_customer_code_from_color_standard,
)

SHEET_ORDER = (
    "Panduan",
    "Ringkasan",
    "Customer",
    "Brand",
    "Buyer",
    "Map_Customer_Brand",
    "Map_Brand_Buyer",
    "CSI",
    "CSI_Baru",
    "Swatch_Tanpa_CSI",
    "Nama_Beda",
    "Item_Kode_Kosong",
    "Item_Tanpa_CSI",
    "Fluorescent_Tanpa_CSI",
    "Alias_Pabrik",
    "Material_Swatch_Tanpa_Item",
    "Item_Tanpa_Swatch",
    "Buyer_Tanpa_Harga",
)

EDITABLE_COLUMNS: dict[str, frozenset[str]] = {
    "Map_Customer_Brand": frozenset({"brand_code"}),
    "Map_Brand_Buyer": frozenset({"buyer_code"}),
    "CSI_Baru": frozenset({"buyer_code", "color_code", "color_name"}),
    "Swatch_Tanpa_CSI": frozenset({"csi_color_code"}),
    "Nama_Beda": frozenset({"keputusan"}),
    "Item_Kode_Kosong": frozenset({"color_code"}),
}

HEADERS: dict[str, tuple[str, ...]] = {
    "Panduan": ("petunjuk",),
    "Ringkasan": ("kode", "jumlah", "keterangan"),
    "Customer": ("customer_id", "code", "name", "is_active"),
    "Brand": ("brand_id", "code", "name"),
    "Buyer": ("buyer_code", "name", "is_active"),
    "Map_Customer_Brand": ("customer_id", "customer_code", "customer_name", "brand_code"),
    "Map_Brand_Buyer": ("brand_id", "brand_code", "brand_name", "buyer_code"),
    "CSI": ("csi_id", "buyer_code", "color_code", "color_name", "is_active"),
    "CSI_Baru": ("buyer_code", "color_code", "color_name"),
    "Swatch_Tanpa_CSI": (
        "recipe_id",
        "factory_key",
        "buyer_code",
        "color_code",
        "color_name",
        "material_name",
        "emboss",
        "csi_color_code",
    ),
    "Nama_Beda": (
        "recipe_id",
        "buyer_code",
        "color_code",
        "swatch_name",
        "csi_id",
        "csi_name",
        "keputusan",
    ),
    "Item_Kode_Kosong": (
        "item_id",
        "item_code",
        "material_name",
        "color",
        "emboss",
        "thickness_range",
        "color_code",
    ),
    "Item_Tanpa_CSI": ("item_id", "item_code", "color_code", "color", "material_name"),
    "Fluorescent_Tanpa_CSI": ("id", "buyer_code", "color_code", "color_name", "season_label"),
    "Alias_Pabrik": ("jenis", "kunci", "kode_customer", "catatan"),
    "Material_Swatch_Tanpa_Item": ("material_name", "canonical_name", "recipe_count"),
    "Item_Tanpa_Swatch": ("item_id", "item_code", "color_code", "color", "material_name"),
    "Buyer_Tanpa_Harga": ("buyer_code", "name"),
    "Match_Customer_Brand": ("customer_code", "customer_name", "brand_code", "brand_name"),
    "Match_Brand_Buyer": ("brand_code", "brand_name", "buyer_code"),
    "Match_Swatch_CSI": ("factory_key", "buyer_code", "color_code", "color_name", "csi_color_code", "csi_color_name"),
    "Match_Nama_Sama": ("buyer_code", "color_code", "swatch_name", "csi_name"),
    "Match_Item_Kode": ("item_code", "color_code", "color", "material_name"),
    "Match_Item_CSI": ("item_code", "color_code", "color", "material_name"),
    "Match_Fluorescent_CSI": ("buyer_code", "color_code", "color_name", "season_label"),
    "Match_Material_Item": ("material_name", "canonical_name", "recipe_count"),
    "Match_Item_Swatch": ("item_code", "color_code", "color", "material_name"),
    "Match_Buyer_Harga": ("buyer_code", "name"),
}

COUNT_NOTES = {
    "customer_tanpa_brand": "Customer belum punya brand. Isi brand_code di Map_Customer_Brand.",
    "brand_tanpa_buyer": "Brand belum punya buyer. Isi buyer_code di Map_Brand_Buyer.",
    "swatch_tanpa_csi": "Resep aktif belum tertaut CSI. Isi csi_color_code, atau tambah CSI di CSI_Baru.",
    "nama_beda": "Nama swatch dan CSI beda jauh. Isi keputusan BIARKAN, IKUT_CSI, atau IKUT_SWATCH.",
    "item_kode_kosong": "Master item belum punya color_code. Isi kolom color_code.",
    "item_tanpa_csi": "Kode warna item tidak ada di CSI buyer mana pun. Tambah di CSI_Baru bila resmi.",
    "fluorescent_tanpa_csi": "Kode fluorescent tidak ada di CSI buyer yang sama. Acuan saja.",
    "material_swatch_tanpa_item": "Material resep tidak ketemu di master item. Acuan saja.",
    "item_tanpa_swatch": "Kode warna item tidak punya resep swatch aktif. Resep baru dibuat di menu Color Swatch.",
    "buyer_tanpa_harga": "Buyer belum punya daftar harga. Harga diisi di menu Buyer Price.",
}

_FILL_EDIT = PatternFill("solid", fgColor="FFF4CC")
_FONT_HEADER = Font(bold=True)
_FLUORESCENT_RE = re.compile(r"\s*\(\s*FLUORESCENT\s*\)|\s+FLUORESCENT\b", re.IGNORECASE)
_SEASON_RE = re.compile(r"\s*S\d{2}$", re.IGNORECASE)
_DECISIONS = {"BIARKAN", "IKUT_CSI", "IKUT_SWATCH"}
_CSI_BARU_BLANK_ROWS = 40
_READONLY_CHECK = (
    "Customer",
    "Brand",
    "Buyer",
    "CSI",
    "Item_Tanpa_CSI",
    "Fluorescent_Tanpa_CSI",
    "Alias_Pabrik",
    "Material_Swatch_Tanpa_Item",
    "Item_Tanpa_Swatch",
    "Buyer_Tanpa_Harga",
)
_READONLY_KEYS = {
    "Customer": "customer_id",
    "Brand": "brand_id",
    "Buyer": "buyer_code",
    "CSI": "csi_id",
    "Item_Tanpa_CSI": "item_id",
    "Fluorescent_Tanpa_CSI": "id",
    "Alias_Pabrik": ("jenis", "kunci", "catatan"),
    "Material_Swatch_Tanpa_Item": "material_name",
    "Item_Tanpa_Swatch": "item_id",
    "Buyer_Tanpa_Harga": "buyer_code",
}


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_name(text: str | None, *, strip_season: bool = False) -> str:
    cleaned = _FLUORESCENT_RE.sub("", text or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip().upper()
    if strip_season:
        cleaned = _SEASON_RE.sub("", cleaned).strip()
    return cleaned


def name_gap_kind(swatch_name: str | None, csi_name: str | None) -> str:
    """same = hanya beda fluorescent/musim, prefix = singkatan, different = nama lain."""
    left = _norm_name(swatch_name)
    right = _norm_name(csi_name, strip_season=True)
    if left == right:
        return "same"
    if left and right and (left.startswith(right) or right.startswith(left)):
        return "prefix"
    return "different"


def _norm_key(name: str | None) -> str:
    return re.sub(r"\s+", " ", (name or "").strip()).upper()


def _active_flag(value: Any) -> str:
    return "1" if bool(value) else "0"


def collect_snapshot(db: Session) -> dict[str, list[dict[str, Any]]]:
    customers = list(db.scalars(select(Customer).order_by(Customer.code.asc(), Customer.id.asc())).all())
    brands = list(db.scalars(select(Brand).order_by(Brand.code.asc(), Brand.id.asc())).all())
    buyers = list(db.scalars(select(BuyerCode).order_by(BuyerCode.code.asc())).all())
    mappings = list(db.scalars(select(CustomerBrandMapping)).all())
    brand_buyers = list(db.scalars(select(BrandBuyerMap)).all())
    csi_rows = list(db.scalars(select(CsiColorMaster).order_by(CsiColorMaster.buyer_code.asc(), CsiColorMaster.color_code.asc())).all())
    swatches = list(
        db.scalars(
            select(ColorSwatchRecipe)
            .where(ColorSwatchRecipe.is_active.is_(True))
            .order_by(ColorSwatchRecipe.id.asc())
        ).all()
    )
    items = list(db.scalars(select(MasterItem).order_by(MasterItem.item_code.asc(), MasterItem.id.asc())).all())
    fluorescents = list(
        db.scalars(
            select(FluorescentColorMaster)
            .where(FluorescentColorMaster.is_active.is_(True))
            .order_by(FluorescentColorMaster.id.asc())
        ).all()
    )
    aliases = list(db.scalars(select(MaterialAlias).where(MaterialAlias.is_active.is_(True))).all())
    price_buyers = {
        _cell(code).upper()
        for code in db.scalars(select(BuyerPriceList.buyer_code)).all()
    }
    customers_by_id = {int(row.id): row for row in customers}
    brands_by_id = {int(row.id): row for row in brands}

    mapped_customers = {int(row.customer_id) for row in mappings}
    brands_with_buyer = {int(row.brand_id) for row in brand_buyers}
    csi_codes = {
        _cell(row.color_code).upper()
        for row in csi_rows
        if row.is_active and _cell(row.color_code)
    }
    csi_by_buyer = {
        (_cell(row.buyer_code).upper(), _cell(row.color_code).upper())
        for row in csi_rows
        if row.is_active and _cell(row.color_code)
    }
    swatch_codes = {_cell(row.color_code).upper() for row in swatches if _cell(row.color_code)}
    alias_map = {_norm_key(row.alias_name): row.canonical_name for row in aliases}
    item_materials = {_norm_key(row.material_name) for row in items if _cell(row.material_name)}

    customer_rows = [
        {
            "customer_id": int(row.id),
            "code": _cell(row.code),
            "name": _cell(row.name),
            "is_active": _active_flag(row.is_active),
        }
        for row in customers
    ]
    brand_rows = [
        {"brand_id": int(row.id), "code": _cell(row.code), "name": _cell(row.name)}
        for row in brands
    ]
    buyer_rows = [
        {
            "buyer_code": _cell(row.code),
            "name": _cell(row.name),
            "is_active": _active_flag(row.is_active),
        }
        for row in buyers
    ]
    map_customer = [
        {
            "customer_id": int(row.id),
            "customer_code": _cell(row.code),
            "customer_name": _cell(row.name),
            "brand_code": "",
        }
        for row in customers
        if int(row.id) not in mapped_customers
    ]
    map_brand = [
        {
            "brand_id": int(row.id),
            "brand_code": _cell(row.code),
            "brand_name": _cell(row.name),
            "buyer_code": "",
        }
        for row in brands
        if int(row.id) not in brands_with_buyer
    ]
    csi_sheet = [
        {
            "csi_id": int(row.id),
            "buyer_code": _cell(row.buyer_code),
            "color_code": _cell(row.color_code),
            "color_name": _cell(row.color_name),
            "is_active": _active_flag(row.is_active),
        }
        for row in csi_rows
    ]
    swatch_gap = [
        {
            "recipe_id": int(row.id),
            "factory_key": _cell(row.factory_key),
            "buyer_code": _cell(row.buyer_code),
            "color_code": _cell(row.color_code),
            "color_name": _cell(row.color_name),
            "material_name": _cell(row.material_name),
            "emboss": _cell(row.emboss),
            "csi_color_code": "",
        }
        for row in swatches
        if row.csi_color_master_id is None
    ]
    csi_by_id = {int(row.id): row for row in csi_rows}
    nama_beda = []
    for row in swatches:
        if not row.csi_color_master_id:
            continue
        csi = csi_by_id.get(int(row.csi_color_master_id))
        if csi is None or name_gap_kind(row.color_name, csi.color_name) != "different":
            continue
        nama_beda.append(
            {
                "recipe_id": int(row.id),
                "buyer_code": _cell(row.buyer_code),
                "color_code": _cell(row.color_code),
                "swatch_name": _cell(row.color_name),
                "csi_id": int(csi.id),
                "csi_name": _cell(csi.color_name),
                "keputusan": "",
            }
        )
    item_kosong = [
        {
            "item_id": int(row.id),
            "item_code": _cell(row.item_code),
            "material_name": _cell(row.material_name),
            "color": _cell(row.color),
            "emboss": _cell(row.emboss),
            "thickness_range": _cell(row.thickness_range),
            "color_code": "",
        }
        for row in items
        if not _cell(row.color_code)
    ]
    item_tanpa_csi = [
        {
            "item_id": int(row.id),
            "item_code": _cell(row.item_code),
            "color_code": _cell(row.color_code),
            "color": _cell(row.color),
            "material_name": _cell(row.material_name),
        }
        for row in items
        if _cell(row.color_code) and _cell(row.color_code).upper() not in csi_codes
    ]
    fluoro_gap = [
        {
            "id": int(row.id),
            "buyer_code": _cell(row.buyer_code),
            "color_code": _cell(row.color_code),
            "color_name": _cell(row.color_name),
            "season_label": _cell(row.season_label),
        }
        for row in fluorescents
        if (_cell(row.buyer_code).upper(), _cell(row.color_code).upper()) not in csi_by_buyer
    ]
    alias_rows: list[dict[str, Any]] = [
        {
            "jenis": "pabrik",
            "kunci": key,
            "kode_customer": code,
            "catatan": "",
        }
        for key, code in sorted(DEFAULT_FACTORY_CUSTOMER_CODE.items())
    ]
    alias_rows.extend(
        {
            "jenis": "alias",
            "kunci": src,
            "kode_customer": dst,
            "catatan": "kode di color_standard",
        }
        for src, dst in sorted(CUSTOMER_CODE_ALIASES.items())
    )
    for row in swatches:
        resolved = extract_customer_code_from_color_standard(row.color_standard)
        customer = customers_by_id.get(int(row.customer_id)) if row.customer_id else None
        stored = _cell(customer.code).upper() if customer else ""
        if resolved and stored and resolved.upper() != stored:
            alias_rows.append(
                {
                    "jenis": "swatch_beda",
                    "kunci": str(int(row.id)),
                    "kode_customer": stored,
                    "catatan": f"{_cell(row.color_standard)} -> {resolved}",
                }
            )
    material_counts: dict[str, dict[str, Any]] = {}
    for row in swatches:
        raw_name = _cell(row.material_name)
        if not raw_name:
            continue
        canonical = alias_map.get(_norm_key(raw_name), raw_name)
        if _norm_key(canonical) in item_materials:
            continue
        bucket = material_counts.setdefault(
            raw_name,
            {"material_name": raw_name, "canonical_name": _cell(canonical), "recipe_count": 0},
        )
        bucket["recipe_count"] += 1
    item_tanpa_swatch = [
        {
            "item_id": int(row.id),
            "item_code": _cell(row.item_code),
            "color_code": _cell(row.color_code),
            "color": _cell(row.color),
            "material_name": _cell(row.material_name),
        }
        for row in items
        if _cell(row.color_code) and _cell(row.color_code).upper() not in swatch_codes
    ]
    buyer_tanpa_harga = [
        {"buyer_code": _cell(row.code), "name": _cell(row.name)}
        for row in buyers
        if _cell(row.code).upper() not in price_buyers
    ]
    match_customer = []
    for link in mappings:
        customer = customers_by_id.get(int(link.customer_id))
        brand = brands_by_id.get(int(link.brand_id))
        if customer is None or brand is None:
            continue
        match_customer.append(
            {
                "customer_code": _cell(customer.code),
                "customer_name": _cell(customer.name),
                "brand_code": _cell(brand.code),
                "brand_name": _cell(brand.name),
            }
        )
    match_customer.sort(key=lambda row: (row["customer_code"], row["brand_code"]))
    match_brand = []
    for link in brand_buyers:
        brand = brands_by_id.get(int(link.brand_id))
        if brand is None:
            continue
        match_brand.append(
            {
                "brand_code": _cell(brand.code),
                "brand_name": _cell(brand.name),
                "buyer_code": _cell(link.buyer_code),
            }
        )
    match_brand.sort(key=lambda row: (row["brand_code"], row["buyer_code"]))
    match_swatch = []
    match_nama = []
    for row in swatches:
        if not row.csi_color_master_id:
            continue
        csi = csi_by_id.get(int(row.csi_color_master_id))
        if csi is None:
            continue
        match_swatch.append(
            {
                "factory_key": _cell(row.factory_key),
                "buyer_code": _cell(row.buyer_code),
                "color_code": _cell(row.color_code),
                "color_name": _cell(row.color_name),
                "csi_color_code": _cell(csi.color_code),
                "csi_color_name": _cell(csi.color_name),
            }
        )
        if name_gap_kind(row.color_name, csi.color_name) != "different":
            match_nama.append(
                {
                    "buyer_code": _cell(row.buyer_code),
                    "color_code": _cell(row.color_code),
                    "swatch_name": _cell(row.color_name),
                    "csi_name": _cell(csi.color_name),
                }
            )
    match_item_kode = [
        {
            "item_code": _cell(row.item_code),
            "color_code": _cell(row.color_code),
            "color": _cell(row.color),
            "material_name": _cell(row.material_name),
        }
        for row in items
        if _cell(row.color_code)
    ]
    match_item_csi = [
        row for row in match_item_kode if row["color_code"].upper() in csi_codes
    ]
    match_item_swatch = [
        row for row in match_item_kode if row["color_code"].upper() in swatch_codes
    ]
    match_fluoro = [
        {
            "buyer_code": _cell(row.buyer_code),
            "color_code": _cell(row.color_code),
            "color_name": _cell(row.color_name),
            "season_label": _cell(row.season_label),
        }
        for row in fluorescents
        if (_cell(row.buyer_code).upper(), _cell(row.color_code).upper()) in csi_by_buyer
    ]
    material_ok: dict[str, dict[str, Any]] = {}
    for row in swatches:
        raw_name = _cell(row.material_name)
        if not raw_name:
            continue
        canonical = alias_map.get(_norm_key(raw_name), raw_name)
        if _norm_key(canonical) not in item_materials:
            continue
        bucket = material_ok.setdefault(
            raw_name,
            {"material_name": raw_name, "canonical_name": _cell(canonical), "recipe_count": 0},
        )
        bucket["recipe_count"] += 1
    match_buyer_harga = [
        {"buyer_code": _cell(row.code), "name": _cell(row.name)}
        for row in buyers
        if _cell(row.code).upper() in price_buyers
    ]
    sheets = {
        "Panduan": [{"petunjuk": line} for line in _guide_lines()],
        "Customer": customer_rows,
        "Brand": brand_rows,
        "Buyer": buyer_rows,
        "Map_Customer_Brand": map_customer,
        "Map_Brand_Buyer": map_brand,
        "CSI": csi_sheet,
        "CSI_Baru": [],
        "Swatch_Tanpa_CSI": swatch_gap,
        "Nama_Beda": nama_beda,
        "Item_Kode_Kosong": item_kosong,
        "Item_Tanpa_CSI": item_tanpa_csi,
        "Fluorescent_Tanpa_CSI": fluoro_gap,
        "Alias_Pabrik": alias_rows,
        "Material_Swatch_Tanpa_Item": sorted(material_counts.values(), key=lambda row: row["material_name"]),
        "Item_Tanpa_Swatch": item_tanpa_swatch,
        "Buyer_Tanpa_Harga": buyer_tanpa_harga,
        "Match_Customer_Brand": match_customer,
        "Match_Brand_Buyer": match_brand,
        "Match_Swatch_CSI": match_swatch,
        "Match_Nama_Sama": match_nama,
        "Match_Item_Kode": match_item_kode,
        "Match_Item_CSI": match_item_csi,
        "Match_Fluorescent_CSI": match_fluoro,
        "Match_Material_Item": sorted(material_ok.values(), key=lambda row: row["material_name"]),
        "Match_Item_Swatch": match_item_swatch,
        "Match_Buyer_Harga": match_buyer_harga,
    }
    counts = {
        "customer_tanpa_brand": len(map_customer),
        "brand_tanpa_buyer": len(map_brand),
        "swatch_tanpa_csi": len(swatch_gap),
        "nama_beda": len(nama_beda),
        "item_kode_kosong": len(item_kosong),
        "item_tanpa_csi": len(item_tanpa_csi),
        "fluorescent_tanpa_csi": len(fluoro_gap),
        "material_swatch_tanpa_item": len(material_counts),
        "item_tanpa_swatch": len(item_tanpa_swatch),
        "buyer_tanpa_harga": len(buyer_tanpa_harga),
        "match_customer_brand": len(match_customer),
        "match_brand_buyer": len(match_brand),
        "match_swatch_csi": len(match_swatch),
        "match_nama_sama": len(match_nama),
        "match_item_kode": len(match_item_kode),
        "match_item_csi": len(match_item_csi),
        "match_fluorescent_csi": len(match_fluoro),
        "match_material_item": len(material_ok),
        "match_item_swatch": len(match_item_swatch),
        "match_buyer_harga": len(match_buyer_harga),
    }
    sheets["Ringkasan"] = [
        {"kode": key, "jumlah": counts[key], "keterangan": COUNT_NOTES[key]}
        for key in COUNT_NOTES
    ]
    return {"counts": counts, "sheets": sheets}


def summary(db: Session) -> dict[str, int]:
    return dict(collect_snapshot(db)["counts"])


WEB_SHEETS = (
    "Map_Customer_Brand",
    "Map_Brand_Buyer",
    "Swatch_Tanpa_CSI",
    "CSI_Baru",
    "Nama_Beda",
    "Item_Kode_Kosong",
)

REFERENCE_SHEETS = (
    "Item_Tanpa_CSI",
    "Fluorescent_Tanpa_CSI",
    "Alias_Pabrik",
    "Material_Swatch_Tanpa_Item",
    "Item_Tanpa_Swatch",
    "Buyer_Tanpa_Harga",
)

MATCHED_SHEETS = (
    "Match_Customer_Brand",
    "Match_Brand_Buyer",
    "Match_Swatch_CSI",
    "Match_Nama_Sama",
    "Match_Item_Kode",
    "Match_Item_CSI",
    "Match_Fluorescent_CSI",
    "Match_Material_Item",
    "Match_Item_Swatch",
    "Match_Buyer_Harga",
)


def workspace(db: Session) -> dict[str, Any]:
    snapshot = collect_snapshot(db)
    sheets = snapshot["sheets"]
    return {
        "counts": snapshot["counts"],
        "brands": [{"code": row["code"], "name": row["name"]} for row in sheets["Brand"] if row["code"]],
        "buyers": [
            {"code": row["buyer_code"], "name": row["name"]}
            for row in sheets["Buyer"]
            if row["is_active"] == "1"
        ],
        "csi": [
            {
                "buyer_code": row["buyer_code"],
                "color_code": row["color_code"],
                "color_name": row["color_name"],
            }
            for row in sheets["CSI"]
            if row["is_active"] == "1" and row["color_code"]
        ],
        "sheets": {name: sheets.get(name) or [] for name in WEB_SHEETS if name != "CSI_Baru"},
    }


def gap_page(db: Session, sheet: str, *, page: int = 1, page_size: int = 50, q: str = "") -> dict[str, Any]:
    if sheet not in REFERENCE_SHEETS and sheet not in MATCHED_SHEETS:
        raise ValueError("Sheet acuan tidak dikenal")
    rows = list(collect_snapshot(db)["sheets"].get(sheet) or [])
    needle = _cell(q).upper()
    if needle:
        rows = [
            row
            for row in rows
            if needle in " ".join(_cell(value).upper() for value in row.values())
        ]
    size = min(max(int(page_size or 50), 1), 100)
    current = max(int(page or 1), 1)
    start = (current - 1) * size
    return {
        "sheet": sheet,
        "total": len(rows),
        "page": current,
        "page_size": size,
        "columns": list(HEADERS[sheet]),
        "rows": rows[start : start + size],
    }


def sheets_from_web(body: dict[str, Any] | None) -> dict[str, list[dict[str, Any]]]:
    parsed: dict[str, list[dict[str, Any]]] = {}
    for name, rows in (body or {}).get("sheets", body or {}).items():
        if name not in WEB_SHEETS or not isinstance(rows, list):
            continue
        built = []
        for index, row in enumerate(rows, start=2):
            if not isinstance(row, dict):
                continue
            record = {header: _cell(row.get(header)) for header in HEADERS[name]}
            record["__row__"] = index
            built.append(record)
        parsed[name] = built
    return parsed


def apply_web(db: Session, body: dict[str, Any] | None) -> dict[str, list[dict[str, Any]]]:
    return apply_parsed(db, sheets_from_web(body))


def _guide_lines() -> list[str]:
    return [
        "Data Intelligent memetakan relasi master dari customer dan buyer sampai item, CSI, dan swatch.",
        "Sel kuning boleh diisi. Sel lain terkunci dan tidak boleh diubah.",
        "Map_Customer_Brand: isi brand_code yang sudah ada untuk customer yang belum punya brand.",
        "Map_Brand_Buyer: isi buyer_code yang sudah ada hanya jika brand belum punya buyer.",
        "Swatch_Tanpa_CSI: isi csi_color_code milik buyer yang sama. Resep baru tidak dibuat di sini.",
        "CSI_Baru: tambah buyer_code, color_code, dan color_name hanya jika pasangan itu belum ada.",
        "Nama_Beda: isi BIARKAN, IKUT_CSI, atau IKUT_SWATCH. Kosong berarti tidak diubah.",
        "Item_Kode_Kosong: isi color_code saja. Material, warna, emboss, dan ketebalan tidak boleh berubah.",
        "Sheet tanpa sel kuning hanya acuan. Harga, identitas item, dan resep swatch baru tidak ditulis dari file ini.",
    ]


def render_workbook(snapshot: dict[str, Any]) -> bytes:
    sheets: dict[str, list[dict[str, Any]]] = snapshot["sheets"]
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name in SHEET_ORDER:
        headers = HEADERS[name]
        rows = list(sheets.get(name) or [])
        if name == "CSI_Baru":
            rows = [{} for _ in range(_CSI_BARU_BLANK_ROWS)]
        _write_sheet(wb, name, headers, rows, EDITABLE_COLUMNS.get(name, frozenset()))
    if "Nama_Beda" in wb.sheetnames:
        ws = wb["Nama_Beda"]
        col = HEADERS["Nama_Beda"].index("keputusan") + 1
        letter = get_column_letter(col)
        validation = DataValidation(type="list", formula1='"BIARKAN,IKUT_CSI,IKUT_SWATCH"', allow_blank=True)
        validation.error = "Isi BIARKAN, IKUT_CSI, atau IKUT_SWATCH"
        validation.errorTitle = "Keputusan"
        ws.add_data_validation(validation)
        last = max(ws.max_row, 2)
        validation.add(f"{letter}2:{letter}{last}")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_sheet(wb: Workbook, name: str, headers: tuple[str, ...], rows: list[dict[str, Any]], editable: frozenset[str]) -> None:
    ws = wb.create_sheet(name)
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = _FONT_HEADER
        cell.protection = Protection(locked=True)
    editable_idx = {index for index, header in enumerate(headers) if header in editable}
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    last_row = max(ws.max_row, 1)
    for cells in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(headers)):
        for cell in cells:
            unlocked = (cell.column - 1) in editable_idx
            cell.protection = Protection(locked=not unlocked)
            if unlocked:
                cell.fill = _FILL_EDIT
    for index, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 14)
        ws.column_dimensions[get_column_letter(index)].width = min(width, 42)
    ws.protection.sheet = True
    ws.protection.enable()
    ws.freeze_panes = "A2"


def export_workbook(db: Session) -> bytes:
    return render_workbook(collect_snapshot(db))


def parse_workbook(raw: bytes) -> dict[str, list[dict[str, Any]]]:
    try:
        wb = load_workbook(BytesIO(raw), data_only=True)
    except Exception as exc:
        raise ValueError("File Excel tidak dapat dibaca") from exc
    parsed: dict[str, list[dict[str, Any]]] = {}
    for name in wb.sheetnames:
        if name not in HEADERS:
            continue
        ws = wb[name]
        headers = [_cell(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        expected = HEADERS[name]
        if tuple(headers[: len(expected)]) != expected:
            parsed[name] = [{"__header_error__": True}]
            continue
        rows = []
        for offset, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            record = {header: _cell(values[index] if index < len(values) else "") for index, header in enumerate(expected)}
            record["__row__"] = offset
            if name != "Panduan" and not any(record[header] for header in expected):
                continue
            rows.append(record)
        parsed[name] = rows
    return parsed


def _reject(bucket: list[dict[str, Any]], sheet: str, row_no: int, reason: str) -> None:
    bucket.append({"sheet": sheet, "row": row_no, "reason": reason})


def _locked_mismatch(record: dict[str, Any], current: dict[str, Any], editable: frozenset[str]) -> bool:
    for key, expected in current.items():
        if key in editable:
            continue
        if _cell(record.get(key)) != _cell(expected):
            return True
    return False


def _row_key(record: dict[str, Any], key: str | tuple[str, ...]) -> str:
    if isinstance(key, tuple):
        return "|".join(_cell(record.get(part)) for part in key)
    return _cell(record.get(key))


def apply_parsed(
    db: Session,
    parsed: dict[str, list[dict[str, Any]]],
    snapshot: dict[str, Any] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    current = snapshot or collect_snapshot(db)
    sheets = current["sheets"]
    applied: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []

    customers = {int(row.id): row for row in db.scalars(select(Customer)).all()}
    brands = list(db.scalars(select(Brand)).all())
    brands_by_id = {int(row.id): row for row in brands}
    brands_by_code = {_cell(row.code).upper(): row for row in brands if _cell(row.code)}
    buyers = {_cell(row.code).upper() for row in db.scalars(select(BuyerCode)).all() if _cell(row.code)}
    pair_keys = {(int(row.customer_id), int(row.brand_id)) for row in db.scalars(select(CustomerBrandMapping)).all()}
    brands_with_buyer = {int(row.brand_id) for row in db.scalars(select(BrandBuyerMap)).all()}
    swatches = {int(row.id): row for row in db.scalars(select(ColorSwatchRecipe)).all()}
    csi_rows = list(db.scalars(select(CsiColorMaster)).all())
    csi_by_id = {int(row.id): row for row in csi_rows}
    csi_index = {
        (_cell(row.buyer_code).upper(), _cell(row.color_code).upper()): row
        for row in csi_rows
        if _cell(row.buyer_code) and _cell(row.color_code)
    }
    items = {int(row.id): row for row in db.scalars(select(MasterItem)).all()}

    _check_readonly(parsed, sheets, rejected)
    _apply_csi_baru(db, parsed.get("CSI_Baru") or [], buyers, csi_index, applied, rejected)
    _apply_customer_brand(
        db,
        parsed.get("Map_Customer_Brand") or [],
        { _cell(row["customer_id"]): row for row in sheets["Map_Customer_Brand"] },
        customers,
        brands_by_code,
        pair_keys,
        applied,
        rejected,
    )
    _apply_brand_buyer(
        db,
        parsed.get("Map_Brand_Buyer") or [],
        {_cell(row["brand_id"]): row for row in sheets["Map_Brand_Buyer"]},
        brands_by_id,
        buyers,
        brands_with_buyer,
        applied,
        rejected,
    )
    _apply_swatch_csi(
        parsed.get("Swatch_Tanpa_CSI") or [],
        {_cell(row["recipe_id"]): row for row in sheets["Swatch_Tanpa_CSI"]},
        swatches,
        csi_index,
        applied,
        rejected,
    )
    _apply_nama_beda(
        parsed.get("Nama_Beda") or [],
        {_cell(row["recipe_id"]): row for row in sheets["Nama_Beda"]},
        swatches,
        csi_by_id,
        applied,
        rejected,
    )
    _apply_item_code(
        parsed.get("Item_Kode_Kosong") or [],
        {_cell(row["item_id"]): row for row in sheets["Item_Kode_Kosong"]},
        items,
        applied,
        rejected,
    )
    return {"applied": applied, "rejected": rejected}


def _header_error(rows: list[dict[str, Any]]) -> bool:
    return bool(rows) and rows[0].get("__header_error__")


def _check_readonly(
    parsed: dict[str, list[dict[str, Any]]],
    sheets: dict[str, list[dict[str, Any]]],
    rejected: list[dict[str, Any]],
) -> None:
    for name in _READONLY_CHECK:
        rows = parsed.get(name)
        if not rows:
            continue
        if _header_error(rows):
            _reject(rejected, name, 1, "Header sheet berubah")
            continue
        key = _READONLY_KEYS[name]
        expected = {_row_key(row, key): row for row in sheets.get(name) or []}
        for record in rows:
            token = _row_key(record, key)
            current = expected.get(token)
            row_no = int(record.get("__row__") or 0)
            if current is None:
                _reject(rejected, name, row_no, "Baris acuan tidak ada di data saat ini")
                continue
            if _locked_mismatch(record, current, frozenset()):
                _reject(rejected, name, row_no, "Kolom terkunci berubah")


def _apply_csi_baru(
    db: Session,
    rows: list[dict[str, Any]],
    buyers: set[str],
    csi_index: dict[tuple[str, str], CsiColorMaster],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "CSI_Baru", 1, "Header sheet berubah")
        return
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        buyer = _cell(record.get("buyer_code")).upper()
        code = _cell(record.get("color_code")).upper()
        name = _cell(record.get("color_name"))
        if not buyer and not code and not name:
            continue
        if not buyer or not code or not name:
            _reject(rejected, "CSI_Baru", row_no, "Buyer, kode warna, dan nama wajib diisi")
            continue
        if buyer not in buyers:
            _reject(rejected, "CSI_Baru", row_no, "Kode buyer tidak ada di master")
            continue
        if len(code) > 64 or len(name) > 255:
            _reject(rejected, "CSI_Baru", row_no, "Kode atau nama warna terlalu panjang")
            continue
        if (buyer, code) in csi_index:
            _reject(rejected, "CSI_Baru", row_no, "CSI buyer dan kode ini sudah ada")
            continue
        created = CsiColorMaster(
            buyer_code=buyer,
            color_code=code,
            color_name=name,
            is_active=True,
            source_file="data-intelligent",
        )
        db.add(created)
        db.flush()
        csi_index[(buyer, code)] = created
        applied.append({"sheet": "CSI_Baru", "row": row_no, "action": "csi_baru"})


def _apply_customer_brand(
    db: Session,
    rows: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    customers: dict[int, Customer],
    brands_by_code: dict[str, Brand],
    pair_keys: set[tuple[int, int]],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "Map_Customer_Brand", 1, "Header sheet berubah")
        return
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        brand_code = _cell(record.get("brand_code"))
        if not brand_code:
            continue
        current = expected.get(_cell(record.get("customer_id")))
        if current is None or _locked_mismatch(record, current, EDITABLE_COLUMNS["Map_Customer_Brand"]):
            _reject(rejected, "Map_Customer_Brand", row_no, "Kolom terkunci berubah")
            continue
        customer = customers.get(int(current["customer_id"]))
        brand = brands_by_code.get(brand_code.upper())
        if customer is None:
            _reject(rejected, "Map_Customer_Brand", row_no, "Customer tidak ditemukan")
            continue
        if brand is None:
            _reject(rejected, "Map_Customer_Brand", row_no, "Kode brand tidak ada di master")
            continue
        pair = (int(customer.id), int(brand.id))
        if pair in pair_keys:
            _reject(rejected, "Map_Customer_Brand", row_no, "Relasi customer-brand sudah ada")
            continue
        db.add(CustomerBrandMapping(customer_id=pair[0], brand_id=pair[1], is_default=False))
        pair_keys.add(pair)
        applied.append({"sheet": "Map_Customer_Brand", "row": row_no, "action": "customer_brand"})


def _apply_brand_buyer(
    db: Session,
    rows: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    brands_by_id: dict[int, Brand],
    buyers: set[str],
    brands_with_buyer: set[int],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "Map_Brand_Buyer", 1, "Header sheet berubah")
        return
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        buyer_code = _cell(record.get("buyer_code")).upper()
        if not buyer_code:
            continue
        brand_token = _cell(record.get("brand_id"))
        if brand_token.isdigit() and int(brand_token) in brands_with_buyer:
            _reject(rejected, "Map_Brand_Buyer", row_no, "Brand sudah punya buyer")
            continue
        current = expected.get(brand_token)
        if current is None or _locked_mismatch(record, current, EDITABLE_COLUMNS["Map_Brand_Buyer"]):
            _reject(rejected, "Map_Brand_Buyer", row_no, "Kolom terkunci berubah")
            continue
        brand_id = int(current["brand_id"])
        if brand_id not in brands_by_id:
            _reject(rejected, "Map_Brand_Buyer", row_no, "Brand tidak ditemukan")
            continue
        if brand_id in brands_with_buyer:
            _reject(rejected, "Map_Brand_Buyer", row_no, "Brand sudah punya buyer")
            continue
        if buyer_code not in buyers:
            _reject(rejected, "Map_Brand_Buyer", row_no, "Kode buyer tidak ada di master")
            continue
        db.add(BrandBuyerMap(brand_id=brand_id, buyer_code=buyer_code))
        brands_with_buyer.add(brand_id)
        applied.append({"sheet": "Map_Brand_Buyer", "row": row_no, "action": "brand_buyer"})


def _apply_swatch_csi(
    rows: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    swatches: dict[int, ColorSwatchRecipe],
    csi_index: dict[tuple[str, str], CsiColorMaster],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "Swatch_Tanpa_CSI", 1, "Header sheet berubah")
        return
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        code = _cell(record.get("csi_color_code")).upper()
        if not code:
            continue
        current = expected.get(_cell(record.get("recipe_id")))
        if current is None or _locked_mismatch(record, current, EDITABLE_COLUMNS["Swatch_Tanpa_CSI"]):
            _reject(rejected, "Swatch_Tanpa_CSI", row_no, "Kolom terkunci berubah")
            continue
        recipe = swatches.get(int(current["recipe_id"]))
        if recipe is None:
            _reject(rejected, "Swatch_Tanpa_CSI", row_no, "Resep tidak ditemukan")
            continue
        if recipe.csi_color_master_id is not None:
            _reject(rejected, "Swatch_Tanpa_CSI", row_no, "Resep sudah tertaut ke CSI")
            continue
        buyer = _cell(recipe.buyer_code).upper()
        csi = csi_index.get((buyer, code))
        if csi is None or not csi.is_active:
            _reject(rejected, "Swatch_Tanpa_CSI", row_no, "Kode CSI tidak ada untuk buyer ini")
            continue
        recipe.csi_color_master_id = int(csi.id)
        applied.append({"sheet": "Swatch_Tanpa_CSI", "row": row_no, "action": "swatch_csi"})


def _apply_nama_beda(
    rows: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    swatches: dict[int, ColorSwatchRecipe],
    csi_by_id: dict[int, CsiColorMaster],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "Nama_Beda", 1, "Header sheet berubah")
        return
    chosen_csi_name: dict[int, str] = {}
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        decision = _cell(record.get("keputusan")).upper()
        current = expected.get(_cell(record.get("recipe_id")))
        if decision and decision not in _DECISIONS:
            _reject(rejected, "Nama_Beda", row_no, "Keputusan harus BIARKAN, IKUT_CSI, atau IKUT_SWATCH")
            continue
        if current is None or _locked_mismatch(record, current, EDITABLE_COLUMNS["Nama_Beda"]):
            _reject(rejected, "Nama_Beda", row_no, "Kolom terkunci berubah")
            continue
        if not decision or decision == "BIARKAN":
            continue
        recipe = swatches.get(int(current["recipe_id"]))
        csi = csi_by_id.get(int(current["csi_id"]))
        if recipe is None or csi is None:
            _reject(rejected, "Nama_Beda", row_no, "Resep atau CSI tidak ditemukan")
            continue
        if decision == "IKUT_CSI":
            recipe.color_name = _cell(csi.color_name)
            applied.append({"sheet": "Nama_Beda", "row": row_no, "action": "ikut_csi"})
            continue
        target = _cell(recipe.color_name)
        previous = chosen_csi_name.get(int(csi.id))
        if previous is not None and previous != target:
            _reject(rejected, "Nama_Beda", row_no, "Nama CSI sudah diubah oleh baris lain pada unggahan ini")
            continue
        csi.color_name = target
        chosen_csi_name[int(csi.id)] = target
        applied.append({"sheet": "Nama_Beda", "row": row_no, "action": "ikut_swatch"})


def _apply_item_code(
    rows: list[dict[str, Any]],
    expected: dict[str, dict[str, Any]],
    items: dict[int, MasterItem],
    applied: list[dict[str, Any]],
    rejected: list[dict[str, Any]],
) -> None:
    if _header_error(rows):
        _reject(rejected, "Item_Kode_Kosong", 1, "Header sheet berubah")
        return
    for record in rows:
        row_no = int(record.get("__row__") or 0)
        code = _cell(record.get("color_code")).upper()
        if not code:
            continue
        current = expected.get(_cell(record.get("item_id")))
        if current is None or _locked_mismatch(record, current, EDITABLE_COLUMNS["Item_Kode_Kosong"]):
            _reject(rejected, "Item_Kode_Kosong", row_no, "Kolom terkunci berubah")
            continue
        item = items.get(int(current["item_id"]))
        if item is None:
            _reject(rejected, "Item_Kode_Kosong", row_no, "Item tidak ditemukan")
            continue
        if _cell(item.color_code):
            _reject(rejected, "Item_Kode_Kosong", row_no, "Kode warna item sudah terisi")
            continue
        if len(code) > 64:
            _reject(rejected, "Item_Kode_Kosong", row_no, "Kode warna terlalu panjang")
            continue
        item.color_code = code
        applied.append({"sheet": "Item_Kode_Kosong", "row": row_no, "action": "item_color_code"})


def apply_workbook(db: Session, raw: bytes) -> dict[str, list[dict[str, Any]]]:
    parsed = parse_workbook(raw)
    return apply_parsed(db, parsed)
