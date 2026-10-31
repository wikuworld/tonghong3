"""Export/import Excel master customer.

Kunci unik impor: code. Baris yang sudah ada diperbarui lewat id atau code.
File yang diunggah ulang tidak menambah duplikat dan tidak menghapus customer
yang tidak ada di file.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Customer, CustomerPic
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "code",
    "name",
    "contact_person",
    "position",
    "email",
    "phone",
    "city",
    "country",
    "tax_id",
    "skep_kb_number",
    "customs_office_code",
    "address",
    "billing_address",
    "remark",
    "is_active",
    "is_export_customer",
    "outside_banten",
    "discount_modifier_enabled",
    "discount_modifier",
    "discount_modifier_type",
    "payment_terms",
    "default_currency",
    "tax_mode",
    "reference_quotation",
    "wilayah_province_code",
    "wilayah_regency_code",
    "wilayah_district_code",
    "wilayah_village_code",
    "created_at",
    "updated_at",
]

TEMPLATE_COLUMNS = [
    name for name in EXPORT_COLUMNS if name not in {"id", "created_at", "updated_at"}
]

_REQUIRED_HEADERS = ("code", "name")

_STRING_FIELDS = (
    "contact_person",
    "position",
    "email",
    "phone",
    "city",
    "country",
    "tax_id",
    "skep_kb_number",
    "customs_office_code",
    "address",
    "billing_address",
    "remark",
    "payment_terms",
    "reference_quotation",
    "wilayah_province_code",
    "wilayah_regency_code",
    "wilayah_district_code",
    "wilayah_village_code",
)

_BOOL_FIELDS = (
    "is_active",
    "is_export_customer",
    "outside_banten",
    "discount_modifier_enabled",
)

_CONTACT_FIELDS = ("contact_person", "position", "phone", "email")

_MAX_LEN = {
    "code": 50,
    "name": 255,
    "contact_person": 255,
    "position": 120,
    "email": 255,
    "phone": 50,
    "city": 120,
    "country": 100,
    "tax_id": 32,
    "skep_kb_number": 80,
    "customs_office_code": 32,
    "payment_terms": 120,
    "reference_quotation": 100,
    "wilayah_province_code": 8,
    "wilayah_regency_code": 16,
    "wilayah_district_code": 24,
    "wilayah_village_code": 32,
}

_HEADER_ALIAS = {
    "id": "id",
    "code": "code",
    "kode": "code",
    "customer_code": "code",
    "name": "name",
    "nama": "name",
    "customer_name": "name",
    "contact_person": "contact_person",
    "pic": "contact_person",
    "contact": "contact_person",
    "position": "position",
    "jabatan": "position",
    "email": "email",
    "phone": "phone",
    "telepon": "phone",
    "telp": "phone",
    "city": "city",
    "kota": "city",
    "country": "country",
    "negara": "country",
    "tax_id": "tax_id",
    "npwp": "tax_id",
    "skep_kb_number": "skep_kb_number",
    "customs_office_code": "customs_office_code",
    "address": "address",
    "alamat": "address",
    "billing_address": "billing_address",
    "remark": "remark",
    "catatan": "remark",
    "is_active": "is_active",
    "active": "is_active",
    "aktif": "is_active",
    "is_export_customer": "is_export_customer",
    "export": "is_export_customer",
    "outside_banten": "outside_banten",
    "discount_modifier_enabled": "discount_modifier_enabled",
    "discount_modifier": "discount_modifier",
    "discount_modifier_type": "discount_modifier_type",
    "payment_terms": "payment_terms",
    "default_currency": "default_currency",
    "currency": "default_currency",
    "tax_mode": "tax_mode",
    "reference_quotation": "reference_quotation",
    "wilayah_province_code": "wilayah_province_code",
    "wilayah_regency_code": "wilayah_regency_code",
    "wilayah_district_code": "wilayah_district_code",
    "wilayah_village_code": "wilayah_village_code",
    "created_at": "created_at",
    "updated_at": "updated_at",
}

_TRUE_VALUES = {"1", "true", "yes", "y", "ya", "aktif", "active"}
_FALSE_VALUES = {"0", "false", "no", "n", "tidak", "nonaktif", "inactive"}
_ADJUST_ALIAS = {
    "nominal": "nominal",
    "amount": "nominal",
    "amt": "nominal",
    "percent": "percent",
    "percentage": "percent",
    "pct": "percent",
    "%": "percent",
}
_TAX_ALIAS = {
    "pkp": "pkp",
    "non_pkp": "non_pkp",
    "non-pkp": "non_pkp",
    "nonpkp": "non_pkp",
    "exempt": "exempt",
    "bebas": "exempt",
    "export": "exempt",
    "ekspor": "exempt",
}


@dataclass
class ImportErrorRow:
    row: int
    code: str | None
    reason: str


@dataclass
class CustomerImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(err) for err in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    text = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in text:
        text = text.replace("__", "_")
    return text


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, bool):
        return "1" if val else "0"
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _blank(raw: str) -> str | None:
    text = _text(raw)
    return text or None


def _parse_int(raw: Any) -> int | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _parse_bool(raw: Any) -> bool | None:
    if isinstance(raw, bool):
        return raw
    text = _cell_str(raw).lower()
    if not text:
        return False
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def _money(raw: Any) -> Decimal | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return Decimal(text).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def _same_money(left: Any, right: Decimal | None) -> bool:
    return _money(left) == right


def _adjust_type(raw: Any) -> str | None:
    text = _text(raw).lower()
    if not text:
        return "nominal"
    return _ADJUST_ALIAS.get(text)


def _tax_mode(raw: Any) -> str | None:
    text = _text(raw).lower().replace(" ", "_")
    if not text:
        return None
    return _TAX_ALIAS.get(text)


def _stamp(value: datetime | None) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return ""


def _xlsx_bytes(headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for r_idx, row in enumerate(data_rows, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _primary_pic(pics: list[CustomerPic]) -> CustomerPic | None:
    if not pics:
        return None
    return sorted(pics, key=lambda pic: pic.id or 0)[0]


def _contact_value(customer: Customer, pic: CustomerPic | None, pic_attr: str, customer_attr: str) -> str:
    if pic is not None and _text(getattr(pic, pic_attr, None)):
        return _text(getattr(pic, pic_attr))
    return _text(getattr(customer, customer_attr, None))


def _export_row(customer: Customer, pic: CustomerPic | None) -> list[Any]:
    discount = (
        float(customer.discount_modifier) if customer.discount_modifier is not None else None
    )
    return [
        customer.id,
        customer.code or "",
        customer.name or "",
        _contact_value(customer, pic, "name", "contact_person"),
        _contact_value(customer, pic, "title", "position"),
        _contact_value(customer, pic, "email", "email"),
        _contact_value(customer, pic, "phone", "phone"),
        customer.city or "",
        customer.country or "",
        customer.tax_id or "",
        customer.skep_kb_number or "",
        customer.customs_office_code or "",
        customer.address or "",
        customer.billing_address or "",
        customer.remark or "",
        1 if customer.is_active else 0,
        1 if customer.is_export_customer else 0,
        1 if customer.outside_banten else 0,
        1 if customer.discount_modifier_enabled else 0,
        discount,
        customer.discount_modifier_type or "",
        customer.payment_terms or "",
        customer.default_currency or "",
        customer.tax_mode or "",
        customer.reference_quotation or "",
        customer.wilayah_province_code or "",
        customer.wilayah_regency_code or "",
        customer.wilayah_district_code or "",
        customer.wilayah_village_code or "",
        _stamp(customer.created_at),
        _stamp(customer.updated_at),
    ]


def export_customers_template_xlsx() -> bytes:
    return _xlsx_bytes(TEMPLATE_COLUMNS, [])


def export_customers_xlsx(db: Session) -> bytes:
    customers = list(db.scalars(select(Customer).order_by(Customer.id.asc())).all())
    pics = list(db.scalars(select(CustomerPic).order_by(CustomerPic.id.asc())).all())
    by_customer: dict[int, list[CustomerPic]] = {}
    for pic in pics:
        if pic.customer_id is None:
            continue
        by_customer.setdefault(pic.customer_id, []).append(pic)
    rows = [_export_row(customer, _primary_pic(by_customer.get(customer.id, []))) for customer in customers]
    return _xlsx_bytes(EXPORT_COLUMNS, rows)


def _read_table(content: bytes) -> tuple[set[str], list[tuple[int, dict[str, str]]], ImportErrorRow | None]:
    if not content:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel kosong")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel tidak bisa dibaca")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel kosong")
    headers = [_HEADER_ALIAS.get(_normalize_header(cell), "") for cell in rows[0]]
    present = {name for name in headers if name}
    if any(name not in present for name in _REQUIRED_HEADERS):
        return set(), [], ImportErrorRow(
            row=1,
            code=None,
            reason="Header wajib memuat kolom code dan name",
        )
    parsed: list[tuple[int, dict[str, str]]] = []
    for offset, raw in enumerate(rows[1:], start=2):
        data: dict[str, str] = {}
        empty = True
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = raw[idx] if idx < len(raw) else None
            text = _cell_str(value)
            if text:
                empty = False
            data[header] = text
        if empty:
            continue
        parsed.append((offset, data))
    return present, parsed, None


def _validate_row(data: dict[str, str], present: set[str], current: Customer | None) -> str | None:
    code = _text(data.get("code"))
    name = _text(data.get("name"))
    if not code:
        return "code wajib diisi"
    if not name:
        return "name wajib diisi"
    for column, limit in _MAX_LEN.items():
        if column in {"code", "name"} or column in present:
            value = code if column == "code" else name if column == "name" else data.get(column, "")
            if len(value) > limit:
                return f"{column} melebihi {limit} karakter"
    for column in _BOOL_FIELDS:
        if column in present and data.get(column) and _parse_bool(data.get(column)) is None:
            return f"Nilai {column} tidak dikenali"
    if "discount_modifier" in present and data.get("discount_modifier") and _money(data.get("discount_modifier")) is None:
        return "Nilai discount_modifier harus angka"
    if "discount_modifier_type" in present and data.get("discount_modifier_type") and _adjust_type(data.get("discount_modifier_type")) is None:
        return "Nilai discount_modifier_type harus nominal atau percent"
    if "tax_mode" in present and data.get("tax_mode") and _tax_mode(data.get("tax_mode")) is None:
        return "Nilai tax_mode harus pkp, non_pkp, atau exempt"
    if "default_currency" in present and data.get("default_currency") and len(_text(data.get("default_currency"))) > 3:
        return "default_currency maksimal 3 karakter"
    kind = (
        _adjust_type(data.get("discount_modifier_type"))
        if "discount_modifier_type" in present
        else ((current.discount_modifier_type if current else None) or "nominal")
    )
    amount = (
        _money(data.get("discount_modifier"))
        if "discount_modifier" in present
        else (current.discount_modifier if current else None)
    )
    if kind == "percent" and amount is not None and (amount < -100 or amount > 100):
        return "Persen discount_modifier harus di antara -100 dan 100"
    return None


def _assign_strings(target: Customer, data: dict[str, str], present: set[str]) -> bool:
    changed = False
    for column in _STRING_FIELDS:
        if column not in present:
            continue
        value = _blank(data.get(column, ""))
        if (getattr(target, column) or None) != value:
            setattr(target, column, value)
            changed = True
    return changed


def _assign_bools(target: Customer, data: dict[str, str], present: set[str], *, creating: bool) -> bool:
    changed = False
    defaults = {
        "is_active": True,
        "is_export_customer": False,
        "outside_banten": False,
        "discount_modifier_enabled": False,
    }
    for column in _BOOL_FIELDS:
        if column not in present:
            if creating:
                setattr(target, column, defaults[column])
            continue
        value = _parse_bool(data.get(column))
        if value is None:
            value = False
        if bool(getattr(target, column)) != value:
            setattr(target, column, value)
            changed = True
    return changed


def _assign_modifiers(target: Customer, data: dict[str, str], present: set[str], *, creating: bool) -> bool:
    changed = False
    if "discount_modifier" in present:
        amount = _money(data.get("discount_modifier"))
        if not _same_money(target.discount_modifier, amount):
            target.discount_modifier = amount
            changed = True
    if "discount_modifier_type" in present or creating:
        kind = _adjust_type(data.get("discount_modifier_type")) if "discount_modifier_type" in present else "nominal"
        if (target.discount_modifier_type or None) != kind:
            target.discount_modifier_type = kind
            changed = True
    if "default_currency" in present:
        currency = _text(data.get("default_currency")).upper()[:3] or None
        if (target.default_currency or None) != currency:
            target.default_currency = currency
            changed = True
    if "tax_mode" in present:
        mode = _tax_mode(data.get("tax_mode"))
        if (target.tax_mode or None) != mode:
            target.tax_mode = mode
            changed = True
    return changed


def _sync_primary_contact(
    db: Session,
    customer: Customer,
    data: dict[str, str],
    present: set[str],
    pics_by_customer: dict[int, list[CustomerPic]],
) -> bool:
    if not any(column in present for column in _CONTACT_FIELDS):
        return False
    bucket = pics_by_customer.setdefault(customer.id, [])
    primary = _primary_pic(bucket)
    name = _blank(data.get("contact_person", "")) if "contact_person" in present else (primary.name if primary else customer.contact_person)
    title = _blank(data.get("position", "")) if "position" in present else (primary.title if primary else customer.position)
    phone = _blank(data.get("phone", "")) if "phone" in present else (primary.phone if primary else customer.phone)
    email = _blank(data.get("email", "")) if "email" in present else (primary.email if primary else customer.email)
    if primary is None:
        if not any((name, title, phone, email)):
            return False
        primary = CustomerPic(customer_id=customer.id, name=name, title=title, phone=phone, email=email)
        db.add(primary)
        bucket.append(primary)
        return True
    changed = False
    for attr, value, column in (
        ("name", name, "contact_person"),
        ("title", title, "position"),
        ("phone", phone, "phone"),
        ("email", email, "email"),
    ):
        if column not in present:
            continue
        if (getattr(primary, attr) or None) != value:
            setattr(primary, attr, value)
            changed = True
    return changed


def import_customers_xlsx(db: Session, content: bytes) -> CustomerImportResult:
    present, parsed, header_error = _read_table(content)
    result = CustomerImportResult()
    if header_error:
        result.errors.append(header_error)
        return result
    result.total_rows = len(parsed)

    customers = list(db.scalars(select(Customer)).all())
    pics = list(db.scalars(select(CustomerPic)).all())
    by_id = {row.id: row for row in customers if row.id is not None}
    by_code: dict[str, Customer] = {}
    for row in customers:
        key = _text(row.code).upper()
        if key and key not in by_code:
            by_code[key] = row
    pics_by_customer: dict[int, list[CustomerPic]] = {}
    for pic in pics:
        if pic.customer_id is None:
            continue
        pics_by_customer.setdefault(pic.customer_id, []).append(pic)

    for excel_row, data in parsed:
        code_label = _blank(data.get("code", ""))
        current_id = _parse_int(data.get("id")) if "id" in present else None
        if "id" in present and data.get("id") and current_id is None:
            result.errors.append(ImportErrorRow(excel_row, code_label, "id harus angka"))
            continue
        target = by_id.get(current_id) if current_id is not None else None
        if current_id is not None and target is None:
            result.errors.append(ImportErrorRow(excel_row, code_label, f"Customer id {current_id} tidak ditemukan"))
            continue
        error = _validate_row(data, present, target)
        if error:
            result.errors.append(ImportErrorRow(excel_row, code_label, error))
            continue

        code = _text(data.get("code"))
        code_key = code.upper()
        owner = by_code.get(code_key)
        if target is None:
            target = owner
        elif owner is not None and owner is not target:
            result.errors.append(ImportErrorRow(excel_row, code_label, "code sudah dipakai customer lain"))
            continue

        creating = target is None
        if creating:
            target = Customer(code=code, name=_text(data.get("name")), is_active=True)
            db.add(target)
            db.flush()
            by_id[target.id] = target
            by_code[code_key] = target
        else:
            old_key = _text(target.code).upper()
            if old_key != code_key:
                by_code.pop(old_key, None)
                target.code = code
                by_code[code_key] = target
            elif (target.code or "") != code and _text(target.code).upper() == code_key:
                pass

        changed = creating
        name = _text(data.get("name"))
        if (target.name or "") != name:
            target.name = name
            changed = True
        changed = _assign_strings(target, data, present) or changed
        changed = _assign_bools(target, data, present, creating=creating) or changed
        changed = _assign_modifiers(target, data, present, creating=creating) or changed
        changed = _sync_primary_contact(db, target, data, present, pics_by_customer) or changed
        if creating:
            result.created += 1
        elif changed:
            target.updated_at = app_now()
            result.updated += 1
        else:
            result.skipped += 1
    return result
