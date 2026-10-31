"""Export/import Excel master brand.

Kunci unik impor: nama brand (disimpan juga sebagai code).
Customer dan buyer pada baris yang sama ikut diperbarui bila kolomnya ada.
Brand yang tidak ada di file tidak dihapus.
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Brand, BrandBuyerMap, Customer, CustomerBrandMapping
from app.services.buyer_registry import ensure_buyer_code
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "name",
    "code",
    "customer_codes",
    "buyer_codes",
    "created_at",
    "updated_at",
]

TEMPLATE_COLUMNS = ["name", "customer_codes", "buyer_codes"]

_REQUIRED_HEADERS = ("name",)

_HEADER_ALIAS = {
    "id": "id",
    "name": "name",
    "nama": "name",
    "brand": "name",
    "brand_name": "name",
    "code": "code",
    "kode": "code",
    "brand_code": "code",
    "customer_codes": "customer_codes",
    "customer_code": "customer_codes",
    "customers": "customer_codes",
    "buyer_codes": "buyer_codes",
    "buyer_code": "buyer_codes",
    "buyers": "buyer_codes",
    "created_at": "created_at",
    "updated_at": "updated_at",
}

_SPLIT_RE = re.compile(r"[,;|\n]+")


@dataclass
class ImportErrorRow:
    row: int
    code: str | None
    reason: str


@dataclass
class BrandImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(err) for err in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    text = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in text:
        text = text.replace("__", "_")
    return text


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, bool):
        return "1" if val else "0"
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _parse_int(raw: Any) -> int | None:
    text = _cell_str(raw)
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        return None


def _stamp(value: datetime | None) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return ""


def _split_tokens(raw: str) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for part in _SPLIT_RE.split(raw or ""):
        token = part.strip()
        if not token:
            continue
        key = token.upper()
        if key in seen:
            continue
        seen.add(key)
        out.append(token)
    return out


def _xlsx_bytes(headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Brands"
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for r_idx, row in enumerate(data_rows, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(r_idx, c_idx, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_brands_template_xlsx() -> bytes:
    return _xlsx_bytes(TEMPLATE_COLUMNS, [])


def export_brands_xlsx(db: Session) -> bytes:
    brands = list(db.scalars(select(Brand).order_by(Brand.id.asc())).all())
    customers = {row.id: (row.code or "") for row in db.scalars(select(Customer)).all()}
    mappings = list(db.scalars(select(CustomerBrandMapping)).all())
    buyers = list(db.scalars(select(BrandBuyerMap)).all())
    customers_by_brand: dict[int, list[str]] = {}
    for mapping in mappings:
        code = customers.get(mapping.customer_id, "")
        if not code:
            continue
        customers_by_brand.setdefault(mapping.brand_id, []).append(code)
    buyers_by_brand: dict[int, list[str]] = {}
    for buyer in buyers:
        if not buyer.buyer_code:
            continue
        buyers_by_brand.setdefault(buyer.brand_id, []).append(buyer.buyer_code)
    rows: list[list[Any]] = []
    for brand in brands:
        rows.append(
            [
                brand.id,
                brand.name or "",
                brand.code or "",
                ", ".join(sorted(set(customers_by_brand.get(brand.id, [])))),
                ", ".join(sorted({code.strip().upper() for code in buyers_by_brand.get(brand.id, []) if code})),
                _stamp(brand.created_at),
                _stamp(brand.updated_at),
            ]
        )
    return _xlsx_bytes(EXPORT_COLUMNS, rows)


def _read_table(content: bytes) -> tuple[set[str], list[tuple[int, dict[str, str]]], ImportErrorRow | None]:
    if not content:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel kosong")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel tidak bisa dibaca")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return set(), [], ImportErrorRow(row=1, code=None, reason="File Excel kosong")
    headers = [_HEADER_ALIAS.get(_normalize_header(cell), "") for cell in rows[0]]
    present = {name for name in headers if name}
    if "name" not in present and "code" in present:
        headers = ["name" if name == "code" else name for name in headers]
        present.add("name")
    if any(name not in present for name in _REQUIRED_HEADERS):
        return set(), [], ImportErrorRow(row=1, code=None, reason="Header wajib memuat kolom name")
    parsed: list[tuple[int, dict[str, str]]] = []
    for offset, raw in enumerate(rows[1:], start=2):
        data: dict[str, str] = {}
        empty = True
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = raw[idx] if idx < len(raw) else None
            text = _cell_str(value)
            if text:
                empty = False
            if header == "name" and header in data and data[header]:
                continue
            data[header] = text
        if empty:
            continue
        parsed.append((offset, data))
    return present, parsed, None


def _sync_customers(
    db: Session,
    brand: Brand,
    tokens: list[str],
    customers_by_code: dict[str, Customer],
    maps_by_brand: dict[int, list[CustomerBrandMapping]],
) -> tuple[bool, str | None]:
    desired: list[int] = []
    for token in tokens:
        customer = customers_by_code.get(token.upper())
        if customer is None:
            return False, f"Customer {token} tidak ditemukan"
        if customer.id not in desired:
            desired.append(customer.id)
    bucket = maps_by_brand.setdefault(brand.id, [])
    current = {mapping.customer_id: mapping for mapping in bucket}
    desired_set = set(desired)
    changed = False
    for customer_id, mapping in list(current.items()):
        if customer_id not in desired_set:
            db.delete(mapping)
            bucket.remove(mapping)
            changed = True
    for customer_id in desired:
        if customer_id in current:
            continue
        mapping = CustomerBrandMapping(customer_id=customer_id, brand_id=brand.id, is_default=False)
        db.add(mapping)
        bucket.append(mapping)
        changed = True
    return changed, None


def _sync_buyers(
    db: Session,
    brand: Brand,
    tokens: list[str],
    maps_by_brand: dict[int, list[BrandBuyerMap]],
) -> bool:
    desired = [token.upper() for token in tokens]
    desired_set = set(desired)
    bucket = maps_by_brand.setdefault(brand.id, [])
    changed = False
    for mapping in list(bucket):
        if _text(mapping.buyer_code).upper() not in desired_set:
            db.delete(mapping)
            bucket.remove(mapping)
            changed = True
    current = {_text(mapping.buyer_code).upper() for mapping in bucket}
    for code in desired:
        if code in current:
            continue
        ensure_buyer_code(db, code)
        mapping = BrandBuyerMap(brand_id=brand.id, buyer_code=code)
        db.add(mapping)
        bucket.append(mapping)
        current.add(code)
        changed = True
    return changed


def import_brands_xlsx(db: Session, content: bytes) -> BrandImportResult:
    present, parsed, header_error = _read_table(content)
    result = BrandImportResult()
    if header_error:
        result.errors.append(header_error)
        return result
    result.total_rows = len(parsed)

    brands = list(db.scalars(select(Brand)).all())
    customers = list(db.scalars(select(Customer)).all())
    customer_maps = list(db.scalars(select(CustomerBrandMapping)).all())
    buyer_maps = list(db.scalars(select(BrandBuyerMap)).all())
    by_id = {row.id: row for row in brands if row.id is not None}
    by_name: dict[str, Brand] = {}
    for row in brands:
        for token in (row.code, row.name):
            key = _text(token).upper()
            if key and key not in by_name:
                by_name[key] = row
    customers_by_code = {_text(row.code).upper(): row for row in customers if _text(row.code)}
    customers_by_brand: dict[int, list[CustomerBrandMapping]] = {}
    for mapping in customer_maps:
        customers_by_brand.setdefault(mapping.brand_id, []).append(mapping)
    buyers_by_brand: dict[int, list[BrandBuyerMap]] = {}
    for mapping in buyer_maps:
        buyers_by_brand.setdefault(mapping.brand_id, []).append(mapping)

    for excel_row, data in parsed:
        name = _text(data.get("name") or data.get("code"))
        label = name or None
        if not name:
            result.errors.append(ImportErrorRow(excel_row, label, "name wajib diisi"))
            continue
        if len(name) > 50:
            result.errors.append(ImportErrorRow(excel_row, label, "name maksimal 50 karakter"))
            continue
        brand_id = _parse_int(data.get("id")) if "id" in present else None
        if "id" in present and data.get("id") and brand_id is None:
            result.errors.append(ImportErrorRow(excel_row, label, "id harus angka"))
            continue
        target = by_id.get(brand_id) if brand_id is not None else None
        if brand_id is not None and target is None:
            result.errors.append(ImportErrorRow(excel_row, label, f"Brand id {brand_id} tidak ditemukan"))
            continue
        owner = by_name.get(name.upper())
        if target is None:
            target = owner
        elif owner is not None and owner is not target:
            result.errors.append(ImportErrorRow(excel_row, label, "name sudah dipakai brand lain"))
            continue

        customer_tokens = _split_tokens(data.get("customer_codes", "")) if "customer_codes" in present else []
        if "customer_codes" in present:
            unknown = next((token for token in customer_tokens if token.upper() not in customers_by_code), None)
            if unknown:
                result.errors.append(ImportErrorRow(excel_row, label, f"Customer {unknown} tidak ditemukan"))
                continue

        creating = target is None
        if creating:
            target = Brand(name=name, code=name)
            db.add(target)
            db.flush()
            by_id[target.id] = target
            by_name[name.upper()] = target
        changed = creating
        if (target.name or "") != name or (target.code or "") != name:
            old_key = _text(target.code).upper()
            if old_key and old_key != name.upper():
                by_name.pop(old_key, None)
            target.name = name
            target.code = name
            by_name[name.upper()] = target
            changed = True
        if "customer_codes" in present:
            customer_changed, customer_error = _sync_customers(
                db, target, customer_tokens, customers_by_code, customers_by_brand
            )
            if customer_error:
                result.errors.append(ImportErrorRow(excel_row, label, customer_error))
                continue
            changed = customer_changed or changed
        if "buyer_codes" in present:
            changed = _sync_buyers(db, target, _split_tokens(data.get("buyer_codes", "")), buyers_by_brand) or changed
        if creating:
            result.created += 1
        elif changed:
            target.updated_at = app_now()
            result.updated += 1
        else:
            result.skipped += 1
    return result
