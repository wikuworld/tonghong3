"""Laporan Pending Payment (Report Mr. Bob) — selaras Excel REPORT TO MR.BOB."""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Invoice
from app.services.accounts_receivable_report import (
    invoice_balance,
    invoice_faktur_date,
    invoice_maturity,
    invoice_remark,
)
from app.services.pending_payment_master import PENDING_PAYMENT_MASTER

PENDING_PAYMENT_HEADERS = [
    "  NO.",
    "客戶 CUSTOMER",
    "品牌 BRAND",
    "付款條件 T/T PAYMENT",
    "每月結賬日 CLOSING DATE",
    "每月付款日 PAYMENT DATE ",
    "月份",
    "應收日",
    "應收金額(USD)",
    "回覆欄 REMARKS",
]


def normalize_customer_key(name: Optional[str]) -> str:
    if not name:
        return ""
    s = str(name).replace("\xa0", " ").upper().strip()
    s = re.sub(r"\s+", " ", s)
    s = s.replace(".,", ".").replace(" ,", ",")
    return s


def _format_report_date(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _format_short_date(d: date) -> str:
    return d.strftime("%d/%m/%y")


def _join_months(months: list[str]) -> Optional[str]:
    uniq: list[str] = []
    for m in months:
        if m and m not in uniq:
            uniq.append(m)
    return ",".join(uniq) if uniq else None


def _join_receivable_dates(dates: list[date]) -> Optional[str]:
    if not dates:
        return None
    seen: set[date] = set()
    ordered: list[date] = []
    for d in sorted(dates):
        if d not in seen:
            seen.add(d)
            ordered.append(d)
    parts = [_format_short_date(d) if d.year < 2000 else _format_report_date(d) for d in ordered]
    return "&".join(parts)


def _is_overdue(inv: Invoice, as_of: date) -> bool:
    if invoice_remark(inv) == "FINISH PAYMENT":
        return False
    if invoice_balance(inv) <= 0:
        return False
    maturity = invoice_maturity(inv)
    if maturity is None:
        return True
    return maturity <= as_of


def _load_overdue_invoices(db: Session, as_of: date) -> list[Invoice]:
    rows = list(db.scalars(select(Invoice).where(Invoice.voided_at.is_(None)).order_by(Invoice.id.asc())).all())
    return [inv for inv in rows if _is_overdue(inv, as_of)]


def _group_invoices_by_customer(invoices: list[Invoice]) -> dict[str, list[Invoice]]:
    grouped: dict[str, list[Invoice]] = {}
    for inv in invoices:
        key = normalize_customer_key(inv.customer_name)
        if not key:
            continue
        grouped.setdefault(key, []).append(inv)
    return grouped


def _aggregate_customer_invoices(invoices: list[Invoice]) -> dict[str, Any]:
    total = Decimal(0)
    months: list[str] = []
    dates: list[date] = []
    remarks: list[str] = []
    for inv in invoices:
        total += invoice_balance(inv)
        faktur = invoice_faktur_date(inv)
        if faktur:
            months.append(f"{faktur.month:02d}")
        maturity = invoice_maturity(inv)
        if maturity:
            dates.append(maturity)
        remark = (inv.payment_remark or "").strip()
        if remark and remark not in remarks:
            remarks.append(remark)
    amount = float(total)
    return {
        "month": _join_months(months),
        "receivable_date": _join_receivable_dates(dates),
        "amount_usd": amount if amount > 0 else 0,
        "remarks": " | ".join(remarks) if remarks else None,
        "invoice_count": len(invoices),
    }


def build_pending_payment_payload(
    db: Session,
    *,
    as_of: Optional[date] = None,
    year: Optional[int] = None,
) -> dict[str, Any]:
    report_date = as_of or date.today()
    report_year = year or report_date.year
    overdue = _load_overdue_invoices(db, report_date)
    grouped = _group_invoices_by_customer(overdue)

    rows: list[dict[str, Any]] = []
    total_amount = Decimal(0)
    overdue_customer_count = 0

    for master in PENDING_PAYMENT_MASTER:
        key = normalize_customer_key(master["customer"])
        invs = grouped.get(key, [])
        agg = _aggregate_customer_invoices(invs) if invs else {
            "month": None,
            "receivable_date": None,
            "amount_usd": 0,
            "remarks": None,
            "invoice_count": 0,
        }
        if agg["amount_usd"] > 0:
            overdue_customer_count += 1
            total_amount += Decimal(str(agg["amount_usd"]))

        rows.append(
            {
                "no": master["no"],
                "customer": master["customer"],
                "brand": master["brand"],
                "payment_term": master["payment_term"],
                "closing_date": master["closing_date"],
                "payment_date": master["payment_date"],
                "month": agg["month"],
                "receivable_date": agg["receivable_date"],
                "amount_usd": agg["amount_usd"],
                "remarks": agg["remarks"],
                "invoice_count": agg["invoice_count"],
            }
        )

    return {
        "report_title": f"印尼東紅逾齡賬款表 Pending Payment {report_year}",
        "updated_label": f"更新  ：{_format_report_date(report_date)}",
        "as_of": report_date.isoformat(),
        "year": report_year,
        "summary": {
            "master_count": len(PENDING_PAYMENT_MASTER),
            "overdue_customer_count": overdue_customer_count,
            "overdue_invoice_count": len(overdue),
            "amount_usd_total": float(total_amount),
        },
        "rows": rows,
    }


def render_pending_payment_xlsx(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = payload.get("as_of", "report").replace("-", "")[-6:] or "Sheet1"

    ws["A1"] = payload.get("report_title") or "Pending Payment"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = payload.get("updated_label") or ""

    for col_idx, header in enumerate(PENDING_PAYMENT_HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = Font(bold=True)

    rows = payload.get("rows") or []
    start_row = 4
    for row_idx, row in enumerate(rows, start=start_row):
        values = [
            row.get("no"),
            row.get("customer"),
            row.get("brand"),
            row.get("payment_term"),
            row.get("closing_date"),
            row.get("payment_date"),
            row.get("month"),
            row.get("receivable_date"),
            row.get("amount_usd") if row.get("amount_usd") else 0,
            row.get("remarks"),
        ]
        for col_idx, val in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    total_row = start_row + len(rows)
    summary = payload.get("summary") or {}
    ws.cell(row=total_row, column=2, value="總計 ：")
    ws.cell(row=total_row, column=9, value=summary.get("amount_usd_total", 0))

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()
