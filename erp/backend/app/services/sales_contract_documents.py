"""Kontrak Perjanjian Jual Beli (KPJB) — PDF & Word, dari multi surat jalan."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from html import escape
from pathlib import Path
from typing import Any, Sequence

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Customer, CustomerPic, DeliveryDetail, DeliveryHeader, PoDetail, ProductionDetail
from app.services.customs_numbering import generate_nomor_kontrak
from app.services.thickness_format import display_po_thickness_mm
from app.timeutil import app_now

TH_NAME = "PT.TONG HONG TANNERY INDONESIA"
TH_NAME_FULL = "PT.Tong Hong Tannery Indonesia"
TH_ADDRESS_L1 = "Kawasan Industri Modern II JL.Modern Industri XIV/W-1"
TH_ADDRESS_L2 = "Ds.Barengkok Kec.Kibin Serang-Banten"
TH_PHONE = "Telp. 0254-8480870 Fax. 0254-8480743"
TH_SIGNATORY = "CHIU PO HAO"
TH_POSITION = "Director"
TH_NPWP_DISPLAY = "21.147.268.3-401.000"
TH_CITY = "Cikande"

_LOGO_PATH = Path(__file__).resolve().parents[3] / "frontend" / "public" / "logo-tonghong.png"
_LOGO_ASPECT = 127 / 258
_LOGO_WIDTH_MM = 28


def _logo_bytes() -> bytes | None:
    try:
        data = _LOGO_PATH.read_bytes()
    except OSError:
        return None
    return data or None


def _fmt_date_long(d: date | None) -> str:
    if not d:
        return "—"
    months = (
        "Jan.",
        "Feb.",
        "Mar.",
        "Apr.",
        "May",
        "Jun.",
        "Jul.",
        "Aug.",
        "Sep.",
        "Oct.",
        "Nov.",
        "Dec.",
    )
    return f"{months[d.month - 1]}{d.day:02d}, {d.year}"


def _dec(v: Decimal | None) -> str:
    if v is None:
        return ""
    return f"{v.quantize(Decimal('0.01')):,}"


def _money2(v: Decimal | None) -> Decimal | None:
    if v is None:
        return None
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _fmt_npwp(raw: str | None) -> str:
    s = "".join(ch for ch in str(raw or "") if ch.isdigit())
    if len(s) >= 15:
        core = s[:15]
        return f"{core[0:2]}.{core[2:5]}.{core[5:8]}.{core[8]}-{core[9:12]}.{core[12:15]}"
    return (raw or "").strip() or "—"


def _buyer_pic(db: Session, customer: Customer | None) -> tuple[str, str]:
    if not customer:
        return "—", "—"
    name = (customer.contact_person or "").strip()
    title = (customer.position or "").strip()
    if name:
        return name, title or "—"
    pics = list(
        db.scalars(
            select(CustomerPic)
            .where(CustomerPic.customer_id == int(customer.id))
            .order_by(CustomerPic.id.asc())
        ).all()
    )
    if pics:
        p0 = pics[0]
        return (p0.name or "—").strip() or "—", (p0.title or "—").strip() or "—"
    return "—", "—"


def _parse_delivery_ids(delivery_ids: Sequence[int] | int) -> list[int]:
    if isinstance(delivery_ids, int):
        raw = [delivery_ids]
    else:
        raw = [int(x) for x in delivery_ids]
    seen: set[int] = set()
    out: list[int] = []
    for i in raw:
        if i in seen:
            continue
        seen.add(i)
        out.append(i)
    if not out:
        raise ValueError("delivery_ids_required")
    return out


def _line_description(pod: PoDetail | None) -> str:
    if not pod:
        return "—"
    thick = display_po_thickness_mm(pod) or ""
    material = (pod.material_name or "").strip()
    emboss = (pod.emboss or "").strip()
    parts = [p for p in [thick, material, emboss] if p]
    return " ".join(parts).strip() or material or "—"


def gather_sales_contract_ctx(
    db: Session,
    delivery_ids: Sequence[int] | int,
) -> dict[str, Any]:
    ids = _parse_delivery_ids(delivery_ids)
    headers: list[DeliveryHeader] = []
    for did in ids:
        dh = db.get(DeliveryHeader, did)
        if not dh:
            raise ValueError("delivery_not_found")
        headers.append(dh)

    customer_ids = {int(h.customer_id) for h in headers if h.customer_id is not None}
    if len(customer_ids) > 1:
        raise ValueError("delivery_customer_mismatch")
    if not customer_ids:
        raise ValueError("delivery_customer_missing")

    customer = db.get(Customer, next(iter(customer_ids)))
    dates = [h.delivery_date for h in headers if h.delivery_date]
    on_date = min(dates) if dates else app_now().date()
    cust_code = (customer.code if customer and customer.code else "").strip()
    kontrak_no = generate_nomor_kontrak(cust_code, on_date)

    buyer_name, buyer_title = _buyer_pic(db, customer)
    cust_name = (customer.name if customer and customer.name else "").strip() or "—"
    cust_addr = ""
    if customer:
        cust_addr = (customer.address or customer.billing_address or "").strip()
    cust_npwp = _fmt_npwp(customer.tax_id if customer else None)

    lines: list[dict[str, Any]] = []
    total_qty = Decimal(0)
    total_amt = Decimal(0)
    for dh in headers:
        sj_no = (dh.delivery_number or str(dh.id)).strip()
        details = list(
            db.scalars(
                select(DeliveryDetail)
                .where(DeliveryDetail.delivery_id == int(dh.id))
                .order_by(DeliveryDetail.id.asc())
            ).all()
        )
        for dd in details:
            pod: PoDetail | None = None
            if dd.po_detail_id:
                pod = db.get(PoDetail, int(dd.po_detail_id))
            if pod is None and dd.production_detail_id:
                pd = db.get(ProductionDetail, int(dd.production_detail_id))
                if pd and pd.po_detail_id:
                    pod = db.get(PoDetail, int(pd.po_detail_id))

            qty = Decimal(str(dd.qty or 0))
            up = pod.unit_price if pod and pod.unit_price is not None else None
            amt = _money2(qty * up) if up is not None else None
            if amt is not None:
                total_amt += amt
            total_qty += qty
            lines.append(
                {
                    "description": _line_description(pod),
                    "color": ((pod.color if pod else None) or "").strip(),
                    "sj": sj_no,
                    "unit": "S/F",
                    "qty": qty,
                    "unit_price": up,
                    "amount": amt,
                }
            )

    return {
        "kontrak_no": kontrak_no,
        "delivery_ids": ids,
        "delivery_numbers": [(h.delivery_number or str(h.id)).strip() for h in headers],
        "order_date": on_date,
        "order_date_fmt": _fmt_date_long(on_date),
        "customer_name": cust_name,
        "customer_address": cust_addr or "—",
        "customer_npwp": cust_npwp,
        "buyer_name": buyer_name,
        "buyer_title": buyer_title,
        "lines": lines,
        "total_qty": total_qty,
        "total_amount": total_amt,
        "currency": (customer.default_currency if customer and customer.default_currency else "USD").upper(),
    }


def _header_letterhead_pdf(center_style: ParagraphStyle, company_style: ParagraphStyle) -> Any:
    """Logo kiri + blok perusahaan (judul + alamat) di kanan."""
    company_block = [
        Paragraph("KONTRAK PERJANJIAN JUAL BELI", company_style),
        Paragraph(escape(TH_NAME), company_style),
        Paragraph(escape(TH_ADDRESS_L1), center_style),
        Paragraph(escape(TH_ADDRESS_L2), center_style),
        Paragraph(escape(TH_PHONE), center_style),
    ]
    logo = _logo_bytes()
    if not logo:
        return company_block

    logo_h = _LOGO_WIDTH_MM * _LOGO_ASPECT * mm
    img = RLImage(io.BytesIO(logo), width=_LOGO_WIDTH_MM * mm, height=logo_h)
    img.hAlign = "LEFT"
    letterhead = Table(
        [[img, company_block]],
        colWidths=[_LOGO_WIDTH_MM * mm + 4 * mm, 150 * mm],
    )
    letterhead.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (0, 0), 6),
                ("RIGHTPADDING", (1, 0), (1, 0), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ("ALIGN", (0, 0), (0, 0), "LEFT"),
                ("ALIGN", (1, 0), (1, 0), "CENTER"),
            ]
        )
    )
    return [letterhead]


def render_sales_contract_pdf(ctx: dict[str, Any]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    company = ParagraphStyle(
        name="kpjbCo",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        alignment=TA_CENTER,
        leading=13,
    )
    center = ParagraphStyle(
        name="kpjbCenter",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_CENTER,
        leading=11,
    )
    body = ParagraphStyle(
        name="kpjbBody",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_JUSTIFY,
        leading=12,
    )
    left = ParagraphStyle(
        name="kpjbLeft",
        parent=styles["Normal"],
        fontSize=9,
        alignment=TA_LEFT,
        leading=12,
    )
    small = ParagraphStyle(
        name="kpjbSmall",
        parent=styles["Normal"],
        fontSize=8,
        alignment=TA_LEFT,
        leading=10,
    )
    right = ParagraphStyle(
        name="kpjbRight",
        parent=styles["Normal"],
        fontSize=8,
        alignment=TA_RIGHT,
        leading=10,
    )

    story: list[Any] = []
    story.extend(_header_letterhead_pdf(center, company))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"<b>{escape(ctx['kontrak_no'])}</b>", center))
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("Yang bertanda tangan dibawah ini:", left))
    story.append(Spacer(1, 2 * mm))

    party1 = (
        f"1.&nbsp;&nbsp;<b>Nama</b>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;:{escape(TH_SIGNATORY)}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Jabatan</b>&nbsp;&nbsp;&nbsp;:{escape(TH_POSITION)}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Perusahaan</b>&nbsp;:{escape(TH_NAME)}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Alamat</b>&nbsp;&nbsp;&nbsp;&nbsp;:{escape(TH_ADDRESS_L1)}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;"
        f"{escape(TH_ADDRESS_L2)}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>NPWP NO.</b>&nbsp;:{escape(TH_NPWP_DISPLAY)}<br/>"
        f"bertindak atas nama {escape(TH_NAME_FULL)}, sebagai pihak yang mengeluarkan barang / Penjual yang "
        f"selanjutnya disebut sebagai pihak I (Pertama)."
    )
    story.append(Paragraph(party1, body))
    story.append(Spacer(1, 3 * mm))

    party2 = (
        f"2.&nbsp;&nbsp;<b>Nama</b>&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;: {escape(ctx['buyer_name'])}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Jabatan</b>&nbsp;&nbsp;&nbsp;: {escape(ctx['buyer_title'])}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Perusahaan</b>&nbsp;: {escape(ctx['customer_name'])}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>Alamat</b>&nbsp;&nbsp;&nbsp;&nbsp;: {escape(ctx['customer_address'])}<br/>"
        f"&nbsp;&nbsp;&nbsp;&nbsp;<b>NPWP NO.</b>&nbsp;: {escape(ctx['customer_npwp'])}<br/>"
        f"bertindak untuk dan atas nama {escape(ctx['customer_name'])} sebagai pihak yang memasukan barang / Pembeli yang "
        f"selanjutnya disebut pihak II (Kedua)."
    )
    story.append(Paragraph(party2, body))
    story.append(Spacer(1, 3 * mm))
    story.append(
        Paragraph(
            "Pihak I (Pertama) dan Pihak II (Kedua) dengan ini mengadakan Kontrak perjanjian jual beli atas barang / "
            "Raw Material Leather sebagai berikut :",
            body,
        )
    )
    story.append(Spacer(1, 3 * mm))

    cur = escape(ctx.get("currency") or "USD")
    head = [
        Paragraph("<b>Description</b>", small),
        Paragraph("<b>Color</b>", small),
        Paragraph("<b>No. SJ</b>", small),
        Paragraph("<b>Unit</b>", small),
        Paragraph("<b>Quantity</b>", small),
        Paragraph(f"<b>Price<br/>({cur})</b>", small),
        Paragraph(f"<b>Amount<br/>({cur})</b>", small),
    ]
    data: list[list[Any]] = [head]
    for ln in ctx["lines"]:
        data.append(
            [
                Paragraph(escape(ln["description"] or ""), small),
                Paragraph(escape(ln["color"] or ""), small),
                Paragraph(escape(ln.get("sj") or ""), small),
                Paragraph(escape(ln["unit"] or "S/F"), small),
                Paragraph(_dec(ln["qty"]), small),
                Paragraph(_dec(ln["unit_price"]) if ln["unit_price"] is not None else "", small),
                Paragraph(_dec(ln["amount"]) if ln["amount"] is not None else "", small),
            ]
        )
    data.append(
        [
            Paragraph("<b>TOTAL</b>", small),
            "",
            "",
            "",
            Paragraph(f"<b>{_dec(ctx['total_qty'])}</b>", small),
            "",
            Paragraph(f"<b>{_dec(ctx['total_amount'])}</b>", small),
        ]
    )
    col_w = [50 * mm, 30 * mm, 35 * mm, 14 * mm, 22 * mm, 20 * mm, 24 * mm]
    tbl = Table(data, colWidths=col_w, repeatRows=1)
    tbl.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.black),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.92, 0.92, 0.92)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(tbl)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Demikian perjanjian kontrak ini dibuat dengan sebenarnya.", left))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"{escape(TH_CITY)},", left))
    story.append(Paragraph(escape(ctx["order_date_fmt"]), left))
    story.append(Spacer(1, 6 * mm))

    sign = Table(
        [
            [
                Paragraph(f"<b>Pihak Pertama</b><br/>{escape(TH_NAME_FULL)}", center),
                Paragraph(f"<b>Pihak Kedua</b><br/>{escape(ctx['customer_name'])}", center),
            ],
            [Spacer(1, 18 * mm), Spacer(1, 18 * mm)],
            [
                Paragraph(f"({escape(TH_SIGNATORY)})", center),
                Paragraph(f"({escape(ctx['buyer_name'])})", center),
            ],
        ],
        colWidths=[90 * mm, 90 * mm],
    )
    sign.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    story.append(sign)
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Page 1 of 1", right))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def render_sales_contract_docx(ctx: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Mm, Pt

    d = Document()

    def set_cell_valign(cell, valign: str = "center") -> None:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        vAlign = OxmlElement("w:vAlign")
        vAlign.set(qn("w:val"), valign)
        tcPr.append(vAlign)

    logo = _logo_bytes()
    head_tbl = d.add_table(rows=1, cols=2)
    head_tbl.autofit = True
    left_cell = head_tbl.rows[0].cells[0]
    right_cell = head_tbl.rows[0].cells[1]
    set_cell_valign(left_cell, "center")
    set_cell_valign(right_cell, "center")

    if logo:
        p_logo = left_cell.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p_logo.add_run()
        run.add_picture(io.BytesIO(logo), width=Mm(_LOGO_WIDTH_MM))

    right_lines = [
        ("KONTRAK PERJANJIAN JUAL BELI", True, 13),
        (TH_NAME, True, 11),
        (TH_ADDRESS_L1, False, 9),
        (TH_ADDRESS_L2, False, 9),
        (TH_PHONE, False, 9),
    ]
    for idx, (text, bold, size) in enumerate(right_lines):
        p = right_cell.paragraphs[0] if idx == 0 else right_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)

    def add_center(text: str, *, bold: bool = False, size: int = 11) -> None:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)

    add_center(ctx["kontrak_no"], bold=True, size=10)
    d.add_paragraph("Yang bertanda tangan dibawah ini:")
    d.add_paragraph(
        f"1. Nama : {TH_SIGNATORY}\n"
        f"   Jabatan : {TH_POSITION}\n"
        f"   Perusahaan : {TH_NAME}\n"
        f"   Alamat : {TH_ADDRESS_L1}, {TH_ADDRESS_L2}\n"
        f"   NPWP NO. : {TH_NPWP_DISPLAY}\n"
        f"bertindak atas nama {TH_NAME_FULL}, sebagai pihak yang mengeluarkan barang / Penjual "
        f"yang selanjutnya disebut sebagai pihak I (Pertama)."
    )
    d.add_paragraph(
        f"2. Nama : {ctx['buyer_name']}\n"
        f"   Jabatan : {ctx['buyer_title']}\n"
        f"   Perusahaan : {ctx['customer_name']}\n"
        f"   Alamat : {ctx['customer_address']}\n"
        f"   NPWP NO. : {ctx['customer_npwp']}\n"
        f"bertindak untuk dan atas nama {ctx['customer_name']} sebagai pihak yang memasukan barang / Pembeli "
        f"yang selanjutnya disebut pihak II (Kedua)."
    )
    d.add_paragraph(
        "Pihak I (Pertama) dan Pihak II (Kedua) dengan ini mengadakan Kontrak perjanjian jual beli atas barang / "
        "Raw Material Leather sebagai berikut :"
    )

    table = d.add_table(rows=1 + len(ctx["lines"]) + 1, cols=7)
    hdr = [
        "Description",
        "Color",
        "No. SJ",
        "Unit",
        "Quantity",
        f"Price ({ctx.get('currency') or 'USD'})",
        f"Amount ({ctx.get('currency') or 'USD'})",
    ]
    for j, h in enumerate(hdr):
        table.rows[0].cells[j].text = h
    for i, ln in enumerate(ctx["lines"], start=1):
        row = table.rows[i].cells
        row[0].text = ln["description"] or ""
        row[1].text = ln["color"] or ""
        row[2].text = ln.get("sj") or ""
        row[3].text = ln["unit"] or "S/F"
        row[4].text = _dec(ln["qty"])
        row[5].text = _dec(ln["unit_price"]) if ln["unit_price"] is not None else ""
        row[6].text = _dec(ln["amount"]) if ln["amount"] is not None else ""
    last = table.rows[len(ctx["lines"]) + 1].cells
    last[0].text = "TOTAL"
    last[4].text = _dec(ctx["total_qty"])
    last[6].text = _dec(ctx["total_amount"])

    d.add_paragraph("Demikian perjanjian kontrak ini dibuat dengan sebenarnya.")
    d.add_paragraph(f"{TH_CITY},")
    d.add_paragraph(ctx["order_date_fmt"])
    d.add_paragraph()
    d.add_paragraph(f"Pihak Pertama\n{TH_NAME_FULL}\n\n\n({TH_SIGNATORY})")
    d.add_paragraph(f"Pihak Kedua\n{ctx['customer_name']}\n\n\n({ctx['buyer_name']})")

    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf.getvalue()
