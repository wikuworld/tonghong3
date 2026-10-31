"""Export/import Excel master Color Swatch.

Kunci unik impor: factory_key + buyer_code + color_code + material_name + emboss + color_standard.
Baris yang sudah ada diperbarui. File yang diunggah ulang tidak menambah duplikat
dan tidak menonaktifkan resep di luar file.
"""

from __future__ import annotations

import io
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Brand, ColorSwatchRecipe, CsiColorMaster, Customer
from app.services.buyer_registry import list_active_buyer_codes
from app.services.color_swatch import normalize_factory_key, suggest_next_color_standard
from app.services.swatch_factory_map import resolve_swatch_customer_id
from app.timeutil import app_now

EXPORT_COLUMNS = [
    "id",
    "factory_key",
    "customer_code",
    "buyer_code",
    "color_code",
    "color_name",
    "emboss",
    "color_standard",
    "backing",
    "material_name",
    "legacy_urut",
    "entered_at",
    "csi_color_code",
    "is_active",
    "created_at",
    "updated_at",
]

TEMPLATE_COLUMNS = [
    "factory_key",
    "buyer_code",
    "color_code",
    "color_name",
    "emboss",
    "color_standard",
    "backing",
    "material_name",
    "legacy_urut",
    "entered_at",
    "customer_code",
    "csi_color_code",
    "is_active",
]

_ROW_FIELDS = (
    "id",
    "factory_key",
    "customer_code",
    "buyer_code",
    "color_code",
    "color_name",
    "emboss",
    "color_standard",
    "backing",
    "material_name",
    "legacy_urut",
    "entered_at",
    "csi_color_code",
    "is_active",
    "created_at",
    "updated_at",
)

_REQUIRED_HEADERS = ("factory_key", "buyer_code", "color_name", "material_name")

_HEADER_ALIAS = {
    "id": "id",
    "factory_key": "factory_key",
    "factory": "factory_key",
    "pabrik": "factory_key",
    "customer": "factory_key",
    "buyer_code": "buyer_code",
    "buyer": "buyer_code",
    "brand": "buyer_code",
    "color_code": "color_code",
    "kode": "color_code",
    "code": "color_code",
    "color_name": "color_name",
    "color": "color_name",
    "warna": "color_name",
    "emboss": "emboss",
    "embos": "emboss",
    "color_standard": "color_standard",
    "color_swatch": "color_standard",
    "swatch": "color_standard",
    "backing": "backing",
    "material_name": "material_name",
    "material": "material_name",
    "legacy_urut": "legacy_urut",
    "no_urut": "legacy_urut",
    "urut": "legacy_urut",
    "entered_at": "entered_at",
    "tggl_masuk": "entered_at",
    "tanggal_masuk": "entered_at",
    "customer_code": "customer_code",
    "kode_pabrik": "customer_code",
    "csi_color_code": "csi_color_code",
    "csi": "csi_color_code",
    "is_active": "is_active",
    "active": "is_active",
    "aktif": "is_active",
    "created_at": "created_at",
    "updated_at": "updated_at",
}

_TRUE_VALUES = {"1", "true", "yes", "y", "aktif", "active"}
_FALSE_VALUES = {"0", "false", "no", "n", "nonaktif", "inactive", "tidak"}


@dataclass
class ImportErrorRow:
    row: int
    color_code: str | None
    reason: str


@dataclass
class ColorSwatchImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    total_rows: int = 0
    errors: list[ImportErrorRow] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "skipped": self.skipped,
            "total_rows": self.total_rows,
            "errors": [asdict(e) for e in self.errors],
        }


def _normalize_header(raw: Any) -> str:
    s = str(raw or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in s:
        s = s.replace("__", "_")
    return s


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val).strip()
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)) and 20000 < float(val) < 80000:
        stamp = datetime(1899, 12, 30) + timedelta(days=float(val))
        return stamp.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _text(raw: Any) -> str:
    return str(raw or "").strip()


def _blank(raw: Any) -> str | None:
    text = _text(raw)
    return text or None


def _business_key(
    factory: str,
    buyer: str,
    color_code: str,
    material: str,
    emboss: str,
    color_standard: str | None,
) -> tuple[str, str, str, str, str, str]:
    return (
        normalize_factory_key(factory),
        _text(buyer).upper(),
        _text(color_code),
        _text(material),
        _text(emboss),
        _text(color_standard),
    )


def _key_from_row(row: ColorSwatchRecipe) -> tuple[str, str, str, str, str, str]:
    return _business_key(
        row.factory_key or "",
        row.buyer_code or "",
        row.color_code or "",
        row.material_name or "",
        row.emboss or "",
        row.color_standard,
    )


def _parse_dt(raw: str) -> datetime | None:
    text = (raw or "").strip().replace("T", " ")
    if not text:
        return None
    if len(text) > 19:
        text = text[:19]
    for fmt, size in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%d/%m/%Y", 10),
    ):
        try:
            return datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
    return None


def _parse_date(raw: str) -> date | None:
    parsed = _parse_dt(raw)
    return parsed.date() if parsed else None


def _format_dt(value: datetime | None) -> str:
    if not value:
        return ""
    return value.strftime("%Y-%m-%d %H:%M:%S")


def _format_date(value: date | datetime | None) -> str:
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.date().strftime("%Y-%m-%d")
    return value.strftime("%Y-%m-%d")


def _parse_active(raw: str) -> bool | None:
    text = (raw or "").strip().lower()
    if not text:
        return True
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def _parse_id(raw: str) -> int | None:
    text = (raw or "").strip()
    if not text:
        return None
    try:
        value = int(float(text)) if "." in text else int(text)
    except ValueError:
        return None
    return value if value > 0 else None


def _workbook_bytes(ws_title: str, headers: list[str], data_rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = ws_title
    for col_idx, col in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=col)
    for row_idx, values in enumerate(data_rows, 2):
        for col_idx, val in enumerate(values, 1):
            ws.cell(row=row_idx, column=col_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_color_swatches_xlsx(db: Session) -> bytes:
    rows = list(
        db.scalars(
            select(ColorSwatchRecipe).order_by(
                ColorSwatchRecipe.factory_key.asc(),
                ColorSwatchRecipe.buyer_code.asc(),
                ColorSwatchRecipe.color_code.asc(),
                ColorSwatchRecipe.material_name.asc(),
                ColorSwatchRecipe.id.asc(),
            )
        ).all()
    )
    customers = {
        int(cust.id): (cust.code or "")
        for cust in db.scalars(select(Customer)).all()
        if getattr(cust, "id", None) is not None
    }
    csi_codes = {
        int(csi.id): (csi.color_code or "")
        for csi in db.scalars(select(CsiColorMaster)).all()
        if getattr(csi, "id", None) is not None
    }
    data = [
        [
            row.id,
            row.factory_key or "",
            customers.get(int(row.customer_id), "") if row.customer_id else "",
            row.buyer_code or "",
            row.color_code or "",
            row.color_name or "",
            row.emboss or "",
            row.color_standard or "",
            row.backing or "",
            row.material_name or "",
            row.legacy_urut or "",
            _format_date(row.entered_at),
            csi_codes.get(int(row.csi_color_master_id), "") if row.csi_color_master_id else "",
            "0" if row.is_active is False else "1",
            _format_dt(row.created_at),
            _format_dt(row.updated_at),
        ]
        for row in rows
    ]
    return _workbook_bytes("Color Swatch", EXPORT_COLUMNS, data)


def export_color_swatches_template_xlsx() -> bytes:
    return _workbook_bytes("Color Swatch", TEMPLATE_COLUMNS, [])


def _parse_rows(content: bytes) -> tuple[list[dict[str, str]], list[ImportErrorRow]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel tidak bisa dibaca")]
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], [ImportErrorRow(row=1, color_code=None, reason="File Excel kosong")]

    col_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row or []):
        key = _HEADER_ALIAS.get(_normalize_header(cell))
        if key and key not in col_map.values():
            col_map[idx] = key

    missing = [name for name in _REQUIRED_HEADERS if name not in col_map.values()]
    if missing:
        return [], [
            ImportErrorRow(
                row=1,
                color_code=None,
                reason="Header wajib memuat kolom factory_key, buyer_code, color_name, dan material_name",
            )
        ]

    parsed: list[dict[str, str]] = []
    for excel_row_num, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {key: "" for key in _ROW_FIELDS}
        for idx, key in col_map.items():
            if idx < len(row) and key in item:
                item[key] = _cell_str(row[idx])
        if not any(item.get(k) for k in ("id", "factory_key", "buyer_code", "color_name", "material_name", "color_code")):
            continue
        item["_row"] = str(excel_row_num)
        item["_present"] = ",".join(col_map.values())
        parsed.append(item)
    return parsed, []


def _brand_id_for_buyer(brands: list[Brand], buyer: str) -> int | None:
    token = buyer.strip().upper()
    if not token:
        return None
    for brand in brands:
        code = _text(getattr(brand, "code", "")).upper()
        name = _text(getattr(brand, "name", "")).upper()
        if token == code or token == name:
            brand_id = getattr(brand, "id", None)
            return int(brand_id) if brand_id is not None else None
    return None


def _allocate_standard(
    db: Session,
    *,
    customer_code: str | None,
    factory_key: str,
    entered: date | None,
    used: set[str],
) -> str | None:
    suggested = suggest_next_color_standard(
        db,
        customer_code=customer_code or None,
        factory_key=factory_key,
        on_date=entered,
    )
    if not suggested:
        return None
    if len(suggested) >= 2 and suggested[-2:].isdigit():
        stem = suggested[:-2]
        seq = int(suggested[-2:])
    else:
        stem = suggested
        seq = 1
    while suggested in used:
        seq += 1
        suggested = f"{stem}{seq:02d}"
    used.add(suggested)
    return suggested


def import_color_swatches_xlsx(db: Session, content: bytes) -> ColorSwatchImportResult:
    result = ColorSwatchImportResult()
    rows, parse_errors = _parse_rows(content)
    result.errors.extend(parse_errors)
    if parse_errors and not rows:
        return result

    known_buyers = {code.strip().upper() for code in list_active_buyer_codes(db) if code}
    existing = list(db.scalars(select(ColorSwatchRecipe)).all())
    customers = [cust for cust in db.scalars(select(Customer)).all() if getattr(cust, "id", None) is not None]
    customers_by_code: dict[str, Customer] = {}
    customers_by_id: dict[int, Customer] = {}
    for cust in customers:
        customers_by_id[int(cust.id)] = cust
        code = _text(getattr(cust, "code", "")).upper()
        if code:
            customers_by_code.setdefault(code, cust)
    csi_by_key: dict[tuple[str, str], CsiColorMaster] = {}
    for csi in db.scalars(select(CsiColorMaster)).all():
        key = (_text(getattr(csi, "buyer_code", "")).upper(), _text(getattr(csi, "color_code", "")).upper())
        if key[0] and key[1]:
            csi_by_key.setdefault(key, csi)
    brands = list(db.scalars(select(Brand)).all())

    by_id = {row.id: row for row in existing if row.id is not None}
    by_key: dict[tuple[str, str, str, str, str, str], ColorSwatchRecipe] = {}
    used_standards: set[str] = set()
    for row in sorted(existing, key=lambda item: item.id or 0):
        by_key.setdefault(_key_from_row(row), row)
        standard = _text(row.color_standard)
        if standard:
            used_standards.add(standard)
    seen_ids: set[int] = set()

    for item in rows:
        result.total_rows += 1
        excel_row = int(item.pop("_row", "0") or 0)
        present = {part for part in (item.pop("_present", "") or "").split(",") if part}
        code_label = _text(item.get("color_code")) or None
        raw_id = _text(item.get("id"))
        row_id = _parse_id(raw_id)
        if raw_id and row_id is None:
            result.skipped += 1
            result.errors.append(ImportErrorRow(row=excel_row, color_code=code_label, reason="id tidak valid"))
            continue
        if row_id is not None and row_id in seen_ids:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason=f"id {row_id} muncul lebih dari sekali di file",
                )
            )
            continue

        factory = normalize_factory_key(item.get("factory_key"))
        buyer = _text(item.get("buyer_code")).upper()
        color_code = _text(item.get("color_code"))
        color_name = _text(item.get("color_name"))
        material = _text(item.get("material_name"))
        emboss = _text(item.get("emboss"))
        file_std = _text(item.get("color_standard"))
        if not factory or not buyer or not color_name or not material:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=code_label,
                    reason="factory_key, buyer_code, color_name, dan material_name wajib diisi",
                )
            )
            continue
        if buyer not in known_buyers:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason=f"Buyer '{buyer}' tidak ditemukan di registry",
                )
            )
            continue

        customer_code = _text(item.get("customer_code")).upper()
        customer_id: int | None
        if customer_code:
            cust = customers_by_code.get(customer_code)
            if cust is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason=f"Customer '{customer_code}' tidak ditemukan",
                    )
                )
                continue
            customer_id = int(cust.id)
        else:
            customer_id = resolve_swatch_customer_id(
                db, factory, color_standard=file_std or None
            )

        csi_id: int | None = None
        csi_specified = "csi_color_code" in present and bool(_text(item.get("csi_color_code")))
        if csi_specified:
            csi_code = _text(item.get("csi_color_code")).upper()
            csi = csi_by_key.get((buyer, csi_code))
            if csi is None or getattr(csi, "id", None) is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason=f"CSI '{csi_code}' tidak ditemukan untuk buyer {buyer}",
                    )
                )
                continue
            csi_id = int(csi.id)

        entered: date | None = None
        if "entered_at" in present and _text(item.get("entered_at")):
            entered = _parse_date(item.get("entered_at") or "")
            if entered is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason="entered_at tidak valid",
                    )
                )
                continue

        is_active = True
        if "is_active" in present:
            parsed_active = _parse_active(item.get("is_active") or "")
            if parsed_active is None:
                result.skipped += 1
                result.errors.append(
                    ImportErrorRow(
                        row=excel_row,
                        color_code=color_code or code_label,
                        reason="is_active tidak valid",
                    )
                )
                continue
            is_active = parsed_active

        created_at = _parse_dt(item.get("created_at") or "") if "created_at" in present else None
        updated_at = _parse_dt(item.get("updated_at") or "") if "updated_at" in present else None
        if "created_at" in present and _text(item.get("created_at")) and created_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=color_code or code_label, reason="created_at tidak valid")
            )
            continue
        if "updated_at" in present and _text(item.get("updated_at")) and updated_at is None:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(row=excel_row, color_code=color_code or code_label, reason="updated_at tidak valid")
            )
            continue

        key = _business_key(factory, buyer, color_code, material, emboss, file_std)
        owner = by_key.get(key)
        target = by_id.get(row_id) if row_id is not None else None
        if owner is not None and target is not None and owner is not target:
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason="kunci resep sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue
        if target is None:
            target = owner

        if target is None and not file_std:
            linked = customers_by_id.get(customer_id) if customer_id is not None else None
            linked_code = _text(getattr(linked, "code", "")) if linked is not None else ""
            suggested = _allocate_standard(
                db,
                customer_code=linked_code or None,
                factory_key=factory,
                entered=entered,
                used=used_standards,
            )
            if suggested:
                file_std = suggested
                key = _business_key(factory, buyer, color_code, material, emboss, file_std)
                sug_owner = by_key.get(key)
                if sug_owner is not None:
                    target = sug_owner

        if row_id is not None:
            seen_ids.add(row_id)
        if file_std:
            used_standards.add(file_std)

        backing = _blank(item.get("backing")) if "backing" in present else None
        legacy_urut = _blank(item.get("legacy_urut")) if "legacy_urut" in present else None
        stored_std = file_std or None
        brand_id = _brand_id_for_buyer(brands, buyer)

        if target is None:
            created = ColorSwatchRecipe(
                factory_key=factory,
                customer_id=customer_id,
                buyer_code=buyer,
                brand_id=brand_id,
                csi_color_master_id=csi_id if csi_specified else None,
                color_code=color_code,
                color_name=color_name,
                emboss=emboss,
                color_standard=stored_std,
                backing=backing,
                material_name=material,
                legacy_urut=legacy_urut,
                entered_at=entered if "entered_at" in present else None,
                is_active=is_active if "is_active" in present else True,
                created_at=created_at or app_now(),
                updated_at=updated_at,
            )
            db.add(created)
            by_key[key] = created
            result.created += 1
            continue

        old_key = _key_from_row(target)
        if old_key != key and by_key.get(key) not in (None, target):
            result.skipped += 1
            result.errors.append(
                ImportErrorRow(
                    row=excel_row,
                    color_code=color_code or code_label,
                    reason="kunci resep sudah ada di master, baris tidak diduplikasi",
                )
            )
            continue

        same_entered = True
        if "entered_at" in present:
            same_entered = target.entered_at == entered
        same_backing = True if "backing" not in present else _text(target.backing) == _text(backing)
        same_legacy = True if "legacy_urut" not in present else _text(target.legacy_urut) == _text(legacy_urut)
        same_active = True if "is_active" not in present else bool(target.is_active) == is_active
        same_csi = True if not csi_specified else target.csi_color_master_id == csi_id
        same_created = created_at is None or _format_dt(target.created_at) == _format_dt(created_at)
        same_updated = updated_at is None or _format_dt(target.updated_at) == _format_dt(updated_at)
        same_fields = (
            _text(target.factory_key) == factory
            and _text(target.buyer_code).upper() == buyer
            and _text(target.color_code) == color_code
            and _text(target.color_name) == color_name
            and _text(target.emboss) == emboss
            and _text(target.color_standard) == _text(stored_std)
            and _text(target.material_name) == material
            and target.customer_id == customer_id
            and target.brand_id == brand_id
            and same_backing
            and same_legacy
            and same_entered
            and same_active
            and same_csi
            and same_created
            and same_updated
        )
        if same_fields:
            result.skipped += 1
            continue

        target.factory_key = factory
        target.buyer_code = buyer
        target.color_code = color_code
        target.color_name = color_name
        target.emboss = emboss
        target.color_standard = stored_std
        target.material_name = material
        target.customer_id = customer_id
        target.brand_id = brand_id
        if "backing" in present:
            target.backing = backing
        if "legacy_urut" in present:
            target.legacy_urut = legacy_urut
        if "entered_at" in present:
            target.entered_at = entered
        if "is_active" in present:
            target.is_active = is_active
        if csi_specified:
            target.csi_color_master_id = csi_id
        if created_at is not None:
            target.created_at = created_at
        target.updated_at = updated_at if updated_at is not None else app_now()
        if old_key != key and by_key.get(old_key) is target:
            by_key.pop(old_key, None)
        by_key[key] = target
        result.updated += 1

    return result
